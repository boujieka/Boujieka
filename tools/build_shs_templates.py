"""Build the Volume 2 Templates Pack (SHS PAYGo): Excel and Word templates in the house style.

Run: python tools/build_shs_templates.py
Writes volumes/02-solar-home-systems/templates/AEF_V2_T0x_*.xlsx / .docx

Excel conventions follow the model: blue font on cream = input; black = formula; every sheet carries a printed footer
with the sheet name and "Page x of y". Word templates carry a contents field, running header and page numbers.
Worked examples are illustrative and use the fictional SolaraPay case where data are needed.
"""
import copy
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import shs_defaults as D  # noqa: E402
from cases import solarapay as SP  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "volumes/02-solar-home-systems/templates"
AUTHOR = "Emmanuel Boujieka Kamga"
HOUSE = "Africa Energy Finance"
VER = "1.0"
GREEN, GOLD, CREAM, GREY = "0B3020", "B07C0F", "FFF8E1", "666666"
FONT = "Arial"

F_IN = Font(name=FONT, size=10, color="0000FF")
F_CALC = Font(name=FONT, size=10, color="000000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, size=9, italic=True, color=GREY)
FILL_IN = PatternFill("solid", fgColor=CREAM)
FILL_HEAD = PatternFill("solid", fgColor=GREEN)
FILL_BAND = PatternFill("solid", fgColor=GREEN)
FILL_SUB = PatternFill("solid", fgColor="F2EEE3")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GOLD_LINE = Border(bottom=Side(style="medium", color=GOLD))
WRAP = Alignment(wrap_text=True, vertical="top")
NUM = '#,##0;(#,##0);0'
NUM1 = '#,##0.0;(#,##0.0);0.0'
PCT = '0.0%;(0.0%);0.0%'
RAT = '0.00"x";(0.00"x");0.00"x"'
DATE = 'mmm yyyy'

BANNED = ["—", "–", " - ", " -- "]


# ------------------------------------------------------------------ Excel helpers
def new_book():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def sheet(wb, name, title, subtitle, code, width_cols=10, landscape=True):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for r in (1, 2, 3):
        for c in range(1, width_cols + 1):
            ws.cell(r, c).fill = FILL_BAND
    ws["A1"] = f"{HOUSE.upper()}  |  VOLUME 2  |  SOLAR HOME SYSTEMS"
    ws["A1"].font = Font(name=FONT, size=8, bold=True, color="E3C27A")
    ws["A2"] = title
    ws["A2"].font = Font(name=FONT, size=15, bold=True, color="FFFFFF")
    ws["A3"] = subtitle
    ws["A3"].font = Font(name=FONT, size=9, italic=True, color="E8E8E8")
    for c in range(1, width_cols + 1):
        ws.cell(4, c).border = GOLD_LINE
    ws["A4"] = f"Template {code}  |  Version {VER}  |  Author & ideation: {AUTHOR}"
    ws["A4"].font = Font(name=FONT, size=8, color=GREY)
    ws.row_dimensions[2].height = 22
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.7
    ws.oddFooter.left.text = f"{HOUSE} | Volume 2 Template {code}"
    ws.oddFooter.center.text = "&A"
    ws.oddFooter.right.text = "Page &P of &N"
    for part in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
        part.size, part.font = 8, "Arial"
    return ws


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[L(i)].width = w


def put(ws, ref, value, kind="calc", fmt=None, bold=False, wrap=False):
    c = ws[ref]
    c.value = value
    if kind == "in":
        c.font, c.fill, c.border = F_IN, FILL_IN, BOX
    elif kind == "head":
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    elif kind == "label":
        c.font = F_BOLD if bold else Font(name=FONT, size=10)
    elif kind == "note":
        c.font = F_NOTE
    else:
        c.font = Font(name=FONT, size=10, bold=bold)
        c.border = BOX
    if fmt:
        c.number_format = fmt
    if wrap or kind == "head":
        c.alignment = Alignment(wrap_text=True, vertical="top" if kind != "head" else "center")
    return c


def section(ws, row, text, ncols):
    for c in range(1, ncols + 1):
        ws.cell(row, c).border = GOLD_LINE
    ws.cell(row, 1, text).font = Font(name=FONT, size=11, bold=True, color=GREEN)


def table_header(ws, row, headers, col0=1, height=30):
    for i, h in enumerate(headers):
        put(ws, f"{L(col0 + i)}{row}", h, "head")
    ws.row_dimensions[row].height = height


def guide_sheet(wb, code, title, purpose, steps, links):
    ws = sheet(wb, "Guide", title, "How to use this template", code, width_cols=6, landscape=False)
    widths(ws, [26, 30, 30, 30, 10, 10])
    section(ws, 6, "Purpose", 6)
    ws.merge_cells("A7:F9")
    put(ws, "A7", purpose, "label", wrap=True)
    section(ws, 11, "How to use it", 6)
    r = 12
    for i, s in enumerate(steps, 1):
        ws.merge_cells(f"B{r}:F{r}")
        put(ws, f"A{r}", f"Step {i}", "label", bold=True)
        put(ws, f"B{r}", s, "label", wrap=True)
        ws.row_dimensions[r].height = 15 * max(1, len(s) // 95 + 1)
        r += 1
    r += 1
    section(ws, r, "Colour code", 6)
    r += 1
    put(ws, f"A{r}", "Input", "in")
    put(ws, f"B{r}", "Blue font on cream: the only cells to edit", "label")
    r += 1
    put(ws, f"A{r}", "Formula")
    put(ws, f"B{r}", "Black font: calculated, do not overwrite", "label")
    r += 2
    section(ws, r, "Links to the Volume 2 model and book", 6)
    r += 1
    for k, v in links:
        put(ws, f"A{r}", k, "label", bold=True)
        ws.merge_cells(f"B{r}:F{r}")
        put(ws, f"B{r}", v, "label", wrap=True)
        r += 1
    r += 1
    ws.merge_cells(f"A{r}:F{r + 1}")
    put(ws, f"A{r}", "Decision support material; not investment, legal, tax or accounting advice. Example values are "
        "illustrative and refer to the fictional SolaraPay case where data are needed. "
        f"(c) {date.today().year} {AUTHOR}.", "note", wrap=True)
    return ws


def finish_book(wb, path, title, subject):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    for b in BANNED:
                        if b in c.value:
                            raise SystemExit(f"dash in {ws.title}!{c.coordinate}: {c.value[:60]}")
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.title = title
    wb.properties.subject = subject
    wb.properties.keywords = "Africa Energy Finance; Volume 2; PAYGo; template"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    _neutral_app(path, b"<Application>.*?</Application>", b"<Application>Africa Energy Finance</Application>")
    print(path.name)


def _neutral_app(path, pat, rep):
    tmp = Path(tempfile.mkstemp(suffix=path.suffix)[1])
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = re.sub(pat, rep, data)
            zout.writestr(item, data)
    shutil.move(tmp, path)
    path.chmod(0o644)


def dv_list(ws, options, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


# ================================================================== T02 Due diligence checklist
DD = [
    ("Commercial", [
        ("Market and segment definition by tier", "Customer survey; sales by segment and region", "Segments defined by product rather than by customer", "Medium", 2),
        ("Sales history by tier, region and channel, 24 months", "Monthly sales file reconciled to revenue", "Growth driven by deposit waivers or price cuts", "High", 5),
        ("Price plan history", "Price lists with effective dates", "Frequent promotions not recorded in the system", "Medium", 4),
        ("Competitive position and alternatives", "Price comparison with grid, generator and competing systems", "No view of competitor pricing", "Low", 2),
        ("Volume plan against distribution capacity", "Agent and installer headcount plan", "Volume growing faster than the agent network", "Medium", 5),
    ]),
    ("Credit and data tape", [
        ("Account level loan tape", "Tape in the Template T06 structure", "Tape does not reconcile to the ledger", "High", 15),
        ("Monthly portfolio data by tier, at least 24 months", "Data in the Credit_Input structure", "Early months missing or restated", "High", 7),
        ("Cohort repayment curves against plan", "Data in the Vintage_Input structure", "Recent cohorts below older cohorts at the same age", "High", 6),
        ("DPD definition and measurement", "Written definition; platform report", "Days without credit used instead of schedule shortfall, undisclosed", "High", 6),
        ("DPD buckets reconcile to gross receivables", "Monthly reconciliation", "Unexplained reconciliation differences", "High", 7),
        ("Collection rate definition", "KPI definitions aligned with PAYGo PERFORM", "Deposits included in the numerator", "High", 7),
        ("Restructuring and re-ageing", "Restructuring log with dates", "Re-ageing clustered before reporting or covenant dates", "High", 15),
        ("Write off policy and history", "Policy with version dates; write off file", "Policy changed during the history period", "High", 8),
        ("Repossession and resale outcomes", "Repossession log; resale proceeds and costs", "Recovery assumptions above observed proceeds", "Medium", 8),
        ("Recalibration of plan curves", "Analyst calibration workpaper", "Plan hazards below observed hazards", "High", 6),
    ]),
    ("Operations and agents", [
        ("Agent register and productivity", "Agent list with sales and cohort quality", "Sales concentrated in a few agents or regions", "Medium", 5),
        ("Commission structure and clawbacks", "Commission plans; clawback records", "No repayment linked component, or clawbacks never applied", "High", 5),
        ("Fraud controls", "Verification call logs; field audit reports", "No fraud losses ever recorded", "High", 5),
        ("Installation and after sales service", "Service tickets and turnaround times", "Rising unresolved tickets", "Medium", 4),
        ("Inventory and supplier terms", "Stock reports; supplier contracts", "Supplier credit shortening", "Medium", 10),
    ]),
    ("Technology and lockout", [
        ("Lockout platform design and access control", "Platform description; user access rights", "Unlock codes issued without audit trail", "High", 15),
        ("Unlocks reconciled to payments", "Unlock log against payment records", "Rising share of unlocks without full payment", "High", 15),
        ("Tampering and bypass rate", "Device telemetry reports", "Bypass rate not measured", "Medium", 15),
        ("Mobile money integration and reconciliation", "Payment platform to ledger reconciliation", "Unreconciled suspense balances", "High", 3),
        ("Data export for lender reporting", "Sample monthly export", "Spreadsheets maintained by hand as the primary record", "Medium", 11),
    ]),
    ("Finance and accounting", [
        ("Audited and management accounts", "Last audited accounts; monthly management accounts", "Audit qualification or late filing", "High", 15),
        ("Revenue recognition policy", "Policy note; auditor correspondence", "Whole contract value booked as revenue at sale", "High", 8),
        ("ECL methodology and staging", "ECL policy and model; auditor view", "Allowance flat while PAR30 rises", "High", 8),
        ("Budget against actual", "Budgets and variance analysis", "Collections systematically over forecast", "Medium", 15),
        ("FX exposure map", "Currency of costs, debt and revenue", "Hard currency debt funding the receivables book, unhedged", "High", 10),
        ("Funding requirement and plan", "Monthly cash forecast", "Plan relies on facilities not yet signed", "High", 11),
        ("Tax position", "Tax filings; open audits", "Disputed VAT or duty exemptions", "Medium", 3),
    ]),
    ("Legal and regulatory", [
        ("Lending and consumer credit licences", "Licences; local counsel opinion", "Lending activity without the required licence", "High", 3),
        ("Pricing disclosure and APR", "Disclosure samples; APR schedule by plan", "APR never computed by the company", "Medium", 3),
        ("Collection conduct and lockout rules", "Collections policy; counsel review", "Lockout practices not reviewed by counsel", "Medium", 3),
        ("Data protection", "Consent forms; registrations", "No consent for the use of data in underwriting", "Medium", 3),
        ("Facility and security documents", "Facility agreements; compliance certificates", "Waivers requested in the last twelve months", "High", 11),
        ("Exchange control and repatriation", "Counsel opinion; investment registrations", "Foreign investment or loans not registered", "High", 3),
        ("True sale and securitisation feasibility, if relevant", "Counsel memo", "No opinion obtained", "Low", 12),
    ]),
    ("ESG and consumer protection", [
        ("Payment burden by tier", "Consumer_Risk based on surveyed incomes", "Illustrative incomes still in use", "High", 2),
        ("Complaints handling", "Complaints log and resolution times", "No complaints log", "Medium", 3),
        ("Hardship and restructuring treatment", "Written policy", "Restructuring used to hide arrears", "Medium", 3),
        ("E-waste and battery take back", "Take back process; recycler contracts", "No end of life process", "Low", 3),
        ("Ownership tracking", "Unlock records at the end of tenor", "Ownership claimed without evidence", "Medium", 13),
    ]),
    ("Management and governance", [
        ("Credit function authority", "Organisation chart; credit committee minutes", "Sales leadership sets credit policy", "High", 15),
        ("Finance team capacity", "Organisation chart; CVs", "Monthly reconciliations not produced", "High", 15),
        ("Board reporting", "Recent board packs", "Board packs without cohort curves", "Medium", 15),
        ("Incentives and key person risk", "Incentive plans; succession", "Incentives on sales volume only", "Medium", 15),
        ("Related party transactions", "Related party register", "Undisclosed related party supply", "Medium", 15),
    ]),
]
STATUS = ["Not started", "Requested", "Received", "Under review", "Issue found", "Closed", "Not applicable"]


def build_t02():
    wb = new_book()
    guide_sheet(wb, "T02", "Investor due diligence checklist",
                "A working checklist for a credit led diligence on a PAYGo solar company: eight workstreams, the evidence to "
                "request, the red flags to look for, and a summary that shows what remains open before the investment committee.",
                ["Assign an owner and a due date to every item on the Checklist sheet.",
                 "Update the status from the drop down list as evidence arrives; record each finding in plain words.",
                 "Mark an item Issue found whenever the red flag is present, and keep it open until the memo explains it.",
                 "Read the Summary sheet before each deal team meeting: high priority items still open block the IC date.",
                 "Copy every unresolved issue into section 11 of the investment memo (Template T01)."],
                [("Book", "Chapter 15 (diligence workstreams, data request, red flags); Annex D"),
                 ("Model", "Credit_Input, Vintage_Input, Consumer_Risk, Investment_Readiness"),
                 ("Related templates", "T01 investment memo; T06 loan tape and data request")])
    ws = sheet(wb, "Checklist", "Investor due diligence checklist", "Status per item; drop down lists in the Priority and Status columns", "T02", 11)
    widths(ws, [7, 22, 40, 36, 38, 10, 14, 36, 14, 12, 9])
    table_header(ws, 6, ["Ref", "Workstream", "Item", "Evidence to request", "Red flag", "Priority", "Status",
                         "Finding", "Owner", "Due date", "Book chapter"])
    r = 7
    for ws_name, items in DD:
        pre = {"Commercial": "CM", "Credit and data tape": "CR", "Operations and agents": "OP", "Technology and lockout": "TL",
               "Finance and accounting": "FI", "Legal and regulatory": "LG", "ESG and consumer protection": "ES",
               "Management and governance": "MG"}[ws_name]
        for i, (item, ev, flag, prio, ch) in enumerate(items, 1):
            put(ws, f"A{r}", f"{pre}{i:02d}")
            put(ws, f"B{r}", ws_name, wrap=True)
            put(ws, f"C{r}", item, wrap=True)
            put(ws, f"D{r}", ev, wrap=True)
            put(ws, f"E{r}", flag, wrap=True)
            put(ws, f"F{r}", prio, "in")
            put(ws, f"G{r}", "Not started", "in")
            put(ws, f"H{r}", None, "in")
            put(ws, f"I{r}", None, "in")
            put(ws, f"J{r}", None, "in", fmt="dd mmm yyyy")
            put(ws, f"K{r}", ch)
            ws.row_dimensions[r].height = 30
            r += 1
    last = r - 1
    dv_list(ws, ["High", "Medium", "Low"], f"F7:F{last}")
    dv_list(ws, STATUS, f"G7:G{last}")
    ws.conditional_formatting.add(f"G7:G{last}", FormulaRule(formula=['$G7="Issue found"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    ws.conditional_formatting.add(f"G7:G{last}", FormulaRule(formula=['OR($G7="Closed",$G7="Not applicable")'], fill=PatternFill("solid", fgColor="D9EAD3")))
    ws.freeze_panes = "D7"
    ws.auto_filter.ref = f"A6:K{last}"
    ws.print_title_rows = "6:6"

    sm = sheet(wb, "Summary", "Diligence status summary", "Counts update from the Checklist sheet", "T02", 8)
    widths(sm, [30, 10, 12, 12, 12, 14, 14, 14])
    table_header(sm, 6, ["Workstream", "Items", "High priority", "Closed or N/A", "Issues found", "High priority open", "Share complete", "Status"])
    rng = lambda col: f"Checklist!${col}$7:${col}${last}"
    r = 7
    for ws_name, _ in DD:
        put(sm, f"A{r}", ws_name, bold=True)
        put(sm, f"B{r}", f'=COUNTIFS({rng("B")},A{r})', fmt=NUM)
        put(sm, f"C{r}", f'=COUNTIFS({rng("B")},A{r},{rng("F")},"High")', fmt=NUM)
        put(sm, f"D{r}", f'=COUNTIFS({rng("B")},A{r},{rng("G")},"Closed")+COUNTIFS({rng("B")},A{r},{rng("G")},"Not applicable")', fmt=NUM)
        put(sm, f"E{r}", f'=COUNTIFS({rng("B")},A{r},{rng("G")},"Issue found")', fmt=NUM)
        put(sm, f"F{r}", f'=C{r}-COUNTIFS({rng("B")},A{r},{rng("F")},"High",{rng("G")},"Closed")-COUNTIFS({rng("B")},A{r},{rng("F")},"High",{rng("G")},"Not applicable")', fmt=NUM)
        put(sm, f"G{r}", f"=IF(B{r}=0,0,D{r}/B{r})", fmt=PCT)
        put(sm, f"H{r}", f'=IF(E{r}>0,"Issues open",IF(F{r}>0,"High items open","Complete"))')
        r += 1
    put(sm, f"A{r}", "Total", bold=True)
    for col in "BCDEF":
        put(sm, f"{col}{r}", f"=SUM({col}7:{col}{r - 1})", fmt=NUM, bold=True)
    put(sm, f"G{r}", f"=IF(B{r}=0,0,D{r}/B{r})", fmt=PCT, bold=True)
    put(sm, f"H{r}", f'=IF(E{r}>0,"Issues open",IF(F{r}>0,"High items open","Complete"))', bold=True)
    tot = r
    r += 2
    put(sm, f"A{r}", "Ready for investment committee", "label", bold=True)
    put(sm, f"B{r}", f'=IF(AND(E{tot}=0,F{tot}=0),"Yes","No: resolve open high priority items and issues first")', bold=True)
    sm.merge_cells(f"B{r}:H{r}")
    for rr in range(7, tot + 1):
        sm.conditional_formatting.add(f"H{rr}", FormulaRule(formula=[f'$H{rr}="Issues open"'], fill=PatternFill("solid", fgColor="F4CCCC")))
        sm.conditional_formatting.add(f"H{rr}", FormulaRule(formula=[f'$H{rr}="Complete"'], fill=PatternFill("solid", fgColor="D9EAD3")))
    finish_book(wb, OUT / "AEF_V2_T02_Due_Diligence_Checklist.xlsx", "AEF Volume 2 Template T02: Investor due diligence checklist",
                "Africa Energy Finance, Volume 2 Templates Pack")
    return last - 6


# ================================================================== T04 Agent economics
def build_t04():
    wb = new_book()
    guide_sheet(wb, "T04", "Agent economics and fully loaded CAC",
                "Builds the customer acquisition cost of a territory on two definitions, narrow (commission and marketing) and "
                "fully loaded (adding the field organisation and fraud controls), ranks agents by the quality of the cohorts "
                "they originate, and shows the cost of opening a new territory.",
                ["Enter the territory inputs on the Inputs sheet (blue cells). Example values are illustrative.",
                 "Read the narrow and fully loaded CAC by tier on the CAC sheet, and the ratio between the two.",
                 "Paste agent level data into Agent_Ranking: units sold, share current at month 3, repayment at month 6, deposit only accounts.",
                 "Use the flags to set commission bands and to decide which agents to retrain or remove.",
                 "Use New_Territory to set a budget and stop criteria before opening a new area.",
                 "Copy the fully loaded CAC per tier into the commission and marketing lines of the model (Products) if you want the unit view to carry field costs, and remove the same costs from fixed costs on Inputs."],
                [("Book", "Chapter 5 (distribution and customer acquisition); Chapter 9 (unit economics); Annex C"),
                 ("Model", "Products (commission, marketing per unit); Inputs (fixed costs); Unit_Economics"),
                 ("Related templates", "T03 credit policy (agent override rules)")])
    ws = sheet(wb, "Inputs", "Territory inputs", "Monthly figures for one territory, local currency", "T04", 8)
    widths(ws, [44, 14, 14, 14, 14, 14, 14, 30])
    section(ws, 6, "Sales force", 8)
    rows = [("Active agents", 50, NUM, "Agents selling in the month"),
            ("Sales per agent per month", 6, NUM1, "All tiers"),
            ("Monthly agent attrition", 0.08, PCT, "Share of agents leaving each month"),
            ("Recruitment and training cost per new agent", 30000, NUM, "Onboarding, training, starter kit")]
    r = 7
    ref = {}
    for lab, v, fmt, note in rows:
        put(ws, f"A{r}", lab, "label")
        put(ws, f"B{r}", v, "in", fmt)
        put(ws, f"H{r}", note, "note")
        ref[lab] = f"Inputs!$B${r}"
        r += 1
    r += 1
    section(ws, r, "Per tier", 8)
    r += 1
    table_header(ws, r, ["Item", "Tier 1", "Tier 2", "Tier 3", "Tier 4", "Tier 5", "", "Notes"])
    r += 1
    P = D.PRODUCTS
    tier_rows = [("Sales mix", [p["mix"] for p in copy.deepcopy(P)], PCT, "Must total 100%"),
                 ("Cash price", [p["price"] for p in P], NUM, "Local currency"),
                 ("Upfront commission per unit", [round(p["comm"] * 0.6) for p in P], NUM, "Paid at activation"),
                 ("Deferred commission per unit", [round(p["comm"] * 0.4) for p in P], NUM, "Paid if the account is current at the checkpoint"),
                 ("Share of accounts qualifying for deferred commission", [0.70, 0.80, 0.85, 0.90, 0.90], PCT, "Credit assumption: current at month 3"),
                 ("Marketing and acquisition per unit", [p["mkt"] for p in P], NUM, "Direct per unit spend")]
    tref = {}
    for lab, vals, fmt, note in tier_rows:
        put(ws, f"A{r}", lab, "label")
        for j, v in enumerate(vals):
            put(ws, f"{L(2 + j)}{r}", v, "in", fmt)
        put(ws, f"H{r}", note, "note")
        tref[lab] = r
        r += 1
    put(ws, f"A{r}", "Check: sales mix totals 100%", "label")
    put(ws, f"B{r}", f'=IF(ABS(SUM(B{tref["Sales mix"]}:F{tref["Sales mix"]})-1)<0.0001,"OK","ERROR")', bold=True)
    mixchk = f"Inputs!$B${r}"
    r += 2
    section(ws, r, "Field organisation and controls (per month)", 8)
    r += 1
    field = [("Field supervisors and area managers", 450000, "5 supervisors at 90,000"),
             ("Transport and fuel", 210000, "Territory budget"),
             ("Verification calls and field audits", 60000, "Territory budget"),
             ("Other field costs", 0, "Demonstration kits, materials")]
    for lab, v, note in field:
        put(ws, f"A{r}", lab, "label")
        put(ws, f"B{r}", v, "in", NUM)
        put(ws, f"H{r}", note, "note")
        ref[lab] = f"Inputs!$B${r}"
        r += 1
    put(ws, f"A{r}", "Fraud and ghost sale losses, share of units", "label")
    put(ws, f"B{r}", 0.01, "in", PCT)
    put(ws, f"H{r}", "From audit sampling; cost valued at upfront commission plus marketing", "note")
    ref["fraud"] = f"Inputs!$B${r}"

    cs = sheet(wb, "CAC", "Customer acquisition cost", "Narrow and fully loaded definitions, by tier", "T04", 8)
    widths(cs, [46, 14, 14, 14, 14, 14, 16, 26])
    table_header(cs, 6, ["Line", "Tier 1", "Tier 2", "Tier 3", "Tier 4", "Tier 5", "Territory total", "Notes"])
    units = f"{ref['Active agents']}*{ref['Sales per agent per month']}"
    lines = [
        ("Units sold per month", lambda c: f"={units}*Inputs!{c}${tref['Sales mix']}", NUM, "Agents x sales per agent x mix"),
        ("Upfront commission", lambda c: f"=Inputs!{c}${tref['Upfront commission per unit']}", NUM, "Per unit"),
        ("Expected deferred commission", lambda c: f"=Inputs!{c}${tref['Deferred commission per unit']}*Inputs!{c}${tref['Share of accounts qualifying for deferred commission']}", NUM, "Accrued at sale on the expected qualification rate"),
        ("Marketing and acquisition", lambda c: f"=Inputs!{c}${tref['Marketing and acquisition per unit']}", NUM, "Per unit"),
        ("Narrow CAC per unit", lambda c: f"=SUM({c}8:{c}10)", NUM, "Commission plus marketing"),
        ("Field organisation per unit", lambda c: f"=IF($G$7=0,0,({ref['Field supervisors and area managers']}+{ref['Transport and fuel']}+{ref['Other field costs']})/$G$7)", NUM, "Allocated equally per unit"),
        ("Recruitment and training per unit", lambda c: f"=IF($G$7=0,0,{ref['Active agents']}*{ref['Monthly agent attrition']}*{ref['Recruitment and training cost per new agent']}/$G$7)", NUM, "Replacing leavers"),
        ("Verification and audits per unit", lambda c: f"=IF($G$7=0,0,{ref['Verification calls and field audits']}/$G$7)", NUM, "Allocated equally per unit"),
        ("Fraud losses per unit", lambda c: f"={ref['fraud']}*({c}8+{c}10)", NUM, "Commission and marketing spent on ghost sales"),
        ("Fully loaded CAC per unit", lambda c: f"={c}11+SUM({c}12:{c}15)", NUM, ""),
        ("Fully loaded ÷ narrow", lambda c: f"=IF({c}11=0,0,{c}16/{c}11)", RAT, ""),
        ("Narrow CAC as share of cash price", lambda c: f"=IF(Inputs!{c}${tref['Cash price']}=0,0,{c}11/Inputs!{c}${tref['Cash price']})", PCT, ""),
        ("Fully loaded CAC as share of cash price", lambda c: f"=IF(Inputs!{c}${tref['Cash price']}=0,0,{c}16/Inputs!{c}${tref['Cash price']})", PCT, ""),
    ]
    for i, (lab, fn, fmt, note) in enumerate(lines):
        r = 7 + i
        put(cs, f"A{r}", lab, "label", bold=lab.startswith(("Narrow CAC per", "Fully loaded CAC per")))
        for j in range(5):
            c = L(2 + j)
            put(cs, f"{c}{r}", fn(c), fmt=fmt, bold=lab.startswith(("Narrow CAC per", "Fully loaded CAC per")))
        put(cs, f"H{r}", note, "note")
    # territory totals (spend per month) and weighted averages
    put(cs, "G7", "=SUM(B7:F7)", fmt=NUM, bold=True)
    for r in range(8, 16):
        put(cs, f"G{r}", f"=IF($G$7=0,0,SUMPRODUCT(B{r}:F{r},$B$7:$F$7)/$G$7)", fmt=NUM)
    put(cs, "G11", "=IF($G$7=0,0,SUMPRODUCT(B11:F11,$B$7:$F$7)/$G$7)", fmt=NUM, bold=True)
    put(cs, "G16", "=IF($G$7=0,0,SUMPRODUCT(B16:F16,$B$7:$F$7)/$G$7)", fmt=NUM, bold=True)
    put(cs, "G17", "=IF(G11=0,0,G16/G11)", fmt=RAT)
    put(cs, "A21", "Monthly territory spend, fully loaded", "label", bold=True)
    put(cs, "B21", "=SUMPRODUCT(B16:F16,B7:F7)", fmt=NUM, bold=True)
    put(cs, "A22", "Check: inputs valid", "label")
    put(cs, "B22", f"={mixchk}", bold=True)
    put(cs, "A24", "Weighted averages in column G use units sold by tier. Field costs are allocated equally per unit; a company "
        "that allocates by time spent per tier should change rows 12 to 14.", "note")

    ar = sheet(wb, "Agent_Ranking", "Agent cohort quality ranking", "Paste agent level data; thresholds on the right", "T04", 12)
    widths(ar, [12, 14, 10, 12, 14, 14, 14, 12, 12, 22, 4, 34])
    table_header(ar, 6, ["Agent ID", "Region", "Tenure (months)", "Units sold, last 6 months", "Share current at month 3",
                         "Repayment rate at month 6", "Deposit only accounts", "Deposit only share", "Repayment rank",
                         "Flag", "", "Thresholds"])
    import random
    rnd = random.Random(20261002)
    n_ag = 30
    for i in range(n_ag):
        r = 7 + i
        put(ar, f"A{r}", f"AG{i + 1:03d}", "in")
        put(ar, f"B{r}", ["North", "South", "East", "West"][i % 4], "in")
        put(ar, f"C{r}", rnd.randint(2, 30), "in", NUM)
        u = rnd.randint(18, 60)
        put(ar, f"D{r}", u, "in", NUM)
        put(ar, f"E{r}", round(rnd.uniform(0.70, 0.95), 3), "in", PCT)
        put(ar, f"F{r}", round(rnd.uniform(0.62, 0.86), 3), "in", PCT)
        put(ar, f"G{r}", rnd.randint(0, max(1, u // 8)), "in", NUM)
        put(ar, f"H{r}", f'=IF(OR(A{r}="",D{r}=0),0,G{r}/D{r})', fmt=PCT)
        put(ar, f"I{r}", f'=IF(A{r}="","",COUNTIFS($F$7:$F${6 + n_ag},">"&F{r})+1)', fmt=NUM)
        put(ar, f"J{r}", f'=IF(A{r}="","",IF(OR(F{r}<$L$9,H{r}>$L$12),"Review",IF(F{r}>=$L$15,"Top band","Standard")))')
    end = 6 + n_ag
    for lab, v, fmt, rr in [("Minimum repayment rate at month 6", 0.70, PCT, 9), ("Maximum deposit only share", 0.10, PCT, 12),
                            ("Repayment rate for top commission band", 0.80, PCT, 15)]:
        put(ar, f"L{rr - 1}", lab, "label", bold=True)
        put(ar, f"L{rr}", v, "in", fmt)
    put(ar, "L18", "Agents flagged for review", "label", bold=True)
    put(ar, "L19", f'=COUNTIFS(J7:J{end},"Review")', fmt=NUM, bold=True)
    put(ar, "L21", "Repayment, top half of agents", "label", bold=True)
    put(ar, "L22", f'=IFERROR(AVERAGEIFS(F7:F{end},I7:I{end},"<="&COUNT(F7:F{end})/2),0)', fmt=PCT)
    put(ar, "L23", "Repayment, bottom half of agents", "label", bold=True)
    put(ar, "L24", f'=IFERROR(AVERAGEIFS(F7:F{end},I7:I{end},">"&COUNT(F7:F{end})/2),0)', fmt=PCT)
    put(ar, "L26", "Example rows are randomly generated illustrations; replace them with company data.", "note")
    ar.conditional_formatting.add(f"J7:J{end}", FormulaRule(formula=['$J7="Review"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    ar.conditional_formatting.add(f"J7:J{end}", FormulaRule(formula=['$J7="Top band"'], fill=PatternFill("solid", fgColor="D9EAD3")))
    ar.freeze_panes = "B7"

    nt = sheet(wb, "New_Territory", "New territory ramp", "Fixed cost per unit during the first twelve months", "T04", 14)
    widths(nt, [40] + [10] * 12 + [24])
    put(nt, "A6", "Fixed territory cost per month (depot, manager, vehicles)", "label")
    put(nt, "B6", 900000, "in", NUM)
    put(nt, "A7", "Mature sales per month", "label")
    put(nt, "B7", 300, "in", NUM)
    table_header(nt, 9, ["Line"] + [f"M{m}" for m in range(1, 13)] + ["Notes"])
    ramp = [0.33, 0.50, 0.67, 0.75, 0.83, 0.90, 0.95, 1, 1, 1, 1, 1]
    put(nt, "A10", "Ramp: share of mature sales", "label")
    put(nt, "A11", "Units sold", "label")
    put(nt, "A12", "Fixed cost per unit", "label")
    put(nt, "A13", "Multiple of mature fixed cost per unit", "label")
    for m in range(12):
        c = L(2 + m)
        put(nt, f"{c}10", ramp[m], "in", PCT)
        put(nt, f"{c}11", f"=$B$7*{c}10", fmt=NUM)
        put(nt, f"{c}12", f"=IF({c}11=0,0,$B$6/{c}11)", fmt=NUM)
        put(nt, f"{c}13", f"=IF($B$7=0,0,{c}12/($B$6/$B$7))", fmt=RAT)
    put(nt, "N10", "Ramp is an input; set it from the last territory opened", "note")
    section(nt, 15, "Stop criteria (set before opening)", 14)
    crit = [("Minimum sales per agent per month by month 6", 4, NUM1),
            ("Maximum gap in repayment at month 3 against established territories (points)", 0.05, PCT),
            ("Maximum fully loaded CAC multiple of mature figure by month 9", 1.3, RAT)]
    for i, (lab, v, fmt) in enumerate(crit):
        put(nt, f"A{16 + i}", lab, "label")
        put(nt, f"B{16 + i}", v, "in", fmt)
    finish_book(wb, OUT / "AEF_V2_T04_Agent_Economics.xlsx", "AEF Volume 2 Template T04: Agent economics and fully loaded CAC",
                "Africa Energy Finance, Volume 2 Templates Pack")


# ================================================================== T05 Lender KPI report
KPI_DEFS = [
    ("Collection rate (month)", "Instalments collected ÷ instalments due in the month, both excluding deposits", "Deposits are collected at sale and would flatter the ratio"),
    ("Collection rate (trailing 3 months)", "Collections ÷ instalments due over the last three months", "Usual covenant measure"),
    ("PAR30", "Gross receivables of accounts more than 30 days past due ÷ gross receivables", "Whole balance of late accounts counts"),
    ("PAR90", "Gross receivables of accounts more than 90 days past due ÷ gross receivables", ""),
    ("PAR30 lagged 3 months", "Balance more than 30 days past due ÷ gross receivables three months earlier", "Corrects for growth in the denominator"),
    ("Receivables at risk (RaR)", "As reported by the company under its stated definition ÷ gross receivables", "Align the definition with the current PAYGo PERFORM documents"),
    ("Write off ratio (trailing 12 months, annualised)", "Write offs ÷ average gross receivables over the same months, scaled to a year", "Depends on the write off policy"),
    ("Recovery rate (trailing 12 months)", "Recoveries net of costs ÷ write offs", ""),
    ("Active ratio", "Accounts with a payment in the last 30 days ÷ accounts not yet paid off or written off", "Optional; leave blank if not reported"),
    ("Eligible receivables", "Current and 1 to 30 DPD balances (eligibility to 30 DPD)", "Change the formula if eligibility differs"),
    ("Borrowing base headroom", "Eligible receivables × advance rate less facility drawn", ""),
    ("Repayment rate (cohort)", "Cumulative collections ÷ cumulative instalments due at a given age", "Compare cohorts at the same age"),
]


def build_t05():
    G, P = copy.deepcopy(D.GENERAL), copy.deepcopy(D.PRODUCTS)
    SP.apply(G, P)
    hist_c, hist_v = SP.history(P, G)
    months = len(next(iter(hist_c.values())))
    agg = []
    for i in range(months):
        rows = [hist_c[j][i] for j in hist_c]
        dep = sum(hist_c[j][i]["orig"] * P[j]["deposit"] for j in hist_c)
        agg.append(dict(date=rows[0]["date"], gross=sum(x["gross"] for x in rows),
                        b=[sum(x[f"b{k}"] for x in rows) for k in range(6)],
                        due=sum(x["due"] for x in rows), coll=sum(x["coll"] for x in rows), dep=dep,
                        wo=sum(x["wo"] for x in rows), rec=sum(x["resale"] - x["rcost"] for x in rows),
                        orig=sum(x["orig"] for x in rows)))
    wb = new_book()
    guide_sheet(wb, "T05", "Lender KPI report",
                "A monthly reporting pack for a PAYGo receivables lender or an investor's portfolio review: portfolio KPIs on "
                "PERFORM style definitions, DPD reconciliation, covenant headroom, cohort repayment against plan and a one page "
                "dashboard. The example data are the synthetic 24 month history of the fictional SolaraPay case (Tiers 1 to 3).",
                ["Replace the example data on Monthly_Input with company data, one column per month (up to 24 months).",
                 "Check the reconciliation row on KPIs: the DPD buckets must add up to reported gross receivables every month.",
                 "Set covenant thresholds and the blended advance rate on Covenants; the sheet shows headroom and breach flags by month.",
                 "Enter cumulative collections and instalments due by cohort on Vintage, with the plan curve, to see cohorts against plan.",
                 "Send the Dashboard page with every monthly report; keep definitions aligned with the current PAYGo PERFORM documents."],
                [("Book", "Chapter 6 (cohorts), Chapter 7 (portfolio KPIs and PERFORM), Chapter 11 (covenants); Annex A"),
                 ("Model", "Credit_Input and Vintage_Input hold the same history; Credit_Portfolio, Covenants, Vintage_Dashboard"),
                 ("Related templates", "T06 loan tape; T07 borrowing base certificate")])

    df = sheet(wb, "Definitions", "KPI definitions", "Definitions used in this report", "T05", 3, landscape=False)
    widths(df, [34, 60, 44])
    table_header(df, 6, ["KPI", "Definition", "Comment"])
    for i, (k, dfn, cm) in enumerate(KPI_DEFS):
        r = 7 + i
        put(df, f"A{r}", k, bold=True, wrap=True)
        put(df, f"B{r}", dfn, wrap=True)
        put(df, f"C{r}", cm, wrap=True)
        df.row_dimensions[r].height = 30

    N = 24
    first, lastc = 3, 3 + N - 1
    MC = [L(c) for c in range(first, lastc + 1)]
    mi = sheet(wb, "Monthly_Input", "Monthly portfolio input", "Local currency; one column per month; month end stocks and monthly flows", "T05", lastc)
    widths(mi, [46, 10] + [12] * N)
    table_header(mi, 6, ["Item", "Unit"] + [f"M{i + 1}" for i in range(N)])
    in_rows = ["Month end date", "Gross receivables (reported)", "Current", "1 to 30 DPD", "31 to 60 DPD", "61 to 90 DPD",
               "91 to 180 DPD", "Over 180 DPD", "Instalments due in month (excluding deposits)",
               "Instalments collected in month (excluding deposits)", "Deposits collected", "Write offs in month",
               "Recoveries net of costs in month", "Units originated", "Receivables at risk (company definition, optional)",
               "Active accounts (optional)", "Open accounts, not paid off or written off (optional)", "Facility drawn"]
    IR = {k: 7 + i for i, k in enumerate(in_rows)}
    for k, r in IR.items():
        put(mi, f"A{r}", k, "label", bold=(k == "Gross receivables (reported)"))
        put(mi, f"B{r}", "date" if k == "Month end date" else ("count" if "accounts" in k.lower() or "Units" in k else "LCY"), "label")
    for i in range(N):
        c = MC[i]
        a = agg[i] if i < len(agg) else None
        y, m = int(a["date"][:4]), int(a["date"][5:7])
        nxt = date(y + (m == 12), m % 12 + 1, 1)
        from datetime import timedelta
        put(mi, f"{c}{IR['Month end date']}", nxt - timedelta(days=1), "in", DATE)
        vals = {"Gross receivables (reported)": a["gross"], "Current": a["b"][0], "1 to 30 DPD": a["b"][1],
                "31 to 60 DPD": a["b"][2], "61 to 90 DPD": a["b"][3], "91 to 180 DPD": a["b"][4], "Over 180 DPD": a["b"][5],
                "Instalments due in month (excluding deposits)": a["due"], "Instalments collected in month (excluding deposits)": a["coll"],
                "Deposits collected": a["dep"], "Write offs in month": a["wo"], "Recoveries net of costs in month": a["rec"],
                "Units originated": a["orig"], "Receivables at risk (company definition, optional)": None,
                "Active accounts (optional)": None, "Open accounts, not paid off or written off (optional)": None,
                "Facility drawn": 0}
        for k, v in vals.items():
            put(mi, f"{c}{IR[k]}", round(v) if isinstance(v, float) else v, "in", NUM)
    put(mi, f"A{7 + len(in_rows) + 1}", "Example: synthetic SolaraPay history, Tiers 1 to 3 combined (fictional). Facility drawn is set to zero "
        "because the facility had not been signed in the history period.", "note")
    mi.freeze_panes = "C7"

    cv = sheet(wb, "Covenants", "Covenant tests", "Thresholds are inputs; headroom and flags by month", "T05", lastc)
    widths(cv, [46, 12] + [12] * N)
    th = [("Minimum trailing 3 month collection rate", 0.70, "min", PCT), ("Maximum RaR", 0.15, "max", PCT),
          ("Maximum PAR30", 0.25, "max", PCT), ("Maximum PAR90", 0.18, "max", PCT),
          ("Minimum borrowing base headroom (LCY)", 0, "min", NUM)]
    section(cv, 6, "Thresholds", lastc)
    for i, (lab, v, _, fmt) in enumerate(th):
        put(cv, f"A{7 + i}", lab, "label")
        put(cv, f"B{7 + i}", v, "in", fmt)
    put(cv, "A12", "Blended advance rate on eligible receivables", "label")
    put(cv, "B12", 0.70, "in", PCT)
    put(cv, "A13", "Watch list when headroom is below this share of the threshold", "label")
    put(cv, "B13", 0.10, "in", PCT)

    k = sheet(wb, "KPIs", "Portfolio KPIs", "Calculated from Monthly_Input; blank where data are not reported", "T05", lastc)
    widths(k, [46, 10] + [12] * N)
    table_header(k, 6, ["KPI", "Unit"] + [f"M{i + 1}" for i in range(N)])
    mref = lambda name, c: f"Monthly_Input!{c}{IR[name]}"
    KR = {}
    kp = [
        ("Month end date", "date", lambda c, i: f"={mref('Month end date', c)}", DATE),
        ("Sum of DPD buckets", "LCY", lambda c, i: f"=SUM(Monthly_Input!{c}{IR['Current']}:{c}{IR['Over 180 DPD']})", NUM),
        ("Reconciliation difference", "LCY", lambda c, i: f"={c}8-{mref('Gross receivables (reported)', c)}", NUM),
        ("Reconciliation check", "flag", lambda c, i: f'=IF({mref("Gross receivables (reported)", c)}="","",IF(ABS({c}9)<=1,"OK","ERROR"))', None),
        ("Collection rate (month)", "%", lambda c, i: f'=IF(N({mref("Instalments due in month (excluding deposits)", c)})=0,"",{mref("Instalments collected in month (excluding deposits)", c)}/{mref("Instalments due in month (excluding deposits)", c)})', PCT),
        ("Collection rate (trailing 3 months)", "%", lambda c, i: "" if i < 2 else
            f'=IF(SUM(Monthly_Input!{MC[i - 2]}{IR["Instalments due in month (excluding deposits)"]}:{c}{IR["Instalments due in month (excluding deposits)"]})=0,"",SUM(Monthly_Input!{MC[i - 2]}{IR["Instalments collected in month (excluding deposits)"]}:{c}{IR["Instalments collected in month (excluding deposits)"]})/SUM(Monthly_Input!{MC[i - 2]}{IR["Instalments due in month (excluding deposits)"]}:{c}{IR["Instalments due in month (excluding deposits)"]}))', PCT),
        ("PAR30", "%", lambda c, i: f'=IF(N({mref("Gross receivables (reported)", c)})=0,"",SUM(Monthly_Input!{c}{IR["31 to 60 DPD"]}:{c}{IR["Over 180 DPD"]})/{mref("Gross receivables (reported)", c)})', PCT),
        ("PAR90", "%", lambda c, i: f'=IF(N({mref("Gross receivables (reported)", c)})=0,"",SUM(Monthly_Input!{c}{IR["91 to 180 DPD"]}:{c}{IR["Over 180 DPD"]})/{mref("Gross receivables (reported)", c)})', PCT),
        ("PAR30 lagged 3 months", "%", lambda c, i: "" if i < 3 else
            f'=IF(N(Monthly_Input!{MC[i - 3]}{IR["Gross receivables (reported)"]})=0,"",SUM(Monthly_Input!{c}{IR["31 to 60 DPD"]}:{c}{IR["Over 180 DPD"]})/Monthly_Input!{MC[i - 3]}{IR["Gross receivables (reported)"]})', PCT),
        ("RaR ratio", "%", lambda c, i: f'=IF(OR({mref("Receivables at risk (company definition, optional)", c)}="",N({mref("Gross receivables (reported)", c)})=0),"",{mref("Receivables at risk (company definition, optional)", c)}/{mref("Gross receivables (reported)", c)})', PCT),
        ("Write off ratio (trailing 12 months, annualised)", "%", lambda c, i:
            f'=IF(AVERAGE(Monthly_Input!{MC[max(0, i - 11)]}{IR["Gross receivables (reported)"]}:{c}{IR["Gross receivables (reported)"]})=0,"",SUM(Monthly_Input!{MC[max(0, i - 11)]}{IR["Write offs in month"]}:{c}{IR["Write offs in month"]})/AVERAGE(Monthly_Input!{MC[max(0, i - 11)]}{IR["Gross receivables (reported)"]}:{c}{IR["Gross receivables (reported)"]})*12/{min(12, i + 1)})', PCT),
        ("Recovery rate (trailing 12 months)", "%", lambda c, i:
            f'=IF(SUM(Monthly_Input!{MC[max(0, i - 11)]}{IR["Write offs in month"]}:{c}{IR["Write offs in month"]})=0,"",SUM(Monthly_Input!{MC[max(0, i - 11)]}{IR["Recoveries net of costs in month"]}:{c}{IR["Recoveries net of costs in month"]})/SUM(Monthly_Input!{MC[max(0, i - 11)]}{IR["Write offs in month"]}:{c}{IR["Write offs in month"]}))', PCT),
        ("Active ratio", "%", lambda c, i: f'=IF(OR({mref("Active accounts (optional)", c)}="",N({mref("Open accounts, not paid off or written off (optional)", c)})=0),"",{mref("Active accounts (optional)", c)}/{mref("Open accounts, not paid off or written off (optional)", c)})', PCT),
        ("Eligible receivables (to 30 DPD)", "LCY", lambda c, i: f"=Monthly_Input!{c}{IR['Current']}+Monthly_Input!{c}{IR['1 to 30 DPD']}", NUM),
        ("Borrowing base", "LCY", lambda c, i: f"={c}20*Covenants!$B$12", NUM),
        ("Borrowing base headroom", "LCY", lambda c, i: f"={c}21-N({mref('Facility drawn', c)})", NUM),
    ]
    for n_, (lab, unit, fn, fmt) in enumerate(kp):
        r = 7 + n_
        KR[lab] = r
        put(k, f"A{r}", lab, "label", bold=lab in ("Collection rate (trailing 3 months)", "PAR30", "Reconciliation check"))
        put(k, f"B{r}", unit, "label")
        for i, c in enumerate(MC):
            v = fn(c, i)
            put(k, f"{c}{r}", v if v != "" else None, fmt=fmt)
    k.conditional_formatting.add(f"C{KR['Reconciliation check']}:{MC[-1]}{KR['Reconciliation check']}",
                                 FormulaRule(formula=[f'C{KR["Reconciliation check"]}="ERROR"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    k.freeze_panes = "C7"

    # covenant flags by month (continued on Covenants)
    table_header(cv, 15, ["Test (1 = breach)", ""] + [f"M{i + 1}" for i in range(N)])
    tests = [("Trailing 3 month collection rate", "Collection rate (trailing 3 months)", 7, "min"),
             ("RaR", "RaR ratio", 8, "max"), ("PAR30", "PAR30", 9, "max"), ("PAR90", "PAR90", 10, "max"),
             ("Borrowing base headroom", "Borrowing base headroom", 11, "min")]
    for t, (lab, kpi, thr, kind) in enumerate(tests):
        r = 16 + t
        put(cv, f"A{r}", lab, "label")
        for c in MC:
            kc = f"KPIs!{c}{KR[kpi]}"
            if lab == "Borrowing base headroom":
                f = f'=IF(N(Monthly_Input!{c}{IR["Facility drawn"]})=0,"",IF({kc}<$B${thr},1,0))'
            else:
                f = f'=IF({kc}="","",IF({kc}{"<" if kind == "min" else ">"}$B${thr},1,0))'
            put(cv, f"{c}{r}", f, fmt=NUM)
        cv.conditional_formatting.add(f"C{r}:{MC[-1]}{r}", FormulaRule(formula=[f"C{r}=1"], fill=PatternFill("solid", fgColor="F4CCCC")))
    put(cv, "A21", "Any covenant in breach", "label", bold=True)
    for c in MC:
        put(cv, f"{c}21", f'=IF(COUNT({c}16:{c}20)=0,"",IF(SUM({c}16:{c}20)>0,1,0))', fmt=NUM, bold=True)
    put(cv, "A22", "Months in breach", "label", bold=True)
    put(cv, "B22", f"=SUM(C21:{MC[-1]}21)", fmt=NUM, bold=True)
    section(cv, 24, "Latest month", lastc)
    table_header(cv, 25, ["Test", "Threshold", "Actual", "Headroom", "Status"], col0=1)
    for t, (lab, kpi, thr, kind) in enumerate(tests):
        r = 26 + t
        put(cv, f"A{r}", lab, "label")
        put(cv, f"B{r}", f"=B{thr}", fmt=PCT if lab != "Borrowing base headroom" else NUM)
        latest = f'IFERROR(INDEX(KPIs!$C${KR[kpi]}:${MC[-1]}${KR[kpi]},MATCH(9.99E+307,KPIs!$C${KR[kpi]}:${MC[-1]}${KR[kpi]})),"")'
        if lab == "Borrowing base headroom":
            fr = f"Monthly_Input!$C${IR['Facility drawn']}:${MC[-1]}${IR['Facility drawn']}"
            latest = f'IF(N(IFERROR(INDEX({fr},MATCH(9.99E+307,{fr})),0))=0,"",{latest})'
        put(cv, f"C{r}", "=" + latest,
            fmt=PCT if lab != "Borrowing base headroom" else NUM)
        hd = f"C{r}-B{r}" if kind == "min" else f"B{r}-C{r}"
        put(cv, f"D{r}", f'=IF(C{r}="","",{hd})', fmt=PCT if lab != "Borrowing base headroom" else NUM)
        put(cv, f"E{r}", f'=IF(C{r}="",IF(A{r}="Borrowing base headroom","No facility drawn","No data"),IF(D{r}<0,"Breach",IF(D{r}<ABS(B{r})*$B$13,"Watch","Compliant")))')
        cv.conditional_formatting.add(f"E{r}", FormulaRule(formula=[f'E{r}="Breach"'], fill=PatternFill("solid", fgColor="F4CCCC")))
        cv.conditional_formatting.add(f"E{r}", FormulaRule(formula=[f'E{r}="Watch"'], fill=PatternFill("solid", fgColor="FCE8B2")))
    put(cv, "A32", "Borrowing base headroom is tested only in months with a facility drawn. RaR is tested only when the company reports it.", "note")

    # vintage
    vt = sheet(wb, "Vintage", "Cohort repayment against plan", "Cumulative collections ÷ instalments due at fixed ages (deposits excluded)", "T05", 22)
    CP = ["M3", "M6", "M12", "M18", "M24"]
    widths(vt, [16] + [13] * 15 + [4] * 6)
    put(vt, "A6", "Tier shown", "label")
    put(vt, "B6", "Tier 2 (example)", "in")
    table_header(vt, 8, ["Cohort (month of sale)"] + [f"Collections {c}" for c in CP] + [f"Due {c}" for c in CP] + [f"Repayment {c}" for c in CP])
    tier_idx = 1
    nco = 18
    for ci in range(1, nco + 1):
        r = 8 + ci
        put(vt, f"A{r}", f"Cohort {ci:02d}", "in")
        v = hist_v[tier_idx][ci]["cp"]
        for kk in range(5):
            cc, dc, rc = L(2 + kk), L(7 + kk), L(12 + kk)
            if kk in v:
                put(vt, f"{cc}{r}", round(v[kk]["coll"]), "in", NUM)
                put(vt, f"{dc}{r}", round(v[kk]["due"]), "in", NUM)
            else:
                put(vt, f"{cc}{r}", None, "in", NUM)
                put(vt, f"{dc}{r}", None, "in", NUM)
            put(vt, f"{rc}{r}", f'=IF(N({dc}{r})=0,"",{cc}{r}/{dc}{r})', fmt=PCT)
    endr = 8 + nco
    put(vt, f"A{endr + 2}", "Average observed", "label", bold=True)
    put(vt, f"A{endr + 3}", "Plan", "label", bold=True)
    put(vt, f"A{endr + 4}", "Gap (points)", "label", bold=True)
    plan = [0.835, 0.803, 0.745, 0.692, 0.644]
    for kk in range(5):
        rc = L(12 + kk)
        put(vt, f"{rc}{endr + 2}", f'=IFERROR(AVERAGE({rc}9:{rc}{endr}),"")', fmt=PCT, bold=True)
        put(vt, f"{rc}{endr + 3}", plan[kk], "in", PCT)
        put(vt, f"{rc}{endr + 4}", f'=IF({rc}{endr + 2}="","",{rc}{endr + 2}-{rc}{endr + 3})', fmt=PCT, bold=True)
    put(vt, f"A{endr + 6}", "Example: Tier 2 cohorts of the synthetic SolaraPay history (fictional); plan curve from the management plan. "
        "Ages beyond the history are left blank, never filled with a proxy.", "note")
    vt.freeze_panes = "B9"
    ch = LineChart()
    ch.title, ch.height, ch.width = "Repayment: average observed against plan", 7, 14
    ch.y_axis.number_format = "0%"
    ch.add_data(Reference(vt, min_col=12, max_col=16, min_row=endr + 2, max_row=endr + 3), from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(vt, min_col=12, max_col=16, min_row=8, max_row=8))
    from openpyxl.chart.series import SeriesLabel
    ch.series[0].tx = SeriesLabel(v="Observed")
    ch.series[1].tx = SeriesLabel(v="Plan")
    vt.add_chart(ch, f"B{endr + 8}")

    # dashboard
    db = sheet(wb, "Dashboard", "Monthly lender dashboard", "Latest month; charts over the reporting period", "T05", 10)
    widths(db, [40, 14, 14, 14, 14, 4, 14, 14, 14, 14])
    table_header(db, 6, ["KPI", "Latest", "3 months earlier", "12 months earlier", "Trend"])
    show = ["Collection rate (month)", "Collection rate (trailing 3 months)", "PAR30", "PAR90", "PAR30 lagged 3 months",
            "Write off ratio (trailing 12 months, annualised)", "Recovery rate (trailing 12 months)", "Eligible receivables (to 30 DPD)"]
    for i, lab in enumerate(show):
        r = 7 + i
        rngk = f"KPIs!$C${KR[lab]}:${MC[-1]}${KR[lab]}"
        fmt = NUM if "Eligible" in lab else PCT
        put(db, f"A{r}", lab, "label")
        put(db, f"B{r}", f'=IFERROR(INDEX({rngk},MATCH(9.99E+307,{rngk})),"")', fmt=fmt, bold=True)
        put(db, f"C{r}", f'=IFERROR(INDEX({rngk},MATCH(9.99E+307,{rngk})-3),"")', fmt=fmt)
        put(db, f"D{r}", f'=IFERROR(INDEX({rngk},MATCH(9.99E+307,{rngk})-12),"")', fmt=fmt)
        better_up = lab.startswith(("Collection", "Recovery", "Eligible"))
        put(db, f"E{r}", f'=IF(OR(B{r}="",C{r}=""),"",IF(B{r}=C{r},"Stable",IF((B{r}>C{r})={"TRUE" if better_up else "FALSE"},"Improving","Deteriorating")))')
        db.conditional_formatting.add(f"E{r}", FormulaRule(formula=[f'E{r}="Deteriorating"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    put(db, "A16", "Reconciliation, latest month", "label", bold=True)
    put(db, "B16", f'=IFERROR(INDEX(KPIs!$C${KR["Reconciliation check"]}:${MC[-1]}${KR["Reconciliation check"]},MATCH("zzz",KPIs!$C${KR["Reconciliation check"]}:${MC[-1]}${KR["Reconciliation check"]})),"")', bold=True)
    put(db, "A17", "Months with any covenant breach", "label", bold=True)
    put(db, "B17", "=Covenants!B22", fmt=NUM, bold=True)
    from openpyxl.chart import Series
    def line_chart(title, rows, anchor):
        ch_ = LineChart()
        ch_.title, ch_.height, ch_.width = title, 7.5, 16
        ch_.y_axis.number_format = "0%"
        for lab_ in rows:
            ch_.series.append(Series(Reference(k, min_col=3, max_col=lastc, min_row=KR[lab_], max_row=KR[lab_]), title=lab_))
        ch_.set_categories(Reference(k, min_col=3, max_col=lastc, min_row=6, max_row=6))
        db.add_chart(ch_, anchor)
    line_chart("Collection rate", ["Collection rate (month)", "Collection rate (trailing 3 months)"], "A20")
    line_chart("PAR30, reported and lagged", ["PAR30", "PAR30 lagged 3 months"], "F20")
    wb.move_sheet("Dashboard", offset=-(len(wb.sheetnames) - 2))
    finish_book(wb, OUT / "AEF_V2_T05_Lender_KPI_Report.xlsx", "AEF Volume 2 Template T05: Lender KPI report",
                "Africa Energy Finance, Volume 2 Templates Pack")
    return agg


# ================================================================== T06 Loan tape and data request
TAPE = [
    ("account_id", "Text", "Unique account identifier", "Yes", "Unique; not blank"),
    ("report_date", "Date", "Month end date of the extract", "Yes", "Same for all rows"),
    ("tier", "Integer", "Product tier 1 to 5 as labelled in the model", "Yes", "1 to 5"),
    ("product_name", "Text", "Commercial product name", "Yes", ""),
    ("origination_date", "Date", "Date of activation after deposit", "Yes", "Not after report date"),
    ("region", "Text", "Sales region", "Yes", ""),
    ("agent_id", "Text", "Originating agent", "Yes", "Links to agent register (T04)"),
    ("channel", "List", "Direct agent, retail partner, other", "Yes", "From list"),
    ("cash_price", "Number", "Cash selling price, local currency", "Yes", "Greater than zero"),
    ("deposit", "Number", "Deposit paid at activation", "Yes", "Not above cash price"),
    ("daily_rate", "Number", "Price of one day of service", "Yes", "Greater than zero"),
    ("tenor_months", "Integer", "Contractual tenor in months", "Yes", "Greater than zero"),
    ("scheduled_instalments", "Number", "Daily rate × 365 ÷ 12 × tenor (calculated)", "Calc", ""),
    ("cumulative_paid_excl_deposit", "Number", "Instalments paid to date, deposit excluded", "Yes", "Not above scheduled instalments"),
    ("instalments_due_to_date", "Number", "Instalments fallen due to date, deposit excluded", "Yes", ""),
    ("days_past_due", "Integer", "Schedule shortfall expressed in days of the daily rate", "Yes", "Zero or more"),
    ("days_without_credit", "Integer", "Consecutive days the device has been locked", "No", "Zero or more"),
    ("status", "List", "Active, Paid off, Defaulted, Repossessed, Written off, Restructured", "Yes", "From list"),
    ("restructure_date", "Date", "Date of the last restructuring, if any", "No", ""),
    ("write_off_amount", "Number", "Amount written off", "No", ""),
    ("write_off_date", "Date", "Date of write off", "No", ""),
    ("repossession_date", "Date", "Date the device was recovered", "No", ""),
    ("resale_proceeds", "Number", "Gross proceeds from resale of the recovered device", "No", ""),
    ("recovery_cost", "Number", "Cost of recovery, refurbishment and resale", "No", ""),
    ("unlock_date", "Date", "Date of permanent unlock", "No", ""),
    ("unlock_type", "List", "Full payment, Settlement, Warranty swap, Goodwill, Other", "No", "Full payment requires paid = scheduled"),
    ("customer_verified", "List", "Yes or No: verification call or visit completed", "Yes", "From list"),
]
STATUS_T = ["Active", "Paid off", "Defaulted", "Repossessed", "Written off", "Restructured"]
UNLOCK_T = ["Full payment", "Settlement", "Warranty swap", "Goodwill", "Other"]
DATA_REQ = [
    ("Account level loan tape (structure on the Tape sheet)", "Excel or CSV", "Since launch"),
    ("Monthly portfolio data by tier: gross receivables, DPD buckets, default exposure, originations, instalments due, collections, write offs, repossessions, resale proceeds and costs, cures, unlocks", "Credit_Input structure", "24 months minimum"),
    ("Cohort observations by tier at fixed ages", "Vintage_Input structure", "All cohorts"),
    ("Reconciliation of DPD buckets and gross receivables to the general ledger", "Excel", "Each month end"),
    ("Credit, write off, restructuring and repossession policies, with the dates each version applied", "PDF", "Current and past versions"),
    ("Price plan history by tier, including deposits, daily rates, tenors and promotions", "Excel", "Since launch"),
    ("Agent register, commission schedules, clawback rules and sales by agent", "Excel", "24 months"),
    ("Unlock logs separated between full payment unlocks and other unlocks", "Excel", "Since launch"),
    ("Audited accounts, monthly management accounts and current budget", "PDF and Excel", "3 years"),
    ("Financing agreements, compliance certificates, waivers and amendments", "PDF", "All current"),
    ("RBF and grant agreements, verification reports and clawback provisions", "PDF", "All current"),
    ("Customer complaint logs and regulatory correspondence", "Excel and PDF", "24 months"),
    ("Hardware supply contracts, warranty terms and FX exposure by currency", "PDF and Excel", "Current"),
]


def build_t06():
    wb = new_book()
    guide_sheet(wb, "T06", "Loan tape specification and data request",
                "The account level data tape a lender, investor or rating analyst needs to rebuild a PAYGo portfolio, with "
                "field definitions, built in validation checks and the standard data request list.",
                ["Send the Schema and Data_Request sheets to the company at the start of diligence.",
                 "Paste the extract into the Tape sheet, one account per row, from row 7 (up to 1,000 rows; extend the formulas for larger books).",
                 "Read the Validation sheet: every check should read zero before the tape is used.",
                 "Track each data request on Data_Request until it is received and reviewed.",
                 "Aggregate the validated tape into monthly and cohort figures for Credit_Input and Vintage_Input in the model, or for Template T05."],
                [("Book", "Chapter 12 (data tape), Chapter 15 (data request list)"),
                 ("Model", "Credit_Input and Vintage_Input are the aggregated form of this tape"),
                 ("Related templates", "T02 due diligence checklist; T05 lender KPI report")])
    sc = sheet(wb, "Schema", "Tape field specification", "Field, type, definition, mandatory flag and validation rule", "T06", 6, landscape=False)
    widths(sc, [6, 30, 10, 54, 10, 34])
    table_header(sc, 6, ["No.", "Field", "Type", "Definition", "Mandatory", "Validation"])
    for i, (f, t, d_, m, v) in enumerate(TAPE):
        r = 7 + i
        put(sc, f"A{r}", i + 1)
        put(sc, f"B{r}", f, bold=True)
        put(sc, f"C{r}", t)
        put(sc, f"D{r}", d_, wrap=True)
        put(sc, f"E{r}", m)
        put(sc, f"F{r}", v, wrap=True)
    NR = 1000
    tp = sheet(wb, "Tape", "Loan tape", "One row per account; three illustrative rows to be replaced", "T06", len(TAPE) + 6)
    hdr = [f for f, *_ in TAPE] + ["chk_present", "chk_missing", "chk_tier", "chk_status", "chk_deposit", "chk_paid",
                                   "chk_origination", "chk_unlock", "chk_dpd", "chk_duplicate"]
    table_header(tp, 6, hdr, height=32)
    widths(tp, [14] * len(hdr))
    col = {f: L(i + 1) for i, f in enumerate(hdr)}
    rd = date(2025, 12, 31)
    ex = [dict(account_id="SP000001", report_date=rd, tier=2, product_name="SHS with TV", origination_date=date(2024, 9, 14),
               region="North", agent_id="AG007", channel="Direct agent", cash_price=39000, deposit=4000, daily_rate=77, tenor_months=24,
               cumulative_paid_excl_deposit=29500, instalments_due_to_date=35130, days_past_due=73, days_without_credit=12,
               status="Active", customer_verified="Yes"),
          dict(account_id="SP000002", report_date=rd, tier=1, product_name="Pico solar kit", origination_date=date(2024, 2, 3),
               region="South", agent_id="AG012", channel="Retail partner", cash_price=6500, deposit=1000, daily_rate=25, tenor_months=12,
               cumulative_paid_excl_deposit=9125, instalments_due_to_date=9125, days_past_due=0, days_without_credit=0,
               status="Paid off", unlock_date=date(2025, 2, 10), unlock_type="Full payment", customer_verified="Yes"),
          dict(account_id="SP000003", report_date=rd, tier=3, product_name="Large SHS with DC fridge", origination_date=date(2024, 5, 20),
               region="East", agent_id="AG021", channel="Direct agent", cash_price=117000, deposit=17550, daily_rate=173, tenor_months=30,
               cumulative_paid_excl_deposit=21000, instalments_due_to_date=99980, days_past_due=456, days_without_credit=380,
               status="Written off", write_off_amount=136862, write_off_date=date(2025, 6, 30), repossession_date=date(2025, 8, 2),
               resale_proceeds=18000, recovery_cost=2700, customer_verified="Yes")]
    for i in range(NR):
        r = 7 + i
        e = ex[i] if i < len(ex) else {}
        for f, t, *_ in TAPE:
            c = f"{col[f]}{r}"
            if f == "scheduled_instalments":
                put(tp, c, f'=IF({col["daily_rate"]}{r}="","",{col["daily_rate"]}{r}*365/12*{col["tenor_months"]}{r})', fmt=NUM)
                continue
            v = e.get(f)
            fmt = "dd mmm yyyy" if t == "Date" else (NUM if t in ("Number", "Integer") else None)
            put(tp, c, v, "in", fmt)
        a = f"{col['account_id']}{r}"
        mand = [f for f, t, d_, m, v in TAPE if m == "Yes" and f != "account_id"]
        put(tp, f"{col['chk_present']}{r}", f'=IF(LEN({a})>0,1,0)', fmt=NUM)
        put(tp, f"{col['chk_missing']}{r}", f'=IF(LEN({a})=0,0,' + "+".join(f'(LEN({col[f]}{r})=0)' for f in mand) + ")", fmt=NUM)
        tc, sc_ = f"{col['tier']}{r}", f"{col['status']}{r}"
        put(tp, f"{col['chk_tier']}{r}", f'=IF(OR(LEN({a})=0,LEN({tc})=0),0,IF(OR({tc}<1,{tc}>5),1,0))', fmt=NUM)
        put(tp, f"{col['chk_status']}{r}", f'=IF(OR(LEN({a})=0,LEN({sc_})=0),0,IF(OR(' + ",".join(f'{sc_}="{x}"' for x in STATUS_T) + '),0,1))', fmt=NUM)
        put(tp, f"{col['chk_deposit']}{r}", f'=IF({a}="",0,IF(N({col["deposit"]}{r})>N({col["cash_price"]}{r}),1,0))', fmt=NUM)
        put(tp, f"{col['chk_paid']}{r}", f'=IF({a}="",0,IF(N({col["cumulative_paid_excl_deposit"]}{r})>N({col["scheduled_instalments"]}{r})+1,1,0))', fmt=NUM)
        put(tp, f"{col['chk_origination']}{r}", f'=IF({a}="",0,IF(N({col["origination_date"]}{r})>N({col["report_date"]}{r}),1,0))', fmt=NUM)
        put(tp, f"{col['chk_unlock']}{r}", f'=IF({a}="",0,IF(AND({col["unlock_type"]}{r}="Full payment",N({col["cumulative_paid_excl_deposit"]}{r})<N({col["scheduled_instalments"]}{r})-1),1,0))', fmt=NUM)
        put(tp, f"{col['chk_dpd']}{r}", f'=IF({a}="",0,IF(N({col["days_past_due"]}{r})<0,1,0))', fmt=NUM)
        put(tp, f"{col['chk_duplicate']}{r}", f'=IF({a}="",0,IF(COUNTIF(${col["account_id"]}$7:${col["account_id"]}${6 + NR},{a})>1,1,0))', fmt=NUM)
    last = 6 + NR
    dv_list(tp, STATUS_T, f"{col['status']}7:{col['status']}{last}")
    dv_list(tp, UNLOCK_T, f"{col['unlock_type']}7:{col['unlock_type']}{last}")
    dv_list(tp, ["Yes", "No"], f"{col['customer_verified']}7:{col['customer_verified']}{last}")
    dv_list(tp, ["Direct agent", "Retail partner", "Other"], f"{col['channel']}7:{col['channel']}{last}")
    tp.freeze_panes = "B7"
    tp.print_title_rows = "6:6"

    va = sheet(wb, "Validation", "Tape validation", "Every check should read zero before the tape is used", "T06", 4, landscape=False)
    widths(va, [56, 14, 14, 30])
    table_header(va, 6, ["Check", "Exceptions", "Status", "Notes"])
    H = lambda c: f"Tape!${col[c]}$7:${col[c]}${last}"
    r = 7
    put(va, f"A{r}", "Accounts in tape", "label", bold=True)
    put(va, f"B{r}", f"=SUM({H('chk_present')})", fmt=NUM, bold=True)
    r += 1
    extra = [("Rows with a missing mandatory field", f'=COUNTIF({H("chk_missing")},">0")'),
             ("Missing mandatory values in total", f"=SUM({H('chk_missing')})"),
             ("Tier outside 1 to 5", f"=SUM({H('chk_tier')})"),
             ("Status not in the list", f"=SUM({H('chk_status')})"),
             ("Deposit above cash price", f"=SUM({H('chk_deposit')})"),
             ("Paid above scheduled instalments", f"=SUM({H('chk_paid')})"),
             ("Origination after report date", f"=SUM({H('chk_origination')})"),
             ("Full payment unlock without full payment", f"=SUM({H('chk_unlock')})"),
             ("Negative days past due", f"=SUM({H('chk_dpd')})"),
             ("Duplicate account identifiers (rows)", f"=SUM({H('chk_duplicate')})")]
    for lab, f in extra:
        put(va, f"A{r}", lab, "label")
        put(va, f"B{r}", f, fmt=NUM)
        put(va, f"C{r}", f'=IF(B{r}=0,"OK","Fix")')
        r += 1
    put(va, f"A{r + 1}", "Tape ready for use", "label", bold=True)
    put(va, f"B{r + 1}", f'=IF(COUNTIF(C8:C{r - 1},"Fix")=0,"Yes","No")', bold=True)
    va.conditional_formatting.add(f"C8:C{r - 1}", FormulaRule(formula=['C8="Fix"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    put(va, f"A{r + 3}", "Each check is computed per account in the chk columns at the right of the Tape sheet; filter a chk column on 1 to find the rows to fix. The mandatory fields are listed on Schema.", "note")

    dr = sheet(wb, "Data_Request", "Data request list", "Issue at the start of diligence", "T06", 9, landscape=True)
    widths(dr, [6, 60, 18, 18, 14, 14, 14, 14, 30])
    table_header(dr, 6, ["Ref", "Item", "Format", "Period", "Owner", "Requested", "Received", "Status", "Comments"])
    for i, (it, fm, per) in enumerate(DATA_REQ):
        r = 7 + i
        put(dr, f"A{r}", f"DR{i + 1:02d}")
        put(dr, f"B{r}", it, wrap=True)
        put(dr, f"C{r}", fm)
        put(dr, f"D{r}", per)
        for c in "EFGI":
            put(dr, f"{c}{r}", None, "in", "dd mmm yyyy" if c in "FG" else None)
        put(dr, f"H{r}", "Not requested", "in")
        dr.row_dimensions[r].height = 30 if len(it) > 60 else 16
    dv_list(dr, ["Not requested", "Requested", "Received", "Reviewed", "Incomplete"], f"H7:H{6 + len(DATA_REQ)}")
    finish_book(wb, OUT / "AEF_V2_T06_Loan_Tape_and_Data_Request.xlsx", "AEF Volume 2 Template T06: Loan tape specification and data request",
                "Africa Energy Finance, Volume 2 Templates Pack")


# ================================================================== T07 Borrowing base certificate and term sheet
TERM = [
    ("Borrower and structure", "Operating company or SPV; security package; true sale or secured loan"),
    ("Facility amount and tenor", "Commitment, availability period, amortisation after availability ends"),
    ("Currency and pricing", "Local or hard currency; margin, fees, hedging cost"),
    ("Borrowing base", "Eligible receivables definition, advance rates by tier, concentration limits by product, region and agent"),
    ("Eligibility criteria", "Maximum DPD, minimum payments made, no restructured accounts, contracts compliant with consumer law"),
    ("Cash management", "Collection accounts, sweep mechanics, priority of payments"),
    ("Portfolio covenants", "Collection rate, RaR, PAR30 and PAR90, write off ratio, with definitions and calculation dates"),
    ("Corporate covenants", "Debt to equity, minimum liquidity, DSCR if any, equity cure rights"),
    ("Cure periods and triggers", "Cure mechanics; triggers for stop of advances and early amortisation"),
    ("Reporting", "Monthly tape and KPI report, audit rights, servicer reports"),
    ("Conditions precedent", "Legal opinions, data tape audit, platform review, insurance"),
    ("Events of default", "Payment default, covenant breach after cure, change of control, licence loss"),
]


def build_t07():
    wb = new_book()
    guide_sheet(wb, "T07", "Borrowing base certificate and term sheet checklist",
                "A monthly borrowing base certificate for a PAYGo receivables facility, with eligibility by DPD bucket, advance "
                "rates by tier, a concentration limit on unproven products and headroom against the drawn balance; and a "
                "checklist of the terms to settle in the facility term sheet.",
                ["Enter gross receivables by tier and DPD bucket at month end on the Certificate sheet (blue cells).",
                 "Set the maximum DPD for eligibility, advance rates, the concentration limit and flag the tiers it applies to.",
                 "Enter ineligible deductions (restructured, deposit only, unverified accounts) and the drawn balance.",
                 "Read the borrowing base, availability and headroom; use the stress block to test a migration between buckets.",
                 "Use Term_Sheet to record each party's position on every term until agreed."],
                [("Book", "Chapter 11 (borrowing base, eligibility, covenants), Chapter 12 (triggers); Annex F"),
                 ("Model", "Products (advance rates), Credit_Assumptions (maximum DPD for eligibility), Credit_Portfolio"),
                 ("Related templates", "T05 lender KPI report; T06 loan tape")])
    ws = sheet(wb, "Certificate", "Borrowing base certificate", "Month end; local currency millions; example from Chapter 11 of the book", "T07", 11)
    widths(ws, [30, 12, 12, 12, 12, 12, 12, 13, 13, 13, 14])
    put(ws, "A6", "Certificate date", "label")
    put(ws, "B6", date(2026, 9, 30), "in", "dd mmm yyyy")
    put(ws, "A7", "Maximum DPD for eligibility", "label")
    put(ws, "B7", 30, "in", NUM)
    put(ws, "A8", "Concentration limit on flagged tiers (share of eligible)", "label")
    put(ws, "B8", 0.12, "in", PCT)
    put(ws, "A9", "Facility limit", "label")
    put(ws, "B9", 2000, "in", NUM)
    put(ws, "A10", "Facility drawn", "label")
    put(ws, "B10", 1500, "in", NUM)
    buckets = [("Current", 0), ("1 to 30", 1), ("31 to 60", 31), ("61 to 90", 61), ("91 to 180", 91), ("Over 180", 181)]
    table_header(ws, 12, ["Tier"] + [f"{b} DPD" if b != "Current" else b for b, _ in buckets] + ["Gross", "Advance rate", "Flag (Yes/No)", "Ineligible deductions"])
    put(ws, "A13", "Lower DPD bound", "label")
    for j, (_, lb) in enumerate(buckets):
        put(ws, f"{L(2 + j)}13", lb, fmt=NUM)
    data = [(120, 20, 15, 0, 25, 0), (930, 150, 110, 0, 140, 0), (700, 100, 60, 0, 70, 0), (290, 30, 10, 0, 10, 0), (145, 15, 5, 0, 5, 0)]
    adv = [0.50, 0.70, 0.70, 0.75, 0.75]
    flag = ["No", "No", "No", "Yes", "Yes"]
    for t in range(5):
        r = 14 + t
        put(ws, f"A{r}", f"Tier {t + 1}", "label", bold=True)
        for j, v in enumerate(data[t]):
            put(ws, f"{L(2 + j)}{r}", v, "in", NUM)
        put(ws, f"H{r}", f"=SUM(B{r}:G{r})", fmt=NUM)
        put(ws, f"I{r}", adv[t], "in", PCT)
        put(ws, f"J{r}", flag[t], "in")
        put(ws, f"K{r}", 0, "in", NUM)
    dv_list(ws, ["Yes", "No"], "J14:J18")
    put(ws, "A19", "Total", "label", bold=True)
    for c in "BCDEFGHK":
        put(ws, f"{c}19", f"=SUM({c}14:{c}18)", fmt=NUM, bold=True)
    put(ws, "A20", "Note: the example merges 31 to 90 DPD into the 31 to 60 column and 91 DPD and over into the 91 to 180 column, as in the book.", "note")

    table_header(ws, 22, ["Tier", "Eligible", "Advance rate", "Base before limits", "Flagged eligible", "", "", "", "", "", ""])
    for t in range(5):
        r = 23 + t
        src = 14 + t
        put(ws, f"A{r}", f"Tier {t + 1}", "label", bold=True)
        put(ws, f"B{r}", f"=MAX(0,SUMPRODUCT((B$13:G$13<=$B$7)*B{src}:G{src})-K{src})", fmt=NUM)
        put(ws, f"C{r}", f"=I{src}", fmt=PCT)
        put(ws, f"D{r}", f"=B{r}*C{r}", fmt=NUM)
        put(ws, f"E{r}", f'=IF(J{src}="Yes",B{r},0)', fmt=NUM)
    put(ws, "A28", "Total", "label", bold=True)
    for c in "BDE":
        put(ws, f"{c}28", f"=SUM({c}23:{c}27)", fmt=NUM, bold=True)

    section(ws, 30, "Borrowing base", 11)
    lines = [("Borrowing base before concentration limits", "=D28", NUM),
             ("Concentration cap on flagged tiers", "=B8*B28", NUM),
             ("Excess flagged receivables", "=MAX(0,E28-B32)", NUM),
             ("Weighted advance rate on flagged tiers", "=IF(E28=0,0,SUMPRODUCT(E23:E27,C23:C27)/E28)", PCT),
             ("Less excess concentration", "=-B33*B34", NUM),
             ("Borrowing base", "=B31+B35", NUM),
             ("Available (lower of limit and borrowing base)", "=MIN(B9,B36)", NUM),
             ("Facility drawn", "=B10", NUM),
             ("Headroom", "=B37-B38", NUM),
             ("Headroom as share of drawn", "=IF(B38=0,0,B39/B38)", PCT),
             ("Effective advance on gross receivables", "=IF(H19=0,0,B36/H19)", PCT),
             ("Status", '=IF(B39<0,"Borrowing base breach: repay or cure",IF(B40<0.05,"Watch: headroom below 5% of drawn","Compliant"))', None)]
    for i, (lab, f, fmt) in enumerate(lines):
        r = 31 + i
        put(ws, f"A{r}", lab, "label", bold=lab in ("Borrowing base", "Headroom", "Status"))
        put(ws, f"B{r}", f, fmt=fmt, bold=lab in ("Borrowing base", "Headroom", "Status"))
    ws.merge_cells("B42:E42")
    ws.conditional_formatting.add("B42", FormulaRule(formula=['LEFT(B42,6)="Borrow"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    ws.conditional_formatting.add("B42", FormulaRule(formula=['LEFT(B42,5)="Watch"'], fill=PatternFill("solid", fgColor="FCE8B2")))

    section(ws, 44, "Stress: migration out of the eligible buckets", 11)
    put(ws, "A45", "Tier migrating", "label")
    put(ws, "B45", 2, "in", NUM)
    put(ws, "A46", "Amount moving from 1 to 30 DPD to 31 to 60 DPD", "label")
    put(ws, "B46", 100, "in", NUM)
    put(ws, "A47", "Eligible receivables after migration", "label")
    put(ws, "B47", "=B28-IF(B7>=1,MIN(B46,INDEX(C14:C18,B45)),0)+IF(B7>=31,MIN(B46,INDEX(C14:C18,B45)),0)", fmt=NUM)
    put(ws, "A48", "Flagged eligible after migration", "label")
    put(ws, "B48", '=E28-IF(INDEX(J14:J18,B45)="Yes",B28-B47,0)', fmt=NUM)
    put(ws, "A49", "Base before limits after migration", "label")
    put(ws, "B49", "=D28-(B28-B47)*INDEX(I14:I18,B45)", fmt=NUM)
    put(ws, "A50", "Excess concentration deduction after migration", "label")
    put(ws, "B50", "=MAX(0,B48-B8*B47)*B34", fmt=NUM)
    put(ws, "A51", "Borrowing base after migration", "label", bold=True)
    put(ws, "B51", "=B49-B50", fmt=NUM, bold=True)
    put(ws, "A52", "Headroom after migration", "label", bold=True)
    put(ws, "B52", "=MIN(B9,B51)-B10", fmt=NUM, bold=True)
    put(ws, "A53", "Headroom lost", "label")
    put(ws, "B53", "=B39-B52", fmt=NUM)

    ts = sheet(wb, "Term_Sheet", "Receivables facility term sheet checklist", "Record each party's position until agreed", "T07", 8)
    widths(ts, [26, 44, 30, 30, 12, 12, 30, 8])
    table_header(ts, 6, ["Term", "Points to settle", "Company position", "Lender position", "Agreed", "Condition precedent", "Notes", "Book"])
    for i, (t, p) in enumerate(TERM):
        r = 7 + i
        put(ts, f"A{r}", t, bold=True, wrap=True)
        put(ts, f"B{r}", p, wrap=True)
        for c in "CDG":
            put(ts, f"{c}{r}", None, "in")
        put(ts, f"E{r}", "Open", "in")
        put(ts, f"F{r}", "No", "in")
        put(ts, f"H{r}", 11)
        ts.row_dimensions[r].height = 30
    dv_list(ts, ["Open", "Agreed", "Dropped"], "E7:E18")
    dv_list(ts, ["Yes", "No"], "F7:F18")
    put(ts, "A20", "Terms agreed", "label", bold=True)
    put(ts, "B20", '=COUNTIF(E7:E18,"Agreed")&" of "&COUNTA(A7:A18)', bold=True)
    finish_book(wb, OUT / "AEF_V2_T07_Borrowing_Base_and_Term_Sheet.xlsx", "AEF Volume 2 Template T07: Borrowing base certificate and term sheet checklist",
                "Africa Energy Finance, Volume 2 Templates Pack")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    n = build_t02()
    build_t04()
    build_t05()
    build_t06()
    build_t07()
    import build_shs_templates_docx as W  # noqa: E402
    W.build_all(OUT)
    print(f"checklist items: {n}")
