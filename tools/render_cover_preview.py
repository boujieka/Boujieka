"""Approximate PNG preview of the Excel Cover sheet (layout check only; Excel's own rendering may differ slightly).

Run: python tools/render_cover_preview.py <workbook.xlsx> <out.png>
"""
import sys

import openpyxl
from openpyxl.utils import get_column_letter as gcl
from PIL import Image, ImageDraw, ImageFont

FORMULA_DISPLAY = {  # default-input values shown for formula cells
    "E35": "OK", "E36": "INVESTMENT READINESS: 3/23 gates met - NOT investment grade", "E37": "Base",
    "E38": "Proxy (model curves)", "E39": "LCY (local), USD for hardware, RBF and term debt"}
wb = openpyxl.load_workbook(sys.argv[1])
ws = wb["Cover"]
S = 1.6
colx = [0]
for c in range(1, 14):
    colx.append(colx[-1] + int(((ws.column_dimensions[gcl(c)].width or 8.43) * 7 + 5) * S))
rowy = [0]
for r in range(1, 49):
    h = ws.row_dimensions[r].height or 15
    rowy.append(rowy[-1] + int(h * 96 / 72 * S))
W, H = colx[-1], rowy[-1]
img = Image.new("RGB", (W, H), "white")
d = ImageDraw.Draw(img)


def font(sz, bold=False, italic=False):
    name = "LiberationSans-" + ("BoldItalic" if bold and italic else "Bold" if bold else "Italic" if italic else "Regular")
    try:
        return ImageFont.truetype(f"/usr/share/fonts/truetype/liberation/{name}.ttf", int(sz * 96 / 72 * S))
    except OSError:
        return ImageFont.truetype("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf", int(sz * 96 / 72 * S))


for r in range(1, 49):
    for c in range(1, 14):
        cell = ws.cell(r, c)
        fg = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
        if fg and isinstance(fg, str):
            d.rectangle([colx[c - 1], rowy[r - 1], colx[c], rowy[r]], fill="#" + fg[-6:])
for img_obj in ws._images:
    anc = img_obj.anchor._from
    logo = Image.open(img_obj.ref if isinstance(img_obj.ref, str) else img_obj.ref).convert("RGBA")
    ext = getattr(img_obj.anchor, "ext", None)
    w_px, h_px = (ext.width / 9525, ext.height / 9525) if ext is not None else (img_obj.width, img_obj.height)
    logo = logo.resize((int(w_px * S), int(h_px * S)))
    img.paste(logo, (colx[anc.col], rowy[anc.row]), logo)
merged = {str(m).split(":")[0]: m for m in ws.merged_cells.ranges}
for r in range(1, 49):
    for c in range(1, 14):
        cell = ws.cell(r, c)
        v = cell.value
        if v is None:
            continue
        if isinstance(v, str) and v.startswith("="):
            v = FORMULA_DISPLAY.get(cell.coordinate, "")
        f = cell.font
        col = "#" + (f.color.rgb[-6:] if f.color is not None and isinstance(f.color.rgb, str) else "000000")
        fnt = font(f.size or 11, f.bold, f.italic)
        y = rowy[r - 1] + (rowy[r] - rowy[r - 1]) // 2
        if cell.coordinate in merged:
            m = merged[cell.coordinate]
            maxw = colx[m.max_col] - colx[c - 1]
            words, line, yy = str(v).split(), "", rowy[r - 1] + 4
            for w_ in words:
                t = (line + " " + w_).strip()
                if d.textlength(t, font=fnt) > maxw:
                    d.text((colx[c - 1], yy), line, font=fnt, fill=col); yy += int(fnt.size * 1.3); line = w_
                else:
                    line = t
            d.text((colx[c - 1], yy), line, font=fnt, fill=col)
        else:
            d.text((colx[c - 1] + 4, y), str(v), font=fnt, fill=col, anchor="lm")
            if f.underline:
                tw = d.textlength(str(v), font=fnt)
                d.line([colx[c - 1] + 4, y + fnt.size // 2 + 2, colx[c - 1] + 4 + tw, y + fnt.size // 2 + 2], fill=col, width=2)
img.save(sys.argv[2])
print(sys.argv[2], img.size)
