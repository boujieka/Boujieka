"""Dependency check for MODEL 7: builds the cell dependency graph from the formulas and reports any cycle.

Run: python3 tools/model/check_cycles.py model/Bankable_Hydro_Model.xlsx
"""
import re,sys
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter as L
wb=load_workbook(sys.argv[1])
g={}
pat=re.compile(r"(?:(?:'([^']+)'|([A-Za-z0-9_]+))!)?(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)")
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str) and v.startswith('='):
                f=re.sub(r'"[^"]*"','',v)
                deps=[]
                for sh,sh2,ref in pat.findall(f):
                    sh=sh or sh2 or ws.title
                    ref=ref.replace('$','')
                    if ':' in ref:
                        a,b,c2,d=range_boundaries(ref)
                        for rr in range(b,d+1):
                            for cc in range(a,c2+1): deps.append((sh,f"{L(cc)}{rr}"))
                    else: deps.append((sh,ref))
                g[(ws.title,c.coordinate)]=deps
sys.setrecursionlimit(100000)
color={}
stack=[]
def dfs(n):
    color[n]=1; stack.append(n)
    for m in g.get(n,[]):
        if m not in g: continue
        if color.get(m)==1:
            i=stack.index(m); print("CYCLE:", stack[i:][:40]); sys.exit()
        if m not in color: dfs(m)
    color[n]=2; stack.pop()
for n in list(g):
    if n not in color: dfs(n)
print("no cycle")
