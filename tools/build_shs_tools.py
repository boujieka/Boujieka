"""Build the Book 2 Decision Tools (SHS PAYGo): six Excel calculators in the house style.

Run: python tools/build_shs_tools.py
Writes volumes/02-solar-home-systems/decision-tools/AEF_V2_D0x_*.xlsx

Default inputs reproduce worked examples of Book 2 or the SolaraPay case, so that each tool can be checked
against a published figure. Formulas avoid array functions so that they evaluate identically in any spreadsheet engine.
"""
import copy
import sys
from pathlib import Path

from openpyxl.chart import LineChart, Reference, Series
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).parent))
import shs_defaults as D  # noqa: E402
from build_shs_templates import (NUM, NUM1, PCT, RAT, dv_list, finish_book, guide_sheet, new_book, put, section,  # noqa: E402
                                 sheet, table_header, widths)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "volumes/02-solar-home-systems/decision-tools"
PCT2 = '0.00%;(0.00%);0.00%'
RED = PatternFill("solid", fgColor="F4CCCC")
AMB = PatternFill("solid", fgColor="FCE8B2")
GRN = PatternFill("solid", fgColor="D9EAD3")
SUBJ = "Africa Energy Finance, Book 2 Decision Tools"
G = D.GENERAL
P = D.PRODUCTS


def inputs_block(ws, row, items, col_label="A", col_val="B", col_note="D"):
    """items: (key, label, value, fmt, note). Returns {key: absolute ref}."""
    ref = {}
    for k, lab, v, fmt, note in items:
        put(ws, f"{col_label}{row}", lab, "label")
        put(ws, f"{col_val}{row}", v, "in", fmt)
        if note:
            put(ws, f"{col_note}{row}", note, "note")
        ref[k] = f"'{ws.title}'!${col_val}${row}"
        row += 1
    return ref, row


def flag_colours(ws, rng, red_text, amber_text=None, green_text=None):
    first = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{red_text}",{first}))'], fill=RED))
    if amber_text:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{amber_text}",{first}))'], fill=AMB))
    if green_text:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{green_text}",{first}))'], fill=GRN))


# ================================================================== D1 Price plan and APR calculator
def build_d1():
    wb = new_book()
    guide_sheet(wb, "D1", "Price plan and APR calculator",
                "Compares up to five PAYGo price plans on the figures that matter to a customer, a regulator and a lender: "
                "instalment, total contract value, PAYGo premium, implied monthly rate, nominal APR, effective annual rate, flat "
                "rate and payment burden; and solves for the daily rate, tenor or price that meets an affordability threshold.",
                ["Enter each plan's cash price, deposit, daily rate and tenor, and the household incomes of its target segment.",
                 "Read the implied rates and the payment burden; the flag shows whether the plan meets the threshold in a lean month.",
                 "Use the solver block to find the daily rate, tenor or price reduction that meets the threshold, or the daily rate for a target APR.",
                 "State which APR convention a disclosure uses; confirm with local counsel which one the rules require."],
                [("Book", "Chapter 2 (payment burden, implied APR), Chapter 3 (APR caps), Chapter 4 (price plan design)"),
                 ("Model", "Products (price plans), Consumer_Risk (burden, APR)"),
                 ("Default inputs", "Plan A is the illustrative Tier 2 plan of Chapters 1 and 2; Plans B to E are the model's default Tiers 1 to 4. Lean month income is set at two thirds of average income for illustration.")])
    ws = sheet(wb, "Price_Plans", "Price plan and APR calculator", "Local currency; five plans side by side", "D1", 8)
    widths(ws, [52, 12, 15, 15, 15, 15, 15, 30])
    table_header(ws, 6, ["Item", "Unit", "Plan A", "Plan B", "Plan C", "Plan D", "Plan E", "Notes"])
    plans = [("Chapter 2 example", 40000, 4000, 72, 24, 18000)] + \
            [(p["name"].split(" - ")[0] + " default", p["price"], p["deposit"], p["daily"], p["tenor"], p["income"]) for p in P[:4]]
    rows_in = [("Plan name", "", None), ("Cash price", "LCY", NUM), ("Deposit", "LCY", NUM), ("Daily rate", "LCY", NUM1),
               ("Tenor", "months", NUM), ("Average monthly household income", "LCY", NUM), ("Lean month household income", "LCY", NUM)]
    for i, (lab, unit, fmt) in enumerate(rows_in):
        r = 7 + i
        put(ws, f"A{r}", lab, "label", bold=(i == 0))
        put(ws, f"B{r}", unit, "label")
        for j, pl in enumerate(plans):
            c = L(3 + j)
            v = [pl[0], pl[1], pl[2], pl[3], pl[4], pl[5], round(pl[5] * 2 / 3)][i]
            put(ws, f"{c}{r}", v, "in", fmt)
    put(ws, "H13", "Illustrative: two thirds of average income", "note")
    section(ws, 15, "Results", 8)
    put(ws, "A16", "Maximum payment burden (threshold)", "label", bold=True)
    put(ws, "C16", G["afford_max"], "in", PCT)
    put(ws, "H16", "Policy choice, not a rule (Chapter 2)", "note")
    put(ws, "A17", "Target nominal APR for the solver", "label", bold=True)
    put(ws, "C17", 0.40, "in", PCT)
    out = [
        ("Amount financed", "LCY", "={c}8-{c}9", NUM),
        ("Monthly instalment", "LCY", "={c}10*365/12", NUM),
        ("Scheduled instalments", "LCY", "={c}20*{c}11", NUM),
        ("Total contract value", "LCY", "={c}9+{c}21", NUM),
        ("PAYGo premium", "LCY", "={c}22-{c}8", NUM),
        ("Premium as share of cash price", "%", "=IF({c}8=0,0,{c}23/{c}8)", PCT),
        ("Implied monthly rate", "%", '=IFERROR(RATE({c}11,-{c}20,{c}19),"")', PCT2),
        ("Nominal APR (12 × monthly rate)", "%", '=IF({c}25="","",12*{c}25)', PCT),
        ("Effective annual rate", "%", '=IF({c}25="","",(1+{c}25)^12-1)', PCT),
        ("Flat rate (premium ÷ amount financed ÷ years)", "%", "=IF({c}19=0,0,{c}23/{c}19/({c}11/12))", PCT),
        ("Payment burden, average income", "%", "=IF({c}12=0,0,{c}20/{c}12)", PCT),
        ("Payment burden, lean month", "%", "=IF({c}13=0,0,{c}20/{c}13)", PCT),
        ("Deposit as share of average monthly income", "%", "=IF({c}12=0,0,{c}9/{c}12)", PCT),
        ("Affordability", "", '=IF({c}30>$C$16,IF({c}29>$C$16,"Above threshold on average","Above threshold in lean months"),"Within threshold")', None),
    ]
    for i, (lab, unit, f, fmt) in enumerate(out):
        r = 19 + i
        put(ws, f"A{r}", lab, "label", bold=lab in ("Nominal APR (12 × monthly rate)", "Affordability"))
        put(ws, f"B{r}", unit, "label")
        for j in range(5):
            c = L(3 + j)
            put(ws, f"{c}{r}", f.format(c=c), fmt=fmt, bold=lab == "Affordability")
    flag_colours(ws, "C32:G32", "Above threshold on", "lean", "Within")
    section(ws, 34, "Solver", 8)
    solve = [
        ("Daily rate that meets the threshold on lean month income", "LCY", "=$C$16*{c}13*12/365", NUM1),
        ("Daily rate that meets the threshold on average income", "LCY", "=$C$16*{c}12*12/365", NUM1),
        ("Daily rate for the target nominal APR", "LCY", "=IF({c}19<=0,0,PMT($C$17/12,{c}11,-{c}19)*12/365)", NUM1),
        ("Tenor that meets the lean month threshold at the same monthly rate", "months", '=IFERROR(NPER({c}25,-$C$16*{c}13,{c}19),"Not reachable")', NUM1),
        ("Reduction in amount financed to meet the lean month threshold, same rate and tenor", "LCY", '=IFERROR(MAX(0,{c}19-PV({c}25,{c}11,-$C$16*{c}13)),"")', NUM),
    ]
    for i, (lab, unit, f, fmt) in enumerate(solve):
        r = 35 + i
        put(ws, f"A{r}", lab, "label", wrap=True)
        put(ws, f"B{r}", unit, "label")
        ws.row_dimensions[r].height = 26
        for j in range(5):
            c = L(3 + j)
            put(ws, f"{c}{r}", f.format(c=c), fmt=fmt)
    put(ws, "A41", "A longer tenor lowers the instalment but extends the months over which default hazard acts (Chapter 4). "
        "\"Not reachable\" means no tenor brings the instalment within the threshold at this monthly rate.", "note")
    finish_book(wb, OUT / "AEF_V2_D1_Price_Plan_and_APR_Calculator.xlsx", "AEF Book 2 Decision Tool D1: Price plan and APR calculator", SUBJ)


# ================================================================== D2 Unit economics calculator and tier dashboard
def build_d2():
    wb = new_book()
    guide_sheet(wb, "D2", "Unit economics calculator",
                "Builds the expected lifetime cash flow of one PAYGo sale month by month from the repayment curve, and derives "
                "expected loss, contribution, LTV to CAC, cash payback, peak funding, unit IRR and unit NPV, with a sensitivity "
                "grid on default hazard and collection rate. A second sheet compares the five default tiers in closed form.",
                ["Enter the price plan, credit behaviour, recovery and cost inputs on Inputs (blue cells).",
                 "Read Results: contribution and LTV to CAC, but also payback, peak funding and unit IRR, which carry the timing.",
                 "Use the sensitivity grid to see how contribution moves with the default hazard and the collection rate.",
                 "Check that the closed form check on Results reads zero difference; it confirms the monthly table.",
                 "Use Tier_Dashboard to compare tiers; replace the default tier inputs with the company's own."],
                [("Book", "Chapter 1 (one Tier 2 sale), Chapter 6 (survival curve), Chapter 9 (unit economics)"),
                 ("Model", "Curves, Unit_Economics"),
                 ("Default inputs", "Reproduce the worked example of Chapter 1: expected collections LCY 33,832, payback in month 14 and lifetime cash of LCY 11,832, with servicing, fees and recoveries set to zero. Add them for a full view.")])
    ws = sheet(wb, "Inputs", "Unit inputs", "One product, local currency", "D2", 5, landscape=False)
    widths(ws, [52, 14, 4, 46, 4])
    items = [("price", "Cash price", 40000, NUM, ""), ("dep", "Deposit", 4000, NUM, ""), ("daily", "Daily rate", 72, NUM1, ""),
             ("T", "Tenor (months)", 24, NUM, "Up to 48 months"), ("h", "Monthly default hazard", 0.026, PCT2, "Share of paying accounts that stop for good each month"),
             ("c", "Collection rate on paying accounts", 0.88, PCT, ""), ("repo", "Repossession rate of defaulted units", 0.0, PCT, "Chapter 1 example: no recoveries"),
             ("lag", "Repossession lag (months)", G["repo_lag"], NUM, ""), ("resale", "Resale value as share of cash price", 0.25, PCT, ""),
             ("rcost", "Recovery cost as share of resale proceeds", G["recov_cost"], PCT, ""), ("hw", "Landed hardware cost", 22000, NUM, ""),
             ("inst", "Installation", 0, NUM, ""), ("war", "Warranty provision, share of landed cost", 0.0, PCT, ""),
             ("cac", "Customer acquisition cost (commission and marketing)", 4000, NUM, ""), ("serv", "Servicing cost per active account per month", 0, NUM, f"Model default {G['cs_cost']}"),
             ("fee", "Payment fees, share of cash collected", 0.0, PCT, f"Model default {G['mm_fee']:.0%}"),
             ("rbf", "RBF per unit (local currency)", 0, NUM, ""), ("rbfm", "Month RBF is received", G["rbf_lag"], NUM, ""),
             ("disc", "Discount rate (annual) for unit NPV", 0.25, PCT, "")]
    section(ws, 6, "Inputs", 5)
    R, r = inputs_block(ws, 7, items)
    r += 1
    section(ws, r, "Derived", 5)
    r += 1
    put(ws, f"A{r}", "Monthly instalment", "label")
    put(ws, f"B{r}", f"={R['daily']}*365/12", fmt=NUM)
    R["instal"] = f"'Inputs'!$B${r}"
    r += 1
    put(ws, f"A{r}", "Scheduled instalments", "label")
    put(ws, f"B{r}", f"={R['instal']}*{R['T']}", fmt=NUM)
    R["sched"] = f"'Inputs'!$B${r}"
    r += 1
    put(ws, f"A{r}", "Day one outlay (hardware, installation, warranty, CAC)", "label")
    put(ws, f"B{r}", f"={R['hw']}*(1+{R['war']})+{R['inst']}+{R['cac']}", fmt=NUM)
    R["outlay"] = f"'Inputs'!$B${r}"
    r += 1
    put(ws, f"A{r}", "Check: tenor plus repossession lag within 60 months", "label")
    put(ws, f"B{r}", f'=IF({R["T"]}+{R["lag"]}<=60,"OK","Reduce tenor or lag")', bold=True)
    R["chk"] = f"'Inputs'!$B${r}"

    cf = sheet(wb, "Unit_Cash_Flow", "Expected cash flow of one unit", "Months since sale; expected values", "D2", 16)
    hdr = ["Month", "Share still paying", "Instalment due", "Collected", "Deposit", "Defaults in month", "Recoveries",
           "Servicing", "Payment fees", "RBF", "Day one outlay", "Net cash", "Cumulative", "Discount factor", "Present value", "Cumulative positive"]
    table_header(cf, 6, hdr)
    widths(cf, [8] + [13] * 15)
    first, last = 7, 67
    for a in range(0, 61):
        r = first + a
        put(cf, f"A{r}", a, fmt=NUM)
        put(cf, f"B{r}", f"=(1-{R['h']})^A{r}", fmt=PCT)
        within = f"AND(A{r}>=1,A{r}<={R['T']})"
        put(cf, f"C{r}", f"=IF({within},{R['instal']},0)", fmt=NUM)
        put(cf, f"D{r}", f"=IF({within},{R['instal']}*{R['c']}*B{r},0)", fmt=NUM)
        put(cf, f"E{r}", f"=IF(A{r}=0,{R['dep']},0)", fmt=NUM)
        put(cf, f"F{r}", f"=IF({within},B{r - 1}-B{r},0)" if a > 0 else "=0", fmt='0.0000')
        put(cf, f"G{r}", f"=IF(A{r}-{R['lag']}>=1,INDEX($F${first}:$F${last},A{r}-{R['lag']}+1)*{R['repo']}*{R['resale']}*{R['price']}*(1-{R['rcost']}),0)", fmt=NUM)
        put(cf, f"H{r}", f"=-IF({within},{R['serv']}*B{r},0)", fmt=NUM)
        put(cf, f"I{r}", f"=-{R['fee']}*(D{r}+E{r}+G{r})", fmt=NUM)
        put(cf, f"J{r}", f"=IF(A{r}={R['rbfm']},{R['rbf']},0)", fmt=NUM)
        put(cf, f"K{r}", f"=IF(A{r}=0,-{R['outlay']},0)", fmt=NUM)
        put(cf, f"L{r}", f"=SUM(D{r}:K{r})-F{r}", fmt=NUM)
        put(cf, f"M{r}", f"=L{r}" if a == 0 else f"=M{r - 1}+L{r}", fmt=NUM)
        put(cf, f"N{r}", f"=(1+{R['disc']})^(-A{r}/12)", fmt='0.0000')
        put(cf, f"O{r}", f"=L{r}*N{r}", fmt=NUM)
        put(cf, f"P{r}", f"=IF(AND(A{r}>=1,M{r}>=0),1,0)", fmt=NUM)
    put(cf, f"A{last + 2}", "Net cash excludes the defaults column, which counts accounts (a share of one unit), not money.", "note")
    cf.freeze_panes = "B7"

    rs = sheet(wb, "Results", "Unit economics results", "Expected values per unit sold", "D2", 6, landscape=False)
    widths(rs, [52, 16, 4, 14, 14, 14])
    rng = lambda col: f"Unit_Cash_Flow!${col}${first}:${col}${last}"
    res = [("Expected instalments collected", f"=SUM({rng('D')})", NUM),
           ("Expected loss rate (missed ÷ scheduled instalments)", f"=IF({R['sched']}=0,0,1-SUM({rng('D')})/{R['sched']})", PCT),
           ("Expected recoveries", f"=SUM({rng('G')})", NUM),
           ("Lifetime contribution (undiscounted)", f"=SUM({rng('L')})", NUM),
           ("Customer acquisition cost", f"={R['cac']}", NUM),
           ("LTV (contribution before CAC)", "=B10+B11", NUM),
           ("LTV to CAC", '=IF(B11=0,"",B12/B11)', RAT),
           ("Cash payback (month)", f'=IFERROR(INDEX({rng("A")},MATCH(1,{rng("P")},0)),"Not within 60 months")', NUM),
           ("Peak funding per unit", f"=-MIN({rng('M')})", NUM),
           ("Unit IRR (annualised)", f'=IFERROR((1+IRR({rng("L")}))^12-1,"")', PCT),
           ("Unit NPV at the discount rate", f"=SUM({rng('O')})", NUM)]
    for i, (lab, f, fmt) in enumerate(res):
        r = 7 + i
        put(rs, f"A{r}", lab, "label", bold=lab.startswith(("Lifetime", "LTV to", "Cash payback", "Unit IRR")))
        put(rs, f"B{r}", f, fmt=fmt, bold=lab.startswith(("Lifetime", "LTV to", "Cash payback", "Unit IRR")))
    # closed form check
    cfm = (lambda hm, cm: f"({R['dep']}+{R['instal']}*MIN(1,{R['c']}*{cm})*(1-{R['h']}*{hm})*(1-(1-{R['h']}*{hm})^{R['T']})/({R['h']}*{hm})"
           f"+(1-(1-{R['h']}*{hm})^{R['T']})*{R['repo']}*{R['resale']}*{R['price']}*(1-{R['rcost']}))*(1-{R['fee']})"
           f"-{R['serv']}*(1-{R['h']}*{hm})*(1-(1-{R['h']}*{hm})^{R['T']})/({R['h']}*{hm})+{R['rbf']}-{R['outlay']}")
    put(rs, "A19", "Closed form contribution (check)", "label")
    put(rs, "B19", "=" + cfm(1, 1), fmt=NUM)
    put(rs, "A20", "Difference with the monthly table", "label")
    put(rs, "B20", f'=IF(ABS(B19-B10)<1,"OK: "&TEXT(B19-B10,"0.00"),"Check inputs: "&TEXT(B19-B10,"0"))', bold=True)
    section(rs, 22, "Sensitivity: lifetime contribution", 6)
    put(rs, "A23", "Default hazard multiplier (rows) × collection rate multiplier (columns)", "note")
    cms = [0.90, 0.95, 1.00]
    hms = [0.50, 0.75, 1.00, 1.30, 1.50, 2.00]
    put(rs, "A24", "Hazard multiplier", "head")
    put(rs, "B24", "Monthly hazard", "head")
    for j, cm in enumerate(cms):
        put(rs, f"{L(4 + j)}24", cm, "in", RAT)
    for i, hm in enumerate(hms):
        r = 25 + i
        put(rs, f"A{r}", hm, "in", RAT)
        put(rs, f"B{r}", f"={R['h']}*A{r}", fmt=PCT2)
        for j in range(3):
            c = L(4 + j)
            put(rs, f"{c}{r}", "=" + cfm(f"$A{r}", f"{c}$24"), fmt=NUM)
    rs.conditional_formatting.add("D25:F30", FormulaRule(formula=["D25<0"], fill=RED))
    put(rs, "A32", "The grid uses the closed form, which matches the monthly table when the check above reads OK. "
        "The Downside scenario of the model raises the hazard by 1.30 times and lowers collections by 0.97 times.", "note")

    # tier dashboard (closed form, default tiers)
    td = sheet(wb, "Tier_Dashboard", "LTV to CAC by tier", "Closed form on the model's default tiers; replace the blue inputs with company data", "D2", 8)
    widths(td, [50, 14, 14, 14, 14, 14, 4, 30])
    table_header(td, 6, ["Item", "Tier 1", "Tier 2", "Tier 3", "Tier 4", "Tier 5", "", "Notes"])
    fx, duty = G["fx0"], G["duty"]
    tin = [("Cash price", [p["price"] for p in P], NUM, ""), ("Deposit", [p["deposit"] for p in P], NUM, ""),
           ("Daily rate", [p["daily"] for p in P], NUM1, ""), ("Tenor (months)", [p["tenor"] for p in P], NUM, ""),
           ("Monthly default hazard", [p["hazard"] for p in P], PCT2, ""), ("Collection rate on paying accounts", [p["coll"] for p in P], PCT, ""),
           ("Repossession rate", [p["repo"] for p in P], PCT, ""), ("Resale value, share of cash price", [p["recov"] for p in P], PCT, ""),
           ("Landed hardware cost", [round(p["hw"] * fx * (1 + duty)) for p in P], NUM, f"USD cost × {fx:.0f} × (1 + {duty:.0%} freight and duty)"),
           ("Installation", [p["install"] for p in P], NUM, ""), ("Warranty, share of landed cost", [p["warranty"] for p in P], PCT, ""),
           ("CAC (commission and marketing)", [p["comm"] + p["mkt"] for p in P], NUM, ""),
           ("RBF per unit (local currency)", [round(p["rbf"] * fx) for p in P], NUM, "USD per unit × opening FX rate")]
    TR = {}
    for i, (lab, vals, fmt, note) in enumerate(tin):
        r = 7 + i
        TR[lab] = r
        put(td, f"A{r}", lab, "label")
        for j, v in enumerate(vals):
            put(td, f"{L(2 + j)}{r}", v, "in", fmt)
        if note:
            put(td, f"H{r}", note, "note")
    put(td, "A20", "Servicing cost per active account per month", "label")
    put(td, "B20", G["cs_cost"], "in", NUM)
    put(td, "A21", "Payment fees, share of cash collected", "label")
    put(td, "B21", G["mm_fee"], "in", PCT)
    put(td, "A22", "Recovery cost, share of resale proceeds", "label")
    put(td, "B22", G["recov_cost"], "in", PCT)
    section(td, 24, "Results", 8)
    g = lambda lab, c: f"{c}${TR[lab]}"
    tout = [
        ("Monthly instalment", lambda c: f"={g('Daily rate', c)}*365/12", NUM),
        ("Sum of survival over the tenor", lambda c: f"=(1-{g('Monthly default hazard', c)})*(1-(1-{g('Monthly default hazard', c)})^{g('Tenor (months)', c)})/{g('Monthly default hazard', c)}", NUM1),
        ("Expected instalments collected", lambda c: f"={c}25*{g('Collection rate on paying accounts', c)}*{c}26", NUM),
        ("Expected loss rate", lambda c: f"=1-{c}27/({c}25*{g('Tenor (months)', c)})", PCT),
        ("Expected recoveries", lambda c: f"=(1-(1-{g('Monthly default hazard', c)})^{g('Tenor (months)', c)})*{g('Repossession rate', c)}*{g('Resale value, share of cash price', c)}*{g('Cash price', c)}*(1-$B$22)", NUM),
        ("Servicing and payment fees", lambda c: f"=-$B$20*{c}26-$B$21*({g('Deposit', c)}+{c}27+{c}29)", NUM),
        ("Day one outlay excluding CAC", lambda c: f"={g('Landed hardware cost', c)}*(1+{g('Warranty, share of landed cost', c)})+{g('Installation', c)}", NUM),
        ("LTV (contribution before CAC)", lambda c: f"={g('Deposit', c)}+{c}27+{c}29+{c}30+{g('RBF per unit (local currency)', c)}-{c}31", NUM),
        ("CAC", lambda c: f"={g('CAC (commission and marketing)', c)}", NUM),
        ("Lifetime contribution", lambda c: f"={c}32-{c}33", NUM),
        ("LTV to CAC", lambda c: f'=IF({c}33=0,"",{c}32/{c}33)', RAT),
        ("Months of collections to recover the day one outlay", lambda c:
            f'=IFERROR(IF(1-({c}31+{c}33-{g("Deposit", c)})*{g("Monthly default hazard", c)}/({c}25*{g("Collection rate on paying accounts", c)}*(1-{g("Monthly default hazard", c)}))<=0,"Beyond tenor",'
            f'IF(LN(1-({c}31+{c}33-{g("Deposit", c)})*{g("Monthly default hazard", c)}/({c}25*{g("Collection rate on paying accounts", c)}*(1-{g("Monthly default hazard", c)})))/LN(1-{g("Monthly default hazard", c)})>{g("Tenor (months)", c)},"Beyond tenor",'
            f'LN(1-({c}31+{c}33-{g("Deposit", c)})*{g("Monthly default hazard", c)}/({c}25*{g("Collection rate on paying accounts", c)}*(1-{g("Monthly default hazard", c)})))/LN(1-{g("Monthly default hazard", c)}))),"")', NUM1),
    ]
    for i, (lab, fn, fmt) in enumerate(tout):
        r = 25 + i
        put(td, f"A{r}", lab, "label", bold=lab in ("LTV to CAC", "Lifetime contribution"))
        for j in range(5):
            c = L(2 + j)
            put(td, f"{c}{r}", fn(c), fmt=fmt, bold=lab in ("LTV to CAC", "Lifetime contribution"))
    put(td, "H36", "Before servicing and fees; deposit counted at sale", "note")
    put(td, "A38", "Closed form on the model's default inputs, without price indexation or timing of recoveries; the model's "
        "Unit_Economics sheet remains the reference. The default tiers have no company history behind them.", "note")
    ch = LineChart()
    from openpyxl.chart import BarChart
    ch = BarChart()
    ch.title, ch.height, ch.width = "LTV to CAC by tier", 7, 14
    ch.series.append(Series(Reference(td, min_col=2, max_col=6, min_row=35, max_row=35), title="LTV to CAC"))
    ch.set_categories(Reference(td, min_col=2, max_col=6, min_row=6, max_row=6))
    td.add_chart(ch, "B40")
    wb.move_sheet("Results", offset=-1)
    finish_book(wb, OUT / "AEF_V2_D2_Unit_Economics_Calculator.xlsx", "AEF Book 2 Decision Tool D2: Unit economics calculator", SUBJ)


# ================================================================== D3 Repayment curve and calibration
def build_d3(observed):
    wb = new_book()
    guide_sheet(wb, "D3", "Repayment curve and cohort calibration",
                "Draws the expected repayment curve of a cohort from a monthly default hazard and a collection rate, estimates "
                "the hazard and collection rate that fit observed cohort repayment most closely, and shows how portfolio growth flatters "
                "the portfolio collection rate.",
                ["On Curve, enter a hazard, a collection rate and a tenor to see survival and cumulative repayment by age.",
                 "On Calibration, enter observed cumulative repayment at M3, M6, M12 and M18 (deposits excluded), with weights.",
                 "Read the closest fitting hazard and collection rate, and compare the fitted curve with the observations.",
                 "Several pairs can fit one checkpoint; use at least three checkpoints and check the fit at each.",
                 "On Portfolio_Effect, see how the same customers produce a higher portfolio collection rate when sales grow."],
                [("Book", "Chapter 6 (survival model, calibration, denominator effect), Chapter 16 (SolaraPay recalibration)"),
                 ("Model", "Curves, Vintage_Input, Vintage_Dashboard; calibrated values go on Products"),
                 ("Default inputs", "Curve: default Tier 2 (2.6% hazard, 88% collection, 24 months). Calibration: average observed Tier 2 repayment of the synthetic SolaraPay cohorts.")])
    cv = sheet(wb, "Curve", "Repayment curve", "Expected behaviour of one cohort by account age", "D3", 7)
    widths(cv, [40, 14, 4, 12, 14, 16, 16])
    R, r = inputs_block(cv, 7, [("h", "Monthly default hazard", 0.026, PCT2, ""), ("c", "Collection rate on paying accounts", 0.88, PCT, ""),
                                ("T", "Tenor (months)", 24, NUM, "")])
    put(cv, "A11", "Expected share of scheduled instalments collected over the tenor", "label", bold=True)
    put(cv, "B11", f"={R['c']}*(1-{R['h']})*(1-(1-{R['h']})^{R['T']})/({R['h']}*{R['T']})", fmt=PCT, bold=True)
    put(cv, "A12", "Share of accounts still paying at the end of the tenor", "label")
    put(cv, "B12", f"=(1-{R['h']})^{R['T']}", fmt=PCT)
    table_header(cv, 6, ["", "", "", "Age (months)", "Share still paying", "Expected payment, share of instalment", "Cumulative repayment"], col0=1)
    for c in "ABC":
        cv[f"{c}6"].fill = PatternFill(fill_type=None)
        cv[f"{c}6"].value = None
    for a in range(1, 61):
        r = 6 + a
        put(cv, f"D{r}", a, fmt=NUM)
        put(cv, f"E{r}", f"=(1-{R['h']})^D{r}", fmt=PCT)
        put(cv, f"F{r}", f"=IF(D{r}<={R['T']},{R['c']}*E{r},0)", fmt=PCT)
        put(cv, f"G{r}", f"=SUM($F$7:F{r})/MIN(D{r},{R['T']})", fmt=PCT)
    ch = LineChart()
    ch.title, ch.height, ch.width = "Cumulative repayment by age", 7, 13
    ch.y_axis.number_format = "0%"
    ch.series.append(Series(Reference(cv, min_col=7, min_row=7, max_row=54), title="Cumulative repayment"))
    ch.series.append(Series(Reference(cv, min_col=5, min_row=7, max_row=54), title="Share still paying"))
    ch.set_categories(Reference(cv, min_col=4, min_row=7, max_row=54))
    cv.add_chart(ch, "I7")

    cb = sheet(wb, "Calibration", "Cohort calibration", "Grid search over hazard and collection rate", "D3", 24)
    widths(cb, [26, 12, 12, 12, 12, 12] + [9] * 18)
    put(cb, "A6", "Checkpoint (months)", "head")
    put(cb, "A7", "Observed cumulative repayment", "label", bold=True)
    put(cb, "A8", "Weight", "label")
    put(cb, "A9", "Fitted", "label")
    put(cb, "A10", "Fit error (points)", "label")
    put(cb, "A11", "Tenor of the product (months)", "label")
    put(cb, "B11", 24, "in", NUM)
    cps = [3, 6, 12, 18]
    for k, m in enumerate(cps):
        c = L(2 + k)
        put(cb, f"{c}6", m, "head")
        put(cb, f"{c}7", round(observed[k], 4) if observed[k] is not None else None, "in", PCT)
        put(cb, f"{c}8", 1, "in", NUM1)
    grid_top = 16
    hz = [round(0.005 + 0.0005 * i, 4) for i in range(151)]
    cs = [round(0.80 + 0.01 * j, 2) for j in range(18)]
    put(cb, f"A{grid_top - 1}", "Hazard (rows) × collection rate (columns): weighted sum of squared errors", "note")
    put(cb, f"A{grid_top}", "Hazard", "head")
    for j, cc in enumerate(cs):
        put(cb, f"{L(2 + j)}{grid_top}", cc, "head")
        cb[f"{L(2 + j)}{grid_top}"].number_format = PCT
    rowmin_col, colidx_col = L(2 + len(cs) + 1), L(2 + len(cs) + 2)
    put(cb, f"{rowmin_col}{grid_top}", "Row minimum", "head")
    put(cb, f"{colidx_col}{grid_top}", "Column of minimum", "head")

    def fit(hc, cc, k):
        m = cps[k]
        ref = f"{L(2 + k)}$6"
        return (f"IF(OR({L(2 + k)}$7=\"\",{L(2 + k)}$8=0),0,{L(2 + k)}$8*({L(2 + k)}$7-{cc}*(1-{hc})*(1-(1-{hc})^MIN({ref},$B$11))/({hc}*MIN({ref},$B$11)))^2)")
    for i, hv in enumerate(hz):
        r = grid_top + 1 + i
        put(cb, f"A{r}", hv, fmt=PCT2)
        for j in range(len(cs)):
            col = L(2 + j)
            put(cb, f"{col}{r}", "=" + "+".join(fit(f"$A{r}", f"{col}${grid_top}", k) for k in range(4)), fmt='0.000000')
        put(cb, f"{rowmin_col}{r}", f"=MIN(B{r}:{L(1 + len(cs))}{r})", fmt='0.000000')
        put(cb, f"{colidx_col}{r}", f"=MATCH({rowmin_col}{r},B{r}:{L(1 + len(cs))}{r},0)", fmt=NUM)
    g0, g1 = grid_top + 1, grid_top + len(hz)
    put(cb, "D11", "Closest fit", "label", bold=True)
    put(cb, "E11", "Hazard", "label")
    put(cb, "F11", f"=INDEX($A${g0}:$A${g1},MATCH(MIN(${rowmin_col}${g0}:${rowmin_col}${g1}),${rowmin_col}${g0}:${rowmin_col}${g1},0))", fmt=PCT2, bold=True)
    put(cb, "E12", "Collection rate on paying accounts", "label")
    put(cb, "F12", f"=INDEX($B${grid_top}:${L(1 + len(cs))}${grid_top},INDEX(${colidx_col}${g0}:${colidx_col}${g1},MATCH(MIN(${rowmin_col}${g0}:${rowmin_col}${g1}),${rowmin_col}${g0}:${rowmin_col}${g1},0)))", fmt=PCT, bold=True)
    put(cb, "E13", "Grid limits", "label")
    put(cb, "F13", f'=IF(OR(F11=$A${g0},F11=$A${g1},F12=$B${grid_top},F12=${L(1 + len(cs))}${grid_top}),"At the edge of the grid: widen it","Inside the grid")')
    for k in range(4):
        c = L(2 + k)
        put(cb, f"{c}9", f"=$F$12*(1-$F$11)*(1-(1-$F$11)^MIN({c}6,$B$11))/($F$11*MIN({c}6,$B$11))", fmt=PCT)
        put(cb, f"{c}10", f'=IF({c}7="","",{c}9-{c}7)', fmt=PCT)
    cb.freeze_panes = f"B{grid_top + 1}"

    pe = sheet(wb, "Portfolio_Effect", "Portfolio growth and the collection rate", "Same customer behaviour, different growth rates", "D3", 9)
    widths(pe, [40, 12, 12, 12, 12, 4, 12, 12, 12])
    put(pe, "A6", "Hazard, collection rate and tenor are linked to the Curve sheet.", "note")
    put(pe, "A8", "Monthly growth in originations", "label", bold=True)
    gr = [0.0, 0.02, 0.05, 0.10]
    for j, gv in enumerate(gr):
        put(pe, f"{L(2 + j)}8", gv, "in", PCT)
    put(pe, "A9", "Portfolio operational collection rate (not a PERFORM KPI), steady growth state", "label", bold=True)
    put(pe, "A10", "Difference with the cohort average (points)", "label")
    table_header(pe, 12, ["Age (months)", "Weight at g1", "Weight at g2", "Weight at g3", "Weight at g4", "", "Payment share", "", ""])
    for a in range(1, 61):
        r = 12 + a
        put(pe, f"A{r}", a, fmt=NUM)
        for j in range(4):
            c = L(2 + j)
            put(pe, f"{c}{r}", f"=IF(A{r}<=Curve!$B$9,(1+{c}$8)^(-A{r}),0)", fmt='0.0000')
        put(pe, f"G{r}", f"=Curve!F{6 + a}", fmt=PCT)
    for j in range(4):
        c = L(2 + j)
        put(pe, f"{c}9", f"=SUMPRODUCT({c}13:{c}72,$G$13:$G$72)/SUM({c}13:{c}72)", fmt=PCT, bold=True)
        put(pe, f"{c}10", f"={c}9-$B$9", fmt=PCT)
    put(pe, "H8", "Chapter 6 reports 64.4%, 68.3% and 71.6% for 0%, 5% and 10% monthly growth on the default Tier 2 curve.", "note")
    finish_book(wb, OUT / "AEF_V2_D3_Repayment_Curve_and_Calibration.xlsx", "AEF Book 2 Decision Tool D3: Repayment curve and cohort calibration", SUBJ)


# ================================================================== D4 Receivables financing calculator
def build_d4():
    wb = new_book()
    guide_sheet(wb, "D4", "Receivables financing calculator",
                "Shows how much cash a sales plan consumes before it returns any, how much of that the receivables facility can "
                "carry, and how much must be equity; and compares local currency and hard currency debt.",
                ["Enter units sold per year, the price plan, credit behaviour, the day one outlay per unit and the facility terms.",
                 "Read the monthly funding path on Funding and the year end summary on Summary.",
                 "Peak equity is the most the plan needs from shareholders to fund the book and central costs, before interest, tax and capex.",
                 "Use Currency to compare a local currency facility with dollar debt plus expected depreciation.",
                 "For a full plan with overheads, tax and interest, use MODEL 2, the PAYGo Company Financial and Investment Model."],
                [("Book", "Chapter 10 (working capital and FX), Chapter 11 (funding the book)"),
                 ("Model", "Financing, FS, Credit_Portfolio, Investment_Summary"),
                 ("Default inputs", "The default Tier 2 product of the model, sold at 10,000, 20,000, 35,000, 50,000 and 50,000 units a year (the volumes of Chapter 10).")])
    ws = sheet(wb, "Inputs", "Financing inputs", "Single product; local currency", "D4", 5, landscape=False)
    widths(ws, [52, 14, 4, 46, 4])
    p2 = P[1]
    landed = round(p2["hw"] * G["fx0"] * (1 + G["duty"]))
    items = [("y1", "Units sold, Year 1", 10000, NUM, ""), ("y2", "Units sold, Year 2", 20000, NUM, ""), ("y3", "Units sold, Year 3", 35000, NUM, ""),
             ("y4", "Units sold, Year 4", 50000, NUM, ""), ("y5", "Units sold, Year 5", 50000, NUM, ""),
             ("price", "Cash price", p2["price"], NUM, ""), ("dep", "Deposit", p2["deposit"], NUM, ""), ("daily", "Daily rate", p2["daily"], NUM1, ""),
             ("T", "Tenor (months)", p2["tenor"], NUM, ""), ("h", "Monthly default hazard", p2["hazard"], PCT2, ""), ("c", "Collection rate on paying accounts", p2["coll"], PCT, ""),
             ("outlay", "Day one outlay per unit (hardware, installation, CAC)", landed + p2["install"] + p2["comm"] + p2["mkt"], NUM, f"Landed {landed:,} + installation {p2['install']:,} + CAC {p2['comm'] + p2['mkt']:,}"),
             ("adv", "Advance rate on eligible receivables", p2["adv"], PCT, ""), ("elig", "Eligible share of paying account balances", 1.0, PCT, "Paying accounts are current or 1 to 30 DPD in the model's proxy"),
             ("limit", "Facility limit", 3_000_000_000, NUM, ""),
             ("start", "Facility available from month", G["rf_start"], NUM, "Model default"),
             ("oh", "Central costs per month (staff, G&A)", G["staff"] + G["ga"], NUM, "Model default; set to zero to isolate the book"),
             ("fx", "Exchange rate (local per USD) for the USD summary", G["fx0"], NUM1, "")]
    section(ws, 6, "Inputs", 5)
    R, r = inputs_block(ws, 7, items)
    put(ws, f"A{r + 1}", "Monthly instalment", "label")
    put(ws, f"B{r + 1}", f"={R['daily']}*365/12", fmt=NUM)
    R["instal"] = f"'Inputs'!$B${r + 1}"
    put(ws, f"A{r + 2}", "PAYGo premium per unit", "label")
    put(ws, f"B{r + 2}", f"={R['dep']}+{R['instal']}*{R['T']}-{R['price']}", fmt=NUM)
    R["prem"] = f"'Inputs'!$B${r + 2}"

    pu = sheet(wb, "Per_Unit", "Per unit curve", "Expected collections and balance of one unit by age", "D4", 5, landscape=False)
    widths(pu, [10, 14, 16, 18, 4])
    table_header(pu, 6, ["Age", "Share still paying", "Collection", "Carrying amount of paying accounts", ""])
    for a in range(0, 61):
        r = 7 + a
        put(pu, f"A{r}", a, fmt=NUM)
        put(pu, f"B{r}", f"=(1-{R['h']})^A{r}", fmt=PCT)
        put(pu, f"C{r}", f"=IF(AND(A{r}>=1,A{r}<={R['T']}),{R['instal']}*{R['c']}*B{r},0)", fmt=NUM)
        put(pu, f"D{r}", f"=IF(A{r}<={R['T']},({R['T']}-A{r})*({R['instal']}-{R['prem']}/{R['T']})*B{r},0)", fmt=NUM)

    mx = sheet(wb, "Cohorts", "Cohort matrix", "Rows: calendar month; columns: month of sale. Upper block balances, lower block collections", "D4", 62)
    widths(mx, [10] + [9] * 61)
    put(mx, "A6", "Month of sale", "head")
    for s in range(1, 61):
        c = L(1 + s)
        put(mx, f"{c}6", s, "head")
        put(mx, f"{c}7", f"=CHOOSE(ROUNDUP({c}6/12,0),{R['y1']},{R['y2']},{R['y3']},{R['y4']},{R['y5']})/12", fmt=NUM)
    put(mx, "A7", "Units sold", "label", bold=True)
    put(mx, "A9", "Balances", "label", bold=True)
    put(mx, "A71", "Collections", "label", bold=True)
    for t in range(1, 61):
        rb, rc = 9 + t, 71 + t
        put(mx, f"A{rb}", t, fmt=NUM)
        put(mx, f"A{rc}", t, fmt=NUM)
        for s in range(1, 61):
            c = L(1 + s)
            if t >= s:
                put(mx, f"{c}{rb}", f"={c}$7*Per_Unit!$D${7 + t - s}", fmt=NUM)
                put(mx, f"{c}{rc}", f"={c}$7*Per_Unit!$C${7 + t - s}", fmt=NUM)

    fd = sheet(wb, "Funding", "Monthly funding path", "Funding of the receivables book and central costs, before tax, interest and capex", "D4", 12)
    hdr = ["Month", "Units sold", "Collections", "Deposits", "Day one outlay and central costs", "Net cash", "Cumulative net cash", "Receivables, carrying amount",
           "Facility capacity", "Funding need", "Facility drawn", "Equity required"]
    table_header(fd, 6, hdr)
    widths(fd, [8] + [15] * 11)
    for t in range(1, 61):
        r = 6 + t
        put(fd, f"A{r}", t, fmt=NUM)
        put(fd, f"B{r}", f"=INDEX(Cohorts!$B$7:$BI$7,A{r})", fmt=NUM)
        put(fd, f"C{r}", f"=SUM(Cohorts!B{71 + t}:BI{71 + t})", fmt=NUM)
        put(fd, f"D{r}", f"=B{r}*{R['dep']}", fmt=NUM)
        put(fd, f"E{r}", f"=-B{r}*{R['outlay']}-{R['oh']}", fmt=NUM)
        put(fd, f"F{r}", f"=C{r}+D{r}+E{r}", fmt=NUM)
        put(fd, f"G{r}", f"=F{r}" if t == 1 else f"=G{r - 1}+F{r}", fmt=NUM)
        put(fd, f"H{r}", f"=SUM(Cohorts!B{9 + t}:BI{9 + t})", fmt=NUM)
        put(fd, f"I{r}", f"=IF(A{r}<{R['start']},0,MIN({R['limit']},H{r}*{R['elig']}*{R['adv']}))", fmt=NUM)
        put(fd, f"J{r}", f"=MAX(0,-G{r})", fmt=NUM)
        put(fd, f"K{r}", f"=MIN(I{r},J{r})", fmt=NUM)
        put(fd, f"L{r}", f"=J{r}-K{r}", fmt=NUM)
    fd.freeze_panes = "B7"

    sm = sheet(wb, "Summary", "Funding summary", "Year end and peak values", "D4", 7, landscape=False)
    widths(sm, [44, 14, 14, 14, 14, 14, 4])
    table_header(sm, 6, ["Item", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5", ""])
    lines = [("Receivables, carrying amount (year end)", "H"), ("Facility drawn (year end)", "K"), ("Equity required (year end)", "L"),
             ("Cumulative net cash (year end)", "G")]
    for i, (lab, col) in enumerate(lines):
        r = 7 + i
        put(sm, f"A{r}", lab, "label")
        for y in range(5):
            put(sm, f"{L(2 + y)}{r}", f"=Funding!{col}{6 + 12 * (y + 1)}", fmt=NUM)
    put(sm, "A12", "Peak funding need", "label", bold=True)
    put(sm, "B12", "=MAX(Funding!J7:J66)", fmt=NUM, bold=True)
    put(sm, "A13", "Peak facility drawn", "label", bold=True)
    put(sm, "B13", "=MAX(Funding!K7:K66)", fmt=NUM, bold=True)
    put(sm, "A14", "Peak equity required", "label", bold=True)
    put(sm, "B14", "=MAX(Funding!L7:L66)", fmt=NUM, bold=True)
    put(sm, "A15", "Month of peak equity", "label")
    put(sm, "B15", "=MATCH(B14,Funding!L7:L66,0)", fmt=NUM)
    put(sm, "A16", "Peak equity required (USD)", "label", bold=True)
    put(sm, "B16", f"=B14/{R['fx']}", fmt=NUM, bold=True)
    put(sm, "A18", "Receivables are the carrying amount of paying accounts (remaining instalments less unearned premium), the basis on "
        "which the facility advances. Defaulted balances are written off as they fall due. Interest, tax and capex are excluded.", "note")
    ch = LineChart()
    ch.title, ch.height, ch.width = "Funding of the receivables book", 7.5, 15
    for col, name in (("H", "Receivables, carrying amount"), ("K", "Facility drawn"), ("L", "Equity required")):
        ci = ord(col) - 64
        ch.series.append(Series(Reference(fd, min_col=ci, min_row=7, max_row=66), title=name))
    ch.set_categories(Reference(fd, min_col=1, min_row=7, max_row=66))
    sm.add_chart(ch, "A20")

    cu = sheet(wb, "Currency", "Local currency or dollar debt", "Cost of dollar debt in local currency terms", "D4", 6, landscape=False)
    widths(cu, [52, 14, 14, 14, 4, 4])
    put(cu, "A6", "Dollar interest rate", "label")
    put(cu, "B6", G["tl_rate"], "in", PCT)
    put(cu, "A7", "Local currency facility rate", "label")
    put(cu, "B7", G["rf_rate"], "in", PCT)
    put(cu, "A8", "Break even depreciation of the local currency, per year", "label", bold=True)
    put(cu, "B8", "=(1+B7)/(1+B6)-1", fmt=PCT2, bold=True)
    table_header(cu, 10, ["Depreciation scenario", "Depreciation per year", "Dollar debt cost in local currency", "Cheaper"])
    for i, (lab, d_) in enumerate([("Base", D.SCENARIOS["dep"][0]), ("Downside", D.SCENARIOS["dep"][1]), ("Severe", D.SCENARIOS["dep"][2])]):
        r = 11 + i
        put(cu, f"A{r}", lab, "label")
        put(cu, f"B{r}", d_, "in", PCT)
        put(cu, f"C{r}", f"=(1+$B$6)*(1+B{r})-1", fmt=PCT)
        put(cu, f"D{r}", f'=IF(C{r}>$B$7,"Local currency","Dollar debt")')
    put(cu, "A15", "The comparison covers the expected cost only. Dollar debt also carries the volatility of the exchange rate and "
        "the translation loss on the principal, which local currency debt does not (Chapter 10).", "note")
    finish_book(wb, OUT / "AEF_V2_D4_Receivables_Financing_Calculator.xlsx", "AEF Book 2 Decision Tool D4: Receivables financing calculator", SUBJ)


# ================================================================== D5 Investment screening scorecard
SC = [
    # label, unit, input type, green rule, amber rule, weight, kill text, default value, note
    ("Months of reconciled portfolio history", "months", "num", (">=", 24), (">=", 12), 2, None, 24, "Below 12 months the credit case rests on proxies"),
    ("DPD buckets reconcile to the ledger", "Yes/No", "yn", None, None, 2, "No", "Yes", "Kill criterion if No"),
    ("Cohort repayment at M12, gap to plan", "points", "num", ("<=", 0.03), ("<=", 0.06), 2, None, 0.051, "SolaraPay Tier 2: 5.1 points"),
    ("PAR30", "%", "num", ("<=", 0.10), ("<=", 0.20), 2, None, 0.17, "SolaraPay latest month: 17.0%"),
    ("Operational collection rate, trailing 12 months (not a PERFORM KPI)", "%", "num", (">=", 0.80), (">=", 0.70), 2, None, 0.693, "SolaraPay latest 12 months: 69.3% (Credit_Portfolio)"),
    ("Positive unit contribution in every tier", "Yes/No", "yn", None, None, 2, "No", "Yes", "Kill criterion if No"),
    ("LTV to CAC, mix weighted", "x", "num", (">=", 3.0), (">=", 2.0), 1, None, 4.8, "SolaraPay calibrated, weighted by planned mix"),
    ("Highest payment burden across tiers", "%", "num", ("<=", 0.10), ("<=", 0.13), 1, None, 0.132, "SolaraPay Tier 3: 13.2%"),
    ("Share of hard currency debt funding the receivables book", "%", "num", ("<=", 0.25), ("<=", 0.50), 1, None, None, "Not assessed in the case"),
    ("Funding runway without new money", "months", "num", (">=", 12), (">=", 6), 1, None, None, "Not assessed in the case"),
    ("Lending licence position confirmed by counsel", "Yes/No", "yn", None, None, 1, "No", None, "Kill criterion if No"),
    ("Credit function independent of sales", "Yes/No", "yn", None, None, 1, None, None, ""),
]


def build_d5():
    wb = new_book()
    guide_sheet(wb, "D5", "Investment screening scorecard",
                "A one page screen that decides whether a PAYGo company deserves full diligence. Twelve criteria are scored "
                "green, amber or red against editable thresholds; three are kill criteria; blank criteria are reported as not "
                "assessed, never scored.",
                ["Enter the value of each criterion from the company's data; leave it blank if it has not been assessed.",
                 "Review the thresholds and weights; they are the author's suggested defaults and should reflect the fund's policy.",
                 "Read the verdict: Decline if a kill criterion fails; Incomplete if criteria are missing; otherwise the weighted score decides.",
                 "Carry every red and amber criterion into the diligence checklist (Template T02) as a priority item."],
                [("Book", "Chapter 15 (red flags, diligence), Chapter 16 (SolaraPay)"),
                 ("Related", "Template T02 checklist; Template T01 memo"),
                 ("Default inputs", "SolaraPay case values where the case reports them; other criteria left blank.")])
    ws = sheet(wb, "Scorecard", "Investment screening scorecard", "Green 2 points, amber 1, red 0; weights apply", "D5", 11)
    widths(ws, [44, 9, 12, 8, 10, 8, 10, 8, 12, 10, 34])
    table_header(ws, 6, ["Criterion", "Unit", "Value", "Green if", "", "Amber if", "", "Weight", "Rating", "Points", "Note"])
    for i, (lab, unit, typ, grn, amb, w, kill, v, note) in enumerate(SC):
        r = 7 + i
        put(ws, f"A{r}", lab, "label", wrap=True)
        put(ws, f"B{r}", unit, "label")
        fmt = PCT if unit in ("%", "points") else (RAT if unit == "x" else (NUM if unit == "months" else None))
        put(ws, f"C{r}", v, "in", fmt)
        if typ == "num":
            put(ws, f"D{r}", grn[0], "in")
            put(ws, f"E{r}", grn[1], "in", fmt)
            put(ws, f"F{r}", amb[0], "in")
            put(ws, f"G{r}", amb[1], "in", fmt)
            cond = lambda op, ref: f'IF(D{r}=">=",C{r}>={ref},C{r}<={ref})' if ref == f"E{r}" else f'IF(F{r}=">=",C{r}>={ref},C{r}<={ref})'
            rating = f'=IF(C{r}="","Not assessed",IF({cond(">=", f"E{r}")},"Green",IF({cond(">=", f"G{r}")},"Amber","Red")))'
            dv_list(ws, [">=", "<="], f"D{r}")
            dv_list(ws, [">=", "<="], f"F{r}")
        else:
            put(ws, f"D{r}", "Yes", "label")
            rating = f'=IF(C{r}="","Not assessed",IF(C{r}="Yes","Green","{"Red: kill" if kill else "Red"}"))'
            dv_list(ws, ["Yes", "No"], f"C{r}")
        put(ws, f"H{r}", w, "in", NUM)
        put(ws, f"I{r}", rating, bold=True)
        put(ws, f"J{r}", f'=IF(I{r}="Green",2,IF(I{r}="Amber",1,0))*H{r}', fmt=NUM)
        put(ws, f"K{r}", note, "note")
        ws.row_dimensions[r].height = 26
    last = 6 + len(SC)
    flag_colours(ws, f"I7:I{last}", "Red", "Amber", "Green")
    r = last + 2
    put(ws, f"A{r}", "Criteria assessed", "label", bold=True)
    put(ws, f"C{r}", f'=COUNTA(A7:A{last})-COUNTIF(I7:I{last},"Not assessed")', fmt=NUM, bold=True)
    put(ws, f"D{r}", f'="of "&COUNTA(A7:A{last})', "label")
    put(ws, f"A{r + 1}", "Weighted score on assessed criteria", "label", bold=True)
    put(ws, f"C{r + 1}", f'=IFERROR(SUM(J7:J{last})/(2*SUMIF(I7:I{last},"<>Not assessed",H7:H{last})),0)', fmt=PCT, bold=True)
    put(ws, f"A{r + 2}", "Kill criteria failed", "label", bold=True)
    put(ws, f"C{r + 2}", f'=COUNTIF(I7:I{last},"Red: kill")', fmt=NUM, bold=True)
    put(ws, f"A{r + 3}", "Score needed to proceed (input)", "label")
    put(ws, f"C{r + 3}", 0.60, "in", PCT)
    put(ws, f"A{r + 4}", "Score needed to proceed without conditions (input)", "label")
    put(ws, f"C{r + 4}", 0.80, "in", PCT)
    put(ws, f"A{r + 6}", "Verdict", "label", bold=True)
    v = (f'=IF(C{r + 2}>0,"Decline: a kill criterion fails",IF(C{r}<COUNTA(A7:A{last}),"Incomplete: "&(COUNTA(A7:A{last})-C{r})&" criteria not assessed; provisional score "&TEXT(C{r + 1},"0%"),'
         f'IF(C{r + 1}>=C{r + 4},"Proceed to diligence",IF(C{r + 1}>=C{r + 3},"Proceed with conditions","Decline"))))')
    put(ws, f"C{r + 6}", v, bold=True)
    ws.merge_cells(f"C{r + 6}:K{r + 6}")
    flag_colours(ws, f"C{r + 6}", "Decline", "Incomplete", "Proceed")
    put(ws, f"A{r + 8}", "Thresholds and weights are suggested defaults reflecting the author's judgement, not sector standards. "
        "Set them to the fund's credit policy before use.", "note")
    finish_book(wb, OUT / "AEF_V2_D5_Investment_Screening_Scorecard.xlsx", "AEF Book 2 Decision Tool D5: Investment screening scorecard", SUBJ)


# ================================================================== D6 Investor returns calculator
def build_d6():
    wb = new_book()
    guide_sheet(wb, "D6", "Investor returns calculator",
                "Translates an entry price and an exit case into the investor's stake, multiple and IRR in US dollars, including "
                "the effect of currency depreciation and of follow on equity; solves for the highest pre money valuation that "
                "meets a target IRR; and tabulates the IRR across exit multiples and depreciation rates.",
                ["Enter the ticket, the pre money valuation, the exit year and the exit case in local currency.",
                 "Enter expected depreciation and any follow on equity calls the investor must fund pro rata.",
                 "Read stake, exit proceeds, multiple and IRR; then the maximum pre money valuation for the target IRR.",
                 "Use the grid to see how much of the return depends on the exit multiple and on the currency."],
                [("Book", "Chapter 15 (valuation, IRR and multiple in USD), Chapter 16 (SolaraPay returns)"),
                 ("Model", "Valuation, Investment_Summary"),
                 ("Default inputs", "SolaraPay calibrated Base: USD 4.0m at USD 8.0m pre money, Year 5 EBITDA KVS 1,418m at 6x, net debt KVS 1,122m, KVS 135 per USD and 5% depreciation a year. The tool returns the case's 29.0% IRR and 3.6x multiple.")])
    ws = sheet(wb, "Returns", "Investor returns", "US dollar returns on a local currency exit", "D6", 9, landscape=False)
    widths(ws, [50, 16, 4, 14, 14, 14, 14, 14, 14])
    items = [("ticket", "Investment (USD)", 4_000_000, NUM, ""), ("pre", "Pre money valuation (USD)", 8_000_000, NUM, ""),
             ("n", "Exit year (1 to 7)", 5, NUM, ""), ("ebitda", "Exit year EBITDA (local currency)", 1_418_022_460, NUM, ""),
             ("mult", "Exit multiple of EBITDA", 6.0, RAT, ""), ("nd", "Net debt at exit (local currency)", 1_121_735_060, NUM, ""),
             ("fx0", "Exchange rate at entry (local per USD)", 135.0, NUM1, ""), ("dep", "Depreciation of the local currency, per year", 0.05, PCT, ""),
             ("target", "Target IRR", 0.25, PCT, "")]
    section(ws, 6, "Inputs", 9)
    R, r = inputs_block(ws, 7, items, col_note="D")
    section(ws, 17, "Follow on equity calls, total company (USD)", 9)
    for y in range(1, 8):
        put(ws, f"{L(1 + y + 1)}18", f"Year {y}", "head")
        put(ws, f"{L(1 + y + 1)}19", 0, "in", NUM)
    put(ws, "A19", "Equity called from shareholders", "label")
    section(ws, 21, "Results", 9)
    res = [("Stake", f"={R['ticket']}/({R['pre']}+{R['ticket']})", PCT),
           ("Exit enterprise value (local currency)", f"={R['ebitda']}*{R['mult']}", NUM),
           ("Exit equity value, 100% (local currency)", f"=MAX(0,B23-{R['nd']})", NUM),
           ("Exchange rate at exit", f"={R['fx0']}*(1+{R['dep']})^{R['n']}", NUM1),
           ("Exit equity value, 100% (USD)", "=B24/B25", NUM),
           ("Investor proceeds at exit (USD)", "=B22*B26", NUM),
           ("Investor follow on equity (USD)", "=B22*SUMPRODUCT((C18:I18<>\"\")*1,C19:I19)", NUM),
           ("Multiple on invested capital", '=IF((B7+B28)=0,"",B27/(' + R["ticket"] + "+B28))", RAT),
           ("IRR (USD)", '=IFERROR(IRR(C33:J33),"Total loss")', PCT),
           ("Maximum pre money valuation for the target IRR, no follow on (USD)",
            f'=IF(B26<=0,"No value at exit",MAX(0,B26*{R["ticket"]}/({R["ticket"]}*(1+{R["target"]})^{R["n"]})-{R["ticket"]}))', NUM)]
    for i, (lab, f, fmt) in enumerate(res):
        r = 22 + i
        put(ws, f"A{r}", lab, "label", bold=lab in ("Multiple on invested capital", "IRR (USD)"))
        put(ws, f"B{r}", f, fmt=fmt, bold=lab in ("Multiple on invested capital", "IRR (USD)"))
    put(ws, "B28", "=B22*(" + "+".join(f"IF({y}<={R['n']},{L(1 + y + 1)}19,0)" for y in range(1, 8)) + ")", fmt=NUM)
    # cash flow row for IRR (years 0 to 7)
    put(ws, "A33", "Investor cash flows (USD), year 0 to year 7", "label")
    for y in range(0, 8):
        c = L(3 + y)
        put(ws, f"{c}32", y, "head")
        if y == 0:
            put(ws, f"{c}33", f"=-{R['ticket']}", fmt=NUM)
        else:
            put(ws, f"{c}33", f"=IF({y}<={R['n']},-$B$22*{L(1 + y + 1)}19,0)+IF({y}={R['n']},$B$27,0)", fmt=NUM)
    put(ws, "A34", "Follow on calls after the exit year are ignored.", "note")
    section(ws, 36, "IRR by exit multiple and depreciation (no follow on equity)", 9)
    deps = [0.0, 0.05, 0.12, 0.20, 0.25]
    mults = [4.0, 5.0, 6.0, 7.0, 8.0]
    put(ws, "A37", "Exit multiple (rows) × depreciation per year (columns)", "head")
    for j, d_ in enumerate(deps):
        put(ws, f"{L(2 + j)}37", d_, "in", PCT)
    for i, m in enumerate(mults):
        r = 38 + i
        put(ws, f"A{r}", m, "in", RAT)
        for j in range(len(deps)):
            c = L(2 + j)
            eq = f"MAX(0,{R['ebitda']}*$A{r}-{R['nd']})/({R['fx0']}*(1+{c}$37)^{R['n']})"
            put(ws, f"{c}{r}", f'=IF({eq}<=0,"Total loss",($B$22*{eq}/{R["ticket"]})^(1/{R["n"]})-1)', fmt=PCT)
    ws.conditional_formatting.add("B38:F42", FormulaRule(formula=['OR(B38="Total loss",B38<0)'], fill=RED))
    ws.conditional_formatting.add("B38:F42", FormulaRule(formula=[f'AND(ISNUMBER(B38),B38>={R["target"]})'], fill=GRN))
    put(ws, "A44", "Exit equity is floored at zero: when net debt exceeds the exit enterprise value, the stake is worth nothing.", "note")
    finish_book(wb, OUT / "AEF_V2_D6_Investor_Returns_Calculator.xlsx", "AEF Book 2 Decision Tool D6: Investor returns calculator", SUBJ)


if __name__ == "__main__":
    import json
    OUT.mkdir(parents=True, exist_ok=True)
    ex = json.load(open(ROOT / "volumes/02-solar-home-systems/case-study/case_exhibits.json"))
    obs = ex["vintage"]["1"]["observed"][:4]
    build_d1()
    build_d2()
    build_d3(obs)
    build_d4()
    build_d5()
    build_d6()
