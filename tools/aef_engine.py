"""Africa Energy Finance (AEF) - shared model-building engine.

Every volume's Excel model is generated from Python with openpyxl so that:
  * the same conventions (colours, timeline, checks, number formats) apply everywhere;
  * shared mechanics (timeline, FX, debt, statements, checks) are written once;
  * every model is reproducible and diff-able in git.

Conventions (see docs/modelling-conventions.md):
  * Monthly time sheets: col A label, col B unit, col C total/notes, col D opening
    balance, cols E.. = month 1..N.
  * Blue font  = hard-coded input; yellow fill = key assumption to review.
  * Black font = formula; green font = link from another sheet.
"""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = "0000FF"
GREEN = "008000"
BLACK = "000000"
NAVY = "1F3864"
GREY = "F2F2F2"
YELLOW = "FFFF00"

FMT_NUM = '#,##0;(#,##0);"-"'
FMT_NUM2 = '#,##0.00;(#,##0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_X = '0.00"x";(0.00"x");"-"'
FMT_DATE = "mmm-yy"
FMT_INT = "0"

FIRST_COL = 5  # column E = month 1


def col(i):
    """Column letter of month i (1-based)."""
    return get_column_letter(FIRST_COL + i - 1)


def q(sheet):
    return f"'{sheet}'" if (" " in sheet or "-" in sheet) else sheet


class ModelBook:
    def __init__(self, months):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.months = months
        self.last = col(months)
        self.rows = {}  # (sheet, key) -> row

    # ---------- sheets ----------
    def sheet(self, name, title, subtitle="", tab=None):
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["A1"] = title
        ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")
        if tab:
            ws.sheet_properties.tabColor = tab
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 16
        return ws

    def time_header(self, ws, row=4):
        """Rows 4-6: month index, period end date, financial year (linked to Timeline)."""
        labels = ["Month #", "Period end", "Year #"]
        for k, lab in enumerate(labels):
            r = row + k
            ws.cell(r, 1, lab).font = Font(name=FONT, bold=True, color="FFFFFF")
            for c in range(1, FIRST_COL + self.months):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=NAVY)
            ws.cell(r, 4, "Opening" if k == 0 else None).font = Font(name=FONT, bold=True, color="FFFFFF")
            for i in range(1, self.months + 1):
                c = col(i)
                cell = ws[f"{c}{r}"]
                if ws.title == "Timeline":
                    continue
                cell.value = f"=Timeline!{c}{r}"
                cell.font = Font(name=FONT, bold=True, color="FFFFFF")
                cell.number_format = FMT_DATE if k == 1 else FMT_INT
                cell.alignment = Alignment(horizontal="right")
        for i in range(1, self.months + 1):
            ws.column_dimensions[col(i)].width = 12
        ws.column_dimensions["D"].width = 12
        ws.freeze_panes = ws[f"{col(1)}7"]

    def section(self, ws, row, text):
        ws.cell(row, 1, text).font = Font(name=FONT, bold=True, color=NAVY, size=11)
        for c in range(1, FIRST_COL + self.months):
            ws.cell(row, c).border = Border(bottom=Side(style="thin", color=NAVY))

    # ---------- monthly rows ----------
    def register(self, sheet, key, row):
        if (sheet, key) in self.rows:
            raise KeyError(f"duplicate row key {sheet}.{key}")
        self.rows[(sheet, key)] = row

    def r(self, sheet, key):
        return self.rows[(sheet, key)]

    def ref(self, sheet, key, c, absolute_row=True, this_sheet=None):
        row = self.r(sheet, key)
        prefix = "" if sheet == this_sheet else f"{q(sheet)}!"
        return f"{prefix}{c}${row}" if absolute_row else f"{prefix}{c}{row}"

    def range_(self, sheet, key, this_sheet=None, absolute=True):
        row = self.r(sheet, key)
        prefix = "" if sheet == this_sheet else f"{q(sheet)}!"
        if absolute:
            return f"{prefix}${col(1)}${row}:${self.last}${row}"
        return f"{prefix}{col(1)}{row}:{self.last}{row}"

    def write_row(self, ws, row, label, unit, fn, fmt=FMT_NUM, bold=False,
                  total=None, opening=None, comment=None, link=False):
        """fn(i, c, p) -> formula string for month i (column c, previous column p)."""
        ws.cell(row, 1, label).font = Font(name=FONT, bold=bold)
        ws.cell(row, 2, unit).font = Font(name=FONT, size=9, color="595959")
        if comment:
            ws.cell(row, 1).comment = Comment(comment, "AEF")
        for i in range(1, self.months + 1):
            c, p = col(i), col(i - 1) if i > 1 else "D"
            cell = ws[f"{c}{row}"]
            cell.value = fn(i, c, p)
            cell.number_format = fmt
            cell.font = Font(name=FONT, bold=bold, color=GREEN if link else BLACK)
        if total == "sum":
            ws.cell(row, 3, f"=SUM({col(1)}{row}:{self.last}{row})")
        elif total == "last":
            ws.cell(row, 3, f"={self.last}{row}")
        elif total == "max":
            ws.cell(row, 3, f"=MAX({col(1)}{row}:{self.last}{row})")
        elif total == "min":
            ws.cell(row, 3, f"=MIN({col(1)}{row}:{self.last}{row})")
        if total:
            ws.cell(row, 3).number_format = fmt
            ws.cell(row, 3).font = Font(name=FONT, bold=True)
        if opening is not None:
            ws.cell(row, 4, opening).number_format = fmt
            ws.cell(row, 4).font = Font(name=FONT)
        if bold:
            for cc in range(1, FIRST_COL + self.months):
                ws.cell(row, cc).border = Border(top=Side(style="thin", color="808080"))


# ---------- static-cell helpers ----------

def label(ws, cell, text, bold=False, size=10, color=BLACK, italic=False):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def put_input(ws, cell, value, fmt=FMT_NUM, key=False, note=None):
    ws[cell] = value
    ws[cell].font = Font(name=FONT, color=BLUE)
    ws[cell].number_format = fmt
    if key:
        ws[cell].fill = PatternFill("solid", fgColor=YELLOW)
    if note:
        ws[cell].comment = Comment(note, "AEF")


def put_calc(ws, cell, formula, fmt=FMT_NUM, bold=False, link=False):
    ws[cell] = formula
    ws[cell].font = Font(name=FONT, bold=bold, color=GREEN if link else BLACK)
    ws[cell].number_format = fmt


def header_row(ws, row, values, start_col=1):
    for k, v in enumerate(values):
        c = ws.cell(row, start_col + k, v)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def set_font_all(ws):
    """Ensure Arial on every populated cell that has no explicit font family."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, bold=f.bold, italic=f.italic, size=f.size, color=f.color)
