"""Compare a workbook recalculated by LibreOffice with the formula-engine value store.

Run: python tools/qa_libreoffice_compare.py <libreoffice_out.xlsx> <values.pkl> <TAG>
The original (formula) workbook is expected beside the out/ folder. Reports error values and differences."""
import openpyxl, pickle, sys, collections, math
lo, pkl, tag = sys.argv[1:4]
d = pickle.load(open(pkl, 'rb'))
wbf = openpyxl.load_workbook(lo.replace('/out/', '/'), read_only=False)   # formulas (original)
wbv = openpyxl.load_workbook(lo, data_only=True)
wcell=None; n = nerr = nmiss = ndiff = 0; errs = collections.Counter(); diffs = []; worst = 0
for ws in wbf.worksheets:
    wv = wbv[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and c.value.startswith('=')):
                continue
            n += 1
            v = wv[c.coordinate].value
            if isinstance(v, str) and (v.startswith('#') or v.startswith('Err')):
                nerr += 1; errs[(ws.title, v)] += 1; continue
            ref = d.get(f"'[{tag}]{ws.title.upper()}'!{c.coordinate}")
            if ref is None or ref == 'empty':
                nmiss += 1; continue
            if isinstance(ref, (int, float)) and not isinstance(ref, bool):
                if v is None: v = 0
                import datetime as _dt
                if isinstance(v, _dt.datetime): v = (v - _dt.datetime(1899, 12, 30)).total_seconds() / 86400
                if isinstance(v, bool): v = float(v)
                if isinstance(v, (int, float)):
                    tol = 1e-6 * max(1, abs(ref))
                    if abs(v - ref) > tol:
                        ndiff += 1; diffs.append((ws.title, c.coordinate, ref, v))
                    r_ = abs(v - ref) / max(1, abs(ref))
                    if r_ > worst: worst, wcell = r_, (ws.title, c.coordinate, ref, v)
                else:
                    ndiff += 1; diffs.append((ws.title, c.coordinate, ref, v))
            else:
                if str(ref) != str(v if v is not None else '') and not (ref in ('', None) and v in ('', None)):
                    ndiff += 1; diffs.append((ws.title, c.coordinate, ref, v))
print(f"{tag}: formulas {n}, LibreOffice errors {nerr}, not in reference {nmiss}, differences {ndiff}, worst relative diff {worst:.2e}")
print('worst cell', wcell if worst else None)
print('errors by sheet:', dict(errs.most_common(10)))
by = collections.Counter(x[0] for x in diffs); print('differences by sheet:', dict(by.most_common(10)))
for x in diffs[:12]: print('  ', x)
