"""Build the user manuals (PDF) from Markdown sources in manuals/src.

Usage: python tools/build_manuals.py [--check-only]

For each manual:
- converts Markdown to HTML (tables, attribute lists), adds a cover, a table of contents with
  page numbers (CSS target-counter) and running headers, then renders a PDF with WeasyPrint;
- applies French typography to French manuals (narrow no-break spaces before ; : ! ? % and inside
  guillemets, and as thousands separators);
- runs editorial checks: no em or en dashes, no stock AI phrases, and every label quoted in the
  manual must exist as a cell text or sheet name in the matching workbook.
"""
import html
import re
import sys
from pathlib import Path

import markdown
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "manuals" / "src"
CSS = ROOT / "manuals" / "manual.css"

MANUALS = [
    # (source, output pdf, workbook, lang, title, subtitle)
    ("p1_fr.md", "product/01-entry-calculator/Manuel_Calculateur_Faisabilite_MiniReseau_FR.pdf",
     "product/01-entry-calculator/Calculateur_Faisabilite_MiniReseau_v1_FR.xlsx", "fr",
     "Calculateur de faisabilité financière de mini-réseau", "Manuel d'utilisation"),
    ("p2_fr.md", "product/02-developer-edition/Manuel_Modele_Financier_Developpeur_FR.pdf",
     "product/02-developer-edition/Modele_Financier_Acces_Energie_Developpeur_v1_FR.xlsx", "fr",
     "Modèle financier de projet d'accès à l'énergie", "Édition Développeur. Manuel d'utilisation"),
    ("p3_fr.md", "product/03-fund-manager/Manuel_Modele_Gestionnaire_Fonds_FR.pdf",
     "product/03-fund-manager/Modele_Gestionnaire_Fonds_Acces_Energie_v1_FR.xlsx", "fr",
     "Modèle du gestionnaire de fonds d'accès à l'énergie", "Édition RBF & Portefeuille. Manuel d'utilisation"),
    ("p1_en.md", "product/01-entry-calculator/Manual_MiniGrid_Feasibility_Calculator_EN.pdf",
     "product/01-entry-calculator/MiniGrid_Feasibility_Calculator_v1.xlsx", "en",
     "Mini-Grid Financial Feasibility Calculator", "User manual"),
    ("p2_en.md", "product/02-developer-edition/Manual_Developer_Edition_EN.pdf",
     "product/02-developer-edition/EnergyAccess_Developer_Model_v1.xlsx", "en",
     "Energy Access Project Financial Model", "Developer Edition. User manual"),
    ("p3_en.md", "product/03-fund-manager/Manual_Fund_Manager_Edition_EN.pdf",
     "product/03-fund-manager/EnergyAccess_Fund_Manager_Model_v1.xlsx", "en",
     "Energy Access Fund Manager Model", "RBF & Portfolio Edition. User manual"),
]

VERSION = {"fr": "Version 1.0, octobre 2026", "en": "Version 1.0, October 2026"}
AUTHOR_LINE = {"fr": "Emmanuel Boujieka Kamga", "en": "Emmanuel Boujieka Kamga"}
TOC_TITLE = {"fr": "Sommaire", "en": "Contents"}
LICENCE = {
    "fr": "Licence : un utilisateur ou une organisation. Reproduction et revente interdites.",
    "en": "Licence: single user or single organisation. No reproduction or resale.",
}

# Phrases that read as machine-written in technical documentation (checked case-insensitively).
BANNED = [
    "—", "–", " -- ", "delve", "tapestry", "testament", "pivotal", "seamless", "vibrant", "showcase", "underscore",
    "landscape", "in today's", "let's", "it's not just", "not only", "in conclusion", "unlock", "game-changer",
    "plongeons", "n'est pas seulement", "il est important de noter", "il convient de noter", "en conclusion",
    "véritable", "incontournable", "révolutionn", "au cœur de", "crucial", "essentiel de", "fluide", "robuste",
    "en somme", "dans un monde", "explorons", "découvrons", "permet de mettre en lumière",
]


def french_typography(text):
    nb = " "
    text = re.sub(r"(\d) (?=\d{3}\b)", lambda m: m.group(1) + nb, text)
    text = re.sub(r"(\d) (?=\d{3}\b)", lambda m: m.group(1) + nb, text)
    text = re.sub(r" ([;:!?%])", nb + r"\1", text)
    text = text.replace("« ", "«" + nb).replace(" »", nb + "»")
    return text


def quoted_labels(md_text, lang):
    if lang == "fr":
        return re.findall(r"«\s*([^»]+?)\s*»", md_text)
    return re.findall(r'"([^"\n]+)"', md_text)


def workbook_texts(path):
    """Cell texts, sheet names and string literals inside formulas (statuses that no example row shows)."""
    texts = set()
    for data_only in (True, False):
        wb = load_workbook(ROOT / path, data_only=data_only)
        texts.update(wb.sheetnames)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if not isinstance(v, str):
                        continue
                    if v.startswith("="):
                        texts.update(t.replace('""', '"') for t in re.findall(r'"((?:[^"]|"")*)"', v))
                    else:
                        texts.add(v.strip())
    return texts


def check(md_text, src, workbook, lang):
    problems = []
    low = md_text.lower()
    for b in BANNED:
        if b.lower() in low:
            problems.append(f"banned phrase {b!r}")
    texts = workbook_texts(workbook)
    for lab in quoted_labels(md_text, lang):
        lab = lab.strip()
        if lab.startswith("=") or lab in texts:
            continue
        # allow quoting the start of a long guide sentence or a verdict family such as "NON - ..."
        if any(t.startswith(lab) for t in texts):
            continue
        problems.append(f"label not found in workbook: {lab!r}")
    return problems


def build_html(md_text, lang, title, subtitle):
    body = markdown.markdown(md_text, extensions=["tables", "attr_list", "toc", "sane_lists"],
                             extension_configs={"toc": {"permalink": False}})
    # table of contents from h2/h3 with ids
    toc_items = re.findall(r'<h([23]) id="([^"]+)">(.*?)</h\1>', body)
    toc = "".join(
        f'<li class="toc-h{lvl}"><a href="#{hid}">{txt}</a></li>' for lvl, hid, txt in toc_items
    )
    if lang == "fr":
        body = french_typography(body)
        toc = french_typography(toc)
        title_t, sub_t = french_typography(title), french_typography(subtitle)
    else:
        title_t, sub_t = title, subtitle
    return f"""<!doctype html>
<html lang="{lang}"><head><meta charset="utf-8"><title>{html.escape(title)}</title></head>
<body>
<section class="cover">
  <div class="cover-band"><p class="cover-kicker">Energy Access Finance Toolkit</p></div>
  <div class="cover-body">
    <h1 class="cover-title">{title_t}</h1>
    <p class="cover-sub">{sub_t}</p>
    <p class="cover-author">{AUTHOR_LINE[lang]}</p>
    <p class="cover-version">{VERSION[lang]}</p>
    <p class="cover-licence">{LICENCE[lang]}</p>
  </div>
</section>
<section class="toc"><h2 class="toc-title">{TOC_TITLE[lang]}</h2><ul>{toc}</ul></section>
<div class="running-title">{title_t}</div>
<main>{body}</main>
</body></html>"""


def main():
    check_only = "--check-only" in sys.argv
    from weasyprint import CSS as WCSS, HTML
    failed = False
    for src, out, workbook, lang, title, subtitle in MANUALS:
        path = SRC / src
        if not path.exists():
            print(f"skip {src} (missing)")
            continue
        md_text = path.read_text(encoding="utf-8")
        problems = check(md_text, src, workbook, lang)
        if problems:
            failed = True
            print(f"{src}: {len(problems)} problem(s)")
            for p in problems:
                print("   ", p)
        else:
            print(f"{src}: checks OK")
        if check_only:
            continue
        doc = build_html(md_text, lang, title, subtitle)
        HTML(string=doc, base_url=str(ROOT)).write_pdf(ROOT / out, stylesheets=[WCSS(filename=str(CSS))])
        print("   wrote", out)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
