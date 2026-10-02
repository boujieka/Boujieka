"""Extract SolaraPay case exhibits (numbers + charts) from the evaluated case workbook and the verified twin.

Run (after qa_evaluate.py on the case workbook):
  python tools/case_exhibits.py <evaluated_case.xlsx>
Writes volumes/02-solar-home-systems/case-study/case_exhibits.json and figures/*.png
"""
import copy
import json
import pickle
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402

sys.path.insert(0, str(Path(__file__).parent))
import shs_defaults as D  # noqa: E402
from cases import solarapay as SP  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUTD = ROOT / "volumes/02-solar-home-systems/case-study"
GREEN, GOLD, GREY, LGREEN = "#0B3020", "#B07C0F", "#8C8C8C", "#4F7F5F"
PAREN = matplotlib.ticker.FuncFormatter(lambda v, _: f"({abs(v):g})" if v < 0 else f"{v:g}")
PAREN_PCT = matplotlib.ticker.FuncFormatter(lambda v, _: f"({abs(v):.0%})" if v < 0 else f"{v:.0%}")
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.spines.top": False, "axes.spines.right": False})

path = sys.argv[1]
vals = pickle.load(open(path + ".pkl", "rb"))
book = Path(path).name.upper()
wb = openpyxl.load_workbook(path)
get = lambda s, c: vals.get(f"'[{book}]{s.upper()}'!{c}")


def _norm(t):
    return t.replace("\u2014", ", ").replace("\u2013", " to ").replace(" - ", ": ").strip()


def row_of(sheet, text, startswith=False):
    ws = wb[sheet]
    text = _norm(text)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and (_norm(v) == text or (startswith and _norm(v).startswith(text))):
            return r
    raise KeyError(text)


G0, P0 = copy.deepcopy(D.GENERAL), copy.deepcopy(D.PRODUCTS)
Gm, Pm = SP.management_inputs(G0, P0)
SP.apply(D.GENERAL, D.PRODUCTS)
from shadow_shs import run  # noqa: E402

X = {"company": SP.CASE["company"], "currency": "KVS", "fx0": D.GENERAL["fx0"]}
cal, man = run(1), run(1, gen=Gm, products=Pm)
X["traj"] = {k: [float(x) for x in cal[k]] for k in ("rev", "ebitda", "ni", "cash", "rf", "cr", "active", "dscr", "fcff")}
X["traj_mgmt"] = {k: [float(x) for x in man[k]] for k in ("rev", "ebitda", "ni", "cr")}
# KPIs from the evaluated workbook (calibrated base)
kp = {}
for key, label in [("units", "Units sold - total"), ("rev_usd", "Revenue (USD, average FX)"), ("gm", "Gross margin"),
                   ("ebitda_m", "EBITDA margin"), ("credit_cost", "Net credit losses / revenue"),
                   ("de", "Debt / book equity"), ("dpd30", "30+ DPD / gross receivables (year end, projection)")]:
    r = row_of("KPIs", label)
    kp[key] = [get("KPIs", f"{c}{r}") for c in "EFGHI"]
X["kpis"] = kp
X["readiness"] = get("Investment_Readiness", "B5")
X["gates_met"] = [wb["Investment_Readiness"].cell(r, 2).value for r in range(7, 30) if get("Investment_Readiness", f"F{r}") == 1]
# unit economics by tier
ue = {}
for key, label in [("price", "Cash price"), ("apr", "Implied annual financing rate (nominal APR)"),
                   ("loss", "Expected loss rate (missed / scheduled instalments)"), ("contrib", "Lifetime contribution per unit"),
                   ("ltv_cac", "LTV / CAC (contribution before CAC / CAC)"), ("payback", "Cash payback (months after sale)"),
                   ("irr", "Unit IRR (annualised, unlevered)")]:
    r = row_of("Unit_Economics", label)
    ue[key] = [get("Unit_Economics", f"{c}{r}") for c in "CDEFG"]
X["unit_econ"] = ue
# consumer risk
cr = {}
for key, label in [("burden", "Payment burden (instalment / income)"), ("apr", "Implied consumer APR")]:
    r = row_of("Consumer_Risk", label)
    cr[key] = [get("Consumer_Risk", f"{c}{r}") for c in "CDEFG"]
X["consumer"] = cr
# history (synthetic) - tiers 1-3
hist_c, hist_v = SP.history(D.PRODUCTS, D.GENERAL)
H = {}
for j, rows in hist_c.items():
    H[j] = dict(cr=[r["coll"] / r["due"] if r["due"] else None for r in rows],
                par30=[(r["b2"] + r["b3"] + r["b4"] + r["b5"]) / r["gross"] if r["gross"] else None for r in rows],
                gross=[r["gross"] for r in rows])
tot_coll = [sum(hist_c[j][i]["coll"] for j in hist_c) for i in range(SP.HIST_MONTHS)]
tot_due = [sum(hist_c[j][i]["due"] for j in hist_c) for i in range(SP.HIST_MONTHS)]
X["hist_cr_last12"] = sum(tot_coll[12:]) / sum(tot_due[12:])
X["hist_cr_first12"] = sum(tot_coll[1:12]) / sum(tot_due[1:12])


def rr_curve(p, months=(3, 6, 12, 18, 24)):
    T, inst, h, c = p["tenor"], p["daily"] * 365 / 12, p["hazard"], p["coll"]
    out = []
    for m in months:
        ages = range(1, min(m, T) + 1)
        out.append(sum(inst * c * (1 - h) ** a for a in ages) / (inst * len(ages)))
    return out


CP = [3, 6, 12, 18, 24]
X["vintage"] = {}
for j in hist_v:
    obs = []
    for k, m in enumerate(CP):
        xs = [v["cp"][k]["coll"] / v["cp"][k]["due"] for v in hist_v[j].values() if k in v["cp"]]
        obs.append(float(np.mean(xs)) if xs else None)
    X["vintage"][j] = dict(tier=D.PRODUCTS[j]["tier"], observed=obs, management=rr_curve(Pm[j]), calibrated=rr_curve(D.PRODUCTS[j]),
                           hazard_mgmt=Pm[j]["hazard"], hazard_cal=D.PRODUCTS[j]["hazard"], coll_mgmt=Pm[j]["coll"], coll_cal=D.PRODUCTS[j]["coll"])
X["own_obs"] = {D.PRODUCTS[j]["tier"]: [v["own"] for v in hist_v[j].values() if v["own"] is not None] for j in hist_v}
X["cases"] = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {}
json.dump(X, open(OUTD / "case_exhibits.json", "w"), indent=1, default=float)

# ---------------- charts
F = OUTD / "figures"
yrs = ["Y1", "Y2", "Y3", "Y4", "Y5"]
fx_avg = [D.GENERAL["fx0"] * 1.05 ** (y - 0.5) for y in range(1, 6)]  # display only (labels in KVS bn)
fig, ax = plt.subplots(figsize=(6.4, 3.0))
x = np.arange(5)
ax.bar(x - 0.2, np.array(X["traj"]["rev"]) / 1e9, 0.4, color=GREEN, label="Revenue")
ax.bar(x + 0.2, np.array(X["traj"]["ebitda"]) / 1e9, 0.4, color=GOLD, label="EBITDA")
ax.plot(x, np.array(X["traj"]["ni"]) / 1e9, color=GREY, marker="o", label="Net income")
ax.axhline(0, color="black", lw=0.6)
ax.set_xticks(x, yrs); ax.set_ylabel("KVS bn"); ax.yaxis.set_major_formatter(PAREN); ax.legend(frameon=False, ncol=3, loc="upper left")
ax.set_title("SolaraPay, calibrated Base case", loc="left", fontsize=10, color=GREEN)
fig.tight_layout(); fig.savefig(F / "fig1_trajectory.png", dpi=200); plt.close(fig)

fig, ax = plt.subplots(figsize=(6.4, 3.0))
m = np.arange(1, SP.HIST_MONTHS + 1)
for j, col_ in zip(H, [GREEN, GOLD, LGREEN]):
    ax.plot(m[1:], H[j]["cr"][1:], color=col_, label=f"Tier {D.PRODUCTS[j]['tier']} collection rate")
ax.axhline(0.70, color="#C00000", ls="--", lw=0.8, label="Covenant minimum 70%")
ax.set_ylim(0.5, 1.0); ax.set_xlabel("History month"); ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
ax.legend(frameon=False, fontsize=7.5, loc="lower left"); ax.set_title("Observed monthly collection rate (synthetic history)", loc="left", fontsize=10, color=GREEN)
fig.tight_layout(); fig.savefig(F / "fig2_history_collections.png", dpi=200); plt.close(fig)

fig, axs = plt.subplots(1, 3, figsize=(7.2, 2.7), sharey=True)
for ax, (j, v) in zip(axs, X["vintage"].items()):
    ax.plot(CP, v["management"], color=GREY, ls="--", marker="o", ms=3, label="Management plan")
    ax.plot(CP, [o if o is not None else np.nan for o in v["observed"]], color=GOLD, marker="s", ms=3, label="Observed cohorts")
    ax.plot(CP, v["calibrated"], color=GREEN, marker="o", ms=3, label="Calibrated")
    ax.set_title(f"Tier {v['tier']}", fontsize=9, color=GREEN); ax.set_xticks(CP)
    ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0)); ax.set_xlabel("Months since sale")
axs[0].set_ylabel("Cumulative repayment rate"); axs[0].legend(frameon=False, fontsize=7)
fig.suptitle("Vintage repayment curves: plan, observed cohorts and calibration", x=0.02, ha="left", fontsize=10, color=GREEN)
fig.tight_layout(); fig.savefig(F / "fig3_vintage_calibration.png", dpi=200); plt.close(fig)

if X["cases"]:
    names = list(X["cases"].keys())
    irr = [X["cases"][n]["irr"] if X["cases"][n]["irr"] is not None else -1.0 for n in names]
    fig, ax = plt.subplots(figsize=(6.4, 3.4))
    colors = [GREEN if v > 0.2 else (GOLD if v > 0 else "#C00000") for v in irr]
    ax.barh(range(len(names)), irr, color=colors)
    for i, (n, v) in enumerate(zip(names, irr)):
        ax.text(max(v, 0) + 0.01, i, "total loss" if X["cases"][n]["irr"] is None else f"{v:.0%}", va="center", fontsize=7.5)
    ax.set_yticks(range(len(names)), [_norm(n).replace(": ", ", ") for n in names], fontsize=7.5); ax.invert_yaxis()
    ax.xaxis.set_major_formatter(PAREN_PCT); ax.axvline(0, color="black", lw=0.6)
    ax.set_xlim(-1.05, 0.5)
    ax.set_title("Investor IRR (USD) by case", loc="left", fontsize=10, color=GREEN)
    fig.tight_layout(); fig.savefig(F / "fig4_irr_cases.png", dpi=200); plt.close(fig)

fig, ax = plt.subplots(figsize=(6.4, 2.8))
ax.plot(x, np.array(X["traj"]["cash"]) / 1e9, color=GOLD, marker="o", label="Closing cash")
ax.bar(x, np.array(X["traj"]["rf"]) / 1e9, 0.5, color=GREEN, label="Receivables facility drawn")
ax.set_xticks(x, yrs); ax.set_ylabel("KVS bn"); ax.legend(frameon=False, loc="upper left")
ax.set_title("Funding: the facility finances the receivables book", loc="left", fontsize=10, color=GREEN)
fig.tight_layout(); fig.savefig(F / "fig5_funding.png", dpi=200); plt.close(fig)
print("exhibits written")
