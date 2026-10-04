"""Assemble Book 7: substitute model numbers, build the generated tables, list sources, lint the style.

python tools/book7/prepare7.py  ->  book7/build/book7_resolved.md
"""
import csv, json, os, re, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
CH = ["ch00_front"] + [f"ch{i:02d}" for i in range(1, 19)] + ["ch99_annexes"]
text = "\n\n".join(open(f"book7/src/{c}.md", encoding="utf8").read().strip() for c in CH)

MAP = json.load(open("model/model_map.json"))
SNAP = json.load(open("model/snapshot_results.json"))
wb = load_workbook("model/Bankable_Hydro_Model.xlsx", data_only=True)
RE_REF = r"'([^']+)'!\$([A-Z]+)\$(\d+)"


def cell(ref):
    m = re.fullmatch(RE_REF, ref)
    return wb[m.group(1)][m.group(2) + m.group(3)].value


M = {k: cell(v) for k, v in MAP["REF"].items() if re.fullmatch(RE_REF, str(v))}

# ---------------- derived values ----------------
X = {}
bs = SNAP["base_sized"]
X["debt_total"] = (bs["debt_c"] or 0) + bs["debt_m"]
for i in range(2, 6):
    X[f"req{i}"] = 112 * (1 + SNAP[f"s{i}"]["req_tariff_flex"])
X["starts_per_close"] = 1 / M["dev_pfc"]
ws = wb["01A_DEVELOPMENT"]
rows = {str(ws.cell(r, 1).value): r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
r_alive = next(r for k, r in rows.items() if k.startswith("Probability the project is still alive"))
r_s3 = rows["Spend: stage 3"]
cols = range(6, 16)
alive = [ws.cell(r_alive, c).value or 0 for c in cols]
first_fs = next(j for j, c in enumerate(cols) if (ws.cell(r_s3, c).value or 0) > 0)
X["palive_fs"] = alive[first_fs]
X["palive_last"] = [a for a in alive if a][-1]
r4 = next(r for k, r in rows.items() if k.startswith("4. ") and isinstance(ws.cell(r, 3).value, float))
X["c4"], X["d4"] = ws.cell(r4, 3).value, ws.cell(r4, 4).value
X["cod_share"] = X["d4"] / X["c4"]
opyr = bs["ts"]["opyr"]
X["xterm_cod"] = bs["ts"]["x_term"][opyr.index(1)]
s_roy, r_roy = MAP["TSROW"]["o_roy"]
s_yr, r_yr = MAP["TSROW"]["opyr"]
k2 = [wb[s_yr].cell(r_yr, 5 + j).value for j in range(40)].index(2)
X["roy_y2"] = wb[s_roy].cell(r_roy, 5 + k2).value
X["be_flow_pct"] = -SNAP["base_locked"]["be_flow_flex"]

MINUS = "−"


def fmt(v, f):
    if v is None:
        return "n/a"
    if isinstance(v, str):
        return v
    neg = lambda s: s.replace("-", MINUS)
    if f == "pct0": return neg(f"{v*100:.0f}%")
    if f == "pct1": return neg(f"{v*100:.1f}%")
    if f == "x1": return neg(f"{v:.1f}x")
    if f == "x2": return neg(f"{v:.2f}x")
    if f == "n0":
        return (MINUS if v < -0.5 else "") + f"{abs(v):,.0f}"
    if f == "n1": return neg(f"{v:,.1f}")
    if f == "n2": return neg(f"{v:.2f}")
    raise ValueError(f)


def sub(m):
    key, f = m.group(1), m.group(2)
    ns, _, rest = key.partition(".")
    v = M[rest] if ns == "m" else X[rest] if ns == "x" else SNAP[ns][rest]
    return fmt(v, f)


text = re.sub(r"\{\{([\w.]+)\|(\w+)\}\}", sub, text)
assert "{{" not in text, re.findall(r"\{\{[^}]*\}\}", text)[:5]

# ---------------- generated tables ----------------
ws = wb["30A_CLOSE_READINESS"]
g = ["Table: Table 17.1. Kasiri River Hydro: financial close readiness gates",
     "| # | Gate | Area | Critical | Status |", "|---|---|---|---|---|"]
r = 5
while isinstance(ws.cell(r, 1).value, int):
    crit = "Yes" if ws.cell(r, 4).value == "Y" else "No"
    test = "" if ws.cell(r, 5).value != "model test" else " (model test)"
    g.append(f"| {ws.cell(r, 1).value} | {ws.cell(r, 2).value} | {ws.cell(r, 3).value} | {crit} | {ws.cell(r, 6).value}{test} |")
    r += 1
g.append("Note: Statuses marked as model tests are computed by the model; the others are evidence statuses entered for the case.")
text = text.replace("%%GATES7", "\n".join(g))

ws = wb["31_CASE_STUDY"]
c = ["Table: Table J.1. Public benchmark cases (as published; nominal, unadjusted)",
     "| Project | Country | Structure | MW | USD/kW | Timeline | Credit support |", "|---|---|---|---|---|---|---|"]
r = MAP["CASE_ROW0"] + 1
while ws.cell(r, 1).value and "FICTIONAL" not in str(ws.cell(r, 1).value):
    v = [ws.cell(r, k).value for k in range(1, 13)]
    ukw = f"{v[6]:,.0f}" if isinstance(v[6], (int, float)) else "n/a"
    row = f"| {v[0]} | {v[1]} | {v[2]} | {v[3]:,.0f} | {ukw} | {v[7]} | {v[9]} |"
    c.append(row.replace("(snippet)", "(search summary)").replace(" - ", ": "))
    r += 1
c.append("Note: Sources for each row are in the research case files. Costs reported in euros are not converted (n/a). PUBLIC DATA NOT FOUND means the research found no public figure.")
text = text.replace("%%CASES7", "\n".join(c))

# ---------------- sources ----------------
SDV = {}
for line in open("research/source_database.md", encoding="utf8"):
    cl = [x.strip() for x in line.strip().strip("|").split("|")]
    if len(cl) >= 8 and re.fullmatch(r"[A-Z]{2}-\d{2}", cl[0]):
        SDV[cl[0]] = cl[-1]
SD = {row["id"]: row for row in csv.DictReader(open("research/source_database.csv", encoding="utf8"))}
CASEFILES = {"A1": "research/case_studies/africa_part1.md", "A2": "research/case_studies/africa_part2.md",
             "INT": "research/case_studies/international_benchmarks.md", "DE": "research/hydro_development_evidence.md"}
CS = {}
for k, p in CASEFILES.items():
    for line in open(p, encoding="utf8"):
        m = re.match(r"^\s*(?:-\s*)?(?:\*\*S(\d+[a-z]?)\*\*|\[S(\d+[a-z]?)\])\s*(.*)$", line)
        if m:
            CS[f"{k}:S{m.group(1) or m.group(2)}"] = m.group(3).strip()
VLAB = {"F": "Full text read", "P": "Landing or summary page read", "S": "Search summary only"}


def clean(t):
    t = t.replace(" - ", ": ").replace("—", ",").replace("–", " to ").replace("**", "").replace("|", "/")
    return re.sub(r"\s+", " ", t).strip()


def entry(cid):
    if cid == "DE:BENCH":
        return "Benchmark table for development model defaults, development evidence file (research note). Records where no public data were found.", "Research note"
    if cid == "CI":
        return "Competitive review of hydropower finance books, models and courses (research note).", "Research note"
    if cid in SD:
        s = SD[cid]
        return clean(f"{s['publisher']} ({s['year']}). {s['title']}. {s['doc_type']}. {s['url']}"), VLAB.get(SDV.get(cid, ""), "Not recorded")
    if cid in CS:
        t = CS[cid]
        low = t.lower()
        st = "Search summary only" if "snippet" in low or "search result" in low else "Page or document opened" if "fetched" in low or "opened" in low else "Opened; see research file"
        t = re.sub(r"\s*\((fetched|snippet[^)]*|opened[^)]*)\)", "", t)
        return clean(t), st
    raise KeyError(cid)


cited = []
for m in re.finditer(r"\[((?:[A-Z]{2,3}(?::S\d+[a-z]?|-\d{2}|:BENCH)|CI)(?:;\s*(?:[A-Z]{2,3}(?::S\d+[a-z]?|-\d{2}|:BENCH)|CI))*)\]", text):
    for i in m.group(1).split(";"):
        i = i.strip()
        if i not in cited:
            cited.append(i)
unknown = [i for i in cited if i not in SD and i not in CS and i not in ("DE:BENCH", "CI")]
assert not unknown, unknown
PREF = {"A1": "African cases, part 1", "A2": "African cases, part 2", "INT": "International benchmarks", "DE": "Development evidence"}


def key(cid):
    a, _, b = cid.replace("-", ":").partition(":")
    n = re.match(r"S?(\d+)([a-z]?)", b)
    return (a, int(n.group(1)) if n else 0, n.group(2) if n else b)


L = ["Identifiers with two letters and a number (such as HY-08) refer to the research source database; identifiers with a prefix and S-number (such as DE:S1) refer to the numbered source lists of the research files: A1 and A2, African case files; INT, international benchmarks; DE, development evidence. Verification status records what the research team saw. Claims resting on a search summary only should be checked against the original before reliance.",
     "", "| ID | Source | Verification |", "|---|---|---|"]
for cid in sorted(cited, key=key):
    e, st = entry(cid)
    L.append(f"| {cid} | {e} | {st} |")
text = text.replace("%%SOURCES7", "\n".join(L))

# ---------------- markdown clean-up for the renderer ----------------
out = []
for line in text.split("\n"):
    if line.startswith("Table: "):
        out += ["", f"**{line[7:].strip()}**", "{: .cap}", ""]
        continue
    if line.startswith("Note: ") and out and out[-1].startswith("|"):
        out += ["", f"*{line}*"]
        continue
    if line.startswith("|") and out and out[-1].strip() and not out[-1].startswith("|"):
        out.append("")
    out.append(line)
text = "\n".join(out)
text = re.sub(r"~([^~\s]+(?: [^~\s]+)*)~", r"<sub>\1</sub>", text)
text = re.sub(r"\^([^\^\s]+)\^", r"<sup>\1</sup>", text)
text = text.replace("](figures/", "](../src/figures/")

# ---------------- style lint ----------------
problems = []
for bad in ["—", "–", " -- ", "delve", "tapestry", "testament", "pivotal", "underscore", "showcase", "vibrant", "intricate",
            "landscape of", "in today's", "it is important to note", "crucial role", "seamless", "Additionally,", "Furthermore,",
            "Moreover,", "robust framework", "navigate the", "leverage", "unlock", "game-changer", "holistic"]:
    for line in text.splitlines():
        if line.startswith("| ") and re.match(r"\| [A-Z0-9:-]+ \| .* \| (Full|Landing|Search|Page|Opened|Not|Research)", line):
            continue
        if bad.lower() in line.lower():
            problems.append((bad, line[:100]))
for p in problems[:40]:
    print("LINT:", p)
os.makedirs("book7/build", exist_ok=True)
open("book7/build/book7_resolved.md", "w", encoding="utf8").write(text)
print("words:", len(re.findall(r"[A-Za-z]+", text)), "sources:", len(cited), "lint issues:", len(problems))
