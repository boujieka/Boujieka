"""Book 7 figures, drawn from MODEL 7 and the full-engine snapshot, in the house colours."""
import json, os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
PRINT = os.environ.get("PRINT") == "1"   # greyscale print figures at 450 ppi for the 7 x 10 in edition
FR = os.environ.get("BOOK_LANG") == "fr"  # French labels, decimal comma; written next to the French sources
OUT = ("book7/build_fr/figures_print" if FR else "book7/build/figures_print") if PRINT else ("book7/src_fr/figures" if FR else "book7/src/figures")
DPI = 450 if PRINT else 220
os.makedirs(OUT, exist_ok=True)
S = json.load(open("model/snapshot_results.json"))
wb = load_workbook("model/Bankable_Hydro_Model.xlsx", data_only=True)

GREEN, GOLD, SAGE, RUST, INK, INK2, GRID = "#0B3020", "#B07C0F", "#6E9A7E", "#A8432A", "#1d1d1d", "#55554f", "#e6e2d6"
NOEV = "#BDBDBD"
if PRINT:
    GREEN, GOLD, SAGE, RUST, INK, INK2, GRID, NOEV = "#1a1a1a", "#9a9a9a", "#c4c4c4", "#5c5c5c", "#000000", "#444444", "#e2e2e2", "#e9e9e9"


def on(fill):
    """Text colour that reads on a fill."""
    h = fill.lstrip("#"); lum = 0.3 * int(h[0:2], 16) + 0.59 * int(h[2:4], 16) + 0.11 * int(h[4:6], 16)
    return "white" if lum < 140 else "black"
plt.rcParams.update({
    "font.family": "Liberation Sans", "font.size": 8, "axes.edgecolor": INK2, "axes.labelcolor": INK2,
    "xtick.color": INK2, "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "axes.axisbelow": True,
    "axes.titlesize": 8.5, "axes.titleweight": "bold", "axes.titlecolor": INK, "legend.frameon": False,
})
W = 5.6
MINUS = "−"


def save(fig, name):
    if FR:  # decimal comma on numeric axes (category axes keep their labels)
        from matplotlib.ticker import FuncFormatter, ScalarFormatter
        for a in fig.axes:
            for axis in (a.xaxis, a.yaxis):
                if isinstance(axis.get_major_formatter(), ScalarFormatter):
                    axis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}".replace(".", ",").replace("-", MINUS)))
    fig.savefig(f"{OUT}/{name}", dpi=DPI, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    if PRINT:
        from PIL import Image
        Image.open(f"{OUT}/{name}").convert("L").save(f"{OUT}/{name}", dpi=(DPI, DPI))


FRL = {
    "Mean river flow (m³/s)": "Débit moyen de la rivière (m³/s)", "Expected energy (GWh per month)": "Énergie attendue (GWh par mois)",
    "Recon.": "Reconn.", "Pre-feas.": "Préfaisab.", "Feasibility": "Faisabilité", "Permitting": "Autorisations", "PPA": "CAE",
    "Financing": "Financement", "USD million, risk-weighted": "M USD, pondéré par le risque",
    "Value of the whole development position at the start of each stage": "Valeur de la position de développement au début de chaque étape",
    "Annual cash flow": "Flux annuel", "Cumulative (right scale)": "Cumul (échelle de droite)", "USD million a year": "M USD par an",
    "USD million, cumulative": "M USD, cumulé",
    "Developer's nominal cash flow on the success path, with a sale of half its stake at COD": "Flux nominal du développeur sur le scénario de succès, avec cession de la moitié de sa part à la COD",
    "Plant cost ±20%": "Coût de la centrale ±20 %", "Tariff ±10%": "Tarif ±10 %", "Generation ±10%": "Production ±10 %",
    "Interest rate +200 bp": "Taux d'intérêt +200 pb", "Operating cost +20%": "Coûts d'exploitation +20 %",
    "P90 one-year in cash flows": "P90 à un an dans les flux", "DEVELOPER": "DÉVELOPPEUR", "LENDERS": "PRÊTEURS", "STATE": "ÉTAT",
    "Three parties must each\naccept the answer": "Trois parties doivent\nchacune accepter la réponse",
    "23\nevidence\ngates": "23\nportes\nde preuve", "Q8\nReady to\nclose?\n\nSTOP\nNOT\nREADY\nCOND.\nGO\nGO": "Q8\nPrêt pour le\nbouclage ?\n\nSTOP\nNOT\nREADY\nCOND.\nGO\nGO",
    "Critical gate": "Porte critique",
    "Q1 Technically viable": "Q1 Techniquement viable", "Q2 Developable": "Q2 Développable", "Q3 Economic for the system": "Q3 Économique pour le système",
    "Q4 Investable for the developer": "Q4 Investissable pour le développeur", "Q5 Bankable": "Q5 Bancable",
    "Q6 Buildable on budget": "Q6 Constructible dans le budget", "Q7 Affordable for buyer and state": "Q7 Abordable pour l'acheteur et l'État",
    "Developer, lenders": "Développeur, prêteurs", "Developer, state": "Développeur, État", "State": "État", "Developer": "Développeur",
    "Lenders": "Prêteurs", "Developer, lenders, state": "Développeur, prêteurs, État",
    "Jan": "jan", "Feb": "fév", "Mar": "mar", "Apr": "avr", "May": "mai", "Jun": "juin", "Jul": "juil", "Aug": "aoû",
    "Sep": "sep", "Oct": "oct", "Nov": "nov", "Dec": "déc",
    "MET": "MET (franchie)", "PARTIAL": "PARTIAL (partielle)", "NOT MET": "NOT MET (non franchie)", "NO EVIDENCE": "NO EVIDENCE (sans preuve)",
}


def T(x):
    return FRL.get(x, x) if FR else x


def num(x):
    """Decimal comma and spaced percent sign in the French edition."""
    return x.replace(".", ",").replace("%", "\u202f%") if FR else x


def pct(v, d=1):
    return num(f"{v*100:.{d}f}%".replace("-", MINUS))


# ---- 2.1 hydrology ----
ws = wb["03_HYDROLOGY"]
hdr = next(r for r in range(1, 60) if ws.cell(r, 1).value == "Month")
rows = {ws.cell(r, 1).value: r for r in range(hdr, hdr + 10)}
months = [T(ws.cell(hdr, c).value) for c in range(5, 17)]
flow = [ws.cell(rows["Mean river flow"], c).value for c in range(5, 17)]
energy = [ws.cell(rows["Expected energy (x availability)"], c).value for c in range(5, 17)]
qd, qe = ws["C5"].value, ws["C9"].value
fig, ax = plt.subplots(1, 2, figsize=(W, 2.3))
ax[0].bar(months, flow, color=SAGE, width=0.62)
ax[0].axhline(qd + qe, color=GOLD, lw=1.4)
ax[0].text(11.4, qd + qe + 1.5, (f"Débit d'équipement plus débit réservé, {qd + qe:.0f}" if FR else f"Design flow plus e-flow, {qd + qe:.0f}"), ha="right", color=INK2, fontsize=6.8)
ax[0].axhline(qe, color=INK2, lw=1, ls=(0, (3, 2)))
ax[0].text(11.4, qe + 1.5, (f"Débit réservé {qe:.0f}" if FR else f"E-flow {qe:.0f}"), ha="right", color=INK2, fontsize=6.8)
ax[0].set_title(T("Mean river flow (m³/s)"), loc="left"); ax[0].tick_params(axis="x", labelsize=6.3)
ax[1].bar(months, energy, color=GREEN, width=0.62)
ax[1].set_title(T("Expected energy (GWh per month)"), loc="left"); ax[1].tick_params(axis="x", labelsize=6.3)
fig.tight_layout(w_pad=2)
save(fig, "fig2_1_hydrology.png")

# ---- 6.1 value of the development position by stage ----
ws = wb["01A_DEVELOPMENT"]
r0 = next(r for r in range(1, ws.max_row) if ws.cell(r, 1).value == "Stage about to start") + 1
labels = [T(x) for x in ["Recon.", "Pre-feas.", "Feasibility", "Permitting", "PPA", "Financing"]]
val = [ws.cell(r0 + k, 5).value for k in range(6)]
pfc = [ws.cell(r0 + k, 2).value for k in range(6)]
fig, ax = plt.subplots(figsize=(W, 2.4))
ax.bar(labels, val, color=[RUST if v < 0 else GREEN for v in val], width=0.58)
for j, v in enumerate(val):
    ax.text(j, v + (0.15 if v >= 0 else -0.15), num(f"{v:.2f}".replace("-", MINUS)), ha="center", va="bottom" if v >= 0 else "top", fontsize=7, color=INK)
ax.set_xticks(range(6)); ax.set_xticklabels([f"{l}\nP(close) {pct(p, 0)}" for l, p in zip(labels, pfc)], fontsize=7)
ax.axhline(0, color=INK, lw=0.8)
ax.set_ylim(min(val) - 1.0, max(val) + 1.0)
ax.set_ylabel(T("USD million, risk-weighted")); ax.grid(axis="x", visible=False)
ax.set_title(T("Value of the whole development position at the start of each stage"), loc="left")
save(fig, "fig6_1_position_value.png")

# ---- 14.1 developer cash flow on the success path ----
ry = next(r for r in range(60, ws.max_row) if ws.cell(r, 1).value == "Year")
yrs, cf = [], []
c = 6
while ws.cell(ry, c).value is not None:
    yrs.append(ws.cell(ry, c).value); cf.append(ws.cell(ry + 1, c).value or 0); c += 1
first = next(i for i, v in enumerate(cf) if abs(v) > 1e-9)
last = max(i for i, v in enumerate(cf) if abs(v) > 1e-9)
yrs, cf = yrs[first:last + 1], cf[first:last + 1]
cum, s = [], 0
for v in cf:
    s += v; cum.append(s)
fig, ax = plt.subplots(figsize=(W, 2.5))
ax.bar(yrs, cf, color=[RUST if v < 0 else GREEN for v in cf], width=0.7, label=T("Annual cash flow"))
ax2 = ax.twinx()
ax2.plot(yrs, cum, color=GOLD, lw=1.8, label=T("Cumulative (right scale)"))
ax2.spines["right"].set_visible(True); ax2.grid(False)
ax.axhline(0, color=INK, lw=0.8)
ax.set_ylabel(T("USD million a year")); ax2.set_ylabel(T("USD million, cumulative"))
h1, l1 = ax.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels()
ax.legend(h1 + h2, l1 + l2, loc="lower right", fontsize=7)
ax.set_title(T("Developer's nominal cash flow on the success path, with a sale of half its stake at COD"), loc="left")
save(fig, "fig14_1_developer_cf.png")

# ---- 16.1 tornado (equity IRR, financing locked) ----
base = S["base_locked"]["kpi_eirr"]
pairs = [("Plant cost ±20%", "capex_m20", "capex_p20"), ("Tariff ±10%", "tar_p10", "tar_m10"), ("Generation ±10%", "gen_p10", "gen_m10"),
         ("Interest rate +200 bp", None, "rate_p200"), ("Operating cost +20%", None, "opex_p20"), ("P90 one-year in cash flows", None, "p90cf")]
rows = []
for lab, up, dn in pairs:
    hi = S[up]["kpi_eirr"] if up else base
    rows.append((T(lab), S[dn]["kpi_eirr"], hi))
rows.sort(key=lambda r: abs(r[2] - r[1]))
fig, ax = plt.subplots(figsize=(W, 2.4))
for i, (lab, lo, hi) in enumerate(rows):
    ax.barh(i, (lo - base) * 100, left=base * 100, color=RUST, height=0.55)
    ax.barh(i, (hi - base) * 100, left=base * 100, color=GREEN, height=0.55)
    ax.text(min(lo, hi) * 100 - 0.3, i, pct(lo), va="center", ha="right", fontsize=6.5, color=INK2)
    if hi != base:
        ax.text(max(lo, hi) * 100 + 0.3, i, pct(hi), va="center", ha="left", fontsize=6.5, color=INK2)
ax.axvline(base * 100, color=GOLD, lw=1.4)
ax.set_yticks(range(len(rows))); ax.set_yticklabels([r[0] for r in rows])
ax.set_xlabel(f"TRI des fonds propres privés, en % (base {pct(base)})" if FR else f"Private equity IRR, percent (base {pct(base)})"); ax.grid(axis="y", visible=False)
lo_all = min(min(r[1], r[2]) for r in rows) * 100; hi_all = max(max(r[1], r[2]) for r in rows) * 100
ax.set_xlim(lo_all - 3, hi_all + 3)
save(fig, "fig16_1_tornado.png")
print("figures written")

# ---- 0.1 framework schematic ----
from matplotlib.patches import FancyBboxPatch
ws = wb["30A_CLOSE_READINESS"]
rf = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("HYDRO READINESS FRAMEWORK")) + 2
QS = [(ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value) for r in range(rf, rf + 7)]
fig, ax = plt.subplots(figsize=(W, 3.7)); ax.set_xlim(0, 100); ax.set_ylim(0, 100); ax.axis("off")


def box(x, y, w, h, txt, fc, tc=None, fs=7, bold=False, ec=None):
    tc = tc or on(fc)
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4,rounding_size=1.2", fc=fc, ec=ec or fc, lw=1))
    ax.text(x + w / 2, y + h / 2, txt, ha="center", va="center", color=tc, fontsize=fs, fontweight="bold" if bold else "normal", wrap=True)


for j, p in enumerate(["DEVELOPER", "LENDERS", "STATE"]):
    box(4 + j * 21, 88, 17, 8, T(p), GOLD, bold=True, fs=6.6 if FR else 7.5)
ax.text(70, 92, T("Three parties must each\naccept the answer"), fontsize=7, color=INK2, va="center")
for i, (q, who, n) in enumerate(QS):
    y = 76 - i * 10.5
    box(2, y, 37, 8, T(f"{q}"), GREEN, fs=6.6)
    ax.text(41, y + 4, f"{T(who)}  |  {n} porte{'s' if n != 1 else ''}" if FR else f"{who}  |  {n} gate{'s' if n != 1 else ''}", fontsize=6, color=INK2, va="center")
    ax.annotate("", xy=(72, 45), xytext=(67, y + 4), arrowprops=dict(arrowstyle="-", color="#c9c3b0", lw=0.8))
box(72, 37, 12, 16, T("23\nevidence\ngates"), SAGE, fs=7.2, bold=True)
ax.annotate("", xy=(87, 45), xytext=(84.5, 45), arrowprops=dict(arrowstyle="->", color=INK, lw=1))
box(87, 27, 12.5, 36, T("Q8\nReady to\nclose?\n\nSTOP\nNOT\nREADY\nCOND.\nGO\nGO"), INK, fs=6, bold=True)
save(fig, "fig0_1_framework.png")

# ---- 17.1 readiness map ----
COL = {"MET": GREEN, "PARTIAL": GOLD, "NOT MET": RUST, "NO EVIDENCE": NOEV}
gates = []
r = 5
while isinstance(ws.cell(r, 1).value, int):
    gates.append((ws.cell(r, 1).value, ws.cell(r, 4).value == "Y", ws.cell(r, 6).value, str(ws.cell(r, 8).value)))
    r += 1
qnames = [q for q, _, _ in QS]
fig, ax = plt.subplots(figsize=(W, 2.9))
for i, qn in enumerate(qnames):
    y = len(qnames) - 1 - i
    gs = [g for g in gates if g[3] == qn]
    for k, (num, crit, st, _) in enumerate(gs):
        ax.add_patch(plt.Rectangle((k * 1.15, y - 0.4), 1.0, 0.8, fc=COL.get(st, "#BDBDBD"), ec=INK if crit else "white", lw=1.6 if crit else 0.5))
        ax.text(k * 1.15 + 0.5, y, str(num), ha="center", va="center", fontsize=7, color=on(COL.get(st, NOEV)), fontweight="bold")
ax.set_yticks(range(len(qnames))); ax.set_yticklabels([T(q) for q in qnames[::-1]], fontsize=7)
ax.set_xlim(-0.2, 7.2); ax.set_ylim(-0.7, len(qnames) - 0.3); ax.set_xticks([]); ax.grid(False); ax.tick_params(axis="y", length=0)
for s in ("left", "bottom"):
    ax.spines[s].set_visible(False)
handles = [plt.Rectangle((0, 0), 1, 1, fc=c) for c in COL.values()] + [plt.Rectangle((0, 0), 1, 1, fc="white", ec=INK, lw=1.6)]
ax.legend(handles, [T(k) for k in COL] + [T("Critical gate")], loc="center left", bbox_to_anchor=(0.62, 0.5), fontsize=7)
ax.set_title(f"Carte de maturité : {sum(1 for g in gates if g[2] == 'MET')} portes franchies sur {len(gates)}" if FR else f"Readiness map: {sum(1 for g in gates if g[2] == 'MET')} of {len(gates)} gates met", loc="left")
save(fig, "fig17_1_readiness_map.png")
print("framework figures written")
