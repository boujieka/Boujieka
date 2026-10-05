"""Build the Book 7 companion materials in the seven-folder structure of the online folder.

Run: python3 tools/book7/build_materials.py   (after prepare7.py, the model run and build_manual7.py)
Writes deliverables/BOOK7_Companion_Materials/ and deliverables/BOOK7_Companion_Materials.zip.
Stand-alone documents (framework, Kasiri case, transaction tools, technical reference, sources) are cut from the
resolved book text, so their numbers are the book's numbers, and rendered in the house style as PDF and Word.
"""
import csv
import os
import re
import shutil
import subprocess
import sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
FR = os.environ.get("BOOK_LANG") == "fr"  # French package: deliverables/BOOK7_Ressources_FR, files suffixed _FR
RESOLVED = "book7/build_fr/book7_resolved.md" if FR else "book7/build/book7_resolved.md"
OUT = "deliverables/BOOK7_Ressources_FR" if FR else "deliverables/BOOK7_Companion_Materials"
WORK = "book7/build_fr/materials" if FR else "book7/build/materials"
FIG = os.path.abspath("book7/src_fr/figures" if FR else "book7/src/figures")
VERSION = "Version 1.0 (release candidate 1)" if FR else "Version 1.0 release candidate 1"
SUB_BASE = ("Ressource d'accompagnement du Livre 7, Développement et financement de l'hydroélectricité : de la rivière au bouclage financier"
            if FR else "Companion to Book 7, Hydropower Development and Finance: From River to Financial Close")
SFX = "_FR" if FR else ""
FOLDERS = ["01_Hydro_Readiness_Framework", "02_Bankable_Hydro_Model", "03_Model_User_Manual", "04_Kasiri_River_Case",
           "05_Transaction_Tools", "06_Technical_Due_Diligence", "07_Sources_and_References"]

lines = open(RESOLVED, encoding="utf8").read().split("\n")


def span(start_pat, end_pat=None):
    i = next(k for k, l in enumerate(lines) if re.match(start_pat, l))
    j = next((k for k in range(i + 1, len(lines)) if end_pat and re.match(end_pat, lines[k])), len(lines))
    return lines[i:j]


def promote(block):
    out, fence = [], False
    for l in block:
        if l.strip().startswith("```"):
            fence = not fence
        elif not fence:
            if l.startswith("### "):
                l = "## " + l[4:]
            elif l.startswith("## "):
                l = "# " + l[3:]
        out.append(l)
    return out


def fix(text):
    return text.replace("](../src/figures/", f"]({FIG}/").replace("](../src_fr/figures/", f"]({FIG}/")


NOTE = ("This document is an extract of Book 7, *Hydropower Development and Finance*. References to chapters and annexes "
        "are to the book. Every Kasiri figure comes from MODEL 7 with its default inputs and is illustrative.")
if FR:
    NOTE = ("Ce document est un extrait du Livre 7, *Développement et financement de l'hydroélectricité*. Les renvois aux chapitres "
            "et aux annexes visent le livre. Tous les chiffres de Kasiri proviennent de MODEL 7 avec ses données par défaut et sont illustratifs.")
    docs = {
        "Etude_de_cas_Kasiri_River_Hydro_FR": ("04_Kasiri_River_Case", "Kasiri River Hydro : le cas du site au bouclage",
            "Un projet au fil de l'eau de 60 MW, République de Navaria (fictive)", span(r"^# Chapitre 18\.", r"^# Annexes")),
        "Outils_de_transaction_Annexes_A_a_I_FR": ("05_Transaction_Tools", "Outils de transaction",
            "Listes de contrôle et modèles, annexes A à I",
            ["# À propos de ces outils", "", NOTE, "", "La version Word de ce document est modifiable : copiez un tableau dans votre propre fichier et complétez-le."]
            + [""] + promote(span(r"^## Annexe A\.", r"^## Annexe J\."))),
        "Referentiel_due_diligence_technique_FR": ("06_Technical_Due_Diligence", "Référentiel de due diligence technique",
            "Hydrologie, génie civil, équipements, construction, exploitation et E&S, annexes N à R",
            promote(span(r"^# Référentiel de due diligence technique"))),
        "Sources_et_statut_de_verification_FR": ("07_Sources_and_References", "Sources et statut de vérification",
            "Annexe L du Livre 7", ["# À propos de cette liste", "", NOTE, ""] + promote(span(r"^## Annexe L\.", r"^## Annexe M\."))),
    }
    fw = span(r"^## Le Hydro Readiness Framework", r"^# Chapitre 1\.")[1:]
    docs = {"Hydro_Readiness_Framework_FR": ("01_Hydro_Readiness_Framework", "Le Hydro Readiness Framework",
            "8 questions, 23 portes, 1 décision de bouclage financier",
            ["# Le Hydro Readiness Framework\u2122", "", NOTE, ""] + [l for l in fw if not l.startswith("La suite du livre")] + [""]
            + span(r"^# Chapitre 17\.", r"^# Chapitre 18\.")), **docs}

docs = docs if FR else {
    "Kasiri_River_Hydro_Case": ("04_Kasiri_River_Case", "Kasiri River Hydro: the case from site to close",
        "A 60 MW run-of-river project, fictional Republic of Navaria",
        span(r"^# Chapter 18\.", r"^# Annexes")),
    "Transaction_Tools_Annexes_A_to_I": ("05_Transaction_Tools", "Transaction tools",
        "Checklists and templates, Annexes A to I",
        ["# About these tools", "", NOTE, "", "The Word version of this document is editable: copy a table into your own file and fill it in."]
        + [""] + promote(span(r"^## Annex A\.", r"^## Annex J\."))),
    "Technical_Due_Diligence_Reference": ("06_Technical_Due_Diligence", "Technical due-diligence reference",
        "Hydrology, civil works, equipment, construction, operation and E&S, Annexes N to R",
        promote(span(r"^# Technical due-diligence reference"))),
    "Sources_and_Verification_Status": ("07_Sources_and_References", "Sources and verification status",
        "Annex L of Book 7",
        ["# About this list", "", NOTE, ""] + promote(span(r"^## Annex L\.", r"^## Annex M\."))),
}
if not FR:
    fw = span(r"^## The Hydro Readiness Framework", r"^# Chapter 1\.")[1:]
    docs = {"Hydro_Readiness_Framework": ("01_Hydro_Readiness_Framework", "The Hydro Readiness Framework",
        "8 Questions, 23 Gates, 1 Financial Close Decision",
        ["# The Hydro Readiness Framework\u2122", "", NOTE, ""] + [l for l in fw if not l.startswith("The rest of the book follows")] + [""] + span(r"^# Chapter 17\.", r"^# Chapter 18\.")), **docs}

if os.path.isdir(OUT):
    shutil.rmtree(OUT)
for f in FOLDERS:
    os.makedirs(os.path.join(OUT, f), exist_ok=True)
os.makedirs(WORK, exist_ok=True)

KICKER = "RESSOURCE DU LIVRE 7" if FR else "BOOK 7 COMPANION"
for name, (folder, title, sub, block) in docs.items():
    md = os.path.join(WORK, f"{name}.md")
    open(md, "w", encoding="utf8").write(fix("\n".join(block)))
    pdf = os.path.abspath(os.path.join(OUT, folder, f"{name}.pdf"))
    r = subprocess.run([sys.executable, os.path.abspath("tools/house/publish_docs.py"), f"{name}.md", pdf, "--title", title,
                        "--subtitle", f"{sub}. {SUB_BASE}", "--kicker", KICKER, "--short", title, "--edition", VERSION,
                        "--case", "cas Kasiri River Hydro" if FR else "Kasiri River Hydro case",
                        "--keywords", "hydroélectricité; financement de projets; Afrique" if FR else "hydropower; project finance; Africa"],
                       cwd=WORK, capture_output=True, text=True)
    print(name, "PDF:", (r.stdout.strip().splitlines() or ["?"])[-1])
    subprocess.run([sys.executable, "tools/book7/build_docx7.py", "--src", md, "--out", os.path.join(OUT, folder, f"{name}.docx"),
                    "--title", title, "--subtitle", f"{sub}. {SUB_BASE}", "--kicker", KICKER, "--edition", VERSION],
                   check=True, capture_output=True)

if FR:
    import glob as _g, json as _j
    TRF = {e["en"]: e["fr"] for f in sorted(_g.glob("model/fr_work/strings_*_fr.json")) for e in _j.load(open(f, encoding="utf8")) if e.get("fr")}
    D = OUT
    shutil.copy("book7/build_fr/Developpement_et_financement_hydroelectricite.pdf",
                f"{D}/01_Hydro_Readiness_Framework/LIVRE7_Developpement_et_financement_hydroelectricite_livre_complet_FR.pdf")
    shutil.copy("book7/build_fr/Developpement_et_financement_hydroelectricite_6x9.pdf",
                f"{D}/01_Hydro_Readiness_Framework/LIVRE7_Developpement_et_financement_hydroelectricite_format_6x9_FR.pdf")
    shutil.copy("model/Bankable_Hydro_Model_FR.xlsx", f"{D}/02_Bankable_Hydro_Model/MODEL7_Modele_financier_hydro_v1.0RC1_FR.xlsx")
    for f in ("MODEL7_TEST_REPORT.md", "MODEL7_GATE_AND_FORMULA_AUDIT.md"):
        shutil.copy(f"docs/{f}", f"{D}/02_Bankable_Hydro_Model/{f}")
    for f in ("MANUAL7_Manuel_utilisation_methodologie.pdf", "MANUAL7_Manuel_utilisation_methodologie.docx"):
        shutil.copy(f"manual/build_fr/{f}", f"{D}/03_Model_User_Manual/{f.replace('.', '_FR.', 1)}")
    for f in ("MODEL7_Video_Guide_FR.mp4", "MODEL7_Video_Guide_FR.fr.srt"):
        shutil.copy(f"course/video/{f}", f"{D}/03_Model_User_Manual/{f}")
    if os.path.exists("course/audio/model7_walkthrough_fr/MODEL7_Guide_audio_FR.mp3"):
        shutil.copy("course/audio/model7_walkthrough_fr/MODEL7_Guide_audio_FR.mp3", f"{D}/03_Model_User_Manual/MODEL7_Guide_audio_FR.mp3")
    wb = load_workbook("model/Bankable_Hydro_Model.xlsx", data_only=True)
    ws = wb["35_BOOK_CHECK"]
    with open(f"{D}/04_Kasiri_River_Case/Kasiri_chiffres_cles_FR.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Indicateur", "Valeur imprimée dans le Livre 7", "Valeur du modèle (données par défaut)", "Résultat", "Où dans le livre"])
        for r in range(5, ws.max_row + 1):
            if ws.cell(r, 1).value and ws.cell(r, 5).value in ("PASS", "CHECK"):
                vals = [ws.cell(r, k).value for k in (1, 2, 3, 5, 6)]
                vals = [TRF.get(v, v) if isinstance(v, str) else (str(v).replace(".", ",") if isinstance(v, float) else v) for v in vals]
                vals[4] = str(vals[4]).replace("Chapters", "Chapitres").replace("Chapter", "Chapitre").replace("Table", "Tableau").replace("Annex", "Annexe")
                w.writerow(vals)
    for f in ("research/source_database.csv", "research/source_database.md", "research/hydro_development_evidence.md",
              "research/book7_positioning_sources.md", "research/case_studies/africa_part1.md", "research/case_studies/africa_part2.md",
              "research/case_studies/international_benchmarks.md", "docs/CASE_DATA_AUDIT.md"):
        shutil.copy(f, f"{D}/07_Sources_and_References/{os.path.basename(f)}")
    open(f"{D}/LISEZMOI.txt", "w", encoding="utf8").write(
        "AFRICA ENERGY FINANCE\nLIVRE 7, DÉVELOPPEMENT ET FINANCEMENT DE L'HYDROÉLECTRICITÉ : ressources d'accompagnement\n" + VERSION + "\n\n"
        "01_Hydro_Readiness_Framework   Le livre complet (PDF A4 et format 6 x 9 pour tablette) et le cadre en document autonome (PDF, Word)\n"
        "02_Bankable_Hydro_Model        Le classeur MODEL 7 en français (il s'ouvre sur sa page de garde) et ses rapports de test (en anglais)\n"
        "03_Model_User_Manual           MANUAL 7 en français (PDF, Word), le guide vidéo en français (MP4, sous-titres) et le guide audio\n"
        "04_Kasiri_River_Case           Le cas Kasiri du chapitre 18 (PDF, Word) et ses chiffres clés (CSV)\n"
        "05_Transaction_Tools           Listes de contrôle et modèles des annexes A à I (PDF ; Word modifiable)\n"
        "06_Technical_Due_Diligence     Annexes N à R (PDF, Word)\n"
        "07_Sources_and_References      Liste des sources avec statut de vérification (PDF, Word) ; fichiers de recherche d'origine (en anglais)\n\n"
        "Commencez par la page de garde du classeur, puis le guide vidéo ou le démarrage rapide du manuel.\n"
        "Téléchargez le classeur avant usage et travaillez sur votre propre copie. Toutes les données par défaut sont illustratives.\n"
        "Les noms des feuilles et les statuts calculés (STOP, MET, ALL OK...) restent en anglais dans le classeur ; le manuel en donne le sens.\n"
        "Outil d'aide à la décision ; ne constitue pas un conseil en investissement, ni un conseil juridique, fiscal ou comptable.\n")
    zip_base = OUT
    if os.path.exists(zip_base + ".zip"):
        os.remove(zip_base + ".zip")
    shutil.make_archive(zip_base, "zip", os.path.dirname(OUT), os.path.basename(OUT))
    n = sum(len(f) for _, _, f in os.walk(OUT))
    print(f"{OUT}: {n} files; {zip_base}.zip {os.path.getsize(zip_base + '.zip') / 1048576:.1f} MB")
    sys.exit(0)

# 01 the full book (A4 PDF) next to the stand-alone framework
shutil.copy("book7/build/Hydropower_Development_and_Finance.pdf", f"{OUT}/01_Hydro_Readiness_Framework/BOOK7_Hydropower_Development_and_Finance_full_book.pdf")
# 02 model and its tests
shutil.copy("model/Bankable_Hydro_Model.xlsx", f"{OUT}/02_Bankable_Hydro_Model/MODEL7_Bankable_Hydro_Model_v1.0RC1.xlsx")
for f in ("MODEL7_TEST_REPORT.md", "MODEL7_GATE_AND_FORMULA_AUDIT.md"):
    shutil.copy(f"docs/{f}", f"{OUT}/02_Bankable_Hydro_Model/{f}")
# 03 manual, audio guide and video guide
for f in ("MANUAL7_User_and_Methodology.pdf", "MANUAL7_User_and_Methodology.docx"):
    shutil.copy(f"manual/build/{f}", f"{OUT}/03_Model_User_Manual/{f}")
shutil.copy("course/audio/model7_walkthrough/MODEL7_Audio_Guide.mp3", f"{OUT}/03_Model_User_Manual/MODEL7_Audio_Guide.mp3")
shutil.copy("course/audio/model7_walkthrough/script.md", f"{OUT}/03_Model_User_Manual/MODEL7_Audio_Guide_script.md")
for f in ("MODEL7_Video_Guide.mp4", "MODEL7_Video_Guide.en.srt", "MODEL7_Video_Guide_FR.mp4", "MODEL7_Video_Guide_FR.fr.srt"):
    shutil.copy(f"course/video/{f}", f"{OUT}/03_Model_User_Manual/{f}")
# 04 Kasiri key figures from the workbook's book check
wb = load_workbook("model/Bankable_Hydro_Model.xlsx", data_only=True)
ws = wb["35_BOOK_CHECK"]
with open(f"{OUT}/04_Kasiri_River_Case/Kasiri_key_figures.csv", "w", newline="", encoding="utf8") as fh:
    w = csv.writer(fh)
    w.writerow(["Figure", "Printed in Book 7", "Model value (default inputs)", "Result", "Where in the book"])
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value and ws.cell(r, 5).value in ("PASS", "CHECK"):
            w.writerow([ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 5).value, ws.cell(r, 6).value])
# 07 research files and audits
for f in ("research/source_database.csv", "research/source_database.md", "research/hydro_development_evidence.md",
          "research/book7_positioning_sources.md", "research/case_studies/africa_part1.md", "research/case_studies/africa_part2.md",
          "research/case_studies/international_benchmarks.md", "docs/CASE_DATA_AUDIT.md", "docs/BOOK7_V04_QA_SUMMARY.md",
          "docs/AUDIT_R1.md", "docs/AUDIT_R2.md", "docs/AUDIT_R3.md", "docs/AUDIT_R4.md"):
    shutil.copy(f, f"{OUT}/07_Sources_and_References/{os.path.basename(f)}")
open(f"{OUT}/README.txt", "w", encoding="utf8").write(
    "AFRICA ENERGY FINANCE\nBOOK 7, HYDROPOWER DEVELOPMENT AND FINANCE: companion materials\n" + VERSION + "\n\n"
    "01_Hydro_Readiness_Framework   The full book (PDF) and the framework as a stand-alone reference (PDF, Word)\n"
    "02_Bankable_Hydro_Model        MODEL 7 workbook (opens on its cover sheet) and its test reports\n"
    "03_Model_User_Manual           MANUAL 7 (PDF, Word), the video guide in English and in French (MP4, subtitles) and the audio guide with its script\n"
    "04_Kasiri_River_Case           The Kasiri case of Chapter 18 (PDF, Word) and its key figures (CSV)\n"
    "05_Transaction_Tools           Checklists and templates of Annexes A to I (PDF; editable Word)\n"
    "06_Technical_Due_Diligence     Annexes N to R (PDF, Word)\n"
    "07_Sources_and_References      Source list with verification status, research files, case-data and source audits\n\n"
    "Start with the workbook cover sheet, then the video guide or the manual's quick start.\n"
    "Download the workbook before use and work on your own copy. All default inputs are illustrative.\n"
    "Decision-support material; not investment, legal, tax or accounting advice.\n")
zip_base = OUT
if os.path.exists(zip_base + ".zip"):
    os.remove(zip_base + ".zip")
shutil.make_archive(zip_base, "zip", os.path.dirname(OUT), os.path.basename(OUT))
n = sum(len(f) for _, _, f in os.walk(OUT))
print(f"{OUT}: {n} files; {zip_base}.zip {os.path.getsize(zip_base + '.zip') / 1048576:.1f} MB")
