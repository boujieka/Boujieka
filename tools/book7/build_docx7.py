"""Word edition of Book 7 in the house style (green and gold, sans serif), from book7/build/book7_resolved.md.

python tools/book7/build_docx7.py  ->  book7/build/Hydropower_Development_and_Finance.docx
"""
import os, re
from datetime import date
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
SRC, OUT = "book7/build/book7_resolved.md", "book7/build/Hydropower_Development_and_Finance.docx"
GREEN, GOLD, INK, MUTED = RGBColor(0x0B, 0x30, 0x20), RGBColor(0xB0, 0x7C, 0x0F), RGBColor(0x1D, 0x1D, 0x1D), RGBColor(0x66, 0x66, 0x66)
FONT = "Arial"
TITLE, SUB = "Hydropower Development and Finance", "Developing, structuring and financing hydropower projects from site to financial close"
AUTHOR = "Emmanuel Boujieka Kamga"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(20); sec.top_margin = Mm(22); sec.bottom_margin = Mm(20)
TEXT_W = 170
st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10); st.font.color.rgb = INK
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
for name, size, color in [("Heading 1", 18, GREEN), ("Heading 2", 13, GREEN), ("Heading 3", 11, GOLD)]:
    h = doc.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = color
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    rf = h.element.rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); h.element.rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    h.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0); h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexfill); tcPr.append(s)


def border_bottom(p, color="B07C0F", sz=12):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "4"), ("w:color", color)):
        e.set(qn(k), v)
    b.append(e); pPr.append(b)


TOK = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|<sub>[^<]+</sub>|<sup>[^<]+</sup>)")


def add_runs(p, text, size=None, color=None, bold=None, italic=None):
    for part in TOK.split(text):
        if not part:
            continue
        b, i, sub, sup = bold, italic, False, False
        if part.startswith("**"): part, b = part[2:-2], True
        elif part.startswith("<sub>"): part, sub = part[5:-6], True
        elif part.startswith("<sup>"): part, sup = part[5:-6], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: part, i = part[1:-1], True
        r = p.add_run(part); r.bold = b; r.italic = i
        r.font.subscript = sub or None; r.font.superscript = sup or None
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


def page_number_footer(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{TITLE}   |   Page "); r.font.size = Pt(8); r.font.color.rgb = MUTED
    r2 = p.add_run()
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r2._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r2._r.append(it)
    r2.font.size = Pt(8); r2.font.color.rgb = MUTED
    h = section.header.paragraphs[0]
    hr = h.add_run("AFRICA ENERGY FINANCE   |   BOOK 7"); hr.font.size = Pt(7.5); hr.bold = True; hr.font.color.rgb = GREEN
    border_bottom(h, sz=6)


# ---- cover ----
if os.path.exists("brand/aef_logo.png"):
    doc.add_picture("brand/aef_logo.png", width=Mm(110))
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(60); add_runs(p, "BOOK 7", 11, GOLD, True)
p = doc.add_paragraph(); add_runs(p, TITLE, 28, GREEN, True); border_bottom(p, sz=18)
p = doc.add_paragraph(); add_runs(p, SUB, 14, INK)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(40); add_runs(p, "AUTHOR & IDEATION", 8.5, GOLD, True)
p = doc.add_paragraph(); add_runs(p, AUTHOR, 14, GREEN, True)
p = doc.add_paragraph(); add_runs(p, f"Africa Energy Finance  |  Business & Financial Models  |  {date.today():%B %Y}", 9.5, MUTED)
p = doc.add_paragraph(); add_runs(p, "First edition, version 0.1 (pre-publication review draft)", 9.5, MUTED)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(80)
add_runs(p, f"© {date.today().year} {AUTHOR}. All rights reserved. Decision support material; not investment, legal, tax or accounting advice. "
            "The default model inputs and the Kasiri River Hydro case are fictional and illustrative.", 8, MUTED)
# ---- contents (Word field, updated on open) ----
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
p = doc.add_paragraph(); add_runs(p, "Contents", 18, GREEN, True)
p = doc.add_paragraph(); r = p.add_run()
for t, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
    if t:
        f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
r.add_text("Right-click and choose Update Field to build the table of contents.")
f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); r._r.append(f)
upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); doc.settings.element.append(upd)
sec.different_first_page_header_footer = True
page_number_footer(sec)

# ---- body ----
lines = open(SRC, encoding="utf8").read().split("\n")
i, first_h1 = 0, True
while i < len(lines):
    L = lines[i]
    if not L.strip() or L.strip() == "{: .cap}":
        i += 1; continue
    if L.startswith("# "):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        h = doc.add_heading(L[2:].strip(), level=1); border_bottom(h)
        i += 1; continue
    if L.startswith("## "):
        doc.add_heading(L[3:].strip(), level=2); i += 1; continue
    if L.startswith("### "):
        doc.add_heading(L[4:].strip(), level=3); i += 1; continue
    m = re.match(r"!\[(.*)\]\((.*)\)", L)
    if m:
        path = os.path.normpath(os.path.join("book7/build", m.group(2)))
        doc.add_picture(path, width=Mm(150)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph(); add_runs(p, m.group(1), 8.5, MUTED, italic=True)
        i += 1; continue
    if L.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.fullmatch(r"\|[\s:|-]+\|?", lines[i].strip()):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        nc = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=nc); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
        lens = [max(len(re.sub(r"<[^>]+>|\*", "", r[j])) if j < len(r) else 0 for r in rows) for j in range(nc)]
        w = [max(6, min(l, 60)) for l in lens]; tot = sum(w)
        for ri, r in enumerate(rows):
            for j in range(nc):
                c = t.cell(ri, j); c.width = Mm(TEXT_W * w[j] / tot)
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
                add_runs(c.paragraphs[0], r[j] if j < len(r) else "", 8, RGBColor(255, 255, 255) if ri == 0 else INK, True if ri == 0 else None)
                if ri == 0: shade(c, "0B3020")
                elif ri % 2 == 0: shade(c, "FAF8F2")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue
    if L.startswith("> "):
        txt = []
        while i < len(lines) and lines[i].startswith(">"):
            if lines[i].strip() != ">": txt.append(lines[i][2:])
            i += 1
        for tline in txt:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Mm(10); add_runs(p, tline, 10, GREEN)
        continue
    m = re.match(r"^(\d+)\.\s+(.*)", L)
    if m:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Mm(7); p.paragraph_format.first_line_indent = Mm(-7)
        p.paragraph_format.space_after = Pt(3); add_runs(p, f"{m.group(1)}.\t" + m.group(2)); i += 1; continue
    if L.startswith("- "):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Mm(7); p.paragraph_format.first_line_indent = Mm(-7)
        p.paragraph_format.space_after = Pt(3); add_runs(p, "\u2022\t" + L[2:]); i += 1; continue
    para = [L.strip()]
    i += 1
    while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|> |- |\d+\.\s|\{:)", lines[i]):
        para.append(lines[i].strip()); i += 1
    txt = " ".join(para)
    p = doc.add_paragraph()
    if txt.startswith("**") and i < len(lines) and lines[i].strip() == "{: .cap}":
        add_runs(p, txt, 9, GREEN); p.paragraph_format.keep_with_next = True
    else:
        add_runs(p, txt)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
cp = doc.core_properties
cp.title, cp.subject, cp.author, cp.keywords = TITLE, SUB, AUTHOR, "hydropower; project development; project finance; Africa"
cp.comments = ""; cp.last_modified_by = AUTHOR
doc.save(OUT)
print("saved", OUT)
