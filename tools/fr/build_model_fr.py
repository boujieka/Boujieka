"""Build the French workbook MODEL 7 (Bankable_Hydro_Model_FR.xlsx) from the English master.

python3 tools/fr/build_model_fr.py
- Text cells: French labels from model/fr_work/strings_*_fr.json.
- Text results written by formulas: translated in the formulas and, identically, in any text cell holding the same words, so that
  every comparison between them still holds; keywords that formulas test with LEFT() (STOP, GO, HIGH, LOW...) stay first.
- Status words entered as evidence and tested by formulas (MET, PARTIAL, NOT MET, NO EVIDENCE, Y, OK, PASS, CHECK...) stay in
  English; sheet names, formulas, number formats and data validations are unchanged.
Then the copy is recalculated with LibreOffice and every numeric cell is compared with the English master: they must be equal.
"""
import glob
import json
import os
import re
import shutil
import subprocess
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
EN = "model/Bankable_Hydro_Model.xlsx"
OUT = "model/Bankable_Hydro_Model_FR.xlsx"
RECALC = os.environ.get("RECALC", "/root/.claude/skills/synced/595554f7-3334-4cb8-90b6-40da4dcb6b84_e238fe12-8a4b-4490-8c46-5bf1323948af/xlsx/scripts/recalc.py")

TRF = {}
for p in sorted(glob.glob("model/fr_work/strings_*_fr.json")):
    for e in json.load(open(p, encoding="utf8")):
        if e.get("fr"):
            TRF[e["en"]] = e["fr"]

# literal strings inside formulas that are results shown to the user (never statuses typed as evidence)
LIT = {
    "STOP: a critical gate is not met": "STOP : une porte critique n'est pas franchie",
    "STOP: critical evidence missing": "STOP : preuve critique manquante",
    "NOT READY: critical gates partly met": "NOT READY : portes critiques en partie franchies",
    "CONDITIONAL GO: all critical gates met": "CONDITIONAL GO : toutes les portes critiques franchies",
    "GO: evidence complete for a close decision": "GO : dossier de preuves complet pour décider du bouclage",
    "LOW additional fiscal pressure": "LOW : faible pression budgétaire supplémentaire",
    "MODERATE additional fiscal pressure": "MODERATE : pression budgétaire supplémentaire modérée",
    "HIGH additional fiscal pressure — escalate to MoF / DSA team": "HIGH : forte pression budgétaire supplémentaire, à remonter au ministère des Finances",
    "NOT BANKABLE — critical gap(s) must be closed": "NOT BANKABLE : lacune(s) critique(s) à combler",
    "NOT YET BANKABLE — development gaps remain": "NOT YET BANKABLE : des lacunes de développement subsistent",
    "BANKABLE SUBJECT TO CONDITIONS": "BANKABLE : bancable sous conditions",
    "READY FOR FINANCIAL CLOSE": "READY : prêt pour le bouclage financier",
    "Structurally under-funded": "Structurellement sous-financé",
    "n/a: negative on the success path": "n/a : négative sur le scénario de succès",
    "No pre-close action required": "Aucune action requise avant le bouclage",
    "Project: ": "Projet : ", "  |  Structure: ": "  |  Structure : ", "  |  Case: ": "  |  Cas : ",
    "  |  Active stresses: ": "  |  Stress actifs : ",
    "9-GATE SCREEN (development stage): ": "FILTRE À 9 PORTES (stade du développement) : ",
    "   |   23-GATE CLOSE READINESS (transaction stage): ": "   |   MATURITÉ À 23 PORTES (stade de la transaction) : ",
    "   |   FISCAL SCREEN: ": "   |   FILTRE BUDGÉTAIRE : ",
    " of 23": " sur 23", " of ": " sur ", " met; ": " franchies ; ", " critical not met": " critiques non franchies",
    " issue(s)": " anomalie(s)", " TO CHECK": " À VÉRIFIER",
    "Agree utility recovery plan (tariff path, loss reduction), escrow of receivables, liquidity facility / LC sizing":
        "Convenir d'un plan de redressement de la compagnie (trajectoire tarifaire, réduction des pertes), séquestre des créances, facilité de liquidité / dimensionnement de la LC",
    "Cap guarantees, replace sovereign guarantees with liquidity instruments, record commitments in fiscal risk statement":
        "Plafonner les garanties, remplacer les garanties souveraines par des instruments de liquidité, inscrire les engagements dans la déclaration des risques budgétaires",
    "Close critical legal/regulatory gaps (tariff pass-through, land, FX) before PPA signature":
        "Combler les lacunes juridiques et réglementaires critiques (répercussion tarifaire, foncier, change) avant la signature du CAE",
    "Commission independent hydrology review; extend record via regional correlation; derive P90 from simulated series":
        "Commander une revue hydrologique indépendante ; allonger la série par corrélation régionale ; calculer le P90 sur séries simulées",
    "Complete RAP implementation milestones before financial close": "Achever les jalons de mise en œuvre du PAR avant le bouclage financier",
    "Complete independent hydrology review and agree P90 basis with lenders' technical adviser":
        "Achever la revue hydrologique indépendante et convenir de la base P90 avec le conseiller technique des prêteurs",
    "Confirm demand absorption with utility dispatch study and anchor-load MoUs":
        "Confirmer l'absorption de la demande par une étude de dispatching de la compagnie et des protocoles d'accord avec des clients d'ancrage",
    "Finalise bankable feasibility study and lenders' technical adviser review":
        "Finaliser l'étude de faisabilité bancable et la revue du conseiller technique des prêteurs",
    "Lock financing terms; run lender stress cases; agree DSCR covenants":
        "Figer les conditions de financement ; passer les cas de stress des prêteurs ; convenir des covenants de DSCR",
    "Lock transmission financing and interface agreement (dates, LDs, deemed energy)":
        "Sécuriser le financement du transport et l'accord d'interface (dates, pénalités, énergie réputée livrée)",
    "Obtain Ministry of Finance fiscal-risk sign-off and disclose commitments":
        "Obtenir l'aval du ministère des Finances sur le risque budgétaire et publier les engagements",
    "Re-phase capacity or secure anchor (mining/export) offtake; update load forecast with utility":
        "Rééchelonner la puissance ou sécuriser un acheteur d'ancrage (mines, export) ; mettre à jour la prévision de charge avec la compagnie",
    "Resolve remaining GAP items via implementation agreement / regulatory decisions":
        "Régler les points GAP restants par la convention de mise en œuvre ou des décisions réglementaires",
    "Restructure financing: more concessional debt, longer tenor, grant/VGF, or tariff adjustment to close gap":
        "Restructurer le financement : plus de dette concessionnelle, maturité plus longue, subvention/VGF ou ajustement tarifaire pour combler l'écart",
    "Secure financing & EPC for evacuation line; align transmission COD with plant COD; agree deemed-energy allocation":
        "Sécuriser le financement et l'EPC de la ligne d'évacuation ; aligner sa mise en service sur celle de la centrale ; convenir de l'énergie réputée livrée",
    "Size LC to ≥6 months, ring-fence collections, monitor utility KPIs":
        "Dimensionner la LC à 6 mois au moins, cantonner les encaissements, suivre les indicateurs de la compagnie",
    "Upgrade ESIA/RAP to lender standards; independent E&S review; riparian notification":
        "Porter l'EIES et le PAR aux normes des prêteurs ; revue E&S indépendante ; notification aux États riverains",
    "Value-engineer layout / re-optimise installed capacity; obtain EPC price discovery":
        "Optimiser la conception et la puissance installée ; obtenir des prix EPC indicatifs",
}


def tr_formula(f):
    def sub(m):
        s = m.group(1)
        return '"' + LIT.get(s, s).replace('"', '""') + '"'
    return re.sub(r'"((?:[^"]|"")*)"', sub, f)


def main():
    wb = load_workbook(EN)
    n_text = n_form = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                if v.startswith("="):
                    nv = tr_formula(v)
                    if nv != v:
                        c.value, n_form = nv, n_form + 1
                elif v in LIT:
                    c.value, n_text = LIT[v], n_text + 1
                elif v in TRF:
                    c.value, n_text = TRF[v], n_text + 1
        if ws.column_dimensions["A"].width and ws.column_dimensions["A"].width < 70:
            ws.column_dimensions["A"].width = round(ws.column_dimensions["A"].width * 1.12, 1)
    wb.properties.title = "MODEL 7 : modèle de développement et de financement de l'hydroélectricité"
    wb.properties.language = "fr-FR"
    wb.save(OUT)
    print(f"{OUT}: {n_text} text cells and {n_form} formulas translated")
    res = json.loads(subprocess.run([sys.executable, RECALC, OUT, "300"], capture_output=True, text=True).stdout)
    print("recalc:", res.get("status"), res.get("total_errors", ""))
    subprocess.run([sys.executable, "tools/model/excel_quote_fix.py", OUT], check=True)
    verify()


def verify():
    a = load_workbook(EN, data_only=True)
    b = load_workbook(OUT, data_only=True)
    diffs, n = [], 0
    for ws in a.worksheets:
        wf = b[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    n += 1
                    w = wf[c.coordinate].value
                    if not isinstance(w, (int, float)) or abs(v - w) > 1e-6 * max(1, abs(v)):
                        diffs.append((ws.title, c.coordinate, v, w))
    print(f"numeric cells compared: {n}; differences: {len(diffs)}")
    for d in diffs[:15]:
        print("  DIFF", d)
    ck, bc = b["33_CHECKS"], b["35_BOOK_CHECK"]
    print("33_CHECKS:", [ck.cell(r, 3).value for r in range(4, 40) if str(ck.cell(r, 1).value or "").upper().startswith(("ALL", "TOUS"))])
    print("35_BOOK_CHECK:", sorted({str(bc.cell(r, 5).value) for r in range(5, 40) if bc.cell(r, 5).value}))
    errs = [(ws.title, c.coordinate, c.value) for ws in b.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper()]
    print("error values:", len(errs), errs[:5])
    return not diffs and not errs


if __name__ == "__main__":
    main()
