"""QA: find formulas that reference EMPTY cells on their own sheet (typical missing-sheet-prefix bug).

Run: python tools/qa_dangling.py <file.xlsx>
"""
import re
import sys
from collections import Counter

import openpyxl
from openpyxl.utils import column_index_from_string as ci

wb = openpyxl.load_workbook(sys.argv[1])
REF = re.compile(r"(?<![A-Za-z_!'\]\.])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
SHEETREF = re.compile(r"(?:'[^']+'|[A-Za-z_][A-Za-z0-9_.]*)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?")
FUNC = re.compile(r"[A-Z][A-Z0-9.]*\(")
hits = Counter()
examples = {}
ALLOW = {"Benchmark_DB"}  # blank FX on count / ratio records is intended and tested with =""
for ws in wb.worksheets:
    if ws.title in ALLOW:
        continue
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            body = re.sub(r'"[^"]*"', '""', v)
            body = SHEETREF.sub("X", body)          # drop cross-sheet refs
            body = FUNC.sub("(", body)               # drop function names
            for m in REF.finditer(body):
                if m.group(3):
                    continue                           # ranges may legitimately include blanks
                col, row_ = ci(m.group(1)), int(m.group(2))
                if col == 4 or row_ < 4:               # column D = opening balances (blank = 0); header rows
                    continue
                if ws.cell(row_, col).value is None:
                    hits[ws.title] += 1
                    examples.setdefault(ws.title, f"{c.coordinate}: {v[:110]}")
print("formulas referencing empty own-sheet cells:", dict(hits) or "none")
for k, e in examples.items():
    print(" ", k, "->", e)
