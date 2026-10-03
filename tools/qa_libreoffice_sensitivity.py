"""Recompute the Sensitivity cases inside the workbook with LibreOffice and compare with the static table.

Run: python tools/qa_libreoffice_sensitivity.py <work_dir> make ; convert <work_dir>/lo/cases/*.xlsx with soffice into <work_dir>/lo/cases_out ;
python tools/qa_libreoffice_sensitivity.py <work_dir> read"""
import openpyxl, subprocess, sys, glob, os
S = sys.argv[1]
SRC = '/home/user/Boujieka/volumes/02-solar-home-systems/model/AEF_SHS_PAYGo_Model_v0.7.xlsx'
CASES = [
 ("Base", {}), ("Downside", {"Inputs!C5": 2}), ("Severe", {"Inputs!C5": 3}),
 ("Base: default hazard x1.5", {"Scenarios!C6": 1.5}),
 ("Base: collection rate x0.95", {"Scenarios!C7": 0.95}),
 ("Base: sales volume -20%", {"Scenarios!C8": 0.8}),
 ("Base: hardware cost +15%", {"Scenarios!C9": 1.15}),
 ("Base: LCY depreciation 20% p.a.", {"Scenarios!C10": 0.20}),
 ("Base: no price increase on new contracts", {"Inputs!C13": 0.0}),
 ("Base: RBF programme off", {"Inputs!C32": 0}),
 ("Base: higher Tier 4-5 mix (20% / 10%)", {"Products!C23": .2, "Products!D23": .3, "Products!E23": .2, "Products!F23": .2, "Products!G23": .1}),
 ("Base: exit at 4.0x EBITDA", {"Inputs!C84": 4.0}),
 ("Base: repayment-linked RBF (mode 2)", {"Inputs!C34": 2}),
 ("Base: securitisation structure", {"Inputs!C60": 2}),
 ("Base: borrowing base up to 90 DPD", {"Credit_Assumptions!C11": 90}),
]
if sys.argv[2] == 'make':
    for i, (name, ch) in enumerate(CASES):
        wb = openpyxl.load_workbook(SRC)
        for ref, v in ch.items():
            sh, c = ref.split('!'); wb[sh][c].value = v
        wb.save(f"{S}/lo/cases/case{i:02d}.xlsx")
    print('made', len(CASES))
else:
    import json
    OUT = [("Peak equity need (USD m)", "KPIs!C57", 1e-6), ("Year 5 revenue (USD m)", "KPIs!I35", 1e-6), ("Year 5 EBITDA margin", "KPIs!I38", 1),
           ("Year 5 collection rate", "KPIs!I19", 1), ("DCF EV (USD m)", "Valuation!C28", 1e-6), ("Investor IRR (USD)", "Valuation!C39", 1),
           ("Investor MOIC", "Valuation!C40", 1), ("Covenant-breach months", "KPIs!C64", 1), ("Master check", "Checks!C30", None)]
    ref = openpyxl.load_workbook(SRC, data_only=False)['Sensitivity']
    res = []
    for i, (name, ch) in enumerate(CASES):
        wv = openpyxl.load_workbook(f"{S}/lo/cases_out/case{i:02d}.xlsx", data_only=True)
        row = [name]
        for label, cell, sc in OUT:
            sh, c = cell.split('!'); v = wv[sh][c].value
            row.append(v * sc if (sc and isinstance(v, (int, float))) else v)
        snap = [ref.cell(6 + i, k).value for k in range(2, 10)]
        res.append((row, snap))
    json.dump(res, open(f"{S}/sens_lo.json", "w"), default=str)
    worst = 0
    for row, snap in res:
        diffs = []
        for a, b in zip(row[1:9], snap):
            if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                d = abs(a - b) / max(1e-9, abs(b)) if abs(b) > 1e-9 else abs(a - b); diffs.append(d); worst = max(worst, d)
            elif str(a) != str(b): diffs.append(f"{a} vs {b}")
        print(f"{row[0][:42]:42s} check={row[9]} IRR wb={row[6] if not isinstance(row[6],float) else round(row[6],4)} snap={snap[5] if not isinstance(snap[5],float) else round(snap[5],4)} maxreldiff={max([d for d in diffs if isinstance(d,float)] or [0]):.2e} {[d for d in diffs if not isinstance(d,float)]}")
    print('worst relative difference', worst)
