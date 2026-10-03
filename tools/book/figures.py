"""Book figures from the companion model and full-engine snapshot runs."""
import json, os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
OUT = "book/figures"
os.makedirs(OUT, exist_ok=True)
S = json.load(open("model/snapshot_results.json"))
MAP = json.load(open("model/model_map.json"))
wb = load_workbook("model/Bankable_Hydro_Model.xlsx", data_only=True)

BLUE, ORANGE, AQUA, YELLOW, VIOLET = "#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#4a3aa7"
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
plt.rcParams.update({
    "font.family": "Liberation Sans", "font.size": 8, "axes.edgecolor": INK2, "axes.labelcolor": INK2,
    "xtick.color": INK2, "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "axes.axisbelow": True,
    "axes.titlesize": 8.5, "axes.titleweight": "bold", "axes.titlecolor": INK, "legend.frameon": False,
})
W = 5.1


def save(fig, name):
    fig.savefig(f"{OUT}/{name}", dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ---- 3.1 hydrology ----
ws = wb["03_HYDROLOGY"]
hdr = next(r for r in range(1, 60) if ws.cell(r, 1).value == "Month")
months = [ws.cell(hdr, c).value for c in range(5, 17)]
flow = [ws.cell(hdr + 2, c).value for c in range(5, 17)]
energy = [ws.cell(hdr + 6, c).value for c in range(5, 17)]
qd = wb["03_HYDROLOGY"]["C5"].value
qe = next(ws.cell(r, 3).value for r in range(4, 20) if ws.cell(r, 1).value == "Environmental flow release")
fig, ax = plt.subplots(1, 2, figsize=(W, 2.2))
ax[0].bar(months, flow, color=BLUE, width=0.62)
ax[0].axhline(qd, color=ORANGE, lw=1.4); ax[0].text(11.4, qd + 8, f"Design flow {qd:.0f}", ha="right", color=INK2, fontsize=7)
ax[0].axhline(qe, color=INK2, lw=1, ls=(0, (3, 2))); ax[0].text(11.4, qe + 8, f"E-flow {qe:.0f}", ha="right", color=INK2, fontsize=7)
ax[0].set_title("Mean river flow (m³/s)", loc="left"); ax[0].tick_params(axis="x", labelsize=6.5)
ax[1].bar(months, energy, color=AQUA, width=0.62)
ax[1].set_title("Expected energy (GWh per month)", loc="left"); ax[1].tick_params(axis="x", labelsize=6.5)
fig.tight_layout(w_pad=2)
save(fig, "fig3_1_hydrology.png")

# ---- 8.1 utility ----
fig, ax = plt.subplots(1, 2, figsize=(W, 2.3), sharey=True)
for a, key, ttl in [(ax[0], "base_locked", "Base case"), (ax[1], "offtaker", "Offtaker stress")]:
    ts = S[key]["ts"]
    idx = [i for i, k in enumerate(ts["opyr"]) if 1 <= k <= 20]
    yr = [ts["year"][i] for i in idx]
    ppa = [ts["u_ppa"][i] for i in idx]; msp = [ts["u_maxppa"][i] for i in idx]
    a.plot(yr, ppa, color=ORANGE, lw=2, label="PPA bill (utility share)")
    a.plot(yr, msp, color=BLUE, lw=2, label="Maximum sustainable payment")
    a.fill_between(yr, msp, ppa, where=[p > m for p, m in zip(ppa, msp)], color=ORANGE, alpha=0.18, lw=0, label="Payment capacity gap")
    top = max(ppa) * 2.6
    a.set_title(ttl, loc="left"); a.set_ylim(0, top)
    if max(msp) > top:
        a.annotate("payment capacity continues to rise", xy=(yr[-1], top * 0.06), ha="right", va="bottom", fontsize=6.5, color=INK2)
ax[0].set_ylabel("USD million")
h, l = ax[1].get_legend_handles_labels()
fig.legend(h, l, loc="lower center", ncol=3, fontsize=7, bbox_to_anchor=(0.5, -0.06))
fig.tight_layout(rect=(0, 0.06, 1, 1))
save(fig, "fig8_1_utility.png")

# ---- 11.1 tornado (equity IRR) ----
base = S["base_locked"]["kpi_eirr"]
pairs = [("CAPEX ±20%", "capex_m20", "capex_p20"), ("Tariff ±10%", "tar_p10", "tar_m10"), ("Generation ±10%", "gen_p10", "gen_m10"),
         ("Commercial rate +200 bp", None, "rate_p200"), ("OPEX +20%", None, "opex_p20"), ("P90 in cash flows", None, "p90cf")]
rows = []
for lab, up, dn in pairs:
    hi = S[up]["kpi_eirr"] if up else base
    lo = S[dn]["kpi_eirr"]
    rows.append((lab, lo, hi))
rows.sort(key=lambda r: abs(r[2] - r[1]))
fig, ax = plt.subplots(figsize=(W, 2.3))
for i, (lab, lo, hi) in enumerate(rows):
    ax.barh(i, (lo - base) * 100, left=base * 100, color=ORANGE, height=0.55)
    ax.barh(i, (hi - base) * 100, left=base * 100, color=BLUE, height=0.55)
    ax.text(min(lo, hi) * 100 - 0.3, i, f"{lo*100:.1f}%", va="center", ha="right", fontsize=6.5, color=INK2)
    if hi != base:
        ax.text(max(lo, hi) * 100 + 0.3, i, f"{hi*100:.1f}%", va="center", ha="left", fontsize=6.5, color=INK2)
ax.axvline(base * 100, color=INK, lw=1)
ax.set_yticks(range(len(rows))); ax.set_yticklabels([r[0] for r in rows])
ax.set_xlabel(f"Private equity IRR, % (base {base*100:.1f}%)"); ax.grid(axis="y", visible=False)
ax.set_xlim(base * 100 - 9, base * 100 + 11)
save(fig, "fig11_1_tornado.png")

# ---- 12.1 exposure composition ----
ws = wb["17A_STRUCTURES"]; C = MAP["CMP_ROWS"]
names = ["Public", "IPP", "PPP", "Hybrid", "Blended"]
pub = [ws.cell(C["pubcap"], 5 + j).value for j in range(5)]
onb = [ws.cell(C["onbud"], 5 + j).value for j in range(5)]
cont = [max(ws.cell(C["guar"], 5 + j).value + ws.cell(C["ppag"], 5 + j).value, ws.cell(C["term"], 5 + j).value) for j in range(5)]
fig, ax = plt.subplots(figsize=(W, 2.4))
b1 = ax.bar(names, pub, color=BLUE, width=0.55, label="Upfront public capital", edgecolor="white", linewidth=1.5)
b2 = ax.bar(names, onb, bottom=pub, color=VIOLET, width=0.55, label="Direct public debt", edgecolor="white", linewidth=1.5)
b3 = ax.bar(names, cont, bottom=[a + b for a, b in zip(pub, onb)], color=ORANGE, width=0.55, label="Contingent (termination or guarantees)", edgecolor="white", linewidth=1.5)
for j in range(5):
    tot = pub[j] + onb[j] + cont[j]
    ax.text(j, tot + 20, f"{tot:,.0f}", ha="center", fontsize=7, color=INK)
ax.set_ylabel("USD million"); ax.set_ylim(0, max(p + o + c for p, o, c in zip(pub, onb, cont)) * 1.18)
ax.legend(loc="upper center", ncol=3, fontsize=6.8, bbox_to_anchor=(0.5, -0.1)); ax.grid(axis="x", visible=False)
save(fig, "fig12_1_exposure.png")

# ---- 13.1 fiscal profile ----
ts = S["base_sized"]["ts"]
idx = [i for i in range(40) if ts["year"][i] <= ts["year"][0] + 37]
yr = [ts["year"][i] for i in idx]
fig, ax = plt.subplots(figsize=(W, 2.3))
ax.bar(yr, [ts["f_net"][i] for i in idx], color=[BLUE if ts["f_net"][i] >= 0 else ORANGE for i in idx], width=0.7, label="Net fiscal cash flow")
ax.plot(yr, [ts["cl_max"][i] for i in idx], color=VIOLET, lw=2, label="Maximum contingent exposure")
ax.axhline(0, color=INK, lw=0.8)
ax.set_ylabel("USD million"); ax.legend(loc="upper right", fontsize=7)
save(fig, "fig13_1_fiscal.png")

# ---- 17.1 fiscal NPV by scenario (central vs consolidated) ----
keys = [("base_locked", "Base"), ("high", "High case"), ("low", "Low case"), ("drought", "Drought"), ("climate", "Climate trend"),
        ("overrun", "CAPEX overrun 27%"), ("overrun96", "CAPEX overrun 96%"), ("delay", "Two-year delay"), ("lowdem", "Low demand"),
        ("rate", "Interest +300 bp"), ("trans", "Transmission 2 yrs late"), ("fx", "FX step 50%"), ("offtaker", "Offtaker stress"),
        ("offtaker_nobs", "Offtaker, no backstop"), ("combined", "Combined")]
c1 = [S[k]["fis_npv"] for k, _ in keys]; c2 = [S[k]["fis_npv_cons"] for k, _ in keys]
fig, ax = plt.subplots(figsize=(W, 3.6))
n = len(keys); y = list(range(n))[::-1]; h = 0.38
ax.barh([v + h / 2 for v in y], c1, height=h, color=BLUE, label="Central government", edgecolor="white", linewidth=0.8)
ax.barh([v - h / 2 for v in y], c2, height=h, color=ORANGE, label="Consolidated (incl. state utility)", edgecolor="white", linewidth=0.8)
ax.set_yticks(y); ax.set_yticklabels([l for _, l in keys]); ax.axvline(0, color=INK, lw=0.8)
ax.set_xlabel("Fiscal NPV, USD million (8%)"); ax.grid(axis="y", visible=False)
ax.legend(loc="upper center", ncol=2, fontsize=7, bbox_to_anchor=(0.45, -0.13))
save(fig, "fig17_1_fiscal_npv.png")
print("figures written")
