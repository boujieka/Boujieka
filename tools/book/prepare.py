"""Assemble the book source: substitute model numbers, resolve citations, build annex tables.

python tools/book/prepare.py  ->  book/build/manuscript_resolved.md
"""
import csv, json, re, sys, os
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
SRC = ["book/src/01_front.md", "book/src/02_part1_3.md", "book/src/03_part4_5.md", "book/src/04_part6_7.md"]
text = "\n\n".join(open(p, encoding="utf8").read() for p in SRC)

# ---------------- numbers ----------------
MAP = json.load(open("model/model_map.json"))
SNAP = json.load(open("model/snapshot_results.json"))
wb = load_workbook("model/Bankable_Hydro_Model.xlsx", data_only=True)


def cell(ref):
    m = re.fullmatch(r"'([^']+)'!\$([A-Z]+)\$(\d+)", ref)
    return wb[m.group(1)][m.group(2) + m.group(3)].value


M = {k: cell(v) for k, v in MAP["REF"].items() if re.fullmatch(r"'([^']+)'!\$([A-Z]+)\$(\d+)", str(v))}
CMP = {}
ws = wb["17A_STRUCTURES"]
for key, row in MAP["CMP_ROWS"].items():
    CMP[key] = {str(j + 1): ws.cell(row, 5 + j).value for j in range(5)}
X = {}
X["debt_total"] = SNAP["base_sized"]["debt_c"] + SNAP["base_sized"]["debt_m"]
pe = [CMP["pubexp"][str(j)] for j in range(1, 6)]
X["pubexp_min"], X["pubexp_max"] = min(pe), max(pe)
yrs = SNAP["base_sized"]["ts"]["opyr"]
k2 = yrs.index(2)
X["blend_y2"] = None
ws15 = wb["15_TARIFF"]
r_bl = MAP["TSROW"]["blend"][1]
X["blend_y2"] = ws15.cell(r_bl, 5 + k2).value
for i in range(2, 6):
    X[f"req{i}"] = X["blend_y2"] * (1 + SNAP[f"s{i}"]["req_tariff_flex"])

STATUS_ABBR = {"READY": "R", "CONDITIONAL": "C", "DEVELOPMENT GAP": "D", "CRITICAL GAP": "X"}


def fmt(v, f):
    if v is None:
        return "n/a"
    if isinstance(v, str):
        if f == "t":
            return STATUS_ABBR.get(v, v)
        if f == "s":
            return v.split(" additional")[0].capitalize()
        return v
    if f == "pct0": return f"{v*100:.0f}%"
    if f == "pct1": return f"{v*100:.1f}%"
    if f == "x1": return f"{v:.1f}x"
    if f == "x2": return f"{v:.2f}x"
    if f == "n0":
        s = f"{abs(v):,.0f}"
        return ("\u2212" if v < -0.5 else "") + s
    if f == "n1": return f"{v:,.1f}"
    if f == "n2": return f"{v:.2f}"
    raise ValueError(f)


def sub(m):
    key, f = m.group(1), m.group(2)
    ns, _, rest = key.partition(".")
    if ns == "m":
        v = M[rest]
    elif ns == "cmp":
        a, b = rest.split(".")
        v = CMP[a][b]
    elif ns == "x":
        v = X[rest]
    else:
        v = SNAP[ns][rest]
    return fmt(v, f)


text = re.sub(r"\{\{([\w.]+)\|(\w+)\}\}", sub, text)
assert "{{" not in text, re.findall(r"\{\{[^}]*\}\}", text)[:5]

# ---------------- red-team chapter ----------------
rt = "book/src/05_redteam.md"
text = text.replace("%%REDTEAM", open(rt, encoding="utf8").read() if os.path.exists(rt) else "(pending)")

# ---------------- annex tables ----------------
ws31 = wb["31_CASE_STUDY"]
rows = ["Table: Table A.1. Public benchmark cases (as published; nominal, unadjusted)",
        "| Project | Country | Structure | MW | USD/kW | Timeline | Credit support |", "|---|---|---|---|---|---|---|"]
r = MAP["CASE_ROW0"] + 1
while ws31.cell(r, 1).value and "FICTIONAL" not in str(ws31.cell(r, 1).value):
    v = [ws31.cell(r, c).value for c in range(1, 13)]
    ukw = f"{v[6]:,.0f}" if isinstance(v[6], (int, float)) else "n/a"
    rows.append(f"| {v[0]} | {v[1]} | {v[2]} | {v[3]:,.0f} | {ukw} | {v[7]} | {v[9]} |")
    r += 1
rows.append("Note: Sources for each row are listed in the case files; costs in EUR are not converted (n/a). 'Snippet' marks a fact seen only in search results.")
text = text.replace("%%CASETABLE", "\n".join(rows))

bench = []
src = open("research/source_database.md", encoding="utf8").read()
sec = src.split("## BENCHMARK RANGES FOR MODEL DEFAULTS")[1].split("###")[0]
bench.append("Table: Table B.1. Benchmark ranges used to set or check model defaults")
bench.append("| Parameter | Low | Central | High | Unit | Source |")
bench.append("|---|---|---|---|---|---|")
for line in sec.splitlines():
    if line.startswith("| ") and not line.startswith("| Parameter") and not line.startswith("|---"):
        c = [x.strip() for x in line.strip().strip("|").split("|")]
        ids = ["[@" + i.strip() + "]" for i in c[5].split(";") if re.fullmatch(r"[A-Z]{2}-\d{2}", i.strip())]
        bench.append(f"| {c[0]} | {c[1]} | {c[2]} | {c[3]} | {c[4]} | {' '.join(ids)} |")
bench.append("Note: Full notes and confidence ratings for each benchmark are in the research database.")
text = text.replace("%%BENCHTABLE", "\n".join(bench))

# ---------------- citations ----------------
SD = {row["id"]: row for row in csv.DictReader(open("research/source_database.csv", encoding="utf8"))}
CASEFILES = {"A1": "research/case_studies/africa_part1.md", "A2": "research/case_studies/africa_part2.md",
             "INT": "research/case_studies/international_benchmarks.md"}
CS = {}
for k, p in CASEFILES.items():
    for line in open(p, encoding="utf8"):
        m = re.match(r"^\s*(?:-\s*)?(?:\*\*S(\d+)\*\*|\[S(\d+)\])\s*(.*)$", line)
        if m:
            n = m.group(1) or m.group(2)
            CS[f"{k}:S{n}"] = m.group(3).strip()

order, refnum = [], {}


def bib(cid):
    if cid == "CI":
        return "Bankable Hydro research team (2026). Competitive intelligence review of hydropower finance books, models, courses and public tools. Research file, companion materials."
    if cid in SD:
        s = SD[cid]
        return f"{s['publisher']} ({s['year']}). *{s['title']}*. {s['doc_type']}. {s['url']}"
    if cid in CS:
        t = CS[cid]
        t = t.replace("—", ".").replace("–", "-").replace("**", "")
        t = re.sub(r"\s*\((fetched|snippet)\)\s*$", lambda mm: " (" + mm.group(1).replace("fetched", "accessed") + ")" if mm.group(1) == "snippet" else "", t)
        t = re.sub(r"\s+", " ", t)
        return t
    raise KeyError(cid)


def cite(m):
    ids = [i.strip().lstrip("@") for i in m.group(1).split(";")]
    nums = []
    for i in ids:
        if i not in refnum:
            bib(i)  # validate
            order.append(i)
            refnum[i] = len(order)
        nums.append(refnum[i])
    nums = sorted(set(nums))
    return "[" + ", ".join(str(n) for n in nums) + "]"


text = re.sub(r"\[(@[^\]]+)\]", cite, text)
refs = "\n".join(f"REF|{refnum[i]}|{bib(i).replace('|', '/')}" for i in order)
text = text.replace("%%REFERENCES", "Numbered in order of first citation. Entries marked 'snippet' or 'search-result only' were seen only in search results and should be verified before reliance.\n\n" + refs)

# ---------------- style lint ----------------
problems = []
for bad in ["—", "–", "delve", "tapestry", "testament", "pivotal", "underscore", "showcase", "vibrant", "intricate", "landscape of",
            "in today's", "it is important to note", "plays a crucial", "crucial role", "seamless", "Additionally,", "Furthermore,", "Moreover,"]:
    for line in text.splitlines():
        if line.startswith("REF|"):
            continue
        if bad.lower() in line.lower():
            problems.append((bad, line[:90]))
if problems:
    for p in problems[:30]:
        print("LINT:", p)
os.makedirs("book/build", exist_ok=True)
open("book/build/manuscript_resolved.md", "w", encoding="utf8").write(text)
print("words:", len(re.findall(r"[A-Za-z]+", text)), "refs:", len(order), "lint issues:", len(problems))
