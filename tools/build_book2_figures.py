"""Figures for Book 2, PAYGo Solar Finance (first edition).

Run: python tools/build_book2_figures.py <model_values.xlsx> <case_values.xlsx>
The two workbooks are the companion model and the SolaraPay case model AFTER a full recalculation (for example a LibreOffice
headless conversion), so that every charted number is a value computed by the workbook itself. Worked examples that live in the
text (Chapter 1.4) are recomputed here from the parameters stated in the text.
Output: volumes/02-solar-home-systems/book/figures/figNN_*.png (300 dpi).
"""
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "volumes/02-solar-home-systems/book/figures"
GREEN, GOLD, INK, MUTED, GRID, PALE = "#0B3020", "#B07C0F", "#222222", "#666666", "#E3E3E3", "#F7F3E8"
TIER = ["#1E8A5C", "#B07C0F", "#3E7CB1", "#A3473A", "#7B5EA7"]  # validated categorical order (dataviz validator, light surface)
plt.rcParams.update({"font.family": "Liberation Sans", "font.size": 9, "axes.edgecolor": "#BBBBBB", "axes.labelcolor": INK,
                     "xtick.color": MUTED, "ytick.color": MUTED, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.titlesize": 10, "axes.titleweight": "bold", "axes.titlecolor": INK, "savefig.dpi": 300})


def save(fig, name):
    OUT.mkdir(parents=True, exist_ok=True)
    fig.savefig(OUT / name, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print(name)


def grid(ax, axis="y"):
    ax.grid(axis=axis, color=GRID, linewidth=0.6)
    ax.set_axisbelow(True)


def box(ax, x, y, w, h, text, fc=PALE, ec=GREEN, color=INK, size=8.5, bold=False):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.012,rounding_size=0.015", fc=fc, ec=ec, lw=1.1))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=size, color=color, wrap=True,
            fontweight="bold" if bold else "normal")


def arrow(ax, x0, y0, x1, y1, color=GREEN):
    ax.add_patch(FancyArrowPatch((x0, y0), (x1, y1), arrowstyle="-|>", mutation_scale=10, color=color, lw=1.1))


def canvas(w=7.2, h=3.6):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    return fig, ax


def row(ws, r, c0, n):
    return [ws.cell(r, c0 + k).value for k in range(n)]


def main(mpath, cpath):
    M = load_workbook(mpath, data_only=True)
    C = load_workbook(cpath, data_only=True)
    tiers = [f"Tier {k}" for k in range(1, 6)]

    # ---- Figure 1: the analytical chain
    fig, ax = canvas(7.4, 3.2)
    links = ["Customer", "Affordability", "Product", "Cohort", "Credit", "Receivables", "Cash", "FX", "Funding", "RBF",
             "Stress\ntesting", "Valuation", "Investment\ncommittee"]
    per_row, w, h = 7, 0.1, 0.2
    for k, t in enumerate(links):
        r_, c_ = divmod(k, per_row)
        x = 0.02 + c_ * 0.14 if r_ == 0 else 0.02 + (per_row - 1 - c_) * 0.14
        y = 0.66 if r_ == 0 else 0.18
        last = k == len(links) - 1
        box(ax, x, y, w, h, t, fc=GREEN if last else PALE, color="white" if last else INK, bold=last)
        if k < len(links) - 1:
            if c_ < per_row - 1:
                if r_ == 0:
                    arrow(ax, x + w, y + h / 2, x + 0.14, y + h / 2)
                else:
                    arrow(ax, x, y + h / 2, x - 0.14 + w, y + h / 2)
            else:
                arrow(ax, x + w / 2, 0.66, x + w / 2, 0.18 + h)
    ax.text(0.02, 0.96, "Each link passes its result to the next; a weakness found early grows on its way down.", fontsize=8.5, color=MUTED)
    save(fig, "fig01_chain.png")

    # ---- Figure 2: three businesses in one legal entity
    fig, ax = canvas(7.2, 3.3)
    cols = [("Retailer", "Hardware margin at sale\nInventory, logistics, agents", "Income statement:\ngross margin"),
            ("Utility like service\nprovider", "Serves the account for years\nPlatform, field service", "Cost to serve per\nactive account"),
            ("Consumer lender", "Funds the instalments\nCredit risk, collections", "Receivables book,\ncredit losses, funding")]
    for k, (t, d, n) in enumerate(cols):
        x = 0.03 + k * 0.33
        box(ax, x, 0.62, 0.28, 0.2, t, fc=GREEN, color="white", bold=True, size=9)
        box(ax, x, 0.33, 0.28, 0.22, d, size=8)
        box(ax, x, 0.05, 0.28, 0.2, n, fc="white", ec=GOLD, size=8)
        arrow(ax, x + 0.14, 0.62, x + 0.14, 0.55)
        arrow(ax, x + 0.14, 0.33, x + 0.14, 0.25)
    ax.text(0.03, 0.9, "One PAYGo company, three businesses: what each one does, and where it shows in the numbers", fontsize=8.5, color=MUTED)
    save(fig, "fig02_three_businesses.png")

    # ---- Figure 3: Chapter 1.4 worked example, cash against two accounting views (parameters from the text)
    inst, cr, h_, T = 2190, 0.88, 0.026, 24
    months = list(range(0, T + 1))
    cash, prof_orig, prof_late = [-22000.0], [18000 - 4000 - 18728.0], [14000.0]
    for a in range(1, T + 1):
        coll = inst * cr * (1 - h_) ** a
        cash.append(cash[-1] + coll)
        prof_orig.append(prof_orig[-1] + 16560 / T)
        prof_late.append(prof_late[-1] + 16560 / T - (inst - coll))
    fig, ax = plt.subplots(figsize=(7.0, 3.4))
    grid(ax)
    for ys, lab, col_, ls in ((cash, "Cumulative cash", GREEN, "-"), (prof_orig, "Cumulative profit, lifetime loss charged at sale (workbook)", GOLD, "-"),
                              (prof_late, "Cumulative profit, losses charged as instalments are missed", "#3E7CB1", "--")):
        ax.plot(months, [v / 1000 for v in ys], color=col_, lw=2, ls=ls, label=lab)
    ax.legend(frameon=False, fontsize=8, loc="lower right")
    ax.axhline(0, color="#999999", lw=0.8)
    pay = next(m for m in months if cash[m] > 0)
    ax.scatter([pay], [cash[pay] / 1000], s=28, color=GREEN, zorder=3)
    ax.annotate(f"Cash turns positive in month {pay}", (pay, cash[pay] / 1000), xytext=(8, -14), textcoords="offset points", ha="left", fontsize=8)
    ax.annotate(f"All three end at LCY {cash[-1]:,.0f}", (T, cash[-1] / 1000), xytext=(23.6, -5.5), textcoords="data", ha="right", fontsize=8,
                arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.7))
    ax.set_xlabel("Months after sale")
    ax.set_ylabel("LCY thousands per unit")
    ax.set_xlim(0, T)
    ax.set_title("One illustrative Tier 2 sale: cash and profit tell different stories", loc="left")
    save(fig, "fig03_unit_cash.png")
    assert pay == 14 and round(cash[-1]) == 11832, (pay, cash[-1])

    # ---- Figure 4: payment burden by tier against the model's policy threshold (Consumer_Risk)
    cr_ws = M["Consumer_Risk"]
    burden = row(cr_ws, 10, 3, 5)
    thr = cr_ws["C4"].value
    fig, ax = plt.subplots(figsize=(6.4, 3.0))
    grid(ax)
    bars = ax.bar(tiers, [b * 100 for b in burden], color=[TIER[k] for k in range(5)], width=0.55)
    for b_, v in zip(bars, burden):
        ax.text(b_.get_x() + b_.get_width() / 2, v * 100 + 0.3, f"{v * 100:.1f}%", ha="center", fontsize=8, color=INK)
    ax.axhline(thr * 100, color=INK, lw=1, ls="--")
    ax.text(4.45, thr * 100 + 0.3, f"Model policy threshold {thr * 100:.0f}%", ha="right", fontsize=8, color=INK)
    ax.set_ylabel("Instalment as % of household income")
    ax.set_title("Payment burden by tier, default model (illustrative incomes)", loc="left")
    save(fig, "fig04_burden.png")

    # ---- Figure 5: cumulative repayment by account age and tier (Curves, default model)
    cv = M["Curves"]
    ends = []
    fig, ax = plt.subplots(figsize=(7.0, 3.4))
    grid(ax)
    for j in range(5):
        c0 = 2 + j * 13
        due = [cv.cell(8 + a, c0 + 1).value or 0 for a in range(61)]
        coll = [cv.cell(8 + a, c0 + 2).value or 0 for a in range(61)]
        cd = cc = 0.0
        xs, ys = [], []
        for a in range(1, 61):
            cd += due[a]
            cc += coll[a]
            if due[a] > 0 or cd > 0:
                xs.append(a)
                ys.append(cc / cd * 100 if cd else None)
        ax.plot(xs, ys, color=TIER[j], lw=2)
        ends.append((ys[-1], j))
    ends.sort()
    placed = []
    for y_, j in ends:
        yl = max(y_, placed[-1] + 1.4) if placed else y_
        placed.append(yl)
        ax.annotate(f"Tier {j + 1}  {y_:.0f}%", (60, y_), xytext=(61, yl), textcoords="data", va="center", fontsize=8, color=INK,
                    annotation_clip=False)
    ax.set_xlim(1, 60)
    ax.set_xlabel("Account age (months)")
    ax.set_ylabel("Cumulative collections ÷ cumulative instalments due (%)")
    ax.set_title("Cohort repayment by account age, default model (Base)", loc="left")
    save(fig, "fig05_repayment_curves.png")

    # ---- Figure 6: SolaraPay actual cohorts at month 12 against the management plan curve
    vd = C["Vintage_Dashboard"]
    actual = [vd.cell(10 + j, 3).value for j in range(3)]
    plan = []
    for j in range(3):
        c0 = 2 + j * 13
        due = sum((cv.cell(8 + a, c0 + 1).value or 0) for a in range(1, 13))
        coll = sum((cv.cell(8 + a, c0 + 2).value or 0) for a in range(1, 13))
        plan.append(coll / due)
    fig, ax = plt.subplots(figsize=(6.4, 3.0))
    grid(ax)
    xs = range(3)
    b1 = ax.bar([x - 0.18 for x in xs], [p * 100 for p in plan], width=0.34, color="#BBBBBB", label="Management plan (default curves)")
    b2 = ax.bar([x + 0.18 for x in xs], [a * 100 for a in actual], width=0.34, color=GREEN, label="SolaraPay cohorts (synthetic history)")
    for bars_, vals in ((b1, plan), (b2, actual)):
        for b_, v in zip(bars_, vals):
            ax.text(b_.get_x() + b_.get_width() / 2, v * 100 + 0.8, f"{v * 100:.1f}%", ha="center", fontsize=8)
    ax.set_xticks(list(xs), tiers[:3])
    ax.set_ylim(0, 112)
    ax.set_ylabel("Repayment at month 12 (%)")
    ax.legend(frameon=False, fontsize=8, loc="upper left", ncol=2)
    ax.set_title("SolaraPay: repayment at month 12, actual cohorts against plan", loc="left")
    save(fig, "fig06_solarapay_vs_plan.png")

    # ---- Figure 7: metric map (PERFORM 2026 KPIs and operational metrics)
    fig, ax = canvas(7.4, 3.6)
    box(ax, 0.02, 0.6, 0.46, 0.3, "PAYGo PERFORM KPIs (June 2026)\nRR paid vs plan · RR paid vs financed\nRR at 90 days · RR at 2x term · Ownership rate at 2x",
        fc=GREEN, color="white", size=8.5)
    box(ax, 0.52, 0.6, 0.46, 0.3, "Operational and lender metrics\nCollection rate · PAR30, PAR90 · Receivables at risk\nWrite off ratio · Active ratio · Enabled rate",
        fc=PALE, size=8.5)
    box(ax, 0.02, 0.12, 0.46, 0.36, "Answers: how well do customers repay,\nand how many come to own the device?\nContract data, daily, payments applied to\ninstalments; subsidies excluded", fc="white", ec=GREEN, size=8)
    box(ax, 0.52, 0.12, 0.46, 0.36, "Answers: how much cash came in this month,\nhow much of the book is late, is the\ncovenant met? Definition printed beside\neach figure; never a substitute for RR", fc="white", ec=GOLD, size=8)
    arrow(ax, 0.25, 0.6, 0.25, 0.48)
    arrow(ax, 0.75, 0.6, 0.75, 0.48, color=GOLD)
    save(fig, "fig07_metric_map.png")

    # ---- Figure 8: growth and the portfolio collection rate (KPIs, default model)
    kp = M["KPIs"]
    cr_y = row(kp, 19, 5, 5)
    units = row(kp, 14, 5, 5)
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(7.2, 2.9))
    years = [f"Y{k}" for k in range(1, 6)]
    for ax, vals, fmt, title, col_ in ((a1, [u / 1000 for u in units], "{:.0f}k", "Units sold per year", "#BBBBBB"),
                                      (a2, [c * 100 for c in cr_y], "{:.1f}%", "Portfolio collection rate", GREEN)):
        grid(ax)
        bars = ax.bar(years, vals, color=col_, width=0.55)
        for b_, v in zip(bars, vals):
            ax.text(b_.get_x() + b_.get_width() / 2, v * 1.01, fmt.format(v), ha="center", fontsize=8)
        ax.set_title(title, loc="left")
    a2.set_ylim(0, 100)
    fig.suptitle("Same customers, same curves: the ratio falls as growth slows (default model, Base)", x=0.02, ha="left", fontsize=9, color=MUTED, y=1.03)
    fig.tight_layout()
    save(fig, "fig08_growth_ratio.png")

    # ---- Figure 9: FX map
    fig, ax = canvas(7.4, 3.4)
    box(ax, 0.02, 0.62, 0.3, 0.28, "Local currency\nPrices, deposits, instalments\nReceivables · Facility\nServicing, staff, overheads", fc=PALE, size=8)
    box(ax, 0.68, 0.62, 0.3, 0.28, "US dollars\nHardware · RBF per unit\nTerm loan · Investor ticket\nand returns", fc=PALE, ec=GOLD, size=8)
    box(ax, 0.35, 0.66, 0.3, 0.2, "Exchange rate path\n(Timeline)", fc=GREEN, color="white", size=8.5)
    arrow(ax, 0.32, 0.76, 0.35, 0.76)
    arrow(ax, 0.68, 0.76, 0.65, 0.76, color=GOLD)
    box(ax, 0.02, 0.08, 0.46, 0.38, "Transaction effects (cash)\nDearer hardware on every new unit\nLarger local repayments and interest on the loan\nMore local currency per USD of RBF", fc="white", ec=GREEN, size=8)
    box(ax, 0.52, 0.08, 0.46, 0.38, "Translation effects (no cash)\nDollar loan restated each month\n(loss booked in the income statement)\nBook and equity worth less in USD", fc="white", ec=GOLD, size=8)
    arrow(ax, 0.5, 0.66, 0.25, 0.46)
    arrow(ax, 0.5, 0.66, 0.75, 0.46, color=GOLD)
    save(fig, "fig09_fx_map.png")

    # ---- Figure 10: funding stack over time (Financing, default model)
    fn = M["Financing"]
    eq = [fn.cell(11, 5 + i).value / 1e9 for i in range(60)]
    tl = [fn.cell(19, 5 + i).value / 1e9 for i in range(60)]
    rf = [fn.cell(23, 5 + i).value / 1e9 for i in range(60)]
    fig, ax = plt.subplots(figsize=(7.0, 3.3))
    grid(ax)
    m_ = list(range(1, 61))
    ax.stackplot(m_, eq, tl, rf, colors=[GREEN, GOLD, "#3E7CB1"], alpha=0.9, edgecolor="white", linewidth=0.5)
    for lab, y_, c_ in (("Cumulative equity invested", eq[-1] / 2, "white"), ("USD term loan (in LCY)", None, INK),
                        ("Receivables facility drawn", eq[-1] + tl[-1] + rf[-1] / 2, "white")):
        if y_ is None:
            ax.text(14, eq[13] + tl[13] / 2, lab, ha="center", va="center", fontsize=8, color="white")
        else:
            ax.text(59, y_, lab, ha="right", va="center", fontsize=8, color=c_)
    ax.set_xlim(1, 60)
    ax.set_xlabel("Month")
    ax.set_ylabel("LCY billions (closing balances)")
    ax.set_title("Funding stack over the plan, default model (Base)", loc="left")
    save(fig, "fig10_funding_stack.png")

    # ---- Figure 11: RBF cycle
    fig, ax = canvas(7.4, 2.6)
    steps = ["Sale and\ninstallation\n(company pays)", "Eligibility\nand claim", "Verification", "Approval", "Disbursement\n(company is paid)"]
    for k, t in enumerate(steps):
        x = 0.02 + k * 0.198
        box(ax, x, 0.42, 0.16, 0.34, t, fc=GREEN if k in (0, 4) else PALE, color="white" if k in (0, 4) else INK, size=8)
        if k < 4:
            arrow(ax, x + 0.16, 0.59, x + 0.198, 0.59)
    ax.plot([0.1, 0.9], [0.22, 0.22], color=GOLD, lw=2)
    ax.text(0.5, 0.08, "Working capital gap: the company finances the subsidy until it is paid (the model uses one verification lag)",
            ha="center", fontsize=8, color=INK)
    save(fig, "fig11_rbf_cycle.png")

    # ---- Figure 12: investor IRR by sensitivity case (Sensitivity, static table recomputed in the workbook)
    sn = M["Sensitivity"]
    cases = []
    for r in range(6, 21):
        name, irr = sn.cell(r, 1).value, sn.cell(r, 7).value
        if name and isinstance(irr, (int, float)):
            cases.append((name.replace("Base: ", ""), irr))
    cases.sort(key=lambda x: x[1])
    fig, ax = plt.subplots(figsize=(7.0, 4.0))
    grid(ax, "x")
    cols_ = [GREEN if n == "Base" else ("#A3473A" if v < 0 else "#BBBBBB") for n, v in cases]
    ax.barh([n for n, _ in cases], [v * 100 for _, v in cases], color=cols_, height=0.6)
    for k, (n, v) in enumerate(cases):
        if v >= 0:
            ax.text(v * 100 + 1, k, f"{v * 100:.1f}%", va="center", ha="left", fontsize=7.5)
        else:
            ax.text(v * 100 + 1, k, f"({-v * 100:.1f}%)", va="center", ha="left", fontsize=7.5, color="white")
    ax.axvline(0, color="#999999", lw=0.8)
    ax.set_xlabel("Investor IRR in USD (%)")
    ax.tick_params(axis="y", labelsize=7.5)
    ax.set_title("Investor IRR by case, default model (Severe has no IRR: no positive flow)", loc="left")
    save(fig, "fig12_irr_cases.png")

    # ---- Figure 13: peak equity by scenario
    pe = [(sn.cell(r, 1).value, sn.cell(r, 2).value) for r in (6, 7, 8)]
    fig, ax = plt.subplots(figsize=(5.6, 2.8))
    grid(ax)
    bars = ax.bar([p[0] for p in pe], [p[1] for p in pe], color=[GREEN, GOLD, "#A3473A"], width=0.5)
    for b_, (_, v) in zip(bars, pe):
        ax.text(b_.get_x() + b_.get_width() / 2, v + 0.8, f"USD {v:.1f}m", ha="center", fontsize=8)
    ax.set_ylabel("Peak equity requirement (USD m)")
    ax.set_title("Peak equity requirement by scenario, default model", loc="left")
    save(fig, "fig13_peak_equity.png")

    # ---- Figure 14: readiness gates, default model and SolaraPay
    gm, gc = M["Investment_Readiness"], C["Investment_Readiness"]
    fig, ax = plt.subplots(figsize=(7.2, 5.6))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(-0.5, 24.5)
    ax.text(0.72, 24, "Default model", ha="center", fontsize=8.5, fontweight="bold")
    ax.text(0.89, 24, "SolaraPay", ha="center", fontsize=8.5, fontweight="bold")
    for k in range(23):
        r = 7 + k
        y = 23 - k
        gname = gm.cell(r, 2).value
        gname = gname if len(gname) <= 62 else gname[:60].rsplit(" ", 1)[0] + "..."
        ax.text(0.0, y, f"{k + 1:>2}  {gname}", va="center", fontsize=7.0, color=INK)
        for x_, ws in ((0.72, gm), (0.89, gc)):
            res = ws.cell(r, 5).value
            typ = ws.cell(r, 3).value
            met = res == "Met"
            ax.add_patch(FancyBboxPatch((x_ - 0.065, y - 0.36), 0.13, 0.72, boxstyle="round,pad=0.01", fc=GREEN if met else ("white" if typ == "Manual" else "#F2DCDB"),
                                        ec=GREEN if met else "#BBBBBB", lw=0.8))
            ax.text(x_, y, "Met" if met else ("Manual: not started" if typ == "Manual" else "Not met"), ha="center", va="center", fontsize=6.5,
                    color="white" if met else INK)
    ax.text(0.0, -0.4, f"Gates met: default {gm['C4'].value} of 23; SolaraPay {gc['C4'].value} of 23. Readiness means ready for independent validation, never 'investment grade'.",
            fontsize=7.5, color=MUTED)
    save(fig, "fig14_readiness.png")

    # ---- Figure 15: investment committee decision flow
    fig, ax = canvas(7.4, 3.0)
    flow = ["Data request\nand cohort tape", "Recalibrate the\nmodel to history", "Stress and\nfunding tests", "Readiness gates\nand red flags"]
    for k, t in enumerate(flow):
        x = 0.02 + k * 0.19
        box(ax, x, 0.55, 0.15, 0.3, t, size=8)
        if k < 3:
            arrow(ax, x + 0.15, 0.7, x + 0.19, 0.7)
    outs = [("Go", GREEN, "white"), ("Conditional go\n(conditions,\ncovenants, caps)", GOLD, "white"), ("Stop", "#A3473A", "white")]
    for k, (t, fc, tc) in enumerate(outs):
        y = 0.74 - k * 0.32
        box(ax, 0.8, y, 0.19, 0.24 if k == 1 else 0.17, t, fc=fc, color=tc, size=7.5, bold=True)
        arrow(ax, 0.74, 0.7, 0.8, y + 0.09)
    ax.text(0.02, 0.2, "Each outcome is tied to evidence: a condition names the gate or the risk it addresses.", fontsize=8, color=MUTED)
    save(fig, "fig15_ic_flow.png")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
