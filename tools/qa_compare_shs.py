"""QA: compare workbook values (evaluated by tools/qa_evaluate.py) with the Python twin.

Run: python tools/qa_compare_shs.py <evaluated.xlsx> <scenario>
Requires <evaluated.xlsx>.pkl produced by qa_evaluate.py.
"""
import pickle
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from shadow_shs import run  # noqa: E402

import json
path, scn = sys.argv[1], int(sys.argv[2])
GEN = json.loads(sys.argv[3]) if len(sys.argv) > 3 else {}
vals = pickle.load(open(path + ".pkl", "rb"))
book = Path(path).name.upper()
wb = openpyxl.load_workbook(path)


def get(sheet, cell):
    return vals.get(f"'[{book}]{sheet.upper()}'!{cell}")


def _norm(t):
    return t.replace("\u2014", ", ").replace("\u2013", " to ").replace(" - ", ": ").strip()


def row_of(sheet, text, col=1):
    ws = wb[sheet]
    text = _norm(text)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and _norm(v) == text:
            return r
    raise KeyError(f"{sheet}: {text}")


def yrs(sheet, text):
    r = row_of(sheet, text)
    return [get(sheet, f"{c}{r}") for c in "EFGHI"]


tw = run(scn, gen=GEN)
rows = [
    ("Revenue", yrs("Annual", "Total revenue"), tw["rev"]),
    ("EBITDA", yrs("Annual", "EBITDA"), tw["ebitda"]),
    ("Net income", yrs("Annual", "Net income"), tw["ni"]),
    ("Closing cash", yrs("Annual", "Closing cash"), tw["cash"]),
    ("Gross receivables", yrs("Annual", "PAYGo receivables - gross"), tw["gross"]),
    ("Facility", yrs("Annual", "Receivables facility"), tw["rf"]),
    ("Collection rate", yrs("KPIs", "Operational collection rate (collected / due, excl. down payments; not a PERFORM KPI)"), tw["cr"]),
    ("DSCR", yrs("KPIs", "DSCR ((CFO + interest) / debt service)"), tw["dscr"]),
    ("FCFF", yrs("Valuation", "Unlevered free cash flow (FCFF)"), tw["fcff"]),
]
singles = [
    ("Peak equity (LCY)", get("KPIs", f"C{row_of('KPIs', 'Peak cumulative equity requirement (LCY)')}"), tw["peak_eq"]),
    ("DCF EV (LCY)", get("Valuation", f"C{row_of('Valuation', 'DCF enterprise value at model start (LCY)')}"), tw["ev"]),
    ("Exit equity (LCY)", get("Valuation", f"C{row_of('Valuation', 'Exit equity value, 100% (LCY)')}"), tw["exit_eq"]),
    ("Investor IRR", get("Valuation", f"C{row_of('Valuation', 'Investor IRR (USD)')}"), tw["irr"]),
    ("Investor MOIC", get("Valuation", f"C{row_of('Valuation', 'Investor MOIC (USD)')}"), tw["moic"]),
    ("Covenant breach months", get("KPIs", f"C{row_of('KPIs', 'Total months with a covenant breach')}"), tw["breach_months"]),
    ("Master check", get("Checks", f"C{row_of('Checks', 'MASTER CHECK')}"), "OK"),
]
worst = 0.0
for name, xl, py in rows:
    diffs = []
    for a, b in zip(xl, py):
        rel = abs(a - b) / max(1.0, abs(b))
        diffs.append(rel)
    worst = max(worst, max(diffs))
    print(f"{name:20s} max rel diff {max(diffs):.2e}   excel Y5 {xl[-1]:,.3f}   twin Y5 {py[-1]:,.3f}")
for name, a, b in singles:
    if isinstance(b, str) or isinstance(a, str):
        ok = str(a) == str(b) or (str(a) == "n/a" and b != b)
        print(f"{name:20s} excel {a}   twin {b}   {'OK' if ok else 'MISMATCH'}")
        continue
    if b != b:  # NaN
        print(f"{name:20s} excel {a}   twin n/a")
        continue
    rel = abs(a - b) / max(1.0, abs(b))
    worst = max(worst, rel)
    print(f"{name:20s} excel {a:,.4f}   twin {b:,.4f}   rel diff {rel:.2e}")
print(f"WORST RELATIVE DIFFERENCE: {worst:.2e}")
