"""Build the Mini-Grid Financial Feasibility Calculator (entry product).

Usage: python tools/build_entry_calculator.py [output.xlsx]
Then recalculate with LibreOffice so cached values exist for previewers.
"""
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import i18n
import print_areas
import branding

LANG, OUT = i18n.setup(sys.argv)
T = i18n.T
if OUT is None:
    OUT = {"en": "product/01-entry-calculator/MiniGrid_Feasibility_Calculator_v1.xlsx",
           "fr": "product/01-entry-calculator/Calculateur_Faisabilite_MiniReseau_v1_FR.xlsx"}[LANG]

MAX_YEARS = 20  # operating years modelled (columns), year 0 = construction
FONT = "Arial"
NAVY = "1F3A5F"
TEAL = "1B7F79"
LIGHT = "EEF3F8"
YELLOW = "FFF2CC"

f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_h1 = Font(name=FONT, size=12, bold=True, color="FFFFFF")
f_h2 = Font(name=FONT, size=10, bold=True, color=NAVY)
f_txt = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_note = Font(name=FONT, size=9, italic=True, color="666666")
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
fill_h1 = PatternFill("solid", fgColor=NAVY)
fill_sec = PatternFill("solid", fgColor=LIGHT)
fill_in = PatternFill("solid", fgColor=YELLOW)
fill_kpi = PatternFill("solid", fgColor="E2F0EE")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
top_line = Border(top=Side(style="thin", color="000000"))

USD = i18n.money('$#,##0;($#,##0);"-"')
USD2 = i18n.money('$#,##0.00;($#,##0.00);"-"')
USD3 = i18n.money('$#,##0.000;($#,##0.000);"-"')
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x";(0.00"x");"-"'

wb = Workbook()
names = {}


def define(name, ws, cell):
    ref = f"'{ws.title}'!${''.join(c for c in cell if c.isalpha())}${''.join(c for c in cell if c.isdigit())}"
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
    names[name] = ref


def header_bar(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text).font = f_h1
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = fill_h1


# ---------------------------------------------------------------- Start Here
ws0 = wb.active
ws0.title = "Start Here"
ws0.sheet_view.showGridLines = False
ws0.column_dimensions["A"].width = 3
ws0.column_dimensions["B"].width = 110
lines = [
    ("Mini-Grid Financial Feasibility Calculator", f_title),
    ("Solar PV + Battery + Diesel hybrid  |  20-year annual model  |  Energy-access edition", f_h2),
    ("", None),
    ("WHAT THIS TOOL ANSWERS", f_h2),
    ("Is this mini-grid financially viable, can it carry debt, and how much subsidy (grant / RBF) closes the gap?", f_txt),
    ("", None),
    ("HOW TO USE IT (10 minutes)", f_h2),
    ("1. Go to 'Inputs'. Edit ONLY the yellow cells with blue text. Every other cell is a formula.", f_txt),
    ("2. Enter your customer segments (households, productive users, commercial, institutions).", f_txt),
    ("3. Check the auto-sizing on 'Inputs' (PV, battery, diesel). Type a value in the Override column to force your own design.", f_txt),
    ("4. Enter unit costs, operating costs and the financing structure (grant, RBF per connection, debt).", f_txt),
    ("5. Read 'Dashboard': IRR, NPV, LCOE, DSCR, payback, and the viability gap (subsidy needed to reach your hurdle rate).", f_txt),
    ("6. Use the Scenario levers on 'Inputs' (tariff, demand, CAPEX multipliers) to stress-test the project.", f_txt),
    ("", None),
    ("COLOUR LEGEND", f_h2),
    ("Blue text on yellow fill = input you may change   |   Black text = formula, do not overwrite   |   Green text = link to another sheet", f_txt),
    ("", None),
    ("SHEETS", f_h2),
    ("Inputs      - all assumptions, auto-sizing and scenario levers", f_txt),
    ("Cash Flow   - 20-year annual engine: demand, generation, revenue, OPEX, tax, debt, CFADS, DSCR, project and equity cash flows", f_txt),
    ("Dashboard   - KPIs, bankability checks, viability gap and charts", f_txt),
    ("Checks      - integrity checks (should all read OK)", f_txt),
    ("", None),
    ("KEY SIMPLIFICATIONS (read before relying on results)", f_h2),
    ("- Annual time step; construction in Year 0, operations from Year 1. No intra-year dispatch: solar share is capped by the target solar fraction (storage limit).", f_txt),
    ("- Night-time solar energy passes through the battery: the round-trip efficiency loss is charged to PV production.", f_txt),
    ("- Tax: flat rate on positive taxable profit, no loss carry-forward. RBF and grants are treated as non-taxable cash inflows. Check local tax treatment.", f_txt),
    ("- Debt: single senior loan, annuity repayment after optional interest-only grace years. No DSRA, no fees, no refinancing.", f_txt),
    ("- PV and battery capacity are fixed after Year 0. If demand keeps growing, diesel fills the gap and fuel costs rise. Plan an expansion or cap growth.", f_txt),
    ("- Battery replacements create negative cash-flow years, so IRR can be misleading. Use NPV and the viability gap as the primary decision metrics.", f_txt),
    ("- Viability gap = upfront subsidy that brings the pre-subsidy project NPV (at the hurdle rate) to zero.", f_txt),
    ("- The example values pre-filled are ILLUSTRATIVE ONLY. They are not market benchmarks. Replace them with your own project data and quotes.", f_txt),
    ("", None),
    ("DISCLAIMER", f_h2),
    ("This workbook is a screening tool for educational and pre-feasibility purposes. It is not investment, legal, tax or engineering advice.", f_note),
    ("Results depend entirely on the inputs you provide. Perform independent technical and financial due diligence before any investment decision.", f_note),
    ("Licence: single user / single organisation. Do not resell or redistribute.", f_note),
]
for i, (t, fnt) in enumerate(lines, start=2):
    c = ws0.cell(row=i, column=2, value=t)
    if fnt:
        c.font = fnt
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------- Inputs
wi = wb.create_sheet("Inputs")
wi.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [44, 14, 14, 14, 14, 14, 3, 60]):
    wi.column_dimensions[col].width = w
wi["A1"] = "INPUTS"
wi["A1"].font = f_title
wi["A2"] = "Edit yellow cells only. Example values are illustrative, not benchmarks."
wi["A2"].font = f_note

r = 4


def section(title):
    global r
    r += 1
    header_bar(wi, r, title, 6)
    r += 1


def inp(label, name, value, fmt, unit="", note="", key=False):
    """Single input: label in A, value in C, unit in D, note in H."""
    global r
    wi.cell(row=r, column=1, value=label).font = f_txt
    c = wi.cell(row=r, column=3, value=value)
    c.font = f_input
    c.fill = fill_in
    c.number_format = fmt
    c.border = box
    if unit:
        wi.cell(row=r, column=4, value=unit).font = f_note
    if note:
        wi.cell(row=r, column=8, value=note).font = f_note
    define(name, wi, f"C{r}")
    r += 1


def calc(label, name, formula, fmt, unit="", note=""):
    global r
    wi.cell(row=r, column=1, value=label).font = f_txt
    c = wi.cell(row=r, column=3, value=formula)
    c.font = f_txt
    c.number_format = fmt
    c.border = box
    if unit:
        wi.cell(row=r, column=4, value=unit).font = f_note
    if note:
        wi.cell(row=r, column=8, value=note).font = f_note
    define(name, wi, f"C{r}")
    r += 1


section("1. PROJECT")
inp("Project name", "ProjectName", "Example village mini-grid", "@")
inp("Country / site", "Site", "Example site", "@")
inp("Currency label", "Currency", "USD", "@", note="Label only. All amounts must be entered in this currency.")
inp("First operating year", "StartYear", 2027, "0", "year")
inp("Project life", "Life", 20, "0", "years", "Maximum 20. Cash flows after this year are zero.")
inp("Discount rate / hurdle rate (project)", "Hurdle", 0.10, PCT, "", "Used for NPV, LCOE and viability gap.", key=True)
inp("Target equity IRR", "TargetEqIRR", 0.15, PCT)
inp("Minimum DSCR required by lender", "MinDSCR", 1.30, MULT)
inp("Tariff escalation", "TariffEsc", 0.03, PCT, "per year")
inp("Cost escalation (OPEX, fuel)", "CostEsc", 0.03, PCT, "per year")
inp("Corporate tax rate", "TaxRate", 0.25, PCT)
inp("Tax depreciation period", "DepYears", 15, "0", "years", "Straight-line on total CAPEX.")

section("2. SCENARIO LEVERS (1.00 = base case)")
inp("Tariff multiplier", "LevTariff", 1.0, MULT)
inp("Demand multiplier", "LevDemand", 1.0, MULT)
inp("CAPEX multiplier", "LevCapex", 1.0, MULT)
inp("OPEX multiplier", "LevOpex", 1.0, MULT)

section("3. CUSTOMERS & DEMAND")
hdr = ["Segment", "Customers", "kWh / month", "Tariff /kWh", "Conn. fee", "Monthly bill"]
for j, h in enumerate(hdr, start=1):
    c = wi.cell(row=r, column=j, value=h)
    c.font = f_h2
    c.fill = fill_sec
    c.alignment = Alignment(horizontal="center", wrap_text=True)
r += 1
segments = [
    ("Households", 800, 15, 0.40, 10),
    ("Productive users (mills, welding, cold rooms)", 60, 250, 0.35, 25),
    ("Commercial (shops, bars, kiosks)", 100, 50, 0.40, 15),
    ("Institutions (health centre, school, water)", 10, 300, 0.35, 50),
    ("Other / custom", 0, 0, 0, 0),
]
SEG_FIRST = r
for i, (lab, n, kwh, tar, fee) in enumerate(segments, start=1):
    wi.cell(row=r, column=1, value=lab).font = f_input
    wi.cell(row=r, column=1).fill = fill_in
    for j, (v, fmt) in enumerate([(n, NUM), (kwh, NUM1), (tar, USD3), (fee, USD)], start=2):
        c = wi.cell(row=r, column=j, value=v)
        c.font = f_input
        c.fill = fill_in
        c.number_format = fmt
        c.border = box
    c = wi.cell(row=r, column=6, value=f"=C{r}*D{r}")
    c.font = f_txt
    c.number_format = USD2
    c.border = box
    define(f"Seg{i}_Name", wi, f"A{r}")
    define(f"Seg{i}_N", wi, f"B{r}")
    define(f"Seg{i}_kWh", wi, f"C{r}")
    define(f"Seg{i}_Tariff", wi, f"D{r}")
    define(f"Seg{i}_Fee", wi, f"E{r}")
    r += 1
NSEG = len(segments)
wi.cell(row=r, column=1, value="Total / weighted average").font = f_bold
wi.cell(row=r, column=2, value=f"=SUM(B{SEG_FIRST}:B{r-1})").number_format = NUM
wi.cell(row=r, column=3, value=f"=IFERROR(SUMPRODUCT(B{SEG_FIRST}:B{r-1},C{SEG_FIRST}:C{r-1})/B{r},0)").number_format = NUM1
wi.cell(row=r, column=4, value=f"=IFERROR(SUMPRODUCT(B{SEG_FIRST}:B{r-1},C{SEG_FIRST}:C{r-1},D{SEG_FIRST}:D{r-1})/SUMPRODUCT(B{SEG_FIRST}:B{r-1},C{SEG_FIRST}:C{r-1}),0)").number_format = USD3
for j in range(1, 7):
    wi.cell(row=r, column=j).font = f_bold
    wi.cell(row=r, column=j).border = top_line
define("TotalCustomers", wi, f"B{r}")
r += 2
inp("Connection ramp - Year 1 (% of customers connected)", "Ramp1", 0.60, PCT)
inp("Connection ramp - Year 2", "Ramp2", 0.85, PCT)
inp("Connection ramp - Year 3 onwards", "Ramp3", 1.00, PCT)
inp("Annual growth in consumption per customer", "DemandGrowth", 0.03, PCT, "per year")
inp("Collection rate (share of bills actually paid)", "Collection", 0.92, PCT, key=True)

section("4. TECHNICAL DESIGN & AUTO-SIZING")
inp("PV specific yield", "SpecYield", 1500, NUM, "kWh/kWp/yr", "Use PVGIS or your design study for the site.")
inp("PV degradation", "Degradation", 0.005, PCT, "per year")
inp("Distribution & conversion losses", "Losses", 0.12, PCT, "", "Energy generated vs energy sold.")
inp("Target solar fraction (max share of generation from PV)", "SolarFraction", 0.90, PCT, "", "Storage-limited cap; the rest comes from diesel.")
inp("PV oversizing factor", "PVOversize", 1.15, MULT, "", "Margin for weather, curtailment and growth.")
inp("Share of daily energy consumed at night", "NightShare", 0.45, PCT)
inp("Battery usable depth of discharge", "DoD", 0.80, PCT)
inp("Battery round-trip efficiency", "RTE", 0.90, PCT, "", "Energy out / energy in. Applies to the solar energy stored for night-time use.")
calc("Storage loss factor on solar energy", "StorageFactor", "=1+NightShare*(1/RTE-1)", "0.000", "",
     "PV must produce this much energy per kWh of solar delivered, because night-time solar passes through the battery.")
inp("Peak-to-average load ratio", "PeakRatio", 2.5, MULT)
inp("Diesel generator efficiency", "DieselEff", 3.0, NUM1, "kWh/litre")
inp("Diesel price at Year 1", "FuelPrice", 1.20, USD2, "per litre")
r += 1
for j, h in zip([1, 2, 3, 4, 5], ["Component", "Auto-size", "Override", "Used", "Unit"]):
    c = wi.cell(row=r, column=j, value=h)
    c.font = f_h2
    c.fill = fill_sec
r += 1
# design demand = Year-3 energy sold at full connection
seg_energy = "+".join(f"Seg{i}_N*Seg{i}_kWh" for i in range(1, NSEG + 1))
wi.cell(row=r, column=1, value="Design energy sold (Year 3, full connection)").font = f_txt
c = wi.cell(row=r, column=2, value=f"=({seg_energy})*12*LevDemand*Ramp3*(1+DemandGrowth)^2")
c.number_format = NUM
wi.cell(row=r, column=5, value="kWh/yr").font = f_note
define("DesignEnergy", wi, f"B{r}")
r += 1
wi.cell(row=r, column=1, value="Design generation (incl. losses)").font = f_txt
c = wi.cell(row=r, column=2, value="=DesignEnergy/(1-Losses)")
c.number_format = NUM
wi.cell(row=r, column=5, value="kWh/yr").font = f_note
define("DesignGen", wi, f"B{r}")
r += 1
sizing = [
    ("PV array", "PV_kWp", "=DesignGen*SolarFraction*StorageFactor/SpecYield*PVOversize", "kWp"),
    ("Battery storage (usable basis -> nameplate)", "BESS_kWh", "=DesignGen/365*NightShare/DoD", "kWh"),
    ("Diesel generator (backup, covers peak)", "Diesel_kW", "=DesignGen/8760*PeakRatio", "kW"),
]
for lab, nm, f, unit in sizing:
    wi.cell(row=r, column=1, value=lab).font = f_txt
    c = wi.cell(row=r, column=2, value=f)
    c.number_format = NUM1
    c.border = box
    c = wi.cell(row=r, column=3)
    c.fill = fill_in
    c.font = f_input
    c.number_format = NUM1
    c.border = box
    c = wi.cell(row=r, column=4, value=f'=IF(ISNUMBER(C{r}),C{r},B{r})')
    c.number_format = NUM1
    c.font = f_bold
    c.border = box
    wi.cell(row=r, column=5, value=unit).font = f_note
    define(nm, wi, f"D{r}")
    r += 1
wi.cell(row=r, column=1, value="Leave Override blank to use the auto-size. Auto-sizing is a screening rule of thumb, not an engineering design (use HOMER or equivalent for final design).").font = f_note
r += 1

section("5. CAPEX (unit costs, Year 0)")
inp("PV modules, structures & inverters", "C_PV", 700, USD, "per kWp")
inp("Battery storage system", "C_BESS", 350, USD, "per kWh")
inp("Diesel generator", "C_Diesel", 400, USD, "per kW")
inp("Distribution network", "C_Network", 450, USD, "per connection")
inp("Smart meters & service drops", "C_Meter", 80, USD, "per connection")
inp("Powerhouse, BOS, civil works (lump sum)", "C_BOS", 60000, USD)
inp("Development costs (studies, permits, community)", "C_DevPct", 0.08, PCT, "of hard CAPEX")
inp("Contingency", "C_ContPct", 0.07, PCT, "of hard CAPEX")
inp("Battery life", "BattLife", 8, "0", "years")
inp("Battery replacement cost", "C_BattRepl", 250, USD, "per kWh", "In real terms at Year 0, escalated with cost escalation.")
inp("Fund replacements via maintenance reserve? (1 = yes, 0 = no)", "UseMRA", 1, "0", "", "1 = annual reserve contributions smooth the battery replacement in CFADS (lender practice).")
r += 1
calc("Hard CAPEX", "HardCapex",
     "=(PV_kWp*C_PV+BESS_kWh*C_BESS+Diesel_kW*C_Diesel+TotalCustomers*(C_Network+C_Meter)+C_BOS)*LevCapex", USD)
calc("TOTAL CAPEX", "TotalCapex", "=HardCapex*(1+C_DevPct+C_ContPct)", USD)
wi.cell(row=r - 1, column=1).font = f_bold
wi.cell(row=r - 1, column=3).font = f_bold
calc("CAPEX per connection", "CapexPerConn", "=IFERROR(TotalCapex/TotalCustomers,0)", USD)

section("6. OPEX (Year-1 values, escalated)")
inp("Fixed O&M", "O_FixedPct", 0.025, PCT, "of hard CAPEX / yr")
inp("Local staff & operator", "O_Staff", 18000, USD, "per year")
inp("Insurance", "O_InsPct", 0.005, PCT, "of total CAPEX / yr")
inp("Licences, land, regulatory fees", "O_Fees", 3000, USD, "per year")
inp("Diesel generator O&M", "O_DieselOM", 0.03, USD3, "per kWh diesel")
inp("Customer management, mobile money fees", "O_CustPct", 0.03, PCT, "of revenue")

section("7. FINANCING & SUBSIDIES")
inp("Upfront CAPEX grant", "Grant", 680000, USD, "", "Capital subsidy received in Year 0.", key=True)
inp("RBF per new connection", "RBFperConn", 250, USD, "per connection", "Paid in the year the connection is made (verified).", key=True)
inp("RBF verification lag", "RBFLag", 0, "0", "years", "0 = paid same year, 1 = paid next year.")
inp("Debt share of CAPEX net of grant", "DebtPct", 0.30, PCT, key=True)
inp("Interest rate", "IntRate", 0.09, PCT)
inp("Loan tenor (incl. grace)", "Tenor", 10, "0", "years")
inp("Grace period (interest only)", "Grace", 1, "0", "years")
r += 1
calc("Debt amount", "DebtAmt", "=MAX(0,(TotalCapex-Grant))*DebtPct", USD)
calc("Equity amount", "EquityAmt", "=MAX(0,TotalCapex-Grant-DebtAmt)", USD)

# data validation: percentages 0..1
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                    showErrorMessage=True, errorTitle=T("Invalid"), error=T("Enter a value between 0% and 100%."))
wi.add_data_validation(dv)
for nm in ["Ramp1", "Ramp2", "Ramp3", "Collection", "SolarFraction", "NightShare", "DoD", "RTE", "Losses", "DebtPct", "TaxRate"]:
    dv.add(names[nm].split("!")[1].replace("$", ""))
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MAX_YEARS),
                     showErrorMessage=True, error=T(f"Project life must be 1-{MAX_YEARS} years."))
wi.add_data_validation(dv2)
dv2.add(names["Life"].split("!")[1].replace("$", ""))
wi.freeze_panes = "A4"

# ---------------------------------------------------------------- Cash Flow
wc = wb.create_sheet("Cash Flow")
wc.sheet_view.showGridLines = False
wc.column_dimensions["A"].width = 46
wc.column_dimensions["B"].width = 14
FC = 3  # first year column (C) = Year 0
LC = FC + MAX_YEARS
for col in range(FC, LC + 1):
    wc.column_dimensions[get_column_letter(col)].width = 11
wc["A1"] = "CASH FLOW ENGINE (annual, nominal)"
wc["A1"].font = f_title
wc["A2"] = "All cells are formulas. Column B = total or NPV where relevant."
wc["A2"].font = f_note

ROW = {}
row_ptr = [4]


def L(y):
    return get_column_letter(FC + y)


def ref(name, y):
    return f"{L(y)}{ROW[name]}"


def line(name, label, fn, fmt=USD, total=None, bold=False, start=0, font=None):
    """fn(y) -> formula string for year y (0..MAX_YEARS). start: first year with formula; earlier years = 0."""
    rr = row_ptr[0]
    ROW[name] = rr
    wc.cell(row=rr, column=1, value=label).font = f_bold if bold else f_txt
    for y in range(0, MAX_YEARS + 1):
        c = wc.cell(row=rr, column=FC + y, value=(fn(y) if y >= start else 0))
        c.number_format = fmt
        c.font = font or (f_bold if bold else f_txt)
        if bold:
            c.border = top_line
    if total == "sum":
        c = wc.cell(row=rr, column=2, value=f"=SUM({L(0)}{rr}:{L(MAX_YEARS)}{rr})")
    elif total == "npv":
        wc.cell(row=rr, column=1).value = label + "  [B = NPV]"
        c = wc.cell(row=rr, column=2, value=f"={L(0)}{rr}+NPV(Hurdle,{L(1)}{rr}:{L(MAX_YEARS)}{rr})")
    else:
        c = None
    if c is not None:
        c.number_format = fmt
        c.font = f_bold
    row_ptr[0] += 1
    return rr


def blank():
    row_ptr[0] += 1


def sec(title):
    rr = row_ptr[0]
    wc.cell(row=rr, column=1, value=title).font = f_h1
    for col in range(1, LC + 1):
        wc.cell(row=rr, column=col).fill = fill_h1
    row_ptr[0] += 1


# timeline
rr = row_ptr[0]
ROW["year"] = rr
wc.cell(row=rr, column=1, value="Project year").font = f_bold
for y in range(0, MAX_YEARS + 1):
    c = wc.cell(row=rr, column=FC + y, value=y)
    c.font = f_bold
    c.fill = fill_sec
    c.alignment = Alignment(horizontal="center")
wc.cell(row=rr, column=2, value="Total / NPV").font = f_bold
wc.cell(row=rr, column=2).fill = fill_sec
row_ptr[0] += 1
line("cal", "Calendar year", lambda y: f"=StartYear-1+{ref('year', y)}", fmt="0")
line("op", "Operating flag", lambda y: f"=IF(AND({ref('year', y)}>=1,{ref('year', y)}<=Life),1,0)", fmt="0")
line("ramp", "Connection ramp", lambda y: f"=IF({ref('year', y)}=1,Ramp1,IF({ref('year', y)}=2,Ramp2,Ramp3))*{ref('op', y)}", fmt=PCT, start=1)
line("tesc", "Tariff index", lambda y: f"=(1+TariffEsc)^({ref('year', y)}-1)", fmt="0.000", start=1)
line("cesc", "Cost index", lambda y: f"=(1+CostEsc)^({ref('year', y)}-1)", fmt="0.000", start=1)
line("dgrow", "Consumption growth index", lambda y: f"=(1+DemandGrowth)^({ref('year', y)}-1)*LevDemand", fmt="0.000", start=1)
blank()

sec("DEMAND & CONNECTIONS")
for i in range(1, NSEG + 1):
    line(f"conn{i}", f"Connections - segment {i}", lambda y, i=i: f"=Seg{i}_N*{ref('ramp', y)}", fmt=NUM, start=1)
    wc.cell(row=ROW[f"conn{i}"], column=1, value=f"=\"Connections - \"&Seg{i}_Name")
line("conn", "Total connections", lambda y: "=" + "+".join(ref(f"conn{i}", y) for i in range(1, NSEG + 1)), fmt=NUM, bold=True)
line("newconn", "New connections in year", lambda y: f"=MAX(0,{ref('conn', y)}-{ref('conn', y-1)})", fmt=NUM, total="sum", start=1)
blank()
for i in range(1, NSEG + 1):
    line(f"kwh{i}", f"Energy sold - segment {i} (kWh)",
         lambda y, i=i: f"={ref(f'conn{i}', y)}*Seg{i}_kWh*12*{ref('dgrow', y)}", fmt=NUM, total="sum", start=1)
    wc.cell(row=ROW[f"kwh{i}"], column=1, value=f"=\"Energy sold - \"&Seg{i}_Name&\" (kWh)\"")
line("kwh", "Total energy sold (kWh)", lambda y: "=" + "+".join(ref(f"kwh{i}", y) for i in range(1, NSEG + 1)), fmt=NUM, total="sum", bold=True)
blank()

sec("GENERATION MIX")
line("gen", "Gross generation required (kWh)", lambda y: f"={ref('kwh', y)}/(1-Losses)", fmt=NUM, total="sum", start=1)
line("pvavail", "PV energy available (kWh)", lambda y: f"=PV_kWp*SpecYield*(1-Degradation)^({ref('year', y)}-1)*{ref('op', y)}", fmt=NUM, total="sum", start=1)
line("solar", "Solar energy delivered (kWh)", lambda y: f"=MIN({ref('pvavail', y)}/StorageFactor,{ref('gen', y)}*SolarFraction)", fmt=NUM, total="sum", start=1)
line("diesel", "Diesel energy delivered (kWh)", lambda y: f"={ref('gen', y)}-{ref('solar', y)}", fmt=NUM, total="sum", start=1)
line("sf", "Actual solar fraction", lambda y: f"=IFERROR({ref('solar', y)}/{ref('gen', y)},0)", fmt=PCT, start=1)
line("litres", "Diesel consumed (litres)", lambda y: f"={ref('diesel', y)}/DieselEff", fmt=NUM, total="sum", start=1)
blank()

sec("REVENUE")
for i in range(1, NSEG + 1):
    line(f"rev{i}", f"Billed revenue - segment {i}",
         lambda y, i=i: f"={ref(f'kwh{i}', y)}*Seg{i}_Tariff*LevTariff*{ref('tesc', y)}", total="sum", start=1)
    wc.cell(row=ROW[f"rev{i}"], column=1, value=f"=\"Billed revenue - \"&Seg{i}_Name")
line("billed", "Total billed energy revenue", lambda y: "=" + "+".join(ref(f"rev{i}", y) for i in range(1, NSEG + 1)), total="sum", bold=True)
line("collected", "Collected energy revenue", lambda y: f"={ref('billed', y)}*Collection", total="sum", start=1)
newconn_fee = lambda y: "=" + "+".join(f"MAX(0,{ref(f'conn{i}', y)}-{ref(f'conn{i}', y-1)})*Seg{i}_Fee" for i in range(1, NSEG + 1))
line("fees", "Connection fees", newconn_fee, total="sum", start=1)
line("rev", "TOTAL OPERATING REVENUE", lambda y: f"={ref('collected', y)}+{ref('fees', y)}", total="npv", bold=True)
blank()

sec("OPERATING COSTS")
line("o_fixed", "Fixed O&M", lambda y: f"=HardCapex*O_FixedPct*{ref('cesc', y)}*LevOpex*{ref('op', y)}", total="sum", start=1)
line("o_staff", "Staff & operator", lambda y: f"=O_Staff*{ref('cesc', y)}*LevOpex*{ref('op', y)}", total="sum", start=1)
line("o_ins", "Insurance", lambda y: f"=TotalCapex*O_InsPct*{ref('cesc', y)}*LevOpex*{ref('op', y)}", total="sum", start=1)
line("o_fees", "Licences & fees", lambda y: f"=O_Fees*{ref('cesc', y)}*LevOpex*{ref('op', y)}", total="sum", start=1)
line("o_cust", "Customer management", lambda y: f"={ref('rev', y)}*O_CustPct", total="sum", start=1)
line("o_dom", "Diesel O&M", lambda y: f"={ref('diesel', y)}*O_DieselOM*{ref('cesc', y)}*LevOpex", total="sum", start=1)
line("o_fuel", "Diesel fuel", lambda y: f"={ref('litres', y)}*FuelPrice*{ref('cesc', y)}", total="sum", start=1)
line("opex", "TOTAL OPEX", lambda y: "=" + "+".join(ref(k, y) for k in ["o_fixed", "o_staff", "o_ins", "o_fees", "o_cust", "o_dom", "o_fuel"]), total="npv", bold=True)
blank()
line("ebitda", "EBITDA", lambda y: f"={ref('rev', y)}-{ref('opex', y)}", total="sum", bold=True)
line("margin", "EBITDA margin", lambda y: f"=IFERROR({ref('ebitda', y)}/{ref('rev', y)},0)", fmt=PCT, start=1)
blank()

sec("CAPEX, REPLACEMENTS & SUBSIDIES")
line("capex", "Initial CAPEX", lambda y: "=TotalCapex" if y == 0 else "=0", total="npv")
line("repl", "Battery replacement",
     lambda y: f"=IF(AND({ref('op', y)}=1,MOD({ref('year', y)},BattLife)=0,{ref('year', y)}<Life),BESS_kWh*C_BattRepl*LevCapex*{ref('cesc', y)},0)",
     total="npv", start=1)
line("grant", "Upfront CAPEX grant received", lambda y: "=Grant" if y == 0 else "=0", total="sum")
line("rbf", "RBF received",
     lambda y: f"=IF({ref('year', y)}-RBFLag>=1,INDEX($C${{n}}:${L(MAX_YEARS)}${{n}},1,{ref('year', y)}-RBFLag+1)*RBFperConn,0)", total="sum", start=1)
line("mra", "Maintenance reserve contribution (battery)",
     lambda y: (f"=IF(AND(UseMRA=1,{ref('op', y)}=1,CEILING({ref('year', y)},BattLife)<Life),"
                f"BESS_kWh*C_BattRepl*LevCapex*(1+CostEsc)^(CEILING({ref('year', y)},BattLife)-1)/BattLife,0)"),
     total="sum", start=1)
line("repl_cfads", "Replacement cost charged to CFADS",
     lambda y: f"=IF(UseMRA=1,{ref('mra', y)},{ref('repl', y)})", total="sum", start=1)
# patch RBF formula row reference to new connections row
for y in range(1, MAX_YEARS + 1):
    c = wc.cell(row=ROW["rbf"], column=FC + y)
    c.value = c.value.replace("{n}", str(ROW["newconn"]))
blank()

sec("TAX (simplified)")
line("dep", "Tax depreciation", lambda y: f"=IF({ref('year', y)}<=MIN(DepYears,Life),TotalCapex/DepYears,0)", total="sum", start=1)
line("tax_u", "Tax - unlevered (for project IRR)", lambda y: f"=MAX(0,{ref('ebitda', y)}-{ref('dep', y)})*TaxRate", total="sum", start=1)
line("tax_l", "Tax - levered (after interest)", lambda y: f"=MAX(0,{ref('ebitda', y)}-{ref('dep', y)}-{{int}})*TaxRate", total="sum", start=1)
blank()

sec("SENIOR DEBT")
line("d_open", "Opening balance", lambda y: f"={{close_prev}}", start=1)
line("d_draw", "Drawdown", lambda y: "=DebtAmt" if y == 0 else "=0", total="sum")
line("d_int", "Interest", lambda y: f"={ref('d_open', y)}*IntRate", total="sum", start=1)
line("d_prin", "Principal repayment",
     lambda y: f"=IF(AND({ref('year', y)}>Grace,{ref('year', y)}<=Tenor),PMT(IntRate,Tenor-Grace,-DebtAmt)-{ref('d_int', y)},0)", total="sum", start=1)
line("d_close", "Closing balance", lambda y: f"={{open}}+{ref('d_draw', y)}-{{prin}}")
line("d_serv", "Total debt service", lambda y: f"={ref('d_int', y)}+{ref('d_prin', y)}", total="sum", bold=True, start=1)
# patch debt cross references
for y in range(0, MAX_YEARS + 1):
    col = FC + y
    if y >= 1:
        wc.cell(row=ROW["d_open"], column=col).value = f"={ref('d_close', y-1)}"
        wc.cell(row=ROW["tax_l"], column=col).value = wc.cell(row=ROW["tax_l"], column=col).value.replace("{int}", ref("d_int", y))
    opening = ref("d_open", y) if y >= 1 else "0"
    prin = ref("d_prin", y) if y >= 1 else "0"
    wc.cell(row=ROW["d_close"], column=col).value = f"={opening}+{ref('d_draw', y)}-{prin}"
blank()

sec("CASH FLOWS")
line("cf_pre", "Project cash flow BEFORE subsidies (unlevered)",
     lambda y: f"=-{ref('capex', y)}+{ref('ebitda', y)}-{ref('repl', y)}-{ref('tax_u', y)}", total="npv", bold=True)
line("cf_post", "Project cash flow AFTER subsidies (unlevered)",
     lambda y: f"={ref('cf_pre', y)}+{ref('grant', y)}+{ref('rbf', y)}", total="npv", bold=True)
line("cfads", "CFADS (cash flow available for debt service)",
     lambda y: f"={ref('ebitda', y)}-{ref('repl_cfads', y)}-{ref('tax_l', y)}+{ref('rbf', y)}", total="sum", start=1)
line("dscr", "DSCR", lambda y: f"=IF({ref('d_serv', y)}>0,{ref('cfads', y)}/{ref('d_serv', y)},\"\")", fmt=MULT, start=1)
line("cf_eq", "Equity cash flow",
     lambda y: (f"=-EquityAmt" if y == 0 else f"={ref('cfads', y)}-{ref('d_serv', y)}"), total="npv", bold=True)
line("cf_eq_cum", "Cumulative equity cash flow",
     lambda y: (f"={ref('cf_eq', 0)}" if y == 0 else f"={ref('cf_eq_cum', y-1)}+{ref('cf_eq', y)}"))
blank()

sec("LCOE INPUTS (discounted at hurdle rate)")
line("lc_cost", "Lifecycle cost (CAPEX + OPEX + replacements)",
     lambda y: f"={ref('capex', y)}+{ref('opex', y)}-{ref('o_cust', y)}+{ref('repl', y)}", total="npv")
line("lc_kwh", "Energy sold (kWh)", lambda y: f"={ref('kwh', y)}", fmt=NUM, total="npv")
line("lc_rev", "Collected energy revenue", lambda y: f"={ref('collected', y)}", total="npv")

# conditional format DSCR below minimum
dscr_rng = f"{L(1)}{ROW['dscr']}:{L(MAX_YEARS)}{ROW['dscr']}"
wc.conditional_formatting.add(dscr_rng, FormulaRule(formula=[f'AND(ISNUMBER({L(1)}{ROW["dscr"]}),{L(1)}{ROW["dscr"]}<MinDSCR)'],
                                                     font=Font(name=FONT, color="C00000", bold=True),
                                                     fill=PatternFill("solid", fgColor="F8D7DA")))
neg_rng = f"{L(0)}{ROW['cf_eq_cum']}:{L(MAX_YEARS)}{ROW['cf_eq_cum']}"
wc.conditional_formatting.add(neg_rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))
wc.freeze_panes = "C4"


def rng(name):
    return f"'Cash Flow'!${L(0)}${ROW[name]}:${L(MAX_YEARS)}${ROW[name]}"


def rng1(name):
    return f"'Cash Flow'!${L(1)}${ROW[name]}:${L(MAX_YEARS)}${ROW[name]}"


def tot(name):
    return f"'Cash Flow'!$B${ROW[name]}"


# ---------------------------------------------------------------- Dashboard
wd = wb.create_sheet("Dashboard", 1)
wd.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJ", [3, 40, 16, 3, 40, 28, 3, 3, 3, 3]):
    wd.column_dimensions[col].width = w
wd["B1"] = "=\"DASHBOARD - \"&ProjectName"
wd["B1"].font = f_title
wd["B2"] = "=Site&\"  |  \"&TotalCustomers&\" customers  |  \"&TEXT(PV_kWp,\"#,##0\")&\" kWp PV / \"&TEXT(BESS_kWh,\"#,##0\")&\" kWh battery / \"&TEXT(Diesel_kW,\"#,##0\")&\" kW diesel  |  amounts in \"&Currency"
wd["B2"].font = f_h2


def kpi_block(col, start_row, title, items):
    lc = get_column_letter(col)
    vc = get_column_letter(col + 1)
    header = wd[f"{lc}{start_row}"]
    header.value = title
    header.font = f_h1
    wd[f"{lc}{start_row}"].fill = fill_h1
    wd[f"{vc}{start_row}"].fill = fill_h1
    rr = start_row + 1
    for label, formula, fmt, nm in items:
        a = wd[f"{lc}{rr}"]
        a.value = label
        a.font = f_txt
        a.border = box
        b = wd[f"{vc}{rr}"]
        b.value = formula
        b.number_format = fmt
        b.font = f_bold
        b.border = box
        b.fill = fill_kpi
        b.alignment = Alignment(horizontal="right")
        if nm:
            define(nm, wd, f"{vc}{rr}")
        rr += 1
    return rr


cfpre, cfpost, cfeq = rng("cf_pre"), rng("cf_post"), rng("cf_eq")
end = kpi_block(2, 4, "PROJECT ECONOMICS", [
    ("Total CAPEX", "=TotalCapex", USD, None),
    ("CAPEX per connection", "=CapexPerConn", USD, None),
    ("Project IRR - before subsidies", f'=IFERROR(IRR({cfpre}),"n/a")', PCT, "IRR_Pre"),
    ("Project IRR - after subsidies", f'=IFERROR(IRR({cfpost}),"n/a")', PCT, "IRR_Post"),
    ("Project NPV @ hurdle - before subsidies", f"={tot('cf_pre')}", USD, "NPV_Pre"),
    ("Project NPV @ hurdle - after subsidies", f"={tot('cf_post')}", USD, "NPV_Post"),
    ("LCOE (cost per kWh sold)", f"=IFERROR({tot('lc_cost')}/{tot('lc_kwh')},0)", USD3, "LCOE"),
    ("Average tariff billed (weighted, Year 1)", f"=IFERROR('Cash Flow'!{L(1)}{ROW['billed']}/'Cash Flow'!{L(1)}{ROW['kwh']},0)", USD3, "AvgTariff"),
    ("Levelized collected revenue per kWh", f"=IFERROR({tot('lc_rev')}/{tot('lc_kwh')},0)", USD3, "LevRev"),
    ("Year-1 EBITDA margin", f"='Cash Flow'!{L(1)}{ROW['margin']}", PCT, None),
    ("Average solar fraction (life)", f"=IFERROR({tot('solar')}/{tot('gen')},0)", PCT, None),
])
end2 = kpi_block(5, 4, "FINANCING & BANKABILITY", [
    ("Debt amount", "=DebtAmt", USD, None),
    ("Equity amount", "=EquityAmt", USD, None),
    ("Total subsidies (grant + RBF, nominal)", f"=Grant+{tot('rbf')}", USD, "TotalSubsidy"),
    ("Subsidy share of CAPEX", "=IFERROR(TotalSubsidy/TotalCapex,0)", PCT, None),
    ("Minimum DSCR", f'=IF(DebtAmt>0,MIN({rng1("dscr")}),"no debt")', MULT, "DSCR_Min"),
    ("Average DSCR", f'=IF(DebtAmt>0,AVERAGE({rng1("dscr")}),"no debt")', MULT, None),
    ("Equity IRR", f'=IFERROR(IRR({cfeq}),"n/a")', PCT, "IRR_Eq"),
    ("Equity payback (years)", f'=IFERROR(MATCH(TRUE,INDEX({rng("cf_eq_cum")}>=0,0),0)-1,"> life")', "0", None),
    ("Lifetime CO2 avoided vs diesel-only (tCO2)", f"={tot('solar')}/DieselEff*CO2perL/1000", NUM, None),
    ("Total energy sold over life (MWh)", f"={tot('kwh')}/1000", NUM, None),
])
# CO2 factor input (placed on Inputs at bottom)
r += 1
section("8. IMPACT")
inp("CO2 emission factor of diesel", "CO2perL", 2.68, "0.00", "kgCO2/litre",
    "Commonly used combustion factor for diesel; verify against the standard your funder uses.")

# Viability gap block
vg_row = max(end, end2) + 1
wd[f"B{vg_row}"] = "VIABILITY GAP & BANKABILITY VERDICT"
wd[f"B{vg_row}"].font = f_h1
for col in "BCDEF":
    wd[f"{col}{vg_row}"].fill = fill_h1
vg_items = [
    ("Viability gap: upfront subsidy needed for NPV = 0 at hurdle", "=MAX(0,-NPV_Pre)", USD, "VGF"),
    ("Viability gap per connection", "=IFERROR(VGF/TotalCustomers,0)", USD, "VGFperConn"),
    ("Subsidies already planned (PV of grant + RBF at hurdle)", f"=NPV_Post-NPV_Pre", USD, "SubsidyPV"),
    ("Remaining funding gap after planned subsidies", "=MAX(0,VGF-SubsidyPV)", USD, "RemainingGap"),
]
rr = vg_row + 1
for label, f, fmt, nm in vg_items:
    wd[f"B{rr}"] = label
    wd[f"B{rr}"].font = f_txt
    wd[f"B{rr}"].border = box
    c = wd[f"C{rr}"]
    c.value = f
    c.number_format = fmt
    c.font = f_bold
    c.fill = fill_kpi
    c.border = box
    define(nm, wd, f"C{rr}")
    rr += 1
checks_v = [
    ("Commercially viable without subsidy?", '=IF(NPV_Pre>=0,"YES","NO - needs subsidy")'),
    ("Viable at hurdle after planned subsidies?", '=IF(NPV_Post>=0,"YES","NO - gap remains")'),
    ("Debt covenant met (min DSCR >= required)?", '=IF(DebtAmt=0,"No debt",IF(DSCR_Min>=MinDSCR,"YES","NO - resize debt"))'),
    ("Equity IRR meets target?", '=IF(ISNUMBER(IRR_Eq),IF(IRR_Eq>=TargetEqIRR,"YES","NO"),"n/a")'),
    ("Collected tariff covers LCOE?", '=IF(LevRev>=LCOE,"YES","NO - tariff below cost")'),
]
rr_v = vg_row + 1
for label, f in checks_v:
    wd[f"E{rr_v}"] = label
    wd[f"E{rr_v}"].font = f_txt
    wd[f"E{rr_v}"].border = box
    c = wd[f"F{rr_v}"]
    c.value = f
    c.font = f_bold
    c.border = box
    c.alignment = Alignment(horizontal="center")
    rr_v += 1
verdict_rng = f"F{vg_row+1}:F{rr_v-1}"
wd.conditional_formatting.add(verdict_rng, FormulaRule(formula=[f'LEFT(F{vg_row+1},LEN("YES"))="YES"'], fill=PatternFill("solid", fgColor="D4EDDA"), font=Font(name=FONT, bold=True, color="155724")))
wd.conditional_formatting.add(verdict_rng, FormulaRule(formula=[f'OR(LEFT(F{vg_row+1},LEN("NO - "))="NO - ",F{vg_row+1}="NO")'], fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(name=FONT, bold=True, color="721C24")))
note_row = max(rr, rr_v) + 1
wd[f"B{note_row}"] = "Viability gap uses the pre-subsidy project cash flow discounted at the hurdle rate. It is the minimum upfront subsidy in present-value terms; RBF paid later must be larger in nominal terms."
wd[f"B{note_row}"].font = f_note

# chart data (hidden helper on dashboard, linked to cash flow)
cd = note_row + 3
wd[f"B{cd}"] = "Chart data (linked)"
wd[f"B{cd}"].font = f_h2
hdrs = ["Year", "Revenue", "OPEX", "CFADS", "Debt service", "Solar kWh", "Diesel kWh"]
for j, h in enumerate(hdrs):
    wd.cell(row=cd + 1, column=12 + j, value=h).font = f_h2
for y in range(1, MAX_YEARS + 1):
    rrr = cd + 1 + y
    wd.cell(row=rrr, column=12, value=f"='Cash Flow'!{L(y)}{ROW['cal']}").font = f_link
    for j, key in enumerate(["rev", "opex", "cfads", "d_serv", "solar", "diesel"], start=1):
        c = wd.cell(row=rrr, column=12 + j, value=f"='Cash Flow'!{L(y)}{ROW[key]}")
        c.font = f_link
        c.number_format = NUM
for j in range(12, 19):
    wd.column_dimensions[get_column_letter(j)].width = 12
wd[f"B{cd}"].value = "Chart data is linked from 'Cash Flow' (columns L-R)."
wd[f"B{cd}"].font = f_note

ch = BarChart()
ch.type = "col"
ch.title = T("Revenue vs OPEX vs CFADS")
ch.y_axis.title = T("Currency")
data = Reference(wd, min_col=13, max_col=15, min_row=cd + 1, max_row=cd + 1 + MAX_YEARS)
cats = Reference(wd, min_col=12, min_row=cd + 2, max_row=cd + 1 + MAX_YEARS)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ln = LineChart()
ln.add_data(Reference(wd, min_col=16, max_col=16, min_row=cd + 1, max_row=cd + 1 + MAX_YEARS), titles_from_data=True)
ch += ln
ch.height = 8
ch.width = 20
wd.add_chart(ch, f"B{note_row+5}")

ch2 = BarChart()
ch2.type = "col"
ch2.grouping = "stacked"
ch2.overlap = 100
ch2.title = T("Generation mix (kWh)")
ch2.add_data(Reference(wd, min_col=17, max_col=18, min_row=cd + 1, max_row=cd + 1 + MAX_YEARS), titles_from_data=True)
ch2.set_categories(cats)
ch2.height = 8
ch2.width = 20
wd.add_chart(ch2, f"B{note_row+22}")

# ---------------------------------------------------------------- Checks
wk = wb.create_sheet("Checks")
wk.sheet_view.showGridLines = False
wk.column_dimensions["A"].width = 70
wk.column_dimensions["B"].width = 14
wk["A1"] = "INTEGRITY CHECKS"
wk["A1"].font = f_title
checks = [
    ("Sources = uses at Year 0 (grant + debt + equity = CAPEX)", "=ABS(Grant+DebtAmt+EquityAmt-MAX(TotalCapex,Grant))<1"),
    ("Debt fully repaid by end of tenor", f"=ABS(INDEX({rng('d_close')},1,MIN(Tenor,{MAX_YEARS})+1))<1"),
    ("Project life within modelled horizon", f"=AND(Life>=1,Life<={MAX_YEARS})"),
    ("Tenor within project life", "=Tenor<=Life"),
    ("Grace shorter than tenor", "=Grace<Tenor"),
    ("Maintenance reserve contributions = replacements (when reserve used)", f"=OR(UseMRA=0,ABS({tot('mra')}-SUM({rng('repl')}))<1)"),
    ("Connection ramp percentages between 0 and 100%", "=AND(Ramp1>=0,Ramp1<=1,Ramp2>=0,Ramp2<=1,Ramp3>=0,Ramp3<=1)"),
    ("Energy balance: solar + diesel = generation", f"=ABS({tot('solar')}+{tot('diesel')}-{tot('gen')})<1"),
    ("At least one customer entered", "=TotalCustomers>0"),
]
for i, (lab, f) in enumerate(checks, start=3):
    wk[f"A{i}"] = lab
    wk[f"A{i}"].font = f_txt
    wk[f"A{i}"].border = box
    c = wk[f"B{i}"]
    c.value = f'=IF({f[1:]},"OK","CHECK")'
    c.font = f_bold
    c.border = box
    c.alignment = Alignment(horizontal="center")
last = 2 + len(checks)
wk[f"A{last+2}"] = "OVERALL"
wk[f"A{last+2}"].font = f_bold
wk[f"B{last+2}"] = f'=IF(COUNTIF(B3:B{last},"OK")={len(checks)},"ALL OK","REVIEW")'
wk[f"B{last+2}"].font = f_bold
define("ChecksOverall", wk, f"B{last+2}")
wk.conditional_formatting.add(f"B3:B{last+2}", FormulaRule(formula=['LEFT(B3,2)="OK"'], fill=PatternFill("solid", fgColor="D4EDDA")))
wk.conditional_formatting.add(f"B3:B{last+2}", FormulaRule(formula=['B3="ALL OK"'], fill=PatternFill("solid", fgColor="D4EDDA")))
wk.conditional_formatting.add(f"B3:B{last+2}", FormulaRule(formula=['OR(B3="CHECK",B3="REVIEW")'], fill=PatternFill("solid", fgColor="F8D7DA")))

wd["F2"] = '="Model checks: "&ChecksOverall'
wd["F2"].font = f_h2

# order sheets & print setup
wb._sheets = [ws0, wi, wd, wc, wk]
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
ws0.sheet_properties.tabColor = NAVY
wi.sheet_properties.tabColor = "FFC000"
wd.sheet_properties.tabColor = TEAL
wc.sheet_properties.tabColor = "7F7F7F"
wk.sheet_properties.tabColor = "7F7F7F"
wb.active = 0
branding.apply(wb, T)
print_areas.apply(wb)
wb.save(OUT)
print("saved", OUT)
i18n.report()
