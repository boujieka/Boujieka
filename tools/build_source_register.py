"""Build the Book 2 source register workbook from tools/source_register_data.py.

Run: python tools/build_source_register.py [out.xlsx]
Default output: output/04_SOURCE_REGISTER_v1.0.xlsx
Sheets: Read_Me (method), Register (every claim, full field list), Book_Claims (book statements and their status),
Conflicts, Legacy_Map (v0.7 register IDs to the new references).
"""
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
import source_register_data as R  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "output/04_SOURCE_REGISTER_v1.0.xlsx"
GREEN, GOLD = "0B3020", "B07C0F"
FILL = {R.VERIFIED: "E2EFDA", R.HISTORICAL: "EDEDED", R.PENDING: "FFF2CC", R.UNVERIFIED: "FCE4D6", R.CONFLICT: "F8CBAD", R.NOT_USED: "F2F2F2"}
COLS = [("Ref", "ref", 8), ("Claim", "claim", 52), ("Value", "value", 14), ("Unit", "unit", 12), ("Period", "period", 16), ("Source", "source", 34),
        ("Publisher", "publisher", 24), ("Publication date", "date", 14), ("Document", "document", 30), ("Page", "page", 14), ("Section", "section", 24),
        ("URL / source location", "url", 34), ("Evidence type", "evidence", 30), ("Grade", "grade", 7), ("Status", "status", 22), ("Book usage", "book", 28),
        ("Model usage", "model", 28), ("Notes", "notes", 50), ("Model key", "key", 26), ("v0.7 ID", "legacy", 8)]


def head(ws, names, widths):
    ws.append(names)
    for i, (c, w) in enumerate(zip(ws[ws.max_row], widths), 1):
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFF", size=9), PatternFill("solid", fgColor=GREEN)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = w


def body(ws, row, status=None):
    ws.append(row)
    for c in ws[ws.max_row]:
        c.font = Font(name="Arial", size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if status in FILL:
            c.fill = PatternFill("solid", fgColor=FILL[status])


def main():
    wb = Workbook()
    rd = wb.active
    rd.title = "Read_Me"
    rd.column_dimensions["A"].width, rd.column_dimensions["B"].width = 30, 110
    lines = [
        ("PAYGo Solar Finance (Book 2) and MODEL 2: source register", ""),
        ("Africa Energy Finance | Business & Financial Models", ""),
        ("", ""),
        ("Purpose", "Every external quantitative or factual claim used in the book or the model, with its source, the evidence actually read, and its status."),
        ("Grade (the source)", "A primary official, regulatory or audited source. B authoritative institutional source, including a company's own unaudited disclosure. "
                               "C reputable secondary source (trade press, news). D unverified or informal."),
        ("Status (what this project has verified)", "VERIFIED: the claim was read in the cited document at the stated page. VERIFIED (HISTORICAL): read, but the document is superseded. "
                                                    "PENDING PRIMARY DOCUMENT: recorded from a secondary account or a client extraction; the primary page has not been read. "
                                                    "UNVERIFIED: no reliable account. CONFLICTING SOURCES: accounts disagree and the conflict is unresolved. NOT USED: held or recorded but not relied on."),
        ("Rule for the model", "Only VERIFIED and VERIFIED (HISTORICAL) claims may feed a model calculation or diagnostic. Other claims may be displayed as context with their status."),
        ("Rule for the book", "Claims that are not VERIFIED are written with their caveat (for example 'reported', 'to be checked against the primary document'). "
                              "CONFLICTING claims are presented as a conflict, never as a single number."),
        ("Pages", "PDF page numbers of the file held by the project, with the printed page in brackets where they differ. 'PAGE TO VERIFY' where the page has not been seen. No page is inferred."),
        ("Grade and status together", "A claim can cite a grade A source and still be PENDING: the grade describes the document, the status describes whether this project has read the claim in it."),
        ("House style", "The brief's status 'VERIFIED, HISTORICAL' is written 'VERIFIED (HISTORICAL)' to keep the publications free of dashes."),
        ("Prepared", "3 October 2026. Version 1.0 of the register format; the claims remain open to verification as listed."),
    ]
    for a, b in lines:
        rd.append([a, b])
    rd["A1"].font = Font(name="Arial", bold=True, size=14, color=GREEN)
    rd["A2"].font = Font(name="Arial", size=10, color=GOLD)
    for row in rd.iter_rows(min_row=4):
        row[0].font = Font(name="Arial", bold=True, size=10)
        row[1].font = Font(name="Arial", size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    cnt = Counter(d["status"] for d in R.REGISTER)
    rd.append([])
    rd.append(["Claims by status", ""])
    rd[f"A{rd.max_row}"].font = Font(name="Arial", bold=True, size=11, color=GREEN)
    for s in (R.VERIFIED, R.HISTORICAL, R.PENDING, R.UNVERIFIED, R.CONFLICT, R.NOT_USED):
        rd.append([s, cnt.get(s, 0)])

    ws = wb.create_sheet("Register")
    head(ws, [c[0] for c in COLS], [c[2] for c in COLS])
    for d in R.REGISTER:
        body(ws, [d[c[1]] for c in COLS], d["status"])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"

    bc = wb.create_sheet("Book_Claims")
    head(bc, ["Where in the book (v0.1)", "Statement", "Register reference", "Status", "Action for v0.2"], [18, 60, 16, 24, 70])
    for row in R.BOOK_CLAIM_CHECKS:
        body(bc, list(row), row[3])
    bc.freeze_panes = "A2"

    cf = wb.create_sheet("Conflicts")
    head(cf, ["Ref", "Claim", "Value", "Unit", "Source", "Notes"], [8, 50, 14, 10, 30, 80])
    for d in R.REGISTER:
        if d["status"] == R.CONFLICT or "CONFLICT" in d["notes"]:
            body(cf, [d["ref"], d["claim"], d["value"], d["unit"], f'{d["publisher"]}: {d["source"]}', d["notes"]], R.CONFLICT)

    lm = wb.create_sheet("Legacy_Map")
    head(lm, ["v0.7 Source_Register ID", "New reference", "Status now"], [22, 16, 26])
    for d in sorted((d for d in R.REGISTER if d["legacy"]), key=lambda x: x["legacy"]):
        body(lm, [d["legacy"], d["ref"], d["status"]], d["status"])

    for s in wb.worksheets:
        s.sheet_view.showGridLines = s.title != "Read_Me"
    wb.properties.creator = "Emmanuel Boujieka Kamga"
    wb.properties.title = "PAYGo Solar Finance: source register"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT, len(R.REGISTER), "claims", dict(cnt))


if __name__ == "__main__":
    main()
