"""Audit of the quantitative claims in the Volume 2 book against the evaluated model and case workbooks.

Run: python tools/audit_book_claims.py <model_values.pkl> <case_values.pkl> <out.xlsx> [key_claims.json]
The value stores map "'[TAG]SHEET'!CELL" to evaluated values (the formula engine run used by tools/qa_evaluate.py).

For every sentence of the book that contains a figure, the script records the chapter, section, sentence and each
figure, classifies the sentence (external source, SolaraPay case, model default, illustrative arithmetic, method)
and looks for cells whose value, rounded to the precision printed in the book, equals the figure. A match is evidence
that a figure is traceable, not proof that the sentence is right: matches on small round numbers are frequent and are
marked as weak.
"""
import bisect
import pickle
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "volumes/02-solar-home-systems/book"
FILES = ["ch00_front.md"] + [f"ch{i:02d}.md" for i in range(1, 17)] + ["ch99_annexes.md"]

EXTERNAL = {
    "E1 ESMAP / World Bank, Off-Grid Solar Market Trends Report 2024": r"ESMAP|Market Trends Report|World Bank",
    "E2 PAYGo PERFORM (2026 current; 2021 historical)": r"PERFORM|CGAP|Lighting Global",
    "E3 Multi Tier Framework (ESMAP 2015)": r"Multi Tier Framework|\bMTF\b|Beyond Connections",
    "E4 M-KOPA FY2024": r"M-KOPA|M KOPA|MKOPA",
    "E5 Sun King": r"Sun King",
    "E6 d.light": r"d\.light",
    "E7 BBOXX": r"BBOXX|Bboxx",
    "E8 GOGLA investment data": r"GOGLA",
    "E9 ZOLA Electric": r"ZOLA",
    "E10 SEforALL Universal Energy Facility": r"Universal Energy Facility|SEforALL|UEF\b",
    "Other company or institution": r"Pawame|TechCabal|Kenyan Wall Street|Citi\b|Citigroup|IFC\b|AFC\b|Proparco|Symbiotics|Engie|Fenix|Zola",
}
CASE_RX = r"SolaraPay|Kivara|\bKVS\b|the case\b|case workbook|history"
MODEL_RX = r"default|Base\b|Downside|Severe|workbook|the model|\bmodel's|sheet|Investment_Summary|Unit_Economics|Sensitivity|Valuation|KPIs|Annual|Covenants|Products|Inputs|Consumer_Risk|Calibration"
ILLUS_RX = r"illustrat|for example|example|suppose|assume|imagine|hypothetical|say,|consider"

NUM = re.compile(r"(?<![\w.])(\d{1,3}(?:,\d{3})+|\d+)(?:\.(\d+))?\s*(%|x\b| times|m\b|bn\b| million| billion)?")
YEAR = re.compile(r"^(19|20)\d{2}$")

PRIORITY = ["INVESTMENT_SUMMARY", "KPIS", "ANNUAL", "VALUATION", "UNIT_ECONOMICS", "SENSITIVITY", "COVENANTS",
            "CREDIT_PORTFOLIO", "VINTAGE_DASHBOARD", "INVESTMENT_READINESS", "PRODUCTS", "CONSUMER_RISK", "INPUTS",
            "CREDIT_ASSUMPTIONS", "SCENARIOS", "CALIBRATION", "MARKET_BENCHMARK", "SOURCE_REGISTER", "FS", "FINANCING",
            "COSTS", "RBF_ENGINE", "CREDIT_ENGINE", "VINTAGE_ENGINE", "OPS", "CURVES", "TIMELINE"]


def load_index(pkl):
    d = pickle.load(open(pkl, "rb"))
    vals = []
    for k, v in d.items():
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            continue
        if v != v or abs(v) > 1e15:
            continue
        m = re.match(r"'\[(.+?)\](.+?)'!([A-Z]+\d+)", k)
        if not m:
            continue
        sheet = m.group(2)
        rank = PRIORITY.index(sheet) if sheet in PRIORITY else len(PRIORITY)
        vals.append((float(v), rank, sheet, m.group(3)))
    vals.sort()
    return vals, [v[0] for v in vals]


def candidates(index, keys, x, decimals, unit):
    """Cells whose value rounds to x at the printed precision, under the plausible scalings of the unit."""
    half = 0.5 * 10 ** (-decimals)
    scales = {"%": [0.01], "m": [1e6, 1.0], "bn": [1e9, 1.0], " million": [1e6, 1.0], " billion": [1e9, 1.0],
              "x": [1.0], " times": [1.0]}.get(unit, [1.0, 1e6, 1e9])
    out = []
    for sc in scales:
        lo, hi = (x - half) * sc, (x + half) * sc
        for sign in (1, -1):
            a, b = (lo, hi) if sign == 1 else (-hi, -lo)
            i, j = bisect.bisect_left(keys, a), bisect.bisect_right(keys, b)
            out += index[i:j][:400]
    out.sort(key=lambda t: t[1])
    return out


def sentences(md):
    md = re.sub(r"\|", " | ", md)
    return [s.strip() for s in re.split(r"(?<=[.!?:])\s+(?=[A-Z(*|])|\n+", md) if s.strip()]


def main(model_pkl, case_pkl, out):
    midx, mkeys = load_index(model_pkl)
    cidx, ckeys = load_index(case_pkl)
    rows = []
    for fname in FILES:
        text = (BOOK / fname).read_text()
        chapter, section = "", ""
        for line in text.splitlines():
            if line.startswith("# "):
                chapter, section = line[2:].strip(), ""
                continue
            if line.startswith("## "):
                section = line[3:].strip()
                continue
            if not line.strip() or line.startswith("![") or set(line.strip()) <= set("|-: "):
                continue
            for s in sentences(line):
                nums = []
                for m in NUM.finditer(s):
                    whole, frac, unit = m.group(1), m.group(2) or "", (m.group(3) or "")
                    if YEAR.match(whole) and not frac and not unit:
                        continue  # calendar years
                    if re.match(r"(Chapter|Tier|Annex|section|Step|Module|T|D|E|SR|Gate)\s*$", s[max(0, m.start() - 9):m.start()]):
                        continue  # references, not figures
                    x = float(whole.replace(",", "") + ("." + frac if frac else ""))
                    nums.append((m.group(0).strip(), x, len(frac), unit.strip() if unit else ""))
                if not nums:
                    continue
                ext = [k for k, rx in EXTERNAL.items() if re.search(rx, s)]
                if ext:
                    cls = "EXTERNAL"
                elif re.search(CASE_RX, s):
                    cls = "SOLARAPAY CASE"
                elif re.search(ILLUS_RX, s, re.I):
                    cls = "ILLUSTRATIVE"
                elif re.search(MODEL_RX, s):
                    cls = "MODEL DEFAULT"
                else:
                    cls = "METHOD / ARITHMETIC"
                for raw, x, dec, unit in nums:
                    weak = (dec == 0 and abs(x) < 100 and unit not in ("%",)) or (unit == "%" and dec == 0 and x in (0, 5, 10, 15, 20, 25, 30, 40, 50, 60, 70, 75, 80, 90, 100))
                    if cls == "SOLARAPAY CASE":
                        order = [("CASE", cidx, ckeys), ("MODEL", midx, mkeys)]
                    else:
                        order = [("MODEL", midx, mkeys), ("CASE", cidx, ckeys)]
                    hit, wb_used = [], ""
                    for tag, idx, keys in order:
                        hit = candidates(idx, keys, x, dec, unit if unit else "")
                        if hit:
                            wb_used = tag
                            break
                    front = bool(hit) and hit[0][1] < PRIORITY.index("FS")
                    if cls == "EXTERNAL":
                        status = "EXTERNAL: see source register"
                    elif hit and not weak and front:
                        status = "CANDIDATE: output or input cell"
                    elif hit and not weak:
                        status = "CANDIDATE: engine cell only"
                    elif hit and weak:
                        status = "WEAK MATCH (round number)"
                    elif cls in ("ILLUSTRATIVE", "METHOD / ARITHMETIC"):
                        status = "TEXT ARITHMETIC: check in place"
                    else:
                        status = "NOT FOUND: reconcile"
                    cells = "; ".join(f"{h[2]}!{h[3]}" for h in hit[:3])
                    rows.append([fname, chapter, section, cls, raw, x, wb_used, cells, "; ".join(ext), status, s[:600]])
    wb = Workbook()
    ws = wb.active
    ws.title = "Cross_Reference"
    head = ["ID", "File", "Chapter", "Section", "Claim class", "Figure as printed", "Value", "Workbook", "Candidate cells (first 3)",
            "External source", "Status", "Sentence"]
    ws.append(head)
    for i, r in enumerate(rows, 1):
        ws.append([f"BC{i:04d}"] + r)
    fill = PatternFill("solid", fgColor="0B3020")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [9, 14, 34, 34, 18, 14, 12, 9, 40, 34, 30, 110]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(rows) + 1}"
    if len(sys.argv) > 4:
        import json
        kc = json.load(open(sys.argv[4]))
        ks = wb.create_sheet("Key_Claims", 0)
        kh = ["Ref", "Book claim", "Where in book", "Class", "Workbook", "Sheet", "Cell / range", "Value recomputed (LibreOffice)",
              "Source / basis", "Status", "Note"]
        ks.append(kh)
        for r in kc:
            ks.append([r.get(k, "") for k in ("ref", "claim", "where", "cls", "wb", "sheet", "cell", "value", "source", "status", "note")])
        for c in ks[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B3020")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate([7, 52, 22, 16, 9, 18, 16, 18, 40, 26, 60], 1):
            ks.column_dimensions[get_column_letter(i)].width = w
        for row in ks.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ks.freeze_panes = "A2"
    summ = wb.create_sheet("Summary", 0)
    counts = defaultdict(int)
    for r in rows:
        counts[(r[3], r[9])] += 1
    summ.append(["Claim class", "Status", "Figures"])
    for (a, b), n in sorted(counts.items()):
        summ.append([a, b, n])
    summ.append([])
    summ.append(["Total figures", "", len(rows)])
    for c in summ[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
    summ.column_dimensions["A"].width, summ.column_dimensions["B"].width = 26, 36
    wb.save(out)
    print(len(rows), "figures")
    for (a, b), n in sorted(counts.items()):
        print(f"{a:22s} {b:34s} {n}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
