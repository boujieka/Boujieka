"""Génère le modèle tarifaire comparatif (Livrable 2) : modele_tarifaire_electricite.xlsx.

Aucune donnée réelle n'est saisie : les feuilles Cameroun et Congo ont leurs entrées vides
(à documenter depuis des sources officielles). La feuille Exemple_fictif contient des valeurs
inventées, uniquement pour vérifier que les formules fonctionnent.

Usage : python build_modele.py [chemin_sortie.xlsx]
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
F_IN = Font(name=FONT, color="0000FF")
F_CALC = Font(name=FONT, color="000000")
F_LINK = Font(name=FONT, color="008000")
F_BOLD = Font(name=FONT, bold=True)
F_TITLE = Font(name=FONT, bold=True, size=14)
F_SEC = Font(name=FONT, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, italic=True, color="666666", size=9)
FILL_IN = PatternFill("solid", fgColor="FFFF00")
FILL_SEC = PatternFill("solid", fgColor="1F3864")
FILL_HEAD = PatternFill("solid", fgColor="D9E1F2")
FILL_WARN = PatternFill("solid", fgColor="F8CBAD")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))

MDS = '#,##0.0;(#,##0.0);"-"'
GWH = '#,##0;(#,##0);"-"'
FCFA = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
NB = '#,##0;(#,##0);"-"'
IDX = '0.0;(0.0);"-"'

YEARS = ["Base"] + [str(y) for y in range(2027, 2036)]  # 10 colonnes : C..L
YC = [L(3 + i) for i in range(len(YEARS))]
N_STEPS = len(YEARS) - 1  # 9 pas de 2027 à 2035

CATS = [
    ("Ménages BT — tranche sociale", 1),
    ("Ménages BT — autres tranches", 0),
    ("Professionnels BT", 0),
    ("Moyenne tension (MT)", 0),
    ("Haute tension — industrie (HT)", 0),
]
IND_CAT = 4  # index de la catégorie industrielle HT

SCEN = ["S0 Statu quo", "S1 Tarifs aux coûts", "S2 Aides ciblées", "S3 Réforme intégrée"]

# Valeurs inventées pour la feuille Exemple_fictif (aucune n'est une donnée réelle).
EX_GEN = {
    "annee": "2025", "inj": 4000, "g": 0.05, "tech": 0.12, "com": 0.13, "reg": 0.15,
    "rec": 0.85, "cmp": 60, "couv": 0.5, "fiab": 0.7, "canal": 1, "benef": 400000,
    "admin": 3000, "incl": 0.25, "cibl": 0.7, "sub_racc": 150000, "ci": 50, "ce": 80,
    "part_ind": 0.15,
}
EX_CATS = [(0.15, 50, 110), (0.30, 80, 110), (0.15, 100, 105), (0.20, 95, 85), (0.20, 70, 60)]
EX_S3 = {"tech": 0.09, "com": 0.06, "rec": 0.95}
EX_RACC = [20000, 0, 50000, 80000]
EX_EQ = {"seuil": 100, "prog": 1,
         "q": [(600000, 0.20, 40), (600000, 0.40, 60), (600000, 0.60, 90),
               (600000, 0.80, 150), (600000, 0.95, 300)]}


def put(ws, ref, value, font=F_CALC, fmt=None, fill=None, comment=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "Modèle")
    return c


def section(ws, row, text, width=14):
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = FILL_SEC
    put(ws, f"A{row}", text, F_SEC)


def inp(ws, ref, value, fmt, example):
    """Cellule d'entrée : bleue ; jaune si elle reste à documenter."""
    if example:
        return put(ws, ref, value, F_IN, fmt)
    return put(ws, ref, None, F_IN, fmt, FILL_IN)


def build_country(wb, name, example):
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 14
    for col in YC + ["M", "N"]:
        ws.column_dimensions[col].width = 13
    put(ws, "A1", f"Modèle tarifaire électricité — {name}", F_TITLE)
    if example:
        put(ws, "A2", "VALEURS FICTIVES — test des formules uniquement, aucune donnée réelle.",
            Font(name=FONT, bold=True, color="C00000"), fill=FILL_WARN)
    else:
        put(ws, "A2", "Cellules jaunes à renseigner depuis des sources officielles "
            "(régulateur, opérateur, ministère des Finances). Citer la source en colonne D.", F_NOTE)
    put(ws, "A3", "Unités : énergie en GWh ; tarifs et coûts en FCFA/kWh ; montants en milliards "
        "de FCFA constants de l'année de base (termes réels).", F_NOTE)

    R = {}
    r = 5
    section(ws, r, "1. HYPOTHÈSES GÉNÉRALES (année de base)")
    r += 1
    for col, h in zip("ABCD", ["Paramètre", "Unité", "Valeur", "Source / note"]):
        put(ws, f"{col}{r}", h, F_BOLD, fill=FILL_HEAD)
    gen = [
        ("annee", "Année de base des données", "année", None,
         "Dernière année avec données vérifiées (période d'analyse 2020–2026)."),
        ("inj", "Énergie injectée dans le réseau", "GWh", GWH, None),
        ("g", "Croissance annuelle de l'énergie injectée", "%/an", PCT,
         "Hypothèse de demande ; pas d'élasticité-prix dans cette version."),
        ("tech", "Pertes techniques", "% injectée", PCT, None),
        ("com", "Pertes commerciales (non techniques)", "% injectée", PCT, None),
        ("reg", "Pertes totales admises par le régulateur", "% injectée", PCT,
         "Au-delà : coût d'inefficience, non imputé à la politique tarifaire."),
        ("rec", "Taux de recouvrement des factures", "% facturé", PCT, None),
        ("cmp", "Coût marginal de production / achat", "FCFA/kWh", FCFA,
         "Valorise les pertes excédentaires."),
        ("couv", "Part de l'écart de sous-tarification que l'État s'engage à compenser", "%", PCT,
         "Selon textes réglementaires / contrat de concession."),
        ("fiab", "Fiabilité de paiement des engagements budgétaires", "% versé", PCT,
         "Versé / dû ; s'applique aux compensations, aides sociales et raccordements."),
        ("canal", "Canal de l'aide sociale (1 = tarif social compensé via la facture ; "
         "0 = transfert direct aux ménages)", "1/0", "0", None),
        ("benef", "Ménages bénéficiaires de l'aide ciblée", "nombre", NB, None),
        ("admin", "Coût administratif du ciblage par bénéficiaire", "FCFA/an", NB, None),
        ("incl", "Erreur d'inclusion (part de l'aide versée à des non-éligibles)", "%", PCT, None),
        ("cibl", "Part de la consommation de la tranche sociale couverte par l'aide", "%", PCT,
         "Les ménages non couverts (erreur d'exclusion) subissent la hausse."),
        ("sub_racc", "Subvention par nouveau raccordement", "FCFA", NB, None),
        ("ci", "CI — coûts indirects liés à la dégradation du service (base)", "Mds FCFA", MDS,
         "Estimation à documenter ; ne pas inclure ce qui figure déjà dans CE."),
        ("ce", "CE — surcoûts des entreprises (autoproduction, délestages) (base)", "Mds FCFA",
         MDS, "Estimation à documenter (enquêtes entreprises)."),
        ("part_ind", "Part de l'électricité dans les coûts de production industriels", "%", PCT,
         None),
    ]
    for key, label, unit, fmt, note in gen:
        r += 1
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", unit, F_NOTE)
        inp(ws, f"C{r}", EX_GEN[key], fmt, example)
        if note:
            put(ws, f"D{r}", note, F_NOTE)
        R[key] = f"$C${r}"
    ws[R["annee"].replace("$", "")].number_format = "@"

    r += 2
    section(ws, r, "2. CATÉGORIES TARIFAIRES (année de base)")
    r += 1
    heads = ["Catégorie", "Sociale (1/0)", "Part des ventes", "Tarif moyen", "Coût efficace",
             "Ventes", "Recettes", "Coût de desserte", "Écart tarif–coût", "Écart / coût"]
    units = ["", "", "% ventes", "FCFA/kWh", "FCFA/kWh", "GWh", "Mds FCFA", "Mds FCFA",
             "Mds FCFA", "%"]
    for i, (h, u) in enumerate(zip(heads, units)):
        put(ws, f"{L(1 + i)}{r}", h, F_BOLD, fill=FILL_HEAD)
        put(ws, f"{L(1 + i)}{r + 1}", u, F_NOTE)
    r += 1
    cat_rows = []
    billed0 = f"{R['inj']}*(1-{R['tech']}-{R['com']})"
    for i, (label, social) in enumerate(CATS):
        r += 1
        cat_rows.append(r)
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", social, F_IN, "0")
        sh, tf, co = EX_CATS[i]
        inp(ws, f"C{r}", sh, PCT, example)
        inp(ws, f"D{r}", tf, FCFA, example)
        inp(ws, f"E{r}", co, FCFA,  example)
        put(ws, f"F{r}", f"={billed0}*C{r}", fmt=GWH)
        put(ws, f"G{r}", f"=F{r}*D{r}/1000", fmt=MDS)
        put(ws, f"H{r}", f"=F{r}*E{r}/1000", fmt=MDS)
        put(ws, f"I{r}", f"=G{r}-H{r}", fmt=MDS)
        put(ws, f"J{r}", f"=IFERROR(D{r}/E{r}-1,0)", fmt=PCT)
    c0, c1 = cat_rows[0], cat_rows[-1]
    r += 1
    put(ws, f"A{r}", "Total / contrôle (la somme des parts doit faire 100 %)", F_BOLD)
    put(ws, f"C{r}", f"=SUM(C{c0}:C{c1})", F_BOLD, PCT)
    for col in "FGHI":
        put(ws, f"{col}{r}", f"=SUM({col}{c0}:{col}{c1})", F_BOLD,
            GWH if col == "F" else MDS)
    put(ws, f"J{r}", f"=IFERROR(G{r}/H{r}-1,0)", F_BOLD, PCT)
    R["share_check"] = f"$C${r}"
    r += 1
    put(ws, f"A{r}", "Avertissement : un écart positif n'est une subvention croisée que si le coût "
        "efficace par catégorie est établi par une étude d'allocation des coûts (tension, "
        "pointe, pertes). Sinon, c'est un simple écart non démontré.", F_NOTE)
    for c in cat_rows:
        ws[f"D{c}"].comment = Comment("Tarif moyen effectif = recettes facturées / kWh "
                                      "(abonnements et primes fixes inclus, taxes exclues).",
                                      "Modèle")
        ws[f"E{c}"].comment = Comment("Coût efficace alloué à la catégorie (étude de coût de "
                                      "service), pertes réglementaires incluses.", "Modèle")

    r += 2
    section(ws, r, "3. SCÉNARIOS DE RÉFORME (paramètres atteints en 2035, trajectoire linéaire)")
    r += 1
    put(ws, f"A{r}", "Scénario actif (choisir dans la liste)", F_BOLD)
    put(ws, f"C{r}", SCEN[3] if example else SCEN[0], F_IN, fill=FILL_IN)
    ws.merge_cells(f"C{r}:D{r}")
    dv = DataValidation(type="list", formula1='"' + ",".join(SCEN) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"C{r}")
    sel = f"$C${r}"
    R["scen"] = sel
    r += 1
    put(ws, f"A{r}", "Paramètre", F_BOLD, fill=FILL_HEAD)
    put(ws, f"B{r}", "Unité", F_BOLD, fill=FILL_HEAD)
    for j, s in enumerate(SCEN):
        put(ws, f"{L(3 + j)}{r}", s, F_BOLD, fill=FILL_HEAD)
    put(ws, f"G{r}", "Actif", F_BOLD, fill=FILL_HEAD)
    head_row = r
    scen_params = [
        ("res_sous", "Résorption de l'écart — catégories sous-tarifées non sociales", "fraction",
         [0, 1, 1, 1], PCT),
        ("res_soc", "Résorption de l'écart — tranche sociale", "fraction", [0, 1, 1, 1], PCT),
        ("res_sur", "Résorption de l'écart — catégories sur-tarifées", "fraction",
         [0, 1, 1, 1], PCT),
        ("aide", "Aide sociale ciblée active", "1/0", [0, 0, 1, 1], "0"),
        ("tech35", "Pertes techniques en 2035", "% injectée", None, PCT),
        ("com35", "Pertes commerciales en 2035", "% injectée", None, PCT),
        ("rec35", "Recouvrement en 2035", "% facturé", None, PCT),
        ("racc", "Raccordements subventionnés par an", "nombre", None, NB),
        ("red_ci", "Réduction de CI atteinte en 2035", "%", [0, 0, 0, None], PCT),
        ("red_ce", "Réduction de CE atteinte en 2035", "%", [0, 0, 0, None], PCT),
    ]
    for key, label, unit, vals, fmt in scen_params:
        r += 1
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", unit, F_NOTE)
        for j in range(4):
            ref = f"{L(3 + j)}{r}"
            if key in ("tech35", "com35", "rec35"):
                base = {"tech35": R["tech"], "com35": R["com"], "rec35": R["rec"]}[key]
                if j < 3:
                    put(ws, ref, f"={base}", fmt=fmt)
                else:
                    inp(ws, ref, EX_S3[key[:-2]], fmt, example)
            elif key == "racc":
                inp(ws, ref, EX_RACC[j], fmt, example)
            elif vals[j] is None:
                inp(ws, ref, {"red_ci": 0.5, "red_ce": 0.6}[key], fmt, example)
            else:
                put(ws, ref, vals[j], F_IN, fmt)
        put(ws, f"G{r}", f"=INDEX(C{r}:F{r},MATCH({sel},$C${head_row}:$F${head_row},0))",
            fmt=fmt)
        R[key] = f"$G${r}"
    r += 1
    put(ws, f"A{r}", "Les valeurs 0/1 de résorption et d'aide définissent les scénarios ; elles "
        "sont modifiables (ex. S3 : ne résorber que les écarts sans justification de coût). "
        "Pertes et recouvrement de S0–S2 = niveau de base (aucune amélioration supposée).", F_NOTE)

    r += 2
    section(ws, r, "4. TRAJECTOIRE 2027–2035 (scénario actif)")
    r += 1
    put(ws, f"A{r}", "Indicateur", F_BOLD, fill=FILL_HEAD)
    put(ws, f"B{r}", "Unité", F_BOLD, fill=FILL_HEAD)
    for col, y in zip(YC, YEARS):
        c = put(ws, f"{col}{r}", y, F_BOLD, fill=FILL_HEAD)
        c.alignment = Alignment(horizontal="right")
    put(ws, f"M{r}", "Cumul 2027–35", F_BOLD, fill=FILL_HEAD)

    T = {}

    def row(key, label, unit, formula, fmt, cumul=False, bold=False):
        nonlocal r
        r += 1
        put(ws, f"A{r}", label, F_BOLD if bold else F_CALC)
        put(ws, f"B{r}", unit, F_NOTE)
        T[key] = r
        for i, col in enumerate(YC):
            put(ws, f"{col}{r}", "=" + formula(i, col), F_BOLD if bold else F_CALC, fmt)
        if cumul:
            put(ws, f"M{r}", f"=SUM({YC[1]}{r}:{YC[-1]}{r})", F_BOLD, fmt)
        for col in YC:
            ws[f"{col}{r}"].border = THIN
        return r

    row("f", "Avancement de la réforme", "fraction", lambda i, c: f"{i}/{N_STEPS}", PCT)
    fr = lambda c: f"{c}${T['f']}"
    row("inj", "Énergie injectée", "GWh",
        lambda i, c: R["inj"] if i == 0 else f"{YC[i - 1]}{T['inj']}*(1+{R['g']})", GWH)
    row("tech", "Pertes techniques", "%",
        lambda i, c: f"{R['tech']}+({R['tech35']}-{R['tech']})*{fr(c)}", PCT)
    row("com", "Pertes commerciales", "%",
        lambda i, c: f"{R['com']}+({R['com35']}-{R['com']})*{fr(c)}", PCT)
    row("rec", "Taux de recouvrement", "%",
        lambda i, c: f"{R['rec']}+({R['rec35']}-{R['rec']})*{fr(c)}", PCT)
    row("billed", "Ventes facturées", "GWh",
        lambda i, c: f"{c}{T['inj']}*(1-{c}{T['tech']}-{c}{T['com']})", GWH)

    tar_rows, sal_rows = [], []
    for k, cr in enumerate(cat_rows):
        res = (f"IF($B${cr}=1,{R['res_soc']},IF($D${cr}<$E${cr},{R['res_sous']},"
               f"{R['res_sur']}))")
        row(f"tar{k}", f"Tarif — {CATS[k][0]}", "FCFA/kWh",
            lambda i, c, cr=cr, res=res: f"$D${cr}+($E${cr}-$D${cr})*{res}*{fr(c)}", FCFA)
        tar_rows.append(T[f"tar{k}"])
    for k, cr in enumerate(cat_rows):
        row(f"sal{k}", f"Ventes — {CATS[k][0]}", "GWh",
            lambda i, c, cr=cr: f"{c}{T['billed']}*$C${cr}", GWH)
        sal_rows.append(T[f"sal{k}"])

    def sp(c, expr):
        """Somme sur les catégories ; expr(tarif, ventes, ligne_categorie)."""
        return "+".join(expr(f"{c}{t}", f"{c}{s}", cr)
                        for t, s, cr in zip(tar_rows, sal_rows, cat_rows))

    row("rev", "Recettes facturées", "Mds FCFA",
        lambda i, c: f"({sp(c, lambda t, s, cr: f'{t}*{s}')})/1000", MDS, True)
    row("cost", "Coût efficace de desserte", "Mds FCFA",
        lambda i, c: f"({sp(c, lambda t, s, cr: f'$E${cr}*{s}')})/1000", MDS, True)
    row("xloss", "Coût des pertes au-delà du niveau admis", "Mds FCFA",
        lambda i, c: f"{c}{T['inj']}*MAX(0,{c}{T['tech']}+{c}{T['com']}-{R['reg']})"
                     f"*{R['cmp']}/1000", MDS, True)
    row("under", "Écart de sous-tarification (Σ coût > tarif)", "Mds FCFA",
        lambda i, c: f"({sp(c, lambda t, s, cr: f'MAX(0,$E${cr}-{t})*{s}')})/1000", MDS, True)
    row("over", "Écart de sur-tarification (Σ tarif > coût)", "Mds FCFA",
        lambda i, c: f"({sp(c, lambda t, s, cr: f'MAX(0,{t}-$E${cr})*{s}')})/1000", MDS, True)
    row("aid_el", "Aide sociale aux éligibles (hausse du tarif social compensée)", "Mds FCFA",
        lambda i, c: f"{R['aide']}*{R['cibl']}*("
                     + sp(c, lambda t, s, cr: f"$B${cr}*MAX(0,{t}-$D${cr})*{s}") + ")/1000",
        MDS, True)
    row("aid_leak", "Fuites (erreur d'inclusion)", "Mds FCFA",
        lambda i, c: f"IFERROR({c}{T['aid_el']}*{R['incl']}/(1-{R['incl']}),0)", MDS, True)
    row("aid_adm", "Coût administratif du ciblage", "Mds FCFA",
        lambda i, c: f"IF({i}=0,0,{R['aide']}*{R['benef']}*{R['admin']}/1E9)", MDS, True)
    row("aid_due", "Aide sociale due (total)", "Mds FCFA",
        lambda i, c: f"{c}{T['aid_el']}+{c}{T['aid_leak']}+{c}{T['aid_adm']}", MDS, True)
    row("aid_paid", "Aide sociale versée", "Mds FCFA",
        lambda i, c: f"{c}{T['aid_due']}*{R['fiab']}", MDS, True)
    row("comp_due", "Compensation tarifaire due à l'opérateur", "Mds FCFA",
        lambda i, c: f"{c}{T['under']}*{R['couv']}", MDS, True)
    row("comp_paid", "Compensation tarifaire versée", "Mds FCFA",
        lambda i, c: f"{c}{T['comp_due']}*{R['fiab']}", MDS, True)
    row("racc", "Subventions au raccordement versées", "Mds FCFA",
        lambda i, c: f"IF({i}=0,0,{R['racc']}*{R['sub_racc']}/1E9*{R['fiab']})", MDS, True)
    row("aid_bill", "Part de l'aide transitant par la facture", "Mds FCFA",
        lambda i, c: f"({c}{T['aid_el']}+{c}{T['aid_leak']})*{R['canal']}", MDS, True)
    row("cash", "Recettes encaissées par l'opérateur (hors compensations)", "Mds FCFA",
        lambda i, c: f"({c}{T['rev']}-{c}{T['aid_bill']})*{c}{T['rec']}"
                     f"+{c}{T['aid_bill']}*{R['fiab']}", MDS, True)
    row("res_op", "Résultat d'exploitation de l'opérateur", "Mds FCFA",
        lambda i, c: f"{c}{T['cash']}+{c}{T['comp_paid']}-{c}{T['cost']}-{c}{T['xloss']}",
        MDS, True, True)
    row("ineff", "  dont coût d'inefficience (pertes excédentaires + non-recouvrement)",
        "Mds FCFA",
        lambda i, c: f"{c}{T['xloss']}+({c}{T['rev']}-{c}{T['aid_bill']})*(1-{c}{T['rec']})",
        MDS, True)

    r += 1
    put(ws, f"A{r + 1}", "Coût total de la politique tarifaire : CTPT = SB + DC + CI + CE", F_BOLD)
    r += 1
    row("SB", "SB — subventions budgétaires versées", "Mds FCFA",
        lambda i, c: f"{c}{T['comp_paid']}+{c}{T['aid_paid']}+{c}{T['racc']}", MDS, True)
    row("DC", "DC — déficit de compensation non financé", "Mds FCFA",
        lambda i, c: f"MAX(0,{c}{T['under']}-{c}{T['comp_paid']})"
                     f"+({c}{T['aid_el']}+{c}{T['aid_leak']})*{R['canal']}*(1-{R['fiab']})",
        MDS, True)
    row("CI", "CI — coûts indirects (dégradation du service)", "Mds FCFA",
        lambda i, c: f"{R['ci']}*(1-{R['red_ci']}*{fr(c)})", MDS, True)
    row("CE", "CE — surcoûts des entreprises", "Mds FCFA",
        lambda i, c: f"{R['ce']}*(1-{R['red_ce']}*{fr(c)})", MDS, True)
    row("CTPT", "CTPT — coût total de la politique tarifaire", "Mds FCFA",
        lambda i, c: f"{c}{T['SB']}+{c}{T['DC']}+{c}{T['CI']}+{c}{T['CE']}", MDS, True, True)
    r += 1
    put(ws, f"A{r}", "Anti-double comptage : DC ne contient que la sous-tarification non "
        "compensée et l'aide due non versée. Les pertes excédentaires et le non-recouvrement "
        "relèvent de la performance de l'opérateur (ligne « inefficience »), pas du CTPT. CI et CE "
        "doivent être estimés sur des périmètres disjoints.", F_NOTE)

    r += 1
    row("ind_idx", "Indice du tarif industriel HT (base = 100)", "indice",
        lambda i, c: f"IFERROR({c}{tar_rows[IND_CAT]}/${YC[0]}${tar_rows[IND_CAT]}*100,0)", IDX)
    row("ind_cost", "Effet sur le coût de production industriel", "% coût",
        lambda i, c: f"{R['part_ind']}*({c}{T['ind_idx']}/100-1)", PCT)
    row("soc_idx", "Indice du tarif de la tranche sociale (base = 100)", "indice",
        lambda i, c: f"IFERROR({c}{tar_rows[0]}/${YC[0]}${tar_rows[0]}*100,0)", IDX)
    row("excl", "Hausse non compensée pour les ménages sociaux exclus du ciblage", "Mds FCFA",
        lambda i, c: f"(1-{R['aide']}*{R['cibl']})*MAX(0,{c}{tar_rows[0]}-$D${cat_rows[0]})"
                     f"*{c}{sal_rows[0]}/1000", MDS, True)
    ws.freeze_panes = "C5"
    return ws, R, T


def build_equity(wb, countries):
    ws = wb.create_sheet("Equite_H2")
    ws.column_dimensions["A"].width = 50
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 15
    put(ws, "A1", "Incidence distributive de la subvention résidentielle (test de H2)", F_TITLE)
    put(ws, "A2", "Source attendue : enquête ménages représentative (consommation, raccordement, "
        "quintile de revenu ou de dépenses). Les ménages non raccordés reçoivent 0.", F_NOTE)
    r = 3
    for name, R, cat_rows_sheet, example in countries:
        r += 2
        section(ws, r, f"{name}", 8)
        r += 1
        put(ws, f"A{r}", "Seuil de la tranche sociale")
        put(ws, f"B{r}", "kWh/mois", F_NOTE)
        inp(ws, f"C{r}", EX_EQ["seuil"], NB, example)
        seuil = f"$C${r}"
        r += 1
        put(ws, f"A{r}", "Structure (1 = progressive : chacun paie la 1re tranche au tarif "
            "social ; 0 = tarif social seulement si conso ≤ seuil)")
        inp(ws, f"C{r}", EX_EQ["prog"], "0", example)
        prog = f"$C${r}"
        r += 1
        s = f"'{name}'!"
        put(ws, f"A{r}", "Tarif social / tarif autres tranches / coût BT résidentiel")
        put(ws, f"B{r}", "FCFA/kWh", F_NOTE)
        put(ws, f"C{r}", f"={s}$D${cat_rows_sheet[0]}", F_LINK, FCFA)
        put(ws, f"D{r}", f"={s}$D${cat_rows_sheet[1]}", F_LINK, FCFA)
        put(ws, f"E{r}", f"={s}$E${cat_rows_sheet[1]}", F_LINK, FCFA)
        put(ws, f"F{r}", "Liés à la feuille pays (base).", F_NOTE)
        t1, t2, cost = f"$C${r}", f"$D${r}", f"$E${r}"
        r += 1
        heads = ["Quintile", "Ménages", "Taux de raccordement", "Conso. raccordés (kWh/mois)",
                 "Subvention / ménage raccordé (FCFA/an)", "Subvention totale (Mds FCFA)",
                 "Part de la subvention", "Part des ménages"]
        for i, h in enumerate(heads):
            c = put(ws, f"{L(1 + i)}{r}", h, F_BOLD, fill=FILL_HEAD)
            c.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[r].height = 42
        q0 = r + 1
        q1 = r + 5
        for q in range(5):
            r += 1
            hh, conn, cons = EX_EQ["q"][q]
            put(ws, f"A{r}", f"Q{q + 1}" + (" (plus pauvres)" if q == 0 else
                                            " (plus riches)" if q == 4 else ""))
            inp(ws, f"B{r}", hh, NB, example)
            inp(ws, f"C{r}", conn, PCT, example)
            inp(ws, f"D{r}", cons, NB, example)
            put(ws, f"E{r}",
                f"=12*IF({prog}=1,MIN(D{r},{seuil})*({cost}-{t1})"
                f"+MAX(0,D{r}-{seuil})*({cost}-{t2}),"
                f"IF(D{r}<={seuil},D{r}*({cost}-{t1}),D{r}*({cost}-{t2})))", fmt=NB)
            put(ws, f"F{r}", f"=B{r}*C{r}*E{r}/1E9", fmt=MDS)
            put(ws, f"G{r}", f"=IFERROR(F{r}/SUM($F${q0}:$F${q1}),0)", fmt=PCT)
            put(ws, f"H{r}", f"=IFERROR(B{r}/SUM($B${q0}:$B${q1}),0)", fmt=PCT)
        r += 1
        put(ws, f"A{r}", "Total", F_BOLD)
        put(ws, f"B{r}", f"=SUM(B{q0}:B{q1})", F_BOLD, NB)
        put(ws, f"F{r}", f"=SUM(F{q0}:F{q1})", F_BOLD, MDS)
        r += 1
        put(ws, f"A{r}", "Part de la subvention captée par Q1+Q2 (40 % les plus pauvres)", F_BOLD)
        put(ws, f"C{r}", f"=G{q0}+G{q0 + 1}", F_BOLD, PCT)
        r += 1
        put(ws, f"A{r}", "Part captée par Q5 (20 % les plus riches)", F_BOLD)
        put(ws, f"C{r}", f"=G{q1}", F_BOLD, PCT)
        r += 1
        put(ws, f"A{r}", "H2 soutenue si la part de Q5 dépasse celle de Q1+Q2 (1 = oui)", F_BOLD)
        put(ws, f"C{r}", f"=IF(F{q1 + 1}=0,\"n.d.\",IF(G{q1}>G{q0}+G{q0 + 1},1,0))", F_BOLD)
        r += 1
        put(ws, f"A{r}", "Lecture : subvention négative = la catégorie paie plus que le coût BT "
            "retenu. Test indicatif, à compléter par des intervalles de confiance d'enquête.",
            F_NOTE)
    return ws


def build_summary(wb, countries):
    ws = wb["Synthese"]
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 12
    for col in "CDE":
        ws.column_dimensions[col].width = 20
    put(ws, "A1", "Synthèse comparative (scénario actif de chaque feuille pays)", F_TITLE)
    put(ws, "A2", "Liens verts vers les feuilles pays. Pour comparer les scénarios, changer le "
        "scénario actif dans chaque feuille pays et relever les résultats.", F_NOTE)
    put(ws, "A4", "Indicateur", F_BOLD, fill=FILL_HEAD)
    put(ws, "B4", "Unité", F_BOLD, fill=FILL_HEAD)
    for j, (name, *_rest) in enumerate(countries):
        put(ws, f"{L(3 + j)}4", name, F_BOLD, fill=FILL_HEAD)
    lines = [
        ("Scénario actif", "", lambda R, T: R["scen"], None),
        ("Contrôle : somme des parts de ventes (doit = 100 %)", "%",
         lambda R, T: R["share_check"], PCT),
        ("Recettes facturées / coût efficace (base)", "%",
         lambda R, T: f"IFERROR({{s}}C{T['rev']}/{{s}}C{T['cost']},0)", PCT),
        ("Écart de sous-tarification (base)", "Mds FCFA", lambda R, T: f"C{T['under']}", MDS),
        ("Écart de sur-tarification (base)", "Mds FCFA", lambda R, T: f"C{T['over']}", MDS),
        ("Résultat d'exploitation opérateur (base)", "Mds FCFA",
         lambda R, T: f"C{T['res_op']}", MDS),
        ("Résultat d'exploitation opérateur (2035)", "Mds FCFA",
         lambda R, T: f"L{T['res_op']}", MDS),
        ("Coût d'inefficience opérateur (2035)", "Mds FCFA", lambda R, T: f"L{T['ineff']}", MDS),
        ("SB cumulées 2027–2035", "Mds FCFA", lambda R, T: f"M{T['SB']}", MDS),
        ("DC cumulé 2027–2035", "Mds FCFA", lambda R, T: f"M{T['DC']}", MDS),
        ("CTPT cumulé 2027–2035", "Mds FCFA", lambda R, T: f"M{T['CTPT']}", MDS),
        ("CTPT 2035 / CTPT base", "%",
         lambda R, T: f"IFERROR({{s}}L{T['CTPT']}/{{s}}C{T['CTPT']},0)", PCT),
        ("Aide sociale due cumulée 2027–2035", "Mds FCFA", lambda R, T: f"M{T['aid_due']}", MDS),
        ("Hausse non compensée des ménages sociaux exclus, cumul", "Mds FCFA",
         lambda R, T: f"M{T['excl']}", MDS),
        ("Indice tarif industriel HT 2035 (base 100)", "indice",
         lambda R, T: f"L{T['ind_idx']}", IDX),
        ("Effet sur le coût de production industriel 2035", "% coût",
         lambda R, T: f"L{T['ind_cost']}", PCT),
    ]
    for i, (label, unit, ref, fmt) in enumerate(lines):
        rr = 5 + i
        put(ws, f"A{rr}", label)
        put(ws, f"B{rr}", unit, F_NOTE)
        for j, (name, R, T) in enumerate((n, R, T) for n, R, _c, _e, T in countries):
            expr = ref(R, T)
            sheet = f"'{name}'!"
            expr = expr.replace("{s}", sheet) if "{s}" in expr else sheet + expr
            put(ws, f"{L(3 + j)}{rr}", "=" + expr, F_LINK, fmt)


def build_readme(wb):
    ws = wb.active
    ws.title = "LISEZMOI"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Modèle financier et tarifaire — électricité (Livrable 2, version de structure)", F_TITLE),
        ("Étude : Le paradoxe des subventions tarifaires en Afrique — Cameroun, République du "
         "Congo et pays dotés d'un pacte national de l'énergie (Mission 300).", F_CALC),
        ("", F_CALC),
        ("STATUT DES DONNÉES", F_BOLD),
        ("• Feuilles « Cameroun » et « Congo » : AUCUNE donnée saisie. Les cellules jaunes sont à "
         "renseigner depuis des sources officielles, avec la source en colonne D.", F_CALC),
        ("• Feuille « Exemple_fictif » : valeurs INVENTÉES servant uniquement à vérifier les "
         "formules. Ne jamais citer ses résultats.", F_CALC),
        ("", F_CALC),
        ("CODE COULEUR", F_BOLD),
        ("• Texte bleu : saisie / levier de scénario.  • Fond jaune : saisie à documenter.  "
         "• Texte noir : formule.  • Texte vert : lien vers une autre feuille.", F_CALC),
        ("", F_CALC),
        ("STRUCTURE (5 modules de la note de cadrage)", F_BOLD),
        ("Module 1 Coûts et Module 2 Tarifs : sections 1–2 des feuilles pays (coût efficace et "
         "tarif moyen par catégorie, écarts).", F_CALC),
        ("Module 3 Équité : feuille Equite_H2 (incidence par quintile) + ligne « ménages sociaux "
         "exclus du ciblage ».", F_CALC),
        ("Module 4 Industrie : indice du tarif HT, effet sur le coût de production, CE.", F_CALC),
        ("Module 5 Réforme : section 3 (S0–S3) et section 4 (trajectoire 2027–2035, CTPT).",
         F_CALC),
        ("", F_CALC),
        ("LIMITES CONNUES DE CETTE VERSION", F_BOLD),
        ("• Pas d'élasticité-prix : une hausse de tarif ne réduit pas la demande ni n'augmente "
         "l'autoproduction.", F_CALC),
        ("• Coûts efficaces constants en termes réels ; pas de programme d'investissement "
         "explicite (H4 non testée quantitativement).", F_CALC),
        ("• Trajectoires linéaires ; un seul scénario actif par feuille à la fois.", F_CALC),
        ("• Le secteur de l'eau n'est pas modélisé (périmètre retenu : électricité).", F_CALC),
        ("• CI et CE sont des estimations exogènes ; leur méthode doit éviter les doubles "
         "comptages.", F_CALC),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = put(ws, f"A{i}", text, font)
        c.alignment = Alignment(wrap_text=True, vertical="top")


def main(out):
    wb = Workbook()
    build_readme(wb)
    wb.create_sheet("Synthese")
    built = []
    for name, ex in [("Cameroun", False), ("Congo", False), ("Exemple_fictif", True)]:
        ws, R, T = build_country(wb, name, ex)
        cat_rows = [r for r in range(1, ws.max_row + 1)
                    if ws[f"A{r}"].value in [c[0] for c in CATS] and ws[f"F{r}"].value]
        built.append((name, R, cat_rows, ex, T))
    build_equity(wb, [(n, R, c, e) for n, R, c, e, _t in built])
    build_summary(wb, built)
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "modele_tarifaire_electricite.xlsx")
