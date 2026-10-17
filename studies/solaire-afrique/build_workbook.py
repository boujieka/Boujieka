"""Génère etude_solaire_afrique.xlsx : gabarit de base de données (volets A, B, C),
journal des observations sourcées, registre des sources et modèle LCOE en formules Excel.

Seules les valeurs vérifiées dans une source citée sont saisies ; tout le reste est vide.
    python build_workbook.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("etude_solaire_afrique.xlsx")
AS_OF = "2026-10-17"

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(italic=True, color="595959")

STATUTS = ["en service", "en construction", "attribué", "annoncé", "à vérifier", "abandonné"]
TYPES_DONNEE = ["observée", "déclarée", "modélisée", "non disponible"]
FIABILITE = ["A - document primaire", "B - institution/développeur", "C - presse spécialisée",
             "D - source secondaire non recoupée"]

PAYS = [
    ("Afrique du Nord", "MAR", "Maroc"), ("Afrique du Nord", "EGY", "Égypte"),
    ("Afrique du Nord", "TUN", "Tunisie"), ("Afrique du Nord", "DZA", "Algérie"),
    ("Afrique du Nord", "MRT", "Mauritanie"),
    ("Afrique de l'Ouest", "SEN", "Sénégal"), ("Afrique de l'Ouest", "CIV", "Côte d'Ivoire"),
    ("Afrique de l'Ouest", "GHA", "Ghana"), ("Afrique de l'Ouest", "BFA", "Burkina Faso"),
    ("Afrique de l'Ouest", "MLI", "Mali"), ("Afrique de l'Ouest", "NER", "Niger"),
    ("Afrique centrale", "CMR", "Cameroun"), ("Afrique centrale", "COG", "Congo"),
    ("Afrique centrale", "TCD", "Tchad"), ("Afrique centrale", "GAB", "Gabon"),
    ("Afrique centrale", "COD", "RDC"),
    ("Afrique de l'Est", "KEN", "Kenya"), ("Afrique de l'Est", "ETH", "Éthiopie"),
    ("Afrique de l'Est", "UGA", "Ouganda"), ("Afrique de l'Est", "RWA", "Rwanda"),
    ("Afrique de l'Est", "TZA", "Tanzanie"),
    ("Afrique australe", "ZAF", "Afrique du Sud"), ("Afrique australe", "ZMB", "Zambie"),
    ("Afrique australe", "NAM", "Namibie"), ("Afrique australe", "ZWE", "Zimbabwe"),
]

# (champ, unité, description) — communs puis spécifiques à chaque volet.
COMMUNS = [
    ("projet_id", "", "Identifiant unique (ex. A-ZMB-001)"),
    ("pays_iso3", "", "Code ISO 3166-1 alpha-3"),
    ("nom", "", "Nom de la centrale"),
    ("localisation", "", "Ville / région / coordonnées"),
    ("developpeur", "", "Développeur ou sponsor principal"),
    ("statut", "", "Voir liste : " + ", ".join(STATUTS)),
    ("annee_attribution", "année", ""),
    ("annee_mise_en_service", "année", "Date de mise en service commerciale (COD)"),
    ("puissance_mwp", "MWp", "Puissance crête DC"),
    ("puissance_mwac", "MWac", "Puissance AC (onduleurs / point d'injection)"),
    ("stockage_mwh", "MWh", "Capacité utile de stockage (0 si aucune)"),
    ("technologie", "", "Modules, onduleurs, fixe / tracker"),
    ("capex_total", "devise d'origine", "Investissement total"),
    ("capex_devise", "ISO 4217", "Devise du CAPEX"),
    ("capex_annee_prix", "année", "Année de valeur du CAPEX"),
    ("capex_usd_kwp", "USD/kWp", "Calculé, USD courants"),
    ("capex_usd_kwac", "USD/kWac", "Calculé, USD courants"),
    ("production_gwh_an", "GWh/an", "Production annuelle (préciser P50 attendue ou mesurée)"),
    ("facteur_charge_ac", "%", "Production / (MWac x 8 760 h)"),
    ("opex_usd_kw_an", "USD/kW/an", ""),
    ("dette_part", "%", "Part de la dette dans le financement"),
    ("dette_taux", "%", "Taux d'intérêt (ou marge + référence)"),
    ("dette_duree", "années", ""),
    ("wacc", "%", "Observé ou modélisé (préciser)"),
    ("tarif", "devise/kWh", "Tarif contractuel année 1"),
    ("tarif_devise", "ISO 4217", ""),
    ("tarif_indexation", "", "Fixe, indexé inflation, part indexée, etc."),
    ("ppa_duree", "années", ""),
    ("procedure", "", "Appel d'offres concurrentiel, FiT, négociation directe, EPC public"),
    ("garanties", "", "Garantie souveraine, IDA/MIGA, lettre de crédit, etc."),
    ("lcoe_usd_mwh", "USD/MWh", "Calculé par le modèle (jamais saisi comme observé)"),
    ("sources", "", "Identifiants du registre Sources, séparés par ;"),
    ("fiabilite_globale", "", "Voir liste"),
    ("notes", "", "Contradictions entre sources, hypothèses"),
]
SPEC_A = [
    ("acheteur", "", "Compagnie nationale ou acheteur"),
    ("raccordement_inclus_capex", "oui/non/inconnu", "Le CAPEX inclut-il ligne et poste ?"),
]
SPEC_B = [
    ("secteur", "", "Mines, forêt, agro-industrie, ciment, ZES, etc."),
    ("modele_commercial", "", "Autoproduction, leasing, PPA privé, ESCO"),
    ("cout_diesel_evite_usd_mwh", "USD/MWh", ""),
    ("taux_couverture_solaire", "%", "Part de la consommation couverte par le solaire"),
    ("tri_projet", "%", ""),
]
SPEC_C = [
    ("nb_clients", "", "Raccordements actifs"),
    ("conso_moyenne_kwh_client_mois", "kWh/mois", ""),
    ("cout_distribution_usd", "USD", "Réseau BT, branchements, compteurs"),
    ("tarif_client", "devise/kWh", "Ou forfait, préciser"),
    ("subvention_usd", "USD", "Subvention d'investissement ou par branchement"),
    ("operateur", "", "Privé, public, communautaire"),
    ("client_ancrage", "oui/non", ""),
]

SOURCES = [
    ("S01", "IRENA", "Renewable Power Generation Costs in 2024 — Summary", 2025,
     "https://www.irena.org/-/media/Files/IRENA/Agency/Publication/2025/Jul/IRENA_TEC_RPGC_in_2024_Summary_2025.pdf",
     "A", "WACC supposé 3,8 % (Europe) à 12 % (Afrique) ; LCOE solaire mondial 0,043 USD/kWh"),
    ("S02", "pv magazine", "Global average solar LCOE stood at $0.043/kWh in 2024, says IRENA", 2025,
     "https://www.pv-magazine.com/2025/07/23/global-average-solar-lcoe-stood-at-0-043-kwh-in-2024-says-irena/",
     "C", "Coût d'installation moyen Afrique 2024 : 1 093 USD/kW (cite IRENA)"),
    ("S03", "Banque mondiale", "Unlocking low-cost, large-scale solar power in Zambia", 2019,
     "https://www.worldbank.org/en/news/feature/2019/05/14/unlocking-low-cost-large-scale-solar-power-in-zambia",
     "B", "Bangweulu 54 MW selon cette source"),
    ("S04", "pv-tech", "IFC Zambia solar auction sees record-low 6¢ price benchmark", 2016,
     "https://www.pv-tech.org/news/ifc-zambia-solar-auction-sees-record-low-6-price-benchmark",
     "C", "Tarifs 6,015 et 7,84 USc/kWh"),
    ("S05", "BloombergNEF Climatescope", "Zambia market profile", 2022,
     "https://2022.global-climatescope.org/markets/zm/",
     "C", "Ngonye 34 MW, Bangweulu 56 MW, mise en service 2019"),
    ("S06", "IFC", "Senegal Scaling Solar PPP brief", 2021,
     "https://www.ifc.org/content/dam/ifc/doc/2021/2021-senegal-scaling-solar-ppp-brief.pdf",
     "B", "Kahone 3,80 et Kael 3,98 c EUR/kWh, appel d'offres avril 2018"),
    ("S07", "IFC", "Scaling Solar: two PV plants bring clean energy to more than 500,000 in Senegal", 2021,
     "https://www.ifc.org/en/pressroom/2021/scaling-solar-two-pv-plants-bring-clean-energy-to-more-than-500000-in-senegal",
     "B", "Mise en service mai 2021, 60 MWac au total, Engie/Meridiam/FONSIS"),
]

# Observations vérifiées : (projet_id, champ, valeur, unité, type, source, note)
OBSERVATIONS = [
    ("A-ZMB-001", "tarif", 0.06015, "USD/kWh", "observée", "S04", "Tarif fixe 25 ans (Neoen/First Solar)"),
    ("A-ZMB-001", "puissance_mwac", 54, "MW", "déclarée", "S03", "S05 indique 56 MW : à réconcilier"),
    ("A-ZMB-001", "annee_mise_en_service", 2019, "année", "déclarée", "S05", ""),
    ("A-ZMB-002", "tarif", 0.0784, "USD/kWh", "observée", "S04", "Enel"),
    ("A-ZMB-002", "puissance_mwac", 34, "MW", "déclarée", "S05", "AC ou DC non précisé"),
    ("A-ZMB-002", "annee_mise_en_service", 2019, "année", "déclarée", "S05", ""),
    ("A-SEN-001", "tarif", 0.0398, "EUR/kWh", "observée", "S06", "Engie/Meridiam/FONSIS"),
    ("A-SEN-001", "annee_mise_en_service", 2021, "année", "déclarée", "S07", ""),
    ("A-SEN-002", "tarif", 0.0380, "EUR/kWh", "observée", "S06", ""),
    ("A-SEN-002", "annee_mise_en_service", 2021, "année", "déclarée", "S07", ""),
]

# Projets de la proposition. Les chiffres ne sont saisis qu'après vérification (onglet Observations).
PROJETS_A = [
    ("A-ZMB-001", "ZMB", "Bangweulu (Scaling Solar)", "Neoen / First Solar", "en service",
     2016, 2019, None, 54, 0.06015, "USD", "Appel d'offres concurrentiel", "S03;S04;S05",
     "Capacité 54 MW (S03) vs 56 MW (S05)"),
    ("A-ZMB-002", "ZMB", "Ngonye (Scaling Solar)", "Enel", "en service",
     2016, 2019, None, 34, 0.0784, "USD", "Appel d'offres concurrentiel", "S04;S05", ""),
    ("A-SEN-001", "SEN", "Kael (Scaling Solar)", "Engie / Meridiam / FONSIS", "en service",
     2018, 2021, None, None, 0.0398, "EUR", "Appel d'offres concurrentiel", "S06;S07",
     "Une source (AE-Africa) place ce tarif à Touba : erreur probable"),
    ("A-SEN-002", "SEN", "Kahone (Scaling Solar)", "Engie / Meridiam / FONSIS", "en service",
     2018, 2021, None, None, 0.0380, "EUR", "Appel d'offres concurrentiel", "S06;S07",
     "Kael + Kahone = 60 MWac au total (S07)"),
    ("A-EGY-001", "EGY", "Benban Solar Park (32 centrales)", "Multiples", "à vérifier",
     None, None, None, None, None, None, "Tarif d'achat garanti (FiT) rounds 1 et 2", "",
     "1 465 MW et ~1,9 Md USD selon presse égyptienne ; tarif round 2 non vérifié. "
     "À éclater en une ligne par centrale"),
    ("A-MAR-001", "MAR", "Noor Ouarzazate IV (PV)", "", "à vérifier",
     None, None, None, None, None, None, "", "", ""),
    ("A-BFA-001", "BFA", "Zagtouli", "", "à vérifier",
     None, None, None, None, None, None, "", "", ""),
    ("A-ZAF-001", "ZAF", "REIPPPP — projets PV (par fenêtre)", "Multiples", "à vérifier",
     None, None, None, None, None, None, "Appel d'offres concurrentiel", "",
     "À éclater en une ligne par projet à partir des documents du DMRE/IPP Office"),
]


def header(ws, cols, row=1):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(name) + 4))
    ws.freeze_panes = f"A{row + 1}"


def add_list(ws, col_letter, values, rows=2000):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{rows}")


def volet_sheet(wb, title, spec):
    ws = wb.create_sheet(title)
    fields = [f for f, _, _ in COMMUNS + spec]
    header(ws, fields)
    add_list(ws, get_column_letter(fields.index("statut") + 1), STATUTS)
    add_list(ws, get_column_letter(fields.index("fiabilite_globale") + 1), FIABILITE)
    return ws, fields


def build() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Lisez-moi"
    lines = [
        "Le coût de l'électricité solaire en Afrique (2016–2026) — base de données et modèle",
        f"Version gabarit générée le {AS_OF} par build_workbook.py. Date d'arrêt de l'étude : 17 octobre 2026.",
        "",
        "Règles :",
        "1. Aucune valeur sans source : chaque chiffre est consigné dans Observations avec type, source et note.",
        "2. Types de donnée : observée (contrat, mesure), déclarée (annonce), modélisée (calcul), non disponible.",
        "3. Ne jamais mélanger USD/kWp (DC) et USD/kWac (AC). Conserver devise et année d'origine.",
        "4. Les projets en construction ou attribués restent hors des statistiques « en service ».",
        "5. PV+stockage et solaire thermodynamique (CSP) sont analysés à part.",
        "6. Le LCOE est toujours calculé (Modele_LCOE), jamais saisi comme donnée observée.",
        "",
        "État actuel : 4 projets partiellement renseignés (Zambie, Sénégal) ; les autres lignes sont des",
        "candidats issus de la proposition, marqués « à vérifier », sans chiffres.",
        "Les cibles de 100 / 75 / 125 projets sont des objectifs de collecte, pas des projets identifiés.",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Dictionnaire")
    header(ws, ["volet", "champ", "unité", "description"])
    for volet, spec in (("commun", COMMUNS), ("A", SPEC_A), ("B", SPEC_B), ("C", SPEC_C)):
        for f, u, d in spec:
            ws.append([volet, f, u, d])
    ws.column_dimensions["D"].width = 70

    ws = wb.create_sheet("Pays")
    header(ws, ["region", "iso3", "pays", "nb_projets_A", "nb_projets_B", "nb_projets_C"])
    for i, (reg, iso, nom) in enumerate(PAYS, 2):
        ws.append([reg, iso, nom] + [f'=COUNTIF(Volet_{v}!B:B,B{i})' for v in "ABC"])

    ws_a, fa = volet_sheet(wb, "Volet_A", SPEC_A)
    for (pid, iso, nom, dev, statut, attr, cod, mwp, mwac, tarif, dev_t, proc, src, note) in PROJETS_A:
        row = dict(projet_id=pid, pays_iso3=iso, nom=nom, developpeur=dev, statut=statut,
                   annee_attribution=attr, annee_mise_en_service=cod, puissance_mwp=mwp,
                   puissance_mwac=mwac, tarif=tarif, tarif_devise=dev_t, procedure=proc,
                   sources=src, notes=note)
        ws_a.append([row.get(f) for f in fa])
    volet_sheet(wb, "Volet_B", SPEC_B)
    volet_sheet(wb, "Volet_C", SPEC_C)

    ws = wb.create_sheet("Observations")
    header(ws, ["projet_id", "champ", "valeur", "unité", "type_donnee", "source_id", "note"])
    for o in OBSERVATIONS:
        ws.append(list(o))
    add_list(ws, "E", TYPES_DONNEE)

    ws = wb.create_sheet("Sources")
    header(ws, ["source_id", "organisme", "titre", "année", "url", "fiabilité", "contenu utilisé"])
    for s in SOURCES:
        ws.append(list(s))
    ws.column_dimensions["E"].width = 60
    ws.column_dimensions["G"].width = 60

    lcoe_sheet(wb)
    verification_sheet(wb)
    return wb


def lcoe_sheet(wb):
    ws = wb.create_sheet("Modele_LCOE")
    ws["A1"] = "Modèle LCOE — cellules jaunes modifiables. Investissement en année 0, flux en années 1..n."
    ws["A1"].font = Font(bold=True)
    inputs = [
        ("Puissance", 50, "MWp"),
        ("CAPEX spécifique", 1000, "USD/kWp"),
        ("Production spécifique année 1", 1800, "kWh/kWp/an"),
        ("OPEX année 1", 15, "USD/kWp/an"),
        ("Indexation OPEX", 0.0, "%/an"),
        ("Dégradation", 0.0, "%/an"),
        ("Durée de vie", 25, "années (max 40)"),
        ("Part de la dette", 0.7, "%"),
        ("Coût de la dette", 0.08, "%"),
        ("Coût des fonds propres", 0.15, "%"),
        ("Taux d'impôt", 0.0, "%"),
        ("WACC calculé", "=B10*B11*(1-B13)+(1-B10)*B12", "%"),
        ("WACC retenu (vide = calculé)", None, "%"),
    ]
    for i, (lab, val, unit) in enumerate(inputs, 3):
        ws.cell(row=i, column=1, value=lab)
        c = ws.cell(row=i, column=2, value=val)
        ws.cell(row=i, column=3, value=unit)
        if not (isinstance(val, str) and val.startswith("=")):
            c.fill = INPUT_FILL
    for r in (7, 8, 10, 11, 12, 13, 14, 15):
        ws.cell(row=r, column=2).number_format = "0.0%"

    ws["A17"], ws["B17"] = "Taux d'actualisation", "=IF(ISBLANK(B15),B14,B15)"
    ws["A18"], ws["B18"] = "CAPEX total (USD)", "=B3*1000*B4"
    ws["A19"], ws["B19"] = "Production année 1 (MWh)", "=B3*B5"
    ws["A20"], ws["B20"] = "OPEX année 1 (USD)", "=B3*1000*B6"
    ws["A21"], ws["B21"] = "Coûts actualisés (USD)", "=B18+SUMPRODUCT(F4:F43,H4:H43)"
    ws["A22"], ws["B22"] = "Production actualisée (MWh)", "=SUMPRODUCT(G4:G43,H4:H43)"
    ws["A23"], ws["B23"] = "LCOE (USD/MWh)", "=B21/B22"
    ws["A24"], ws["B24"] = "Facteur de charge DC", "=B5/8760"
    ws["B17"].number_format = ws["B24"].number_format = "0.0%"
    ws["B23"].number_format = "0.0"
    ws["A23"].font = ws["B23"].font = Font(bold=True)

    ws["E3"], ws["F3"], ws["G3"], ws["H3"] = "année", "OPEX (USD)", "production (MWh)", "facteur d'actualisation"
    for c in ("E3", "F3", "G3", "H3"):
        ws[c].font = Font(bold=True)
    for k, r in enumerate(range(4, 44), 1):
        ws[f"E{r}"] = k
        ws[f"F{r}"] = f"=IF(E{r}<=$B$9,$B$20*(1+$B$7)^(E{r}-1),0)"
        ws[f"G{r}"] = f"=IF(E{r}<=$B$9,$B$19*(1-$B$8)^(E{r}-1),0)"
        ws[f"H{r}"] = f"=1/(1+$B$17)^E{r}"

    ws["J3"] = "Sensibilité au WACC (sans dégradation ni indexation : (CAPEX x CRF + OPEX) / E)"
    ws["J3"].font = Font(bold=True)
    ws["J4"], ws["K4"], ws["L4"] = "WACC", "LCOE (USD/MWh)", "écart vs 4 %"
    for i, r in enumerate((0.04, 0.06, 0.08, 0.10, 0.12, 0.14), 5):
        ws[f"J{i}"] = r
        ws[f"J{i}"].number_format = "0%"
        ws[f"K{i}"] = f"=($B$18*J{i}*(1+J{i})^$B$9/((1+J{i})^$B$9-1)+$B$20)/$B$19"
        ws[f"K{i}"].number_format = "0.0"
        ws[f"L{i}"] = f"=K{i}/$K$5-1"
        ws[f"L{i}"].number_format = "0.0%"
    ws["J12"] = ("Valeurs par défaut = cas illustratif de la proposition (50 MWp, 50 M USD, 90 GWh, "
                 "0,75 M USD/an, 25 ans). Résultats illustratifs, pas des coûts observés.")
    ws["J12"].font = NOTE
    ws.column_dimensions["A"].width = 30
    for col in "EFGH":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 16


def verification_sheet(wb):
    ws = wb.create_sheet("Verif_proposition")
    header(ws, ["affirmation de la proposition", "statut", "constat", "source"])
    rows = [
        ("WACC jusqu'à 12 % en Afrique dans l'analyse IRENA 2024", "confirmé",
         "IRENA : hypothèses de WACC de 3,8 % (Europe) à 12 % (Afrique), conditions 2023", "S01"),
        ("LCOE solaire mondial 2024 : 43 USD/MWh", "confirmé", "0,043 USD/kWh", "S01;S02"),
        ("CAPEX Afrique 2024 : 1 093 USD/kW", "corroboré (secondaire)",
         "Cité par pv magazine d'après IRENA ; non retrouvé dans le résumé IRENA", "S02"),
        ("LCOE Afrique 2024 : 74 USD/MWh", "non vérifié",
         "Absent du résumé IRENA ; à contrôler dans le rapport complet ou la base IRENA", ""),
        ("CAPEX Afrique 2015 : 3 437 USD/kW ; LCOE 2015 : 225 USD/MWh", "non vérifié",
         "Une source secondaire citant l'édition suivante d'IRENA donne 3 520 USD/kW et 222 USD/MWh "
         "pour 2015 : les séries historiques sont révisées d'une édition à l'autre", ""),
        ("Tableau WACC 4–14 % (43,9 à 89,3 USD/MWh)", "confirmé (calcul)",
         "Reproduit : 43,9 / 51,8 / 60,4 / 69,5 / 79,2 / 89,2 ; écarts d'arrondi ≤ 0,1", "lcoe.py"),
        ("+80 % de LCOE entre WACC 4 % et 12 %", "confirmé (calcul)", "+80,4 %", "lcoe.py"),
        ("Ngonye / Bangweulu", "précisé",
         "Bangweulu = 6,015 USc/kWh (Neoen/First Solar) ; Ngonye = 7,84 USc/kWh (Enel)", "S04;S05"),
        ("Kael et Kahone", "confirmé", "3,98 et 3,80 c EUR/kWh, mise en service mai 2021", "S06;S07"),
    ]
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["C"].width = 80


if __name__ == "__main__":
    build().save(OUT)
    print(f"écrit : {OUT}")
