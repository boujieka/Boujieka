"""The live formula workbook must reproduce the engine after recalculation in LibreOffice. Skipped when soffice is
not installed."""
from __future__ import annotations

import shutil
import subprocess

import numpy as np
import openpyxl
import pytest

from app.engine.demo import demo_inputs
from app.engine.inputs import GridArrival, ProjectInputs
from app.engine.model import run_model
from app.services.formula_xlsx import FIRST_COL, build_with_rows, model_rows

pytestmark = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")


def _cases():
    wc = demo_inputs()
    wc.opex.receivable_days, wc.opex.payable_days = 30, 45
    tr = demo_inputs()
    tr.grid_arrival = GridArrival(enabled=True, operating_year=7, treatment="asset_transfer", compensation_pct_nbv=0.6, regulatory_source="test")
    sc = demo_inputs()
    sc.financing.debt[0].amortisation = "sculpted"
    sc.financing.debt[0].share_of_capex = 0.8
    sc.financing.debt[1].share_of_capex = 0.1
    return {"demo": demo_inputs(), "working_capital": wc, "asset_transfer": tr, "sculpted": ProjectInputs.model_validate(sc.model_dump())}


@pytest.mark.parametrize("name", list(_cases()))
def test_formula_workbook_matches_engine(name, tmp_path):
    p = _cases()[name]
    r = run_model(p)
    sculpt = {t["name"]: {"commitment": t["commitment"], "ds": t["interest"] + t["principal"]} for t in r.tranches if t.get("sculpted")}
    data, rows = build_with_rows(p, {"case": name}, sculpt)
    src = tmp_path / f"{name}.xlsx"
    src.write_bytes(data)
    out = tmp_path / "out"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/lo", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=900)
    wb = openpyxl.load_workbook(out / f"{name}.xlsx", data_only=True)
    m = wb["Model"]
    for key, sk in model_rows().items():
        xl = [m.cell(rows[key], FIRST_COL + i).value for i in range(r.T)]
        assert not any(isinstance(v, str) for v in xl), f"{key}: error value {[v for v in xl if isinstance(v, str)][:1]}"
        xl = np.array([0.0 if v is None else float(v) for v in xl])
        eng = r.series[sk]
        scale = max(float(np.max(np.abs(eng))), 1.0)
        assert float(np.max(np.abs(xl - eng))) <= 1e-6 * scale, f"{key} differs from engine series {sk}"
    s = wb["Summary"]
    vals = {s.cell(i, 1).value: s.cell(i, 2).value for i in range(1, 60)}
    assert vals["Project IRR after subsidy (annual)"] == pytest.approx(r.metrics["project_irr_post_subsidy"], abs=1e-6)
    assert vals["Equity IRR (annual)"] == pytest.approx(r.metrics["equity_irr"], abs=1e-6)
    assert vals["Project NPV before subsidy"] == pytest.approx(r.metrics["project_npv_pre_subsidy"], rel=1e-9)
    assert abs(vals["Maximum absolute balance check"]) < 1e-3
    errors = [c.coordinate for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
    assert not errors, errors[:5]
