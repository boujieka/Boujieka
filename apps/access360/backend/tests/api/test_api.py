from __future__ import annotations

import io
import uuid

import pyotp
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app.engine.demo import demo_inputs

PW = "Correct-horse-42-battery"


def new_client() -> TestClient:
    from app.main import app
    c = TestClient(app)
    c.headers.update({"x-requested-with": "access360"})
    return c


def register(c: TestClient, kind="developer", org="Org") -> dict:
    email = f"u{uuid.uuid4().hex[:10]}@example.com"
    r = c.post("/api/auth/register", json={"email": email, "password": PW, "full_name": "Test User", "org_name": org,
                                           "org_kind": kind, "accept_terms": True})
    assert r.status_code == 201, r.text
    return {**r.json(), "password": PW}


def org_id(u: dict) -> str:
    return u["organizations"][0]["id"]


def make_project(c, oid) -> str:
    r = c.post(f"/api/orgs/{oid}/projects", json={"name": "P1", "country": "Kenya", "latitude": 0.1, "longitude": 35.0})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    r = c.put(f"/api/orgs/{oid}/projects/{pid}/inputs", json=demo_inputs().model_dump(mode="json"))
    assert r.status_code == 200, r.text
    return pid


def test_health(client):
    assert client.get("/api/health").json()["status"] == "ok"


def test_register_login_me_logout(db_ready):
    c = new_client()
    u = register(c)
    assert u["organizations"][0]["role"] == "owner"
    assert c.get("/api/auth/me").status_code == 200
    c.post("/api/auth/logout")
    c.cookies.clear()
    assert c.get("/api/auth/me").status_code == 401
    r = c.post("/api/auth/login", json={"email": u["email"], "password": "wrong-Password-1"})
    assert r.status_code == 401
    r = c.post("/api/auth/login", json={"email": u["email"], "password": PW})
    assert r.status_code == 200 and "token" not in r.json()


def test_weak_password_and_duplicate_rejected(db_ready):
    c = new_client()
    r = c.post("/api/auth/register", json={"email": "weak@example.com", "password": "short", "full_name": "x", "org_name": "o",
                                           "org_kind": "developer", "accept_terms": True})
    assert r.status_code == 422
    u = register(c)
    r = c.post("/api/auth/register", json={"email": u["email"], "password": PW, "full_name": "x", "org_name": "o",
                                           "org_kind": "developer", "accept_terms": True})
    assert r.status_code == 409


def test_csrf_header_required_for_cookie_writes(db_ready):
    c = new_client()
    u = register(c)
    del c.headers["x-requested-with"]
    r = c.post(f"/api/orgs/{org_id(u)}/projects", json={"name": "x", "country": "Ghana"})
    assert r.status_code == 403


def test_full_project_workflow(db_ready):
    c = new_client()
    u = register(c)
    oid = org_id(u)
    pid = make_project(c, oid)
    r = c.post(f"/api/orgs/{oid}/projects/{pid}/runs")
    assert r.status_code == 201, r.text
    run = r.json()
    assert run["status"] == "warning"  # integrity checks pass; synthetic and unreviewed inputs keep it below validated
    assert all(ch["passed"] for ch in run["checks"])
    assert run["is_demo"]
    rid = run["id"]
    assert len(c.get(f"/api/orgs/{oid}/projects/{pid}/runs").json()) == 1
    sens = c.post(f"/api/orgs/{oid}/projects/{pid}/sensitivity").json()
    assert len(sens["rows"]) >= 19 and all(row["checks_passed"] for row in sens["rows"])
    ms = c.post(f"/api/orgs/{oid}/projects/{pid}/minimum-subsidy", json={"min_dscr": 1.3}).json()
    assert ms["feasible"] and ms["dscr_min"] >= 1.3
    c.post(f"/api/orgs/{oid}/projects/{pid}/milestones", json={"title": "Feasibility study", "due_date": "2027-03-31"})
    x = c.get(f"/api/orgs/{oid}/projects/{pid}/runs/{rid}/export.xlsx")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Summary", "Inputs", "Model", "Run_Metrics", "Run_Annual", "Run_Quarterly", "Run_Checks"} <= set(wb.sheetnames)
    assert str(wb["Model"]["E20"].value).startswith("=")  # live formulas, not values
    cover = {wb["Summary"].cell(i, 1).value: wb["Summary"].cell(i, 2).value for i in range(1, 25)}
    assert cover["Model run"] == rid and cover["Inputs SHA-256"] == run["inputs_sha256"]
    p = c.get(f"/api/orgs/{oid}/projects/{pid}/runs/{rid}/teaser.pdf")
    assert p.status_code == 200 and p.content[:4] == b"%PDF"
    docs = c.get(f"/api/orgs/{oid}/documents").json()
    assert {d["kind"] for d in docs} == {"model_xlsx", "teaser_pdf"}
    assert all(d["inputs_sha256"] == run["inputs_sha256"] for d in docs)
    dash = c.get(f"/api/orgs/{oid}/dashboard").json()
    assert dash["real"]["projects"] == 1  # this org is not a demo org; the run is flagged demo by its inputs


def test_cross_tenant_access_blocked_at_api(db_ready):
    a, b = new_client(), new_client()
    ua, ub = register(a), register(b)
    pid = make_project(a, org_id(ua))
    assert b.get(f"/api/orgs/{org_id(ua)}/projects").status_code == 404
    assert b.get(f"/api/orgs/{org_id(ua)}/projects/{pid}").status_code == 404
    # using its own org id with the other tenant's project id must not leak either
    assert b.get(f"/api/orgs/{org_id(ub)}/projects/{pid}").status_code == 404


def test_row_level_security_in_the_database(db_ready):
    import os
    a = new_client()
    ua = register(a)
    make_project(a, org_id(ua))
    eng = create_engine(os.environ["A360_DATABASE_URL"])  # app role, not the owner
    with eng.connect() as conn:
        n_none = conn.execute(text("select count(*) from projects")).scalar()
        conn.execute(text("select set_config('app.current_org', :v, true)"), {"v": str(uuid.uuid4())})
        n_other = conn.execute(text("select count(*) from projects")).scalar()
        conn.rollback()
    with eng.connect() as conn:
        conn.execute(text("select set_config('app.current_org', :v, true)"), {"v": org_id(ua)})
        n_own = conn.execute(text("select count(*) from projects")).scalar()
        with pytest.raises(Exception):
            conn.execute(text("insert into projects (id, org_id, name, country, technology, stage, inputs_version, is_demo, marketplace_visible, created_at, updated_at) "
                              "values (:i, :o, 'x', 'y', 't', 'screening', 0, false, false, now(), now())"),
                         {"i": str(uuid.uuid4()), "o": str(uuid.uuid4())})
        conn.rollback()
    assert n_none == 0 and n_other == 0 and n_own >= 1


def test_roles_viewer_cannot_write(db_ready):
    owner, viewer = new_client(), new_client()
    uo, uv = register(owner), register(viewer)
    oid = org_id(uo)
    # adding members is privileged: requires MFA
    r = owner.post(f"/api/orgs/{oid}/members", json={"email": uv["email"], "role": "viewer"})
    assert r.status_code == 403 and "Multi factor" in r.text
    sec = owner.post("/api/auth/mfa/setup").json()["secret"]
    assert owner.post("/api/auth/mfa/verify", json={"code": pyotp.TOTP(sec).now()}).status_code == 200
    r = owner.post(f"/api/orgs/{oid}/members", json={"email": uv["email"], "role": "viewer"})
    assert r.status_code == 201, r.text
    assert viewer.get(f"/api/orgs/{oid}/projects").status_code == 200
    assert viewer.post(f"/api/orgs/{oid}/projects", json={"name": "x", "country": "Ghana"}).status_code == 403


def test_accreditation_needs_admin_verification(db_ready):
    c = new_client()
    u = register(c, kind="accredited", org="Verifier Ltd")
    assert u["organizations"][0]["accreditation_status"] == "pending"
    # an ordinary user cannot verify
    assert c.post(f"/api/admin/organizations/{org_id(u)}/accreditation", json={"status": "verified"}).status_code == 403
    # a platform administrator with MFA can
    admin = new_client()
    ua = register(admin)
    eng = create_engine(__import__("os").environ["A360_DATABASE_OWNER_URL"])
    with eng.begin() as conn:
        conn.execute(text("update users set is_superadmin = true where email = :e"), {"e": ua["email"]})
    assert admin.post(f"/api/admin/organizations/{org_id(u)}/accreditation", json={"status": "verified"}).status_code == 403
    sec = admin.post("/api/auth/mfa/setup").json()["secret"]
    admin.post("/api/auth/mfa/verify", json={"code": pyotp.TOTP(sec).now()})
    r = admin.post(f"/api/admin/organizations/{org_id(u)}/accreditation", json={"status": "verified"})
    assert r.status_code == 200 and r.json()["accreditation_status"] == "verified"


def test_dataset_upload_validation(db_ready):
    c = new_client()
    u = register(c)
    oid = org_id(u)
    meta = {"title": "Health posts", "layer_type": "health", "source_institution": "User survey", "publication_date": "2026",
            "geographic_coverage": "District X", "resolution": "Point", "licence": "Internal", "retrieval_date": "2026-10-01",
            "quality": "user_assumption", "limitations": "Partial survey"}
    good = {"type": "FeatureCollection", "features": [{"type": "Feature", "geometry": {"type": "Point", "coordinates": [30, -2]}, "properties": {"n": 1}}]}
    r = c.post(f"/api/orgs/{oid}/datasets", json={**meta, "geojson": good})
    assert r.status_code == 201 and r.json()["feature_count"] == 1
    bad = {"type": "FeatureCollection", "features": [{"type": "Feature", "geometry": {"type": "Circle", "coordinates": [0, 0]}}]}
    assert c.post(f"/api/orgs/{oid}/datasets", json={**meta, "geojson": bad}).status_code == 422
    missing = {k: v for k, v in meta.items() if k != "licence"}
    assert c.post(f"/api/orgs/{oid}/datasets", json={**missing, "geojson": good}).status_code == 422
    import json
    csv_bytes = b"name,lat,lon\nA,-1.5,30.1\nB,-1.6,30.2\n"
    r = c.post(f"/api/orgs/{oid}/datasets/csv", data={"meta": json.dumps(meta)}, files={"file": ("p.csv", csv_bytes, "text/csv")})
    assert r.status_code == 201, r.text
    assert r.json()["feature_count"] == 2


def test_optimizer_endpoint(db_ready):
    c = new_client()
    register(c)
    demo = c.get("/api/optimizer/demo").json()
    r = c.post("/api/optimizer/compare", json={k: demo[k] for k in ("settlement", "costs", "weights")})
    assert r.status_code == 200
    out = r.json()
    assert out["recommended"] and len(out["options"]) == 7


def test_seed_demo_workspace(db_ready, monkeypatch):
    monkeypatch.setenv("A360_DEMO_PASSWORD", "Demo-password-2027")
    from app.seed import DEMO_EMAIL, seed_demo
    seed_demo()
    c = new_client()
    r = c.post("/api/auth/login", json={"email": DEMO_EMAIL, "password": "Demo-password-2027"})
    assert r.status_code == 200
    oid = [o for o in r.json()["organizations"] if o["is_demo"]][0]["id"]
    dash = c.get(f"/api/orgs/{oid}/dashboard").json()
    assert dash["demo"]["projects"] == 3 and dash["real"]["projects"] == 0
