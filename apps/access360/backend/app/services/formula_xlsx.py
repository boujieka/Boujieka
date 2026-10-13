"""Live formula Excel model of one project.

Every calculation row of the engine is written as Excel formulas over quarterly columns, so the user can change an
input in the workbook and see every statement and metric recompute. The workbook is generated from the same inputs
as a model run and is verified against the engine after recalculation in LibreOffice (tests/engine/test_formula_xlsx.py).

Formula rules (as in docs/modelling-conventions.md): inputs in blue on the Inputs sheet; one formula per row, the
same in every period; interest on opening balances; no circular references; functions available in Excel 2010 and
LibreOffice only. The one exception is a sculpted tranche: its debt service schedule depends on CFADS through tax
and is written as values from the engine run (labelled), because a live formula would need iteration.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engine.inputs import CapexCategory, ProjectInputs

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD = Font(bold=True, color="FFFFFF")
NAVY = PatternFill("solid", fgColor="0B1F3A")
GREY = PatternFill("solid", fgColor="F2F4F7")
FIRST_COL = 5  # column E holds period 0
NUM = '#,##0;(#,##0);"-"'
DEC = '#,##0.0000;(#,##0.0000);"-"'


class Book:
    def __init__(self, p: ProjectInputs):
        self.p = p
        self.wb = Workbook()
        self.inp = self.wb.active
        self.inp.title = "Inputs"
        self.m = self.wb.create_sheet("Model")
        self.cells: dict[str, str] = {}  # input name -> absolute reference
        self.rows: dict[str, int] = {}
        self.r_in = 3
        self.r_m = 6
        self.pending: list = []
        self.C = p.construction_quarters
        self.T = self.C + p.operating_years * 4

    # ---------------------------------------------------------------- inputs
    def section(self, title: str) -> None:
        self.r_in += 1
        c = self.inp.cell(self.r_in, 1, title)
        c.font, c.fill = HEAD, NAVY
        self.inp.cell(self.r_in, 2).fill = NAVY
        self.inp.cell(self.r_in, 3).fill = NAVY
        self.r_in += 1

    def put(self, name: str, label: str, value, unit: str = "") -> str:
        self.inp.cell(self.r_in, 1, label)
        v = self.inp.cell(self.r_in, 2, value)
        v.font = BLUE
        self.inp.cell(self.r_in, 3, unit)
        ref = f"Inputs!$B${self.r_in}"
        self.cells[name] = ref
        self.r_in += 1
        return ref

    # ---------------------------------------------------------------- model rows
    def col(self, i: int) -> str:
        return get_column_letter(FIRST_COL + i)

    def ref(self, name: str, i: int) -> str:
        """Reference to row `name` in period i (relative to the period being written)."""
        return f"{self.col(i)}{self.rows[name]}"

    def rng(self, name: str, a: int | None = None, b: int | None = None, absolute: bool = True) -> str:
        a = 0 if a is None else a
        b = self.T - 1 if b is None else b
        r = self.rows[name]
        d = "$" if absolute else ""
        return f"{d}{self.col(a)}{d}{r}:{d}{self.col(b)}{d}{r}"

    def header(self, title: str) -> None:
        self.r_m += 1
        c = self.m.cell(self.r_m, 1, title)
        c.font, c.fill = HEAD, NAVY
        for j in range(2, FIRST_COL + self.T):
            self.m.cell(self.r_m, j).fill = NAVY
        self.r_m += 1

    def row(self, name: str, label: str, unit: str, fn, total: str | None = "sum", fmt: str = NUM, values: bool = False) -> None:
        """Register a row; formulas are written by `flush` once every row has its number (rows may refer forward)."""
        self.rows[name] = self.r_m
        self.pending.append((self.r_m, name, label, unit, fn, total, fmt, values))
        self.r_m += 1

    def flush(self) -> None:
        for args in self.pending:
            self._write(*args)
        self.pending = []

    def _write(self, r: int, name: str, label: str, unit: str, fn, total, fmt, values) -> None:
        self.m.cell(r, 1, label)
        self.m.cell(r, 2, unit)
        for i in range(self.T):
            v = fn(i)
            c = self.m.cell(r, FIRST_COL + i, v if values else (v if isinstance(v, (int, float)) else f"={v}"))
            c.number_format = fmt
            if values:
                c.font = BLUE
        if total == "sum":
            self.m.cell(r, 3, f"=SUM({self.rng(name, absolute=False)})").number_format = fmt
        elif total == "last":
            self.m.cell(r, 3, f"={self.ref(name, self.T - 1)}").number_format = fmt


def build(p: ProjectInputs, run_meta: dict | None = None, sculpt_schedule: dict | None = None) -> bytes:
    """Return the workbook bytes. `sculpt_schedule` maps a sculpted tranche name to its engine debt service list."""
    return build_with_rows(p, run_meta, sculpt_schedule)[0]


def build_with_rows(p: ProjectInputs, run_meta: dict | None = None, sculpt_schedule: dict | None = None) -> tuple[bytes, dict[str, int]]:
    b = Book(p)
    C, T = b.C, b.T
    tech, opx, mac, fin, sub, ga = p.technical, p.opex, p.macro, p.financing, p.subsidies, p.grid_arrival
    I = b.cells

    # ================================================================ Inputs sheet
    b.inp["A1"] = f"ACCESS360 AFRICA: live formula model, {p.project_name}"
    b.inp["A1"].font = Font(bold=True, size=14, color="0B1F3A")
    b.inp["A2"] = ("Blue cells are inputs. Change them and recalculate: every row on the Model sheet and every "
                   "metric on Summary follows. This is an analytical modelling treatment and does not constitute a "
                   "determination of the applicable accounting treatment under IFRS.")
    b.inp["A2"].alignment = Alignment(wrap_text=True)
    b.section("Timeline and macro")
    b.put("C", "Construction quarters", C, "quarters")
    b.put("infl", "Inflation", mac.inflation_pa, "per year")
    b.put("tidx", "Tariff indexation (share of inflation)", mac.tariff_indexation)
    b.put("fx0", "FX at financial close", mac.fx_lcy_per_usd, "LCY per USD")
    b.put("fxdep", "LCY depreciation", mac.fx_depreciation_pa, "per year")
    b.put("tax", "Tax rate", mac.tax_rate)
    b.put("gtax", "Grant income taxable (1 = yes)", 1 if mac.grant_income_taxable else 0)
    b.put("disc", "Project discount rate", mac.discount_rate_pa, "per year")
    b.put("resid", "Residual value at the horizon (share of NBV)", p.residual_value_pct_nbv)
    b.section("Technical (from the design study)")
    b.put("deliv", "Deliverable energy", tech.deliverable_energy_kwh_pa, "kWh a year")
    b.put("tl", "Technical loss", tech.technical_loss)
    b.put("cl", "Commercial loss", tech.commercial_loss)
    b.put("dshare", "Diesel share of generation", tech.diesel_share)
    b.put("dcost", "Diesel fuel cost at COD", tech.diesel_cost_per_kwh, "LCY per kWh")
    b.put("bkwh", "Battery capacity", tech.battery_kwh, "kWh")
    b.put("blife", "Battery life", tech.battery_life_years, "years")
    b.put("busd", "Battery replacement cost at COD", tech.battery_replacement_usd_per_kwh, "USD per kWh")
    b.put("bdec", "Battery cost decline", tech.battery_cost_decline_pa, "per year")
    b.section("Operating costs and collection")
    b.put("ofix", "Fixed opex at COD prices", opx.fixed_pa, "LCY a year")
    b.put("ocon", "Opex per connection at COD prices", opx.per_connection_pa, "LCY a year")
    b.put("oins", "Insurance (share of initial capex)", opx.insurance_pct_capex, "per year")
    b.put("eff", "Collection efficiency", opx.collection_efficiency)
    b.put("rdays", "Receivable days", opx.receivable_days, "days")
    b.put("pdays", "Payable days", opx.payable_days, "days")
    b.section("Customer classes")
    for k, c in enumerate(p.customers):
        b.put(f"c{k}_n", f"{c.name}: connections target", c.connections_target)
        b.put(f"c{k}_s", f"{c.name}: start (operating quarter)", c.start_quarter)
        b.put(f"c{k}_r", f"{c.name}: ramp", c.ramp_quarters, "quarters")
        b.put(f"c{k}_kwh", f"{c.name}: consumption at COD", c.kwh_per_connection_month, "kWh a month")
        b.put(f"c{k}_g", f"{c.name}: consumption growth", c.demand_growth_pa, "per year")
        b.put(f"c{k}_t", f"{c.name}: tariff at COD", c.tariff_per_kwh, "LCY per kWh")
        b.put(f"c{k}_f", f"{c.name}: fixed charge", c.fixed_charge_month, "LCY a month")
        b.put(f"c{k}_fee", f"{c.name}: connection fee", c.connection_fee, "LCY")
    b.section("Capital expenditure")
    for j, it in enumerate(p.capex):
        b.put(f"k{j}_a", f"{it.name}: amount", it.amount, it.currency)
        b.put(f"k{j}_usd", f"{it.name}: in USD (1) or LCY (0)", 1 if it.currency == "USD" else 0)
        b.put(f"k{j}_life", f"{it.name}: life", it.asset_life_years if it.category != CapexCategory.STORAGE else tech.battery_life_years, "years")
        b.put(f"k{j}_rep", f"{it.name}: replaced at end of life (1 = yes)", 1 if it.replace_at_end_of_life else 0)
        b.put(f"k{j}_rs", f"{it.name}: replacement cost share", it.replacement_cost_share)
        b.put(f"k{j}_rp", f"{it.name}: replacement price change", it.real_price_change_pa, "per year")
    b.section("Financing")
    b.put("basis", "Debt sized net of capital grant (1) or on gross capex (0)", 1 if fin.debt_basis == "capex_net_of_grant" else 0)
    b.put("lock", "Lock up DSCR", fin.lockup_dscr, "x")
    b.put("minc", "Minimum cash", fin.min_cash, "LCY")
    b.put("mra", "Maintenance reserve (1 = yes)", 1 if fin.maintenance_reserve else 0)
    for j, tr in enumerate(fin.debt):
        b.put(f"d{j}_sh", f"{tr.name}: share of capex basis", tr.share_of_capex)
        b.put(f"d{j}_usd", f"{tr.name}: in USD (1) or LCY (0)", 1 if tr.currency == "USD" else 0)
        b.put(f"d{j}_r", f"{tr.name}: interest rate", tr.rate_pa, "per year")
        b.put(f"d{j}_ten", f"{tr.name}: tenor from COD", tr.tenor_years, "years")
        b.put(f"d{j}_gr", f"{tr.name}: grace", tr.grace_quarters, "quarters")
        b.put(f"d{j}_fee", f"{tr.name}: upfront fee", tr.upfront_fee)
        b.put(f"d{j}_am", f"{tr.name}: amortisation (1 annuity, 2 straight line, 3 sculpted)",
              {"annuity": 1, "straight_line": 2, "sculpted": 3}[tr.amortisation])
    b.section("Subsidies")
    b.put("grant", "Capital grant", sub.capital_grant, "LCY")
    ms = sub.grant_milestones or []
    for j, m_ in enumerate(ms):
        b.put(f"g{j}_q", f"Grant milestone {j + 1}: model quarter", m_.quarter)
        b.put(f"g{j}_sh", f"Grant milestone {j + 1}: share", m_.share)
    b.put("rbf", "RBF per verified connection", sub.rbf_per_connection, "LCY")
    b.put("rlag", "RBF verification lag", sub.rbf_verification_lag_quarters, "quarters")
    b.put("rcap", "RBF programme cap (blank = none)", sub.rbf_cap if sub.rbf_cap is not None else 1e30, "LCY")
    b.put("ts", "Tariff support", sub.tariff_support_per_kwh, "LCY per kWh")
    b.put("tsy", "Tariff support duration", sub.tariff_support_years, "years")
    b.section("Grid arrival")
    b.put("ga_on", "Grid arrival modelled (1 = yes)", 1 if ga.enabled else 0)
    b.put("ga_y", "Grid arrives in operating year", ga.operating_year)
    b.put("ga_tr", "Treatment (1 interconnect, 2 asset transfer, 3 stranded)",
          {"interconnect": 1, "asset_transfer": 2, "stranded": 3}[ga.treatment])
    b.put("ga_comp", "Compensation (share of NBV)", ga.compensation_pct_nbv)
    b.put("ga_tf", "Tariff factor after interconnection", ga.tariff_factor_after)
    b.put("ga_df", "Demand factor after interconnection", ga.demand_factor_after)
    b.inp.column_dimensions["A"].width = 58
    b.inp.column_dimensions["B"].width = 16

    # ================================================================ Model sheet
    m = b.m
    m["A1"] = "Model: quarterly calculations (all formulas). Column C = total or closing value."
    m["A1"].font = BOLD
    m["A3"], m["A4"], m["A5"] = "Period (t)", "Operating quarter (q)", "Operating year"
    for i in range(T):
        m.cell(3, FIRST_COL + i, i)
        m.cell(4, FIRST_COL + i, f"={b.col(i)}3-{I['C']}")
        m.cell(5, FIRST_COL + i, f"=IF({b.col(i)}4>=0,INT({b.col(i)}4/4)+1,0)")
    m.column_dimensions["A"].width = 46
    m.freeze_panes = "E6"
    rows = b.rows
    rows["t"], rows["q"], rows["yr"] = 3, 4, 5

    def t_(i):
        return b.ref("t", i)

    def q_(i):
        return b.ref("q", i)

    def prev(name, i):
        return b.ref(name, i - 1) if i > 0 else "0"

    b.header("Timeline")
    b.row("t_arr", "Grid arrival period (closure or interconnection)", "t", lambda i: f"IF({I['ga_on']}=1,{I['C']}+({I['ga_y']}-1)*4,999999)", total=None)
    b.row("closes", "Operations end at grid arrival (1 = yes)", "flag", lambda i: f"IF(AND({I['ga_on']}=1,{I['ga_tr']}>1,{b.ref('t_arr', i)}<{T}),1,0)", total=None)
    b.row("op", "Operating period (1 = yes)", "flag", lambda i: f"IF({q_(i)}>=0,1,0)", total=None)
    b.row("act", "Site operating (1 = yes)", "flag", lambda i: f"IF(AND({b.ref('op', i)}=1,OR({b.ref('closes', i)}=0,{t_(i)}<{b.ref('t_arr', i)})),1,0)", total=None)
    b.row("ic", "After interconnection (1 = yes)", "flag", lambda i: f"IF(AND({I['ga_on']}=1,{I['ga_tr']}=1,{t_(i)}>={b.ref('t_arr', i)},{b.ref('op', i)}=1),1,0)", total=None)
    b.row("fx", "Exchange rate", "LCY per USD", lambda i: f"{I['fx0']}*(1+{I['fxdep']})^({t_(i)}/4)", total="last", fmt=DEC)
    b.row("cidx", "Cost index", "index", lambda i: f"(1+{I['infl']})^(MAX({q_(i)},0)/4)", total=None, fmt=DEC)
    b.row("tix", "Tariff index", "index", lambda i: f"(1+{I['infl']}*{I['tidx']})^(MAX({q_(i)},0)/4)", total=None, fmt=DEC)
    b.row("tf", "Tariff factor", "factor", lambda i: f"IF({b.ref('ic', i)}=1,{I['ga_tf']},1)", total=None, fmt=DEC)
    b.row("df", "Demand factor", "factor", lambda i: f"IF({b.ref('ic', i)}=1,{I['ga_df']},1)", total=None, fmt=DEC)

    b.header("Customers and energy")
    for k, c in enumerate(p.customers):
        b.row(f"conn{k}", f"{c.name}: connections", "number",
              lambda i, k=k: f"IF({b.ref('act', i)}=1,{I[f'c{k}_n']}*MIN(1,MAX(0,({q_(i)}-{I[f'c{k}_s']}+1)/{I[f'c{k}_r']})),0)", total="last")
        b.row(f"new{k}", f"{c.name}: new connections", "number", lambda i, k=k: f"MAX(0,{b.ref(f'conn{k}', i)}-{prev(f'conn{k}', i)})")
        b.row(f"dem{k}", f"{c.name}: demand", "kWh",
              lambda i, k=k: f"{b.ref(f'conn{k}', i)}*{I[f'c{k}_kwh']}*3*(1+{I[f'c{k}_g']})^(MAX({q_(i)},0)/4)*{b.ref('df', i)}")
    K = range(len(p.customers))
    b.row("conn", "Connections", "number", lambda i: "+".join(b.ref(f"conn{k}", i) for k in K) or "0", total="last")
    b.row("new", "New connections", "number", lambda i: "+".join(b.ref(f"new{k}", i) for k in K) or "0")
    b.row("dem", "Demand", "kWh", lambda i: "+".join(b.ref(f"dem{k}", i) for k in K) or "0")
    b.row("cap", "Capacity to customers", "kWh", lambda i: f"IF({b.ref('act', i)}=1,{I['deliv']}/4*(1-{I['tl']}),0)")
    b.row("sup", "Energy supplied", "kWh", lambda i: f"MIN({b.ref('dem', i)},{b.ref('cap', i)})")
    b.row("gen", "Energy generated", "kWh", lambda i: f"IF({b.ref('act', i)}=1,{b.ref('sup', i)}/(1-{I['tl']}),0)")
    b.row("uns", "Unserved demand", "kWh", lambda i: f"{b.ref('dem', i)}-{b.ref('sup', i)}")
    b.row("bkwh", "Energy billed", "kWh", lambda i: f"{b.ref('sup', i)}*(1-{I['cl']})")
    b.row("shr", "Share of demand served", "share", lambda i: f"IF({b.ref('dem', i)}>0,{b.ref('sup', i)}/{b.ref('dem', i)},0)", total=None, fmt=DEC)

    b.header("Revenue and operating costs")
    b.row("erev", "Energy revenue", "LCY", lambda i: "+".join(
        f"{b.ref(f'dem{k}', i)}*{b.ref('shr', i)}*(1-{I['cl']})*{I[f'c{k}_t']}*{b.ref('tix', i)}*{b.ref('tf', i)}" for k in K) or "0")
    b.row("frev", "Fixed charges", "LCY", lambda i: "+".join(
        f"{b.ref(f'conn{k}', i)}*{I[f'c{k}_f']}*3*{b.ref('tix', i)}*{b.ref('tf', i)}" for k in K) or "0")
    b.row("fee", "Connection fees", "LCY", lambda i: "+".join(f"{b.ref(f'new{k}', i)}*{I[f'c{k}_fee']}*{b.ref('tix', i)}" for k in K) or "0")
    b.row("rev", "Revenue", "LCY", lambda i: f"{b.ref('erev', i)}+{b.ref('frev', i)}+{b.ref('fee', i)}")
    b.row("colb", "Collectable billing", "LCY", lambda i: f"({b.ref('erev', i)}+{b.ref('frev', i)})*{I['eff']}")
    b.row("bad", "Bad debt", "LCY", lambda i: f"{b.ref('erev', i)}+{b.ref('frev', i)}-{b.ref('colb', i)}")
    b.row("tsup", "Tariff support", "LCY", lambda i: f"IF(AND({b.ref('act', i)}=1,{q_(i)}<{I['tsy']}*4),{b.ref('bkwh', i)}*{I['ts']},0)")
    b.row("capex", "Initial capex", "LCY", lambda i: "+".join(
        f"IF({t_(i)}<{I['C']},{I[f'k{j}_a']}/{I['C']}*IF({I[f'k{j}_usd']}=1,{b.ref('fx', i)},1),0)" for j in range(len(p.capex))) or "0")
    b.row("ofix", "Fixed opex", "LCY", lambda i: f"IF({b.ref('act', i)}=1,{I['ofix']}/4*{b.ref('cidx', i)},0)")
    b.row("ocon", "Opex per connection", "LCY", lambda i: f"IF({b.ref('act', i)}=1,{I['ocon']}/4*{b.ref('conn', i)}*{b.ref('cidx', i)},0)")
    b.row("oins", "Insurance", "LCY", lambda i: f"IF({b.ref('act', i)}=1,{I['oins']}*$C${rows['capex']}/4*{b.ref('cidx', i)},0)")
    b.row("fuel", "Diesel fuel", "LCY", lambda i: f"{b.ref('gen', i)}*{I['dshare']}*{I['dcost']}*{b.ref('cidx', i)}")
    b.row("opex", "Operating costs", "LCY", lambda i: f"{b.ref('ofix', i)}+{b.ref('ocon', i)}+{b.ref('oins', i)}+{b.ref('fuel', i)}")
    b.row("ebitda", "EBITDA", "LCY", lambda i: f"{b.ref('rev', i)}+{b.ref('tsup', i)}-{b.ref('opex', i)}-{b.ref('bad', i)}")
    b.row("recv", "Receivables", "LCY", lambda i: f"IF({b.ref('act', i)}=1,{b.ref('colb', i)}*{I['rdays']}/(365/4),0)", total="last")
    b.row("payb", "Payables", "LCY", lambda i: f"IF({b.ref('act', i)}=1,{b.ref('opex', i)}*{I['pdays']}/(365/4),0)", total="last")
    b.row("ccol", "Cash collections", "LCY", lambda i: f"{b.ref('colb', i)}-({b.ref('recv', i)}-{prev('recv', i)})+{b.ref('fee', i)}")
    b.row("copex", "Cash operating costs", "LCY", lambda i: f"{b.ref('opex', i)}-({b.ref('payb', i)}-{prev('payb', i)})")

    b.header("Replacements and maintenance reserve")
    # battery replacement every battery life after COD while operating
    b.row("brep", "Battery replacement", "LCY", lambda i: (
        f"IF(AND({b.ref('act', i)}=1,{I['bkwh']}>0,{I['busd']}>0,{q_(i)}>0,MOD({q_(i)},MAX(1,ROUND({I['blife']}*4,0)))=0),"
        f"{I['bkwh']}*{I['busd']}*(1-{I['bdec']})^({q_(i)}/4)*{b.ref('fx', i)},0)"))
    for j, it in enumerate(p.capex):
        b.row(f"rep{j}", f"{it.name}: replacement", "LCY", lambda i, j=j: (
            f"IF(AND({I[f'k{j}_rep']}=1,{b.ref('act', i)}=1,{q_(i)}>0,MOD({q_(i)},MAX(1,ROUND({I[f'k{j}_life']}*4,0)))=0),"
            f"{I[f'k{j}_a']}*{I[f'k{j}_rs']}*(1+{I[f'k{j}_rp']})^({q_(i)}/4)*IF({I[f'k{j}_usd']}=1,{b.ref('fx', i)},1),0)"))
    b.row("repl", "Replacements", "LCY", lambda i: "+".join([b.ref("brep", i)] + [b.ref(f"rep{j}", i) for j in range(len(p.capex))]))
    # helper: next replacement period (right to left) and start of the funding window (left to right)
    b.row("nxt", "Next replacement period (helper)", "t", lambda i: (
        f"IF({b.ref('repl', i + 1)}>0,{t_(i + 1)},{b.ref('nxt', i + 1)})" if i < T - 1 else "-1"), total=None)
    b.row("wst", "Reserve funding window start (helper)", "t", lambda i: (
        f"IF({b.ref('repl', i - 1)}>0,{t_(i)},{b.ref('wst', i - 1)})" if i > C else f"{I['C']}"), total=None)
    b.row("mrain", "Reserve contributions", "LCY", lambda i: (
        f"IF(AND({I['mra']}=1,{b.ref('op', i)}=1,{b.ref('repl', i)}=0,{b.ref('nxt', i)}>=0),"
        f"INDEX({b.rng('repl')},1,{b.ref('nxt', i)}+1)/MAX(1,{b.ref('nxt', i)}-{b.ref('wst', i)}),0)"))
    b.row("mraout", "Reserve releases", "LCY", lambda i: f"IF({I['mra']}=1,{b.ref('repl', i)},0)")
    b.row("mrab", "Maintenance reserve", "LCY", lambda i: f"{prev('mrab', i)}+{b.ref('mrain', i)}-{b.ref('mraout', i)}", total="last")

    b.header("Debt")
    b.row("gr_tot", "Capital grant used for debt sizing", "LCY", lambda i: f"{I['grant']}", total=None)
    b.row("dbase", "Debt basis", "LCY", lambda i: f"MAX($C${rows['capex']}-IF({I['basis']}=1,{b.ref('gr_tot', i)},0),0)", total=None)
    b.row("spsh", "Share of capex spent", "share", lambda i: f"IF($C${rows['capex']}>0,{b.ref('capex', i)}/$C${rows['capex']},0)", total=None, fmt=DEC)
    tr_rows = []
    for j, tr in enumerate(p.financing.debt):
        pre = f"d{j}"
        rate_q = f"((1+{I[f'{pre}_r']})^0.25-1)"
        fxr = lambda i, pre=pre: f"IF({I[f'{pre}_usd']}=1,{b.ref('fx', i)},1)"
        n_am = f"MAX(1,ROUND({I[f'{pre}_ten']}*4,0)-{I[f'{pre}_gr']})"
        first = f"({I['C']}+{I[f'{pre}_gr']})"
        mat = f"({I['C']}+ROUND({I[f'{pre}_ten']}*4,0)-1)"
        sculpt = sculpt_schedule.get(tr.name) if (sculpt_schedule and tr.amortisation == "sculpted") else None
        if sculpt is not None:
            b.row(f"{pre}_cmt", f"{tr.name}: commitment (engine, sculpted)", "LCY", lambda i, v=sculpt["commitment"]: v, total=None, values=True)
            b.row(f"{pre}_sds", f"{tr.name}: sculpted debt service (engine values)", "LCY", lambda i, v=sculpt["ds"]: float(v[i]), values=True)
        else:
            b.row(f"{pre}_cmt", f"{tr.name}: commitment", "LCY", lambda i, pre=pre: f"{I[f'{pre}_sh']}*{b.ref('dbase', i)}", total=None)
        b.row(f"{pre}_drw", f"{tr.name}: draws", "LCY", lambda i, pre=pre: f"{b.ref(f'{pre}_cmt', i)}*{b.ref('spsh', i)}")
        b.row(f"{pre}_ob", f"{tr.name}: opening balance (tranche currency)", "CCY", lambda i, pre=pre: prev(f"{pre}_cb", i), total=None)
        b.row(f"{pre}_int", f"{tr.name}: interest (tranche currency)", "CCY", lambda i, pre=pre, rq=rate_q: f"{b.ref(f'{pre}_ob', i)}*{rq}")
        b.row(f"{pre}_b0", f"{tr.name}: balance at end of grace (helper)", "CCY",
              lambda i, pre=pre, first=first: f"INDEX({b.rng(f'{pre}_ob')},1,{first}+1)", total=None)
        b.row(f"{pre}_pay", f"{tr.name}: level instalment (helper)", "CCY",
              lambda i, pre=pre, rq=rate_q, n=n_am: f"IF({I[f'{pre}_am']}=2,{b.ref(f'{pre}_b0', i)}/{n},IF({rq}=0,{b.ref(f'{pre}_b0', i)}/{n},{b.ref(f'{pre}_b0', i)}*{rq}/(1-(1+{rq})^(-{n}))))",
              total=None)
        if sculpt is not None:
            prin = lambda i, pre=pre, mat=mat, first=first, fxr=fxr: (
                f"IF(AND({b.ref('closes', i)}=1,{t_(i)}={b.ref('t_arr', i)}),0,IF(AND({t_(i)}>={first},{b.ref(f'{pre}_ob', i)}>0.000001),"
                f"IF({t_(i)}<{mat},{b.ref(f'{pre}_sds', i)}/{fxr(i)}-{b.ref(f'{pre}_int', i)},{b.ref(f'{pre}_ob', i)}),0))")
        else:
            prin = lambda i, pre=pre, mat=mat, first=first: (
                f"IF(AND({b.ref('closes', i)}=1,{t_(i)}={b.ref('t_arr', i)}),0,IF(AND({t_(i)}>={first},{b.ref(f'{pre}_ob', i)}>0.000001),"
                f"IF({t_(i)}<{mat},MIN({b.ref(f'{pre}_ob', i)},IF({I[f'{pre}_am']}=2,{b.ref(f'{pre}_pay', i)},{b.ref(f'{pre}_pay', i)}-{b.ref(f'{pre}_int', i)})),"
                f"{b.ref(f'{pre}_ob', i)}),0))")
        b.row(f"{pre}_pr", f"{tr.name}: principal (tranche currency)", "CCY", prin)
        b.row(f"{pre}_pp", f"{tr.name}: prepayment at closure (tranche currency)", "CCY",
              lambda i, pre=pre: f"IF(AND({b.ref('closes', i)}=1,{t_(i)}={b.ref('t_arr', i)}),{b.ref(f'{pre}_ob', i)},0)")
        b.row(f"{pre}_cb", f"{tr.name}: closing balance (tranche currency)", "CCY",
              lambda i, pre=pre, fxr=fxr: (f"IF({b.ref(f'{pre}_ob', i)}+{b.ref(f'{pre}_drw', i)}/{fxr(i)}-{b.ref(f'{pre}_pr', i)}-{b.ref(f'{pre}_pp', i)}<0.000001,0,"
                                           f"{b.ref(f'{pre}_ob', i)}+{b.ref(f'{pre}_drw', i)}/{fxr(i)}-{b.ref(f'{pre}_pr', i)}-{b.ref(f'{pre}_pp', i)})"), total="last")
        b.row(f"{pre}_intl", f"{tr.name}: interest", "LCY", lambda i, pre=pre, fxr=fxr: f"{b.ref(f'{pre}_int', i)}*{fxr(i)}")
        b.row(f"{pre}_prl", f"{tr.name}: principal", "LCY", lambda i, pre=pre, fxr=fxr: f"{b.ref(f'{pre}_pr', i)}*{fxr(i)}")
        b.row(f"{pre}_ppl", f"{tr.name}: prepayment", "LCY", lambda i, pre=pre, fxr=fxr: f"{b.ref(f'{pre}_pp', i)}*{fxr(i)}")
        b.row(f"{pre}_bl", f"{tr.name}: closing balance", "LCY", lambda i, pre=pre, fxr=fxr: f"{b.ref(f'{pre}_cb', i)}*{fxr(i)}", total="last")
        b.row(f"{pre}_fxl", f"{tr.name}: FX loss on revaluation", "LCY",
              lambda i, pre=pre, fxr=fxr: f"{b.ref(f'{pre}_ob', i)}*({fxr(i)}-{fxr(i - 1) if i > 0 else fxr(i)})")
        b.row(f"{pre}_feel", f"{tr.name}: upfront fee", "LCY", lambda i, pre=pre: f"IF({t_(i)}=0,{b.ref(f'{pre}_cmt', i)}*{I[f'{pre}_fee']},0)")
        tr_rows.append(pre)

    def tsum(suffix, i):
        return "+".join(b.ref(f"{pre}_{suffix}", i) for pre in tr_rows) or "0"
    b.row("draws", "Debt draws", "LCY", lambda i: tsum("drw", i))
    b.row("intc", "Interest during construction (capitalised)", "LCY", lambda i: f"IF({t_(i)}<{I['C']},{tsum('intl', i)},0)")
    b.row("into", "Interest expense (operations)", "LCY", lambda i: f"IF({t_(i)}>={I['C']},{tsum('intl', i)},0)")
    b.row("prin", "Principal", "LCY", lambda i: tsum("prl", i))
    b.row("prep", "Prepayment", "LCY", lambda i: tsum("ppl", i))
    b.row("debt", "Debt balance", "LCY", lambda i: tsum("bl", i), total="last")
    b.row("fxl", "FX loss on debt", "LCY", lambda i: tsum("fxl", i))
    b.row("fees", "Financing fees (capitalised)", "LCY", lambda i: tsum("feel", i))
    b.row("ds", "Scheduled debt service", "LCY",
          lambda i: f"IF(AND({b.ref('closes', i)}=1,{t_(i)}={b.ref('t_arr', i)}),0,{b.ref('into', i)})+{b.ref('prin', i)}")
    nq = fin.dsra_quarters
    b.row("dsrat", "DSRA target", "LCY", lambda i: (
        f"IF(OR({nq}=0,AND({b.ref('closes', i)}=1,{t_(i)}>={b.ref('t_arr', i)})),0,"
        + (f"SUM({b.col(i + 1)}{rows['ds']}:{b.col(min(i + nq, T - 1))}{rows['ds']})" if i < T - 1 and nq > 0 else "0") + ")"), total=None)

    b.header("Grants and RBF")
    b.row("gcash", "Capital grant receipts", "LCY", lambda i: (
        "+".join(f"IF({t_(i)}={I[f'g{j}_q']},IF(OR({t_(i)}<{I['C']},{b.ref('act', i)}=1,{b.ref('closes', i)}=0),{I['grant']}*{I[f'g{j}_sh']},0),0)"
                 for j in range(len(ms))) if ms else f"IF({t_(i)}={min(C, T - 1)},{I['grant']},0)"))
    b.row("rbfdue", "RBF due on verified connections", "LCY", lambda i: (
        f"IF(AND({b.ref('act', i)}=1,{t_(i)}-{I['rlag']}>=0),INDEX({b.rng('new')},1,{t_(i)}-{I['rlag']}+1)*{I['rbf']},0)"))
    b.row("rbf", "RBF receipts", "LCY", lambda i: (
        f"MAX(0,MIN({b.ref('rbfdue', i)},{I['rcap']}-{('SUM($E$' + str(rows.get('rbf', 0)) + ':' + b.col(i - 1) + '$' + str(rows.get('rbf', 0)) + ')') if i > 0 else '0'}))"))

    b.header("Fixed assets and depreciation")
    n_life = lambda key: f"MAX(1,ROUND({key}*4,0))"
    b.row("finv", "Capitalised interest and fees", "LCY", lambda i: f"{b.ref('intc', i)}+{b.ref('fees', i)}")
    wl_num = "+".join(f"{I[f'k{j}_a']}*IF({I[f'k{j}_usd']}=1,{I['fx0']},1)*{I[f'k{j}_life']}" for j in range(len(p.capex)))
    wl_den = "+".join(f"{I[f'k{j}_a']}*IF({I[f'k{j}_usd']}=1,{I['fx0']},1)" for j in range(len(p.capex)))
    b.row("wlife", "Weighted asset life", "years", lambda i: f"({wl_num})/MAX({wl_den},0.000001)", total=None, fmt=DEC)

    def dep_terms(i, physical_only=False):
        terms = []
        for j in range(len(p.capex)):  # initial items in service from COD
            amt = f"{I[f'k{j}_a']}*SUMPRODUCT(({b.rng('t')}<{I['C']})*({b.rng('fx')}^{I[f'k{j}_usd']}))/{I['C']}"
            terms.append(f"IF(AND({t_(i)}>={I['C']},{t_(i)}<{I['C']}+{n_life(I[f'k{j}_life'])}),{amt}/{n_life(I[f'k{j}_life'])},0)")
            terms.append(f"SUMPRODUCT({b.rng(f'rep{j}')}*({b.rng('t')}+1<={t_(i)})*({b.rng('t')}+1+{n_life(I[f'k{j}_life'])}>{t_(i)}))/{n_life(I[f'k{j}_life'])}")
        terms.append(f"SUMPRODUCT({b.rng('brep')}*({b.rng('t')}+1<={t_(i)})*({b.rng('t')}+1+{n_life(I['blife'])}>{t_(i)}))/{n_life(I['blife'])}")
        if not physical_only:
            terms.append(f"IF(AND({t_(i)}>={I['C']},{t_(i)}<{I['C']}+{n_life(b.ref('wlife', i))}),SUM({b.rng('finv')})/{n_life(b.ref('wlife', i))},0)")
        return "+".join(terms)

    b.row("depsch", "Scheduled depreciation", "LCY", lambda i: f"IF({b.ref('act', i)}=1,{dep_terms(i)},0)")
    b.row("udep", "Depreciation for unlevered tax (physical assets)", "LCY", lambda i: f"IF({b.ref('act', i)}=1,{dep_terms(i, True)},0)")
    b.row("add", "Additions", "LCY", lambda i: f"{b.ref('capex', i)}+{b.ref('repl', i)}+{b.ref('finv', i)}")
    b.row("dep", "Depreciation", "LCY", lambda i: f"MIN({b.ref('depsch', i)},{prev('nbv', i)}+{b.ref('add', i)})")
    b.row("nbvpre", "NBV before disposal (helper)", "LCY", lambda i: f"{prev('nbv', i)}+{b.ref('add', i)}-{b.ref('dep', i)}", total=None)
    b.row("dprc", "Disposal proceeds", "LCY", lambda i: (
        f"IF(AND({b.ref('closes', i)}=1,{t_(i)}={b.ref('t_arr', i)}),IF({I['ga_tr']}=2,{I['ga_comp']}*{b.ref('nbvpre', i)},0),"
        + (f"IF(AND({I['resid']}>0,{b.ref('nbvpre', i)}>0.000001,{b.ref('closes', i)}=0),{I['resid']}*{b.ref('nbvpre', i)},0)" if i == T - 1 else "0") + ")"))
    b.row("disp", "Disposed NBV", "LCY", lambda i: (
        f"IF(AND({b.ref('closes', i)}=1,{t_(i)}={b.ref('t_arr', i)}),{b.ref('nbvpre', i)},"
        + (f"IF(AND({I['resid']}>0,{b.ref('nbvpre', i)}>0.000001,{b.ref('closes', i)}=0),{b.ref('nbvpre', i)},0)" if i == T - 1 else "0") + ")"))
    b.row("dgain", "Gain (loss) on disposal", "LCY", lambda i: f"{b.ref('dprc', i)}-{b.ref('disp', i)}")
    b.row("nbv", "Net book value", "LCY", lambda i: f"{b.ref('nbvpre', i)}-{b.ref('disp', i)}", total="last")

    b.header("Deferred grant")
    b.row("gbase", "Deferred grant before release (helper)", "LCY", lambda i: f"{prev('dgr', i)}+{b.ref('gcash', i)}+{b.ref('rbf', i)}", total=None)
    b.row("ginc", "Grant income released", "LCY", lambda i: (
        f"IF(OR({b.ref('disp', i)}>0,AND({b.ref('closes', i)}=1,{t_(i)}={b.ref('t_arr', i)})),{b.ref('gbase', i)},"
        f"IF(AND({t_(i)}>={I['C']},{b.ref('gbase', i)}>0.000000000001),{b.ref('gbase', i)}/MAX(ROUND({b.ref('wlife', i)}*4,0)-({t_(i)}-{I['C']}),1),0))"))
    b.row("dgr", "Deferred grant", "LCY", lambda i: f"{b.ref('gbase', i)}-{b.ref('ginc', i)}", total="last")

    b.header("Tax")
    b.row("ebt", "Earnings before tax", "LCY", lambda i: f"{b.ref('ebitda', i)}+{b.ref('ginc', i)}-{b.ref('dep', i)}-{b.ref('into', i)}-{b.ref('fxl', i)}+{b.ref('dgain', i)}")
    b.row("txb", "Taxable income", "LCY", lambda i: f"{b.ref('ebt', i)}-IF({I['gtax']}=1,0,{b.ref('ginc', i)})")

    def loss_pool(base, pool):
        b.row(pool, f"Tax loss pool ({pool})", "LCY", lambda i: f"MAX(0,{prev(pool, i)}-{b.ref(base, i)})", total="last")

    loss_pool("txb", "lp")
    b.row("taxp", "Tax", "LCY", lambda i: f"MAX(0,{b.ref('txb', i)}-{prev('lp', i)})*{I['tax']}")
    b.row("ni", "Net income", "LCY", lambda i: f"{b.ref('ebt', i)}-{b.ref('taxp', i)}")
    b.row("utb_pre", "Unlevered taxable income before subsidy", "LCY", lambda i: f"{b.ref('ebitda', i)}-{b.ref('tsup', i)}-{b.ref('udep', i)}+{b.ref('dgain', i)}")
    loss_pool("utb_pre", "lpu1")
    b.row("utax_pre", "Unlevered tax before subsidy", "LCY", lambda i: f"MAX(0,{b.ref('utb_pre', i)}-{prev('lpu1', i)})*{I['tax']}")
    b.row("utb_post", "Unlevered taxable income after subsidy", "LCY",
          lambda i: f"{b.ref('ebitda', i)}-{b.ref('udep', i)}+{b.ref('dgain', i)}+IF({I['gtax']}=1,{b.ref('ginc', i)},0)")
    loss_pool("utb_post", "lpu2")
    b.row("utax_post", "Unlevered tax after subsidy", "LCY", lambda i: f"MAX(0,{b.ref('utb_post', i)}-{prev('lpu2', i)})*{I['tax']}")

    b.header("Cash waterfall")
    b.row("cops", "Cash from operations before tax", "LCY", lambda i: f"{b.ref('ccol', i)}+{b.ref('tsup', i)}-{b.ref('copex', i)}")
    b.row("cfads", "CFADS", "LCY", lambda i: f"{b.ref('cops', i)}-{b.ref('taxp', i)}-{b.ref('repl', i)}+{b.ref('mraout', i)}-{b.ref('mrain', i)}")
    b.row("dmv", "DSRA move to target (helper)", "LCY", lambda i: f"{b.ref('dsrat', i)}-{prev('dsra', i)}", total=None)
    b.row("floor", "Cash floor (helper)", "LCY", lambda i: f"IF({t_(i)}={I['C']}-1,{I['minc']},IF(AND({t_(i)}>={I['C']},{t_(i)}<{T - 1}),{I['minc']},0))", total=None)
    b.row("eqc", "Equity contributions (construction)", "LCY", lambda i: (
        f"IF({t_(i)}<{I['C']},MAX(0,{b.ref('capex', i)}+{b.ref('intc', i)}+{b.ref('fees', i)}+{b.ref('dmv', i)}+{b.ref('taxp', i)}"
        f"-{b.ref('draws', i)}-{b.ref('gcash', i)}-{prev('cash', i)}+{b.ref('floor', i)}),0)"))
    b.row("pre0", "Cash before reserve movements (helper)", "LCY", lambda i: (
        f"IF({t_(i)}<{I['C']},0,{prev('cash', i)}+{b.ref('cfads', i)}+{b.ref('gcash', i)}+{b.ref('rbf', i)}+{b.ref('dprc', i)}+{b.ref('draws', i)}"
        f"-{b.ref('into', i)}-{b.ref('prin', i)}-{b.ref('prep', i)}-{b.ref('capex', i)})"), total=None)
    b.row("topup", "DSRA funding (+) or release (-)", "LCY", lambda i: (
        f"IF({t_(i)}<{I['C']},{b.ref('dmv', i)},IF({b.ref('dmv', i)}<=0,{b.ref('dmv', i)},MIN({b.ref('dmv', i)},MAX(0,{b.ref('pre0', i)}-{b.ref('floor', i)}))))"))
    b.row("pre1", "Cash after reserve funding (helper)", "LCY", lambda i: f"{b.ref('pre0', i)}-IF({t_(i)}<{I['C']},0,{b.ref('topup', i)})", total=None)
    b.row("ddraw", "DSRA draw", "LCY", lambda i: (
        f"IF({t_(i)}<{I['C']},0,IF(AND({b.ref('pre1', i)}<{b.ref('floor', i)},{prev('dsra', i)}+{b.ref('topup', i)}>0),"
        f"MIN({prev('dsra', i)}+{b.ref('topup', i)},{b.ref('floor', i)}-{b.ref('pre1', i)}),0))"))
    b.row("dsra", "DSRA balance", "LCY", lambda i: f"{prev('dsra', i)}+{b.ref('topup', i)}-{b.ref('ddraw', i)}", total="last")
    b.row("supp", "Equity support", "LCY", lambda i: f"IF({t_(i)}<{I['C']},0,MAX(0,{b.ref('floor', i)}-{b.ref('pre1', i)}-{b.ref('ddraw', i)}))")
    b.row("dscrq", "DSCR (quarter)", "x", lambda i: f"IF({b.ref('ds', i)}>0.000001,{b.ref('cfads', i)}/{b.ref('ds', i)},\"\")", total=None, fmt=DEC)
    b.row("dist", "Distributions", "LCY", lambda i: (
        f"IF({t_(i)}<{I['C']},0,IF(AND({t_(i)}<{T - 1},OR(AND({b.ref('ds', i)}>0.000001,N({b.ref('dscrq', i)})<{I['lock']}),"
        f"{b.ref('dsra', i)}<{b.ref('dsrat', i)}-0.000001)),0,MAX(0,{b.ref('pre1', i)}+{b.ref('ddraw', i)}+{b.ref('supp', i)}-{b.ref('floor', i)})))"))
    b.row("cash", "Cash", "LCY", lambda i: (
        f"IF({t_(i)}<{I['C']},{prev('cash', i)}+{b.ref('draws', i)}+{b.ref('gcash', i)}+{b.ref('eqc', i)}-({b.ref('capex', i)}+{b.ref('intc', i)}+{b.ref('fees', i)}+{b.ref('dmv', i)}+{b.ref('taxp', i)}),"
        f"{b.ref('pre1', i)}+{b.ref('ddraw', i)}+{b.ref('supp', i)}-{b.ref('dist', i)})"), total="last")

    b.header("Balance sheet and checks")
    b.row("assets", "Total assets", "LCY", lambda i: f"{b.ref('cash', i)}+{b.ref('dsra', i)}+{b.ref('mrab', i)}+{b.ref('nbv', i)}+{b.ref('recv', i)}", total="last")
    b.row("liab", "Total liabilities", "LCY", lambda i: f"{b.ref('debt', i)}+{b.ref('dgr', i)}+{b.ref('payb', i)}", total="last")
    b.row("equity", "Total equity", "LCY", lambda i: (
        f"SUM($E${rows['eqc']}:{b.col(i)}{rows['eqc']})+SUM($E${rows['supp']}:{b.col(i)}{rows['supp']})"
        f"+SUM($E${rows['ni']}:{b.col(i)}{rows['ni']})-SUM($E${rows['dist']}:{b.col(i)}{rows['dist']})"), total="last")
    b.row("bschk", "Balance check (assets - liabilities - equity)", "LCY", lambda i: f"{b.ref('assets', i)}-{b.ref('liab', i)}-{b.ref('equity', i)}", total="last", fmt=DEC)

    b.header("Project and equity cash flows")
    b.row("basecf", "Project cash flow before tax and subsidy (helper)", "LCY", lambda i: (
        f"IF({t_(i)}<{I['C']},-{b.ref('capex', i)},{b.ref('ccol', i)}-{b.ref('copex', i)}-{b.ref('repl', i)})+{b.ref('dprc', i)}"))
    b.row("pcf_pre", "Project cash flow before subsidy", "LCY", lambda i: f"{b.ref('basecf', i)}-{b.ref('utax_pre', i)}")
    b.row("pcf_post", "Project cash flow after subsidy", "LCY",
          lambda i: f"{b.ref('basecf', i)}+{b.ref('tsup', i)}+{b.ref('gcash', i)}+{b.ref('rbf', i)}-{b.ref('utax_post', i)}")
    b.row("ecf", "Equity cash flow", "LCY", lambda i: f"-{b.ref('eqc', i)}-{b.ref('supp', i)}+{b.ref('dist', i)}")

    b.flush()

    # ================================================================ Summary
    s = b.wb.create_sheet("Summary", 0)
    s["A1"] = f"ACCESS360 AFRICA: {p.project_name}"
    s["A1"].font = Font(bold=True, size=14, color="0B1F3A")
    meta = run_meta or {}
    lines = [("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))] + [(k, str(v)) for k, v in meta.items()]
    for n, (k, v) in enumerate(lines, start=3):
        s.cell(n, 1, k).font = BOLD
        s.cell(n, 2, v)
    r0 = 4 + len(lines)
    s.cell(r0, 1, "Metric (live formulas)").font = BOLD
    rq = f"((1+{I['disc']})^0.25-1)"
    yrs = p.operating_years
    met = [
        ("Project IRR before subsidy (annual)", f"=(1+IRR(Model!{b.rng('pcf_pre')},0.02))^4-1", "0.0%"),
        ("Project IRR after subsidy (annual)", f"=(1+IRR(Model!{b.rng('pcf_post')},0.02))^4-1", "0.0%"),
        ("Equity IRR (annual)", f"=(1+IRR(Model!{b.rng('ecf')},0.02))^4-1", "0.0%"),
        ("Project NPV before subsidy", f"=NPV({rq},Model!{b.rng('pcf_pre')})", NUM),
        ("Project NPV after subsidy", f"=NPV({rq},Model!{b.rng('pcf_post')})", NUM),
        ("Initial capex", f"=Model!C{rows['capex']}", NUM),
        ("Equity support (total)", f"=Model!C{rows['supp']}", NUM),
        ("Maximum absolute balance check", f"=MAX(MAX(Model!{b.rng('bschk')}),-MIN(Model!{b.rng('bschk')}))", DEC),
    ]
    for n, (k, f, fmt) in enumerate(met, start=r0 + 1):
        s.cell(n, 1, k)
        c = s.cell(n, 2, f)
        c.number_format = fmt
    ra = r0 + len(met) + 2
    s.cell(ra, 1, "Operating year").font = BOLD
    s.cell(ra, 2, "CFADS").font = BOLD
    s.cell(ra, 3, "Debt service").font = BOLD
    s.cell(ra, 4, "DSCR").font = BOLD
    for y in range(1, yrs + 1):
        rr = ra + y
        s.cell(rr, 1, y)
        s.cell(rr, 2, f"=SUMIF(Model!{b.rng('yr')},{y},Model!{b.rng('cfads')})").number_format = NUM
        s.cell(rr, 3, f"=SUMIF(Model!{b.rng('yr')},{y},Model!{b.rng('ds')})").number_format = NUM
        s.cell(rr, 4, f'=IF(C{rr}>0.000001,B{rr}/C{rr},"")').number_format = "0.00x"
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 22
    b.wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    b.wb.save(buf)
    return buf.getvalue(), dict(b.rows)


def model_rows() -> dict[str, str]:
    """Mapping from workbook row keys to engine series keys, used by the verification test."""
    return {"conn": "connections", "dem": "demand_kwh", "sup": "supplied_kwh", "bkwh": "billed_kwh", "rev": "revenue",
            "bad": "bad_debt", "tsup": "tariff_support", "opex": "opex", "ebitda": "ebitda", "recv": "receivables",
            "payb": "payables", "ccol": "collections", "capex": "capex", "repl": "replacement_capex", "mrab": "mra",
            "debt": "debt_balance", "into": "interest_expense", "prin": "principal", "fxl": "fx_loss", "ds": "debt_service",
            "gcash": "grant_receipts", "rbf": "rbf_receipts", "dep": "depreciation", "nbv": "nbv", "ginc": "grant_income",
            "dgr": "deferred_grant", "taxp": "tax", "ni": "net_income", "cfads": "cfads", "dsra": "dsra",
            "supp": "equity_support", "eqc": "equity_contributions", "dist": "distributions", "cash": "cash",
            "pcf_pre": "project_cf_pre_subsidy", "pcf_post": "project_cf_post_subsidy", "ecf": "equity_cf"}
