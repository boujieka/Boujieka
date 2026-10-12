"""Excel model export and one page investment teaser. Both are built only from a stored model run, and both carry
the engine version, the inputs hash, the run identifier and the generation date."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

NAVY = "0B1F3A"
ACCOUNTING_NOTE = ("This is an analytical modelling treatment and does not constitute a determination of the applicable "
                   "accounting treatment under IFRS.")
STATUS_TEXT = {"validated": "Validated: integrity checks passed, no warnings, assumptions reviewed by a named person",
               "warning": "Warning: integrity checks passed, but warnings remain or the assumptions are not reviewed",
               "failed_checks": "Failed checks: do not rely on the outputs",
               "incomplete_inputs": "Incomplete inputs: outputs are not meaningful"}


def _flatten(d, prefix=""):
    out = []
    if isinstance(d, dict):
        for k, v in d.items():
            out += _flatten(v, f"{prefix}{k}.")
    elif isinstance(d, list):
        for i, v in enumerate(d):
            out += _flatten(v, f"{prefix}{i}.")
    else:
        out.append((prefix[:-1], d))
    return out


def _stamp(project, run) -> list[tuple[str, str]]:
    return [("Project", project.name), ("Country", project.country), ("Model status", STATUS_TEXT.get(run.status, run.status)),
            ("Engine version", run.engine_version), ("Model run", str(run.id)), ("Inputs version", str(run.inputs_version)),
            ("Inputs SHA-256", run.inputs_sha256), ("Run date", run.created_at.strftime("%Y-%m-%d %H:%M UTC")),
            ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
            ("Data", "SYNTHETIC DEMO DATA: not a real project" if run.is_demo else "User supplied inputs (see provenance)")]


def model_xlsx(project, run) -> bytes:
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY)

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font, c.fill = head, fill

    ws = wb.active
    ws.title = "Cover"
    ws["A1"] = "ACCESS360 AFRICA: ACCESS FINANCE model export"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    for i, (k, v) in enumerate(_stamp(project, run), start=3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    r0 = ws.max_row + 2
    notes = ["This workbook holds the values of one engine run. Formulas are not exported in this version: the calculation "
             "logic is documented in docs/financial-engine-spec.md and reproduced by the platform for any inputs.",
             ACCOUNTING_NOTE,
             "Nothing here is investment, legal, tax or accounting advice."]
    for i, n in enumerate(notes):
        ws.cell(r0 + i, 1, n).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 90

    ws = wb.create_sheet("Inputs")
    header(ws, ["Input", "Value"])
    for k, v in _flatten(run.inputs_snapshot):
        ws.append([k, v if not isinstance(v, bool) else str(v)])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 55, 30

    ws = wb.create_sheet("Metrics")
    header(ws, ["Metric", "Value"])
    for k, v in run.metrics.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 40

    ws = wb.create_sheet("Sources_Uses")
    header(ws, ["Item", "LCY"])
    for k, v in run.sources_uses.items():
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34

    ws = wb.create_sheet("Annual")
    keys = list(run.annual.keys())
    header(ws, keys)
    for i in range(len(run.annual["year"])):
        ws.append([run.annual[k][i] for k in keys])

    ws = wb.create_sheet("Quarterly")
    keys = list(run.quarterly.keys())
    header(ws, keys)
    for i in range(len(run.quarterly["date"])):
        ws.append([run.quarterly[k][i] for k in keys])
    ws.freeze_panes = "B2"

    ws = wb.create_sheet("Checks")
    header(ws, ["Code", "Check", "Passed", "Detail"])
    for c in run.checks:
        ws.append([c["code"], c["label"], "OK" if c["passed"] else "FAIL", c["detail"]])
    header(ws, ["Code", "Warning", "", "Detail"])
    for c in run.warnings:
        ws.append([c["code"], c["label"], "", c["detail"]])
    for m in run.missing_inputs:
        ws.append(["MISSING", m, "", ""])
    for wsx in wb.worksheets[1:]:
        for col in range(1, wsx.max_column + 1):
            if wsx.column_dimensions[get_column_letter(col)].width in (None, 13):
                wsx.column_dimensions[get_column_letter(col)].width = 16
        for row in wsx.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00;(#,##0.00);\"-\"" if abs(c.value) < 10 else "#,##0;(#,##0);\"-\""
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _money(x, ccy="USD"):
    if x is None:
        return "Not available"
    if abs(x) >= 1e6:
        return f"{ccy} {x / 1e6:,.2f}m"
    return f"{ccy} {x:,.0f}"


def _pct(x):
    return "Not available" if x is None else f"{x:.1%}"


def teaser_pdf(project, run, milestones) -> bytes:
    inp, M = run.inputs_snapshot, run.metrics
    fx = inp["macro"]["fx_lcy_per_usd"]
    ccy = inp["currency"]
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7.5, leading=9.5)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=8.5, leading=11)
    h = ParagraphStyle("h", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#" + NAVY), spaceBefore=4, spaceAfter=2)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=10 * mm,
                            title=f"Investment teaser: {project.name}")
    el = []
    el.append(Paragraph(f"<b>{project.name}</b>", ParagraphStyle("t", parent=styles["Title"], fontSize=16, alignment=0,
                                                                   textColor=colors.HexColor("#" + NAVY))))
    if run.is_demo:
        el.append(Paragraph("<font color='#B00020'><b>SYNTHETIC DEMO DATA: not a real project and not an offer</b></font>", body))
    el.append(Paragraph(f"Investment teaser | Model status: <b>{STATUS_TEXT.get(run.status, run.status)}</b>", small))
    ok_outputs = run.status in ("validated", "warning")
    tech = inp["technical"]
    conns = M.get("connections_target")
    yr = run.annual
    demand5 = sum(yr["energy_sold_kwh"][:5]) / 5 if yr.get("energy_sold_kwh") else None
    debt = inp["financing"]["debt"]
    su = run.sources_uses
    fin = "; ".join(f"{d['name']} ({d['currency']}, {d['share_of_capex']:.0%} of capex basis, {d['rate_pa']:.1%}, {d['tenor_years']:g} years)" for d in debt) or "No debt"
    grants = inp["subsidies"]
    rows = [
        ["Location", f"{project.country}" + (f" ({project.latitude:.4f}, {project.longitude:.4f})" if project.latitude is not None else "")],
        ["Technology", f"Solar PV {tech['pv_kwp']:g} kWp, battery {tech['battery_kwh']:g} kWh" + (f", diesel {tech['diesel_kw']:g} kW" if tech['diesel_kw'] else "") + (f", network {tech['network_km']:g} km" if tech['network_km'] else "")],
        ["Customers", ", ".join(f"{c['name']} ({c['connections_target']})" for c in inp["customers"]) + f"; {conns:,.0f} connections"],
        ["Expected demand", f"{demand5:,.0f} kWh sold a year, average of years 1 to 5" if demand5 else "Not available"],
        ["Investment cost", f"{_money(M['capex_total_usd'])} ({_money(M['capex_total'], ccy)}); {_money(M.get('cost_per_connection_usd'))} per connection"],
        ["Financing structure", f"Debt {_money(su['debt'] / fx)}; grants in construction {_money(su['grants_in_construction'] / fx)}; equity {_money(su['equity'] / fx)}. {fin}"],
        ["Subsidies", f"Capital grant {_money(grants['capital_grant'] / fx)}; RBF {_money(grants['rbf_per_connection'] / fx)} per verified connection" + (f" (cap {_money(grants['rbf_cap'] / fx)})" if grants.get('rbf_cap') else "")],
        ["Viability gap", f"{_money((M.get('viability_gap_per_connection') or 0) / fx)} per connection (present value, before subsidies)"],
    ]
    if ok_outputs:
        rows += [["Expected returns", f"Project IRR {_pct(M.get('project_irr_post_subsidy'))} after subsidies ({_pct(M.get('project_irr_pre_subsidy'))} before); "
                                      f"equity IRR {_pct(M.get('equity_irr'))}; minimum annual DSCR {M['dscr_min']:.2f}x" if M.get('dscr_min') is not None else
                                      f"Project IRR {_pct(M.get('project_irr_post_subsidy'))} after subsidies; equity IRR {_pct(M.get('equity_irr'))}"]]
    else:
        rows += [["Expected returns", "Not shown: the model has failed checks or incomplete inputs"]]
    rows += [["Development status", project.stage.replace("_", " ").capitalize()]]
    t = Table([[Paragraph(f"<b>{a}</b>", body), Paragraph(str(b), body)] for a, b in rows], colWidths=[38 * mm, 144 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D5DD")),
                           ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F4F7"))]))
    el += [Spacer(1, 3 * mm), t]
    el.append(Paragraph("Key risks", h))
    risks = [w["label"] + (f" ({w['detail']})" if w.get("detail") else "") for w in run.warnings if w["code"] not in ("W_PROV", "W_REVIEW")]
    if any(d["currency"] == "USD" for d in debt):
        risks.append("Hard currency debt against local currency revenue")
    if not inp["grid_arrival"]["enabled"]:
        risks.append("Grid arrival not modelled in the base case: see sensitivity")
    risks.append("Demand ramp and collection performance against plan")
    el.append(Paragraph("<br/>".join("&bull; " + r for r in risks[:7]), body))
    el.append(Paragraph("Next investment milestones", h))
    open_ms = [m for m in milestones if m.status != "done"][:5]
    el.append(Paragraph("<br/>".join(f"&bull; {m.title}" + (f" (due {m.due_date.isoformat()})" if m.due_date else "") for m in open_ms)
                        or "No milestones recorded", body))
    el.append(Paragraph("Sources and assumptions", h))
    prov = inp.get("provenance", "user_estimate")
    stamp = "; ".join(f"{k}: {v}" for k, v in _stamp(project, run) if k in ("Engine version", "Model run", "Inputs SHA-256", "Generated"))
    el.append(Paragraph(f"Input provenance: {prov.replace('_', ' ')}. Figures are outputs of the ACCESS FINANCE engine run identified below; "
                        f"USD figures at the financial close rate of {fx:,.2f} {ccy} per USD. {ACCOUNTING_NOTE} This teaser is information, "
                        f"not an offer of securities. {stamp}.", small))
    doc.build(el)
    return buf.getvalue()
