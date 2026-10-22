"""Documented official-data import workflow (Master prompt §7, §13).

For sources with no machine access (AfDB portal when unreachable, national budget documents,
ministry reports, institutional uploads) an authorised Data Contributor uploads a CSV or Excel
file built from an official document. Every row MUST cite its source document (URL or document
reference) and reporting date; rows are staged as source records and go through the same
pipeline, validation and review as automated connectors.

Templates: ap2i/data_templates/projects_import_template.csv and budget_lines_template.csv.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from decimal import Decimal

from app.connectors.base import (
    Connector,
    Money,
    NormalizedAgreement,
    NormalizedRecord,
    RawRecord,
    SyncParams,
    parse_date,
    parse_decimal,
)

PROJECT_COLUMNS = [
    "country_iso3", "source_project_id", "title", "title_language", "sector_code", "description", "objectives",
    "executing_agency", "implementing_agency", "status", "approval_date", "start_date", "planned_end_date",
    "actual_end_date", "original_cost", "revised_cost", "cost_currency", "physical_progress_pct",
    "physical_progress_date", "operational_status", "funder_code", "financing_instrument", "financing_amount",
    "financing_currency", "financing_date", "cumulative_disbursed", "cumulative_disbursed_as_of",
    "latitude", "longitude", "location_name", "source_url", "source_document", "reporting_date", "notes",
]
REQUIRED_PROJECT = ["country_iso3", "source_project_id", "title", "source_url", "reporting_date"]

BUDGET_COLUMNS = [
    "country_iso3", "budget_year", "budget_line_code", "project_identifier", "project_label",
    "administrative_classification", "program_classification", "funding_source", "funding_type", "currency",
    "approved_appropriation", "revised_appropriation", "commitment", "payment", "source_url",
    "source_document", "reporting_date",
]
REQUIRED_BUDGET = ["country_iso3", "budget_year", "budget_line_code", "currency", "source_url", "reporting_date"]

VALID_STATUS = {"pipeline", "approved", "active", "completed", "closed", "suspended", "cancelled", "dropped", "unknown"}
VALID_OPS = {"operational", "partially_operational", "not_operational", "unknown"}


class UploadValidationError(Exception):
    def __init__(self, errors: list[dict]):
        super().__init__(f"{len(errors)} invalid row(s)")
        self.errors = errors


def read_table(content: bytes, filename: str) -> list[dict]:
    """CSV (UTF-8, ; or ,) or Excel (.xlsx). Values returned as stripped strings."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for r in rows[1:]:
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            out.append({header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header)})
        return out
    if name.endswith(".csv") or name.endswith(".txt"):
        text = content.decode("utf-8-sig")
        sample = text[:2048]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        return [{(k or "").strip(): (v or "").strip() for k, v in row.items()}
                for row in csv.DictReader(io.StringIO(text), delimiter=delimiter)]
    raise UploadValidationError([{"row": 0, "error": "Unsupported file type (use .csv or .xlsx)"}])


def validate_project_rows(rows: list[dict]) -> list[dict]:
    errors = []
    for i, r in enumerate(rows, start=2):
        for c in REQUIRED_PROJECT:
            if not r.get(c):
                errors.append({"row": i, "field": c, "error": "required"})
        if r.get("status") and r["status"] not in VALID_STATUS:
            errors.append({"row": i, "field": "status", "error": f"must be one of {sorted(VALID_STATUS)}"})
        if r.get("operational_status") and r["operational_status"] not in VALID_OPS:
            errors.append({"row": i, "field": "operational_status", "error": f"must be one of {sorted(VALID_OPS)}"})
        for c in ("original_cost", "revised_cost", "financing_amount", "cumulative_disbursed", "physical_progress_pct",
                  "latitude", "longitude"):
            if r.get(c) and parse_decimal(r[c]) is None:
                errors.append({"row": i, "field": c, "error": "not a number"})
        for c in ("approval_date", "start_date", "planned_end_date", "actual_end_date", "financing_date",
                  "cumulative_disbursed_as_of", "reporting_date", "physical_progress_date"):
            if r.get(c) and parse_date(r[c]) is None:
                errors.append({"row": i, "field": c, "error": "date must be YYYY-MM-DD"})
        pp = parse_decimal(r.get("physical_progress_pct"))
        if pp is not None and not (Decimal(0) <= pp <= Decimal(100)):
            errors.append({"row": i, "field": "physical_progress_pct", "error": "must be between 0 and 100"})
        if r.get("financing_amount") and not (r.get("financing_currency") and r.get("funder_code")):
            errors.append({"row": i, "field": "financing_currency", "error": "amount needs funder_code and currency"})
        if r.get("original_cost") and not r.get("cost_currency"):
            errors.append({"row": i, "field": "cost_currency", "error": "cost needs a currency"})
    return errors


class ManualUploadConnector(Connector):
    key = "manual_upload"
    name = "Official-data import workflow (CSV/Excel built from official documents)"
    partner_code = None
    access_method = "manual_upload"
    usage_terms = "Contributor certifies that each row reproduces an official document cited in source_url."
    update_frequency = "On upload"
    default_schedule = None

    def __init__(self, config=None, client=None, rows: list[dict] | None = None, uploaded_by: str | None = None):
        super().__init__(config, client)
        self.rows = rows or []
        self.uploaded_by = uploaded_by

    def fetch(self, params: SyncParams) -> Iterator[RawRecord]:
        for r in self.rows:
            yield RawRecord(
                source_identifier=f"upload:{r['country_iso3'].upper()}:{r['source_project_id']}",
                record_type="upload_project_row",
                country_iso3=r["country_iso3"].upper(),
                payload={**r, "_uploaded_by": self.uploaded_by},
                source_url=r.get("source_url"),
                source_last_updated=r.get("reporting_date"),
            )

    def normalize(self, raw: RawRecord) -> NormalizedRecord | None:
        r = raw.payload
        rec = NormalizedRecord(
            country_iso3=r["country_iso3"].upper(),
            title=r["title"],
            title_language=r.get("title_language") or None,
            description=r.get("description") or None,
            objectives=r.get("objectives") or None,
            sector_code=(r.get("sector_code") or "").upper() or None,
            status=r.get("status") or "unknown",
            status_label=r.get("status") or None,
            executing_agency=r.get("executing_agency") or None,
            implementing_agency=r.get("implementing_agency") or None,
            approval_date=parse_date(r.get("approval_date")),
            start_date=parse_date(r.get("start_date")),
            planned_end_date=parse_date(r.get("planned_end_date")),
            actual_end_date=parse_date(r.get("actual_end_date")),
            original_cost=parse_decimal(r.get("original_cost")),
            cost_currency=(r.get("cost_currency") or "").upper() or None,
        )
        if r.get("revised_cost"):
            rec.observations.append(("revised_estimated_cost", "revised_cost", parse_decimal(r["revised_cost"])))
        if r.get("physical_progress_pct"):
            rec.observations.append(("physical_progress_pct", "physical_progress_pct", parse_decimal(r["physical_progress_pct"])))
            rec.observations.append(("physical_progress_date", "physical_progress_date",
                                     parse_date(r.get("physical_progress_date")) or parse_date(r.get("reporting_date"))))
        if r.get("operational_status"):
            rec.observations.append(("operational_status", "operational_status", r["operational_status"]))
        if r.get("source_document"):
            rec.documents.append({"title": r["source_document"], "url": r.get("source_url"), "doc_type": "official_document",
                                  "date": parse_date(r.get("reporting_date")), "language": r.get("title_language") or None})
        if r.get("latitude") and r.get("longitude"):
            rec.locations.append({"name": r.get("location_name"), "lat": parse_decimal(r["latitude"]),
                                  "lon": parse_decimal(r["longitude"]), "precision": "as stated in source document"})
        if r.get("funder_code"):
            amount = parse_decimal(r.get("financing_amount"))
            cur = (r.get("financing_currency") or "").upper() or None
            fdate = parse_date(r.get("financing_date"))
            agr = NormalizedAgreement(
                partner_code=r["funder_code"].upper(),
                source_identifier=f"upload:{r['source_project_id']}:{r['funder_code'].upper()}",
                name=r["title"], instrument_code=(r.get("financing_instrument") or "").upper() or None,
                instrument_label=r.get("financing_instrument") or None, approval_date=fdate,
                status=r.get("status") or "unknown", currency=cur, amount_original=amount,
                amount_label="Financing amount stated in the cited official document" if amount is not None else None,
            )
            if amount is not None and cur:
                agr.commitments.append(Money(amount, cur, fdate, f"{raw.source_identifier}:financing", "original",
                                             "uploaded financing amount"))
            if r.get("cumulative_disbursed"):
                agr.reported_cumulative_disbursed = parse_decimal(r["cumulative_disbursed"])
                agr.reported_cumulative_as_of = parse_date(r.get("cumulative_disbursed_as_of")) or parse_date(r.get("reporting_date"))
                agr.reported_cumulative_label = "Cumulative disbursement stated in the cited official document"
            rec.agreements.append(agr)
        return rec
