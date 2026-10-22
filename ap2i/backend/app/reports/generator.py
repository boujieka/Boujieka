"""Reporting engine (Master prompt §27): PDF, Excel and CSV from database queries only.

Every report carries: AP²I logo (PDF/Excel), title, reporting period, generation date, data
coverage, sources, methodology, limitations and a unique report identifier.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.config import get_settings
from app.finance.indicators import agreement_position
from app.finance.portfolio import compute_portfolio
from app.models import (
    Country,
    FinancingAgreement,
    FinancingPartner,
    Project,
    ProjectAlert,
    RecoveryPlan,
    Report,
    Sector,
    SourceConnector,
)
from app.performance.quality import dashboard as dq_dashboard

LOGO = Path(__file__).resolve().parents[3] / "brand" / "ap2i-logo.png"
NAVY, GREEN, GOLD = "082B50", "087F5B", "D9A323"

REPORT_TYPES = {
    "country_investment": "Country Investment Report",
    "sector_investment": "Sector Investment Report",
    "annual_portfolio": "Annual Portfolio Report",
    "donor_portfolio": "Donor Portfolio Report",
    "disbursement": "Disbursement Report",
    "project_performance": "Project Performance Report",
    "project_rescue": "Project Rescue Report",
    "operational_infrastructure": "Operational Infrastructure Report",
    "data_quality": "Data Quality Report",
}
METHODOLOGY = ("Figures come only from records imported into AP²I from official sources, with provenance. Amounts are "
               "kept in original currencies; USD equivalents use the annual average rate of each transaction's year "
               "(World Bank WDI PA.NUS.FCRF; IMF ER for SDR). Projects are counted once per category per year. "
               "See the public methodology page for the full rules (methodology version 1.0).")
LIMITATIONS = ["Coverage is limited to the connected sources; this is not the complete national investment portfolio.",
               "Missing values are shown as empty — never replaced by zero.",
               "Alerts are signals for verification; they do not establish wrongdoing."]


@dataclass
class Table:
    title: str
    headers: list[str]
    rows: list[list]


@dataclass
class ReportData:
    report_type: str
    title: str
    period: str
    coverage: str
    sources: list[str]
    tables: list[Table] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    report_uid: str = field(default_factory=lambda: "RPT-" + uuid.uuid4().hex[:12].upper())
    generated_at: datetime = field(default_factory=lambda: datetime.now(timezone.utc))


def _sources(db: Session) -> list[str]:
    return [f"{c.name} — {c.base_url or ''} (last success {c.last_success_at:%Y-%m-%d})"
            for c in db.scalars(select(SourceConnector).where(SourceConnector.last_success_at.is_not(None)))]


def _projects(db: Session, country: Country, tenant_cond, sector: str | None = None):
    q = select(Project).where(Project.country_id == country.id, Project.merged_into_id.is_(None), tenant_cond).options(
        selectinload(Project.agreements).selectinload(FinancingAgreement.partner), selectinload(Project.sector))
    if sector:
        q = q.join(Sector, Sector.id == Project.sector_id).where(Sector.code == sector)
    return list(db.scalars(q.order_by(Project.ap2i_id)).unique())


def build(db: Session, report_type: str, params: dict, principal) -> ReportData:
    from app.security.deps import tenant_filter

    tenant_cond = tenant_filter(Project, principal)
    plan_cond = tenant_filter(RecoveryPlan, principal)
    if report_type not in REPORT_TYPES:
        raise ValueError(f"Unknown report type {report_type}")
    iso3 = (params.get("country") or get_settings().default_country).upper()
    country = db.scalar(select(Country).where(Country.iso3 == iso3))
    if country is None:
        raise ValueError(f"Unknown country {iso3}")
    y0, y1 = int(params.get("year_from") or 2011), int(params.get("year_to") or 2026)
    sector = params.get("sector")
    title = f"{REPORT_TYPES[report_type]} — {country.name_en}"
    projects = _projects(db, country, tenant_cond, sector)
    data = ReportData(report_type, title, f"{y0}–{y1}",
                      f"{len(projects)} documented projects for {country.name_en} in AP²I (connected sources only)", _sources(db))

    def project_rows(items):
        rows = []
        for p in items:
            for a in p.agreements or [None]:
                pos = agreement_position(db, a) if a else None
                rows.append([p.ap2i_id, p.title, p.sector.name_en if p.sector else "", p.status,
                             a.partner.short_name if a and a.partner else "", a.source_identifier if a else "",
                             str(a.approval_date or "") if a else "", pos.currency if pos else "",
                             str(pos.committed) if pos and pos.committed is not None else "",
                             str(pos.disbursed) if pos and pos.disbursed is not None else "",
                             str(pos.disbursed_as_of or "") if pos else ""])
        return rows

    headers = ["AP²I ID", "Title", "Sector", "Status", "Partner", "Source ID", "Approval", "Currency", "Committed", "Disbursed", "Disbursed as of"]
    if report_type in ("country_investment", "sector_investment", "project_performance"):
        data.tables.append(Table("Projects and financing", headers, project_rows(projects)))
    if report_type in ("country_investment", "annual_portfolio", "sector_investment", "disbursement", "donor_portfolio"):
        partner_id = None
        if report_type == "donor_portfolio" and params.get("partner"):
            partner_id = db.scalar(select(FinancingPartner.id).where(FinancingPartner.code == params["partner"].upper()))
            data.title += f" — {params['partner'].upper()}"
        sector_id = db.scalar(select(Sector.id).where(Sector.code == sector)) if sector else None
        pf = compute_portfolio(db, iso3, y0, y1, sector_id=sector_id, partner_id=partner_id)
        rows = []
        for y in pf["years"]:
            c = y["counts"]
            rows.append([y["year"], c["approved"], c["active"], c["completed"] + c["completed_closing_date_proxy"],
                         "; ".join(f"{k} {v}" for k, v in y["commitments"]["by_currency"].items()),
                         y["commitments"]["usd_equivalent"] or "", "yes" if y["commitments"]["usd_complete"] else "no",
                         "; ".join(f"{k} {v}" for k, v in y["disbursements"]["by_currency"].items()),
                         y["disbursements"]["usd_equivalent"] or "", "yes" if y["disbursements"]["usd_complete"] else "no"])
        data.tables.append(Table("Annual portfolio", ["Year", "Approved", "Active", "Completed", "Commitments (original currencies)",
                                                      "Commitments USD eq.", "USD complete", "Disbursements (original currencies)",
                                                      "Disbursements USD eq.", "USD complete"], rows))
        if report_type == "donor_portfolio":
            items = [p for p in projects if partner_id is None or any(a.partner_id == partner_id for a in p.agreements)]
            data.tables.append(Table("Projects", headers, project_rows(items)))
    if report_type in ("project_performance", "country_investment"):
        alerts = db.scalars(select(ProjectAlert).where(ProjectAlert.is_active.is_(True),
                                                       ProjectAlert.project_id.in_([p.id for p in projects]))).all()
        by_id = {p.id: p for p in projects}
        data.tables.append(Table("Active alerts", ["Alert ID", "AP²I ID", "Rule", "Level", "Observed", "Threshold", "Verification"],
                                 [[a.public_id, by_id[a.project_id].ap2i_id, a.rule_code, a.level, a.observed_value, a.threshold,
                                   a.recommended_verification] for a in alerts]))
    if report_type == "operational_infrastructure":
        done = [p for p in projects if p.status in ("completed", "closed")]
        data.tables.append(Table("Completed projects — operational status", ["AP²I ID", "Title", "Sector", "Status", "Operational status", "Commissioning"],
                                 [[p.ap2i_id, p.title, p.sector.name_en if p.sector else "", p.status, p.operational_status,
                                   str(p.commissioning_date or "")] for p in done]))
        data.notes.append(f"Operational status undocumented for {sum(1 for p in done if p.operational_status == 'unknown')} of {len(done)} completed projects.")
    if report_type == "project_rescue":
        plans = db.scalars(select(RecoveryPlan).where(RecoveryPlan.project_id.in_([p.id for p in projects]), plan_cond)).all() if projects else []
        rows = []
        for pl in plans:
            for a in pl.actions:
                rows.append([pl.id, db.get(Project, pl.project_id).ap2i_id, a.problem, a.recommended_action, a.responsible_institution or "",
                             a.priority, a.approval_status, a.implementation_status, str(a.due_date or ""), a.contributing_factors])
        data.tables.append(Table("Recovery actions", ["Plan", "AP²I ID", "Problem", "Action", "Responsible", "Priority", "Approval",
                                                      "Implementation", "Due", "Contributing factors"], rows))
    if report_type == "data_quality":
        dq = dq_dashboard(db, country.id)
        data.tables.append(Table("Summary", ["Measure", "Value"], [[k, str(v)] for k, v in dq.items() if not isinstance(v, (dict, list))]))
        data.tables.append(Table("Missing information (projects)", ["Field", "Projects missing it"], [[k, v] for k, v in dq["missing_information"].items()]))
        data.tables.append(Table("Connectors", ["Key", "Status", "Enabled", "Last success", "Last failure", "Records"],
                                 [[c["key"], c["status"], c["enabled"], c["last_success_at"] or "", c["last_failure_at"] or "", c["records_stored"]]
                                  for c in dq["connectors"]]))
    return data


def header_lines(d: ReportData) -> list[tuple[str, str]]:
    return [("Report", d.title), ("Report ID", d.report_uid), ("Reporting period", d.period),
            ("Generated", d.generated_at.strftime("%Y-%m-%d %H:%M UTC")), ("Data coverage", d.coverage),
            ("Methodology", METHODOLOGY), ("Limitations", " ".join(LIMITATIONS + d.notes)),
            ("Sources", " | ".join(d.sources) or "No synchronised source")]


def render_csv(d: ReportData) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    for k, v in header_lines(d):
        w.writerow([f"# {k}", v])
    for t in d.tables:
        w.writerow([])
        w.writerow([f"## {t.title}"])
        w.writerow(t.headers)
        w.writerows(t.rows)
    return buf.getvalue().encode("utf-8-sig")


def render_xlsx(d: ReportData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    row = 1
    if LOGO.exists():
        img = XLImage(str(LOGO))
        img.width, img.height = 120, 120
        ws.add_image(img, "A1")
        row = 8
    for k, v in header_lines(d):
        ws.cell(row, 1, k).font = Font(bold=True, color=NAVY)
        c = ws.cell(row, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    for t in d.tables:
        sh = wb.create_sheet(t.title[:31])
        for j, h in enumerate(t.headers, 1):
            cell = sh.cell(1, j, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            sh.column_dimensions[cell.column_letter].width = min(60, max(12, len(h) + 4))
        for i, r in enumerate(t.rows, 2):
            for j, v in enumerate(r, 1):
                sh.cell(i, j, v)
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_pdf(d: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1 * cm,
                            bottomMargin=1 * cm, title=d.title, author="AP²I")
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=7, leading=8.5)
    story = []
    if LOGO.exists():
        story.append(Image(str(LOGO), width=3 * cm, height=3 * cm))
    story.append(Paragraph(f"<font color='#{NAVY}'>{_esc(d.title)}</font>", ss["Title"]))
    for k, v in header_lines(d)[1:]:
        story.append(Paragraph(f"<b>{_esc(k)}:</b> {_esc(v)}", small))
    for t in d.tables:
        story.append(Spacer(1, 0.4 * cm))
        story.append(Paragraph(f"<font color='#{GREEN}'>{_esc(t.title)}</font>", ss["Heading2"]))
        if not t.rows:
            story.append(Paragraph("No records.", small))
            continue
        data = [[Paragraph(f"<b>{_esc(h)}</b>", small) for h in t.headers]] + [
            [Paragraph(_esc(str(v))[:600], small) for v in r] for r in t.rows[:2000]]
        tbl = LongTable(data, repeatRows=1)
        tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                 ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
                                 ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                 ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")])]))
        story.append(tbl)
        if len(t.rows) > 2000:
            story.append(Paragraph(f"{len(t.rows) - 2000} more rows in the Excel/CSV version.", small))

    def footer(canvas, _doc):
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor(f"#{NAVY}"))
        canvas.drawString(1.2 * cm, 0.6 * cm, f"AP²I — {d.report_uid} — generated {d.generated_at:%Y-%m-%d} — page {_doc.page}")

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _esc(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


RENDERERS = {"pdf": (render_pdf, "application/pdf"),
             "xlsx": (render_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
             "csv": (render_csv, "text/csv; charset=utf-8")}


def generate(db: Session, report_type: str, fmt: str, params: dict, principal, user_id: int | None,
             tenant_id: int | None) -> tuple[Report, bytes, str]:
    if fmt not in RENDERERS:
        raise ValueError("format must be pdf, xlsx or csv")
    data = build(db, report_type, params, principal)
    render, mime = RENDERERS[fmt]
    content = render(data)
    folder = Path(get_settings().document_store_dir) / "reports"
    folder.mkdir(parents=True, exist_ok=True)
    key = folder / f"{data.report_uid}.{fmt}"
    key.write_bytes(content)
    rep = Report(report_uid=data.report_uid, report_type=report_type, title=data.title, fmt=fmt, params=params,
                 period=data.period, generated_by=user_id, tenant_id=tenant_id, sha256=hashlib.sha256(content).hexdigest(),
                 size_bytes=len(content), storage_key=str(key))
    db.add(rep)
    db.commit()
    return rep, content, mime


__all__ = ["generate", "build", "REPORT_TYPES", "func", "os"]
