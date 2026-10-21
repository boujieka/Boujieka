"""PDF (ReportLab), Word (python-docx) and Excel (openpyxl) exporters."""

from __future__ import annotations

import io

from src.reports.charts import chart
from src.reports.content import build_sections, report_title

FOOTER = "Industrial Energy AI Africa - MVP 1.0 - preliminary screening study"


# ------------------------------------------------------------------------------------ PDF

def to_pdf(snap: dict, lang: str = "fr") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title=report_title(snap, lang), author="Industrial Energy AI Africa")
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("small", parent=body, fontSize=7, leading=9)
    warn = ParagraphStyle("warn", parent=body, textColor=colors.HexColor("#9A3412"), fontName="Helvetica-Bold")
    story = [Paragraph(escape(report_title(snap, lang)), ss["Title"]),
             Paragraph(escape(f"{snap['meta']['generated_at']} - v{snap['meta']['app_version']}"), small), Spacer(1, 8)]
    if snap["data_flags"].get("project_synthetic"):
        story.append(Paragraph(escape("DONNÉES SYNTHÉTIQUES / SYNTHETIC DATA"), warn))
    width = A4[0] - 3.6 * cm
    for i, s in enumerate(build_sections(snap, lang), start=1):
        story.append(Paragraph(escape(f"{i}. {s.title}"), ss["Heading2"]))
        for ptxt in s.paragraphs:
            story.append(Paragraph(escape(ptxt), body))
        if s.bullets:
            story.append(ListFlowable([ListItem(Paragraph(escape(b), body)) for b in s.bullets], bulletType="bullet", leftIndent=12))
        for tbl in s.tables:
            story.append(Paragraph(escape(tbl.caption), ss["Heading4"]))
            data = [[Paragraph(escape(str(c)), small) for c in tbl.header]] + [[Paragraph(escape(str(c)), small) for c in r] for r in tbl.rows]
            t = Table(data, repeatRows=1, colWidths=[width / len(tbl.header)] * len(tbl.header))
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3EEEA")),
                                   ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B9C6C1")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 6)]
        for ck in s.charts:
            png = chart(ck, snap, lang)
            if png:
                img = Image(io.BytesIO(png), width=width, height=width * 0.42)
                story += [img, Spacer(1, 6)]
        if i in (1, 8):
            story.append(PageBreak())

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(1.8 * cm, 1.0 * cm, FOOTER)
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.0 * cm, str(canvas.getPageNumber()))
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


# ----------------------------------------------------------------------------------- Word

def to_docx(snap: dict, lang: str = "fr") -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt

    d = Document()
    d.add_heading(report_title(snap, lang), 0)
    d.add_paragraph(f"{snap['meta']['generated_at']} - v{snap['meta']['app_version']}")
    if snap["data_flags"].get("project_synthetic"):
        r = d.add_paragraph().add_run("DONNÉES SYNTHÉTIQUES / SYNTHETIC DATA")
        r.bold = True
    for i, s in enumerate(build_sections(snap, lang), start=1):
        d.add_heading(f"{i}. {s.title}", 1)
        for ptxt in s.paragraphs:
            d.add_paragraph(ptxt)
        for b in s.bullets:
            d.add_paragraph(b, style="List Bullet")
        for tbl in s.tables:
            d.add_paragraph(tbl.caption).runs[0].bold = True
            t = d.add_table(rows=1, cols=len(tbl.header))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(tbl.header):
                t.rows[0].cells[j].text = str(h)
            for row in tbl.rows:
                cells = t.add_row().cells
                for j, c in enumerate(row):
                    cells[j].text = str(c)
            for row in t.rows:
                for cell in row.cells:
                    for par in cell.paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(8)
        for ck in s.charts:
            png = chart(ck, snap, lang)
            if png:
                d.add_picture(io.BytesIO(png), width=Cm(16))
    d.add_paragraph(FOOTER)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------- Excel

def to_xlsx(snap: dict, lang: str = "fr", hourly: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="E3EEEA")

    def sheet(title, header, rows):
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col) + 2))
        return ws

    ws = wb.active
    ws.title = "Summary"
    ws.append([report_title(snap, lang)])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Generated", snap["meta"]["generated_at"], "Version", snap["meta"]["app_version"]])
    if snap["data_flags"].get("project_synthetic"):
        ws.append(["SYNTHETIC DATA - demonstration only"])
    sections = build_sections(snap, lang)
    for tbl in sections[0].tables:
        ws.append([])
        ws.append(tbl.header)
        for r in tbl.rows:
            ws.append(r)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 45

    # Scenarios (LP and verification side by side)
    rows = []
    for key, sc in (snap.get("scenarios") or {}).items():
        for method in ("lp", "check"):
            r = sc[method]
            m, c, e = r.get("metrics") or {}, r.get("capacities") or {}, r.get("annual_energy") or {}
            rows.append([key, r.get("method"), r.get("status"), c.get("pv_kwp"), c.get("battery_kwh"), c.get("battery_kw"),
                         c.get("diesel_existing_kw"), c.get("diesel_new_kw"), c.get("biomass_kw"), e.get("pv_used"), e.get("discharge"),
                         e.get("diesel"), e.get("biomass"), e.get("grid"), e.get("unserved"), m.get("fuel_l"), m.get("renewable_share"),
                         m.get("annual_operating_cost"), m.get("annual_cost"), m.get("lcoe_per_kwh"), m.get("total_investment"),
                         f"scenarios.{key}.{method}"])
    ws_s = sheet("Scenarios", ["scenario", "method", "status", "pv_kwp", "battery_kwh", "battery_kw", "diesel_existing_kw", "diesel_new_kw",
                               "biomass_kw", "pv_used_kwh", "bess_discharge_kwh", "diesel_kwh", "biomass_kwh", "grid_kwh", "unserved_kwh",
                               "fuel_l", "renewable_share", "annual_operating_cost", "equivalent_annual_cost", "lcoe_per_kwh",
                               "investment", "source_path"], rows)
    if rows:
        ch = BarChart()
        ch.title = "Equivalent annual cost"
        ch.add_data(Reference(ws_s, min_col=19, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws_s, min_col=1, min_row=2, max_row=len(rows) + 1))
        ws_s.add_chart(ch, "B" + str(len(rows) + 4))

    fin = snap.get("finance")
    if fin:
        cols = ["years", "phase", "capex", "replacements", "revenue", "opex", "ebitda", "depreciation", "tax_unlevered", "project_cf",
                "debt_drawdown", "idc", "interest", "principal", "debt_service", "debt_balance_end", "tax_levered", "cfads", "equity_cf",
                "dscr", "energy_kwh"]
        n = len(fin["years"])
        ws_f = sheet("Cash flow", cols, [[fin[c][i] for c in cols] for i in range(n)])
        ch = BarChart()
        ch.title = "Project cash flow"
        ch.add_data(Reference(ws_f, min_col=10, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws_f, min_col=1, min_row=2, max_row=n + 1))
        ws_f.add_chart(ch, f"B{n + 4}")
        ind = ["npv_project", "npv_equity", "irr_project", "irr_equity", "payback_years", "discounted_payback_years", "min_dscr", "avg_dscr",
               "llcr", "lcoe_per_kwh", "total_financing_requirement", "debt_amount", "equity_amount", "cumulative_savings", "real_discount_rate"]
        sheet("Indicators", ["indicator", "value", "source_path"], [[k, fin.get(k), f"finance.{k}"] for k in ind])
    if snap.get("structures"):
        sheet("Structures", ["structure", "investor", "operator", "client_npv", "investor_irr", "terms_missing", "payment_basis"],
              [[r["label_en"], r["investor"], r["operator"], r["client_npv"], r["investor_irr"], r["terms_missing"], r["payment_basis"]] for r in snap["structures"]])
    sens = snap.get("sensitivity") or {}
    if sens.get("tornado"):
        sheet("Sensitivity", ["parameter", "low_multiplier", "high_multiplier", "low", "high", "swing"],
              [[r["parameter"], r["low_multiplier"], r["high_multiplier"], r["low"], r["high"], r["swing"]] for r in sens["tornado"]["rows"]])
    if sens.get("matrix"):
        mx = sens["matrix"]
        sheet("Sensitivity matrix", [f"{mx['row_parameter']} \\ {mx['col_parameter']}"] + [str(c) for c in mx["cols"]],
              [[r] + list(vals) for r, vals in zip(mx["rows"], mx["values"])])
    a = snap.get("assumptions", {})
    arows = []
    for group, vals in a.items():
        if isinstance(vals, dict):
            for k, v in vals.items():
                arows.append([group, k, str(v) if isinstance(v, (list, dict)) else v])
        else:
            arows.append(["general", group, vals])
    sheet("Assumptions", ["group", "parameter", "value"], arows)
    if hourly:
        keys = [k for k in ("load", "pv_available", "pv_used", "pv_curtailed", "charge", "discharge", "soc", "diesel", "diesel_excess",
                            "biomass", "grid", "unserved") if k in hourly]
        n = len(hourly[keys[0]])
        sheet("Hourly dispatch", ["hour"] + keys, [[i] + [float(hourly[k][i]) for k in keys] for i in range(n)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
