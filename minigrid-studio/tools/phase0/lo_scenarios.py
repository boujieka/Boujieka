"""Recalculate reference scenarios of the CrossBoundary workbook in headless
LibreOffice, replaying the workbook's three VBA macros in Python.

This is a CROSS-CHECK oracle, not the reference: reference values must come
from Microsoft Excel (see docs/reference_scenarios.md). LibreOffice is used
because it reproduces the Excel values cached in the published file once the
negative-width OFFSET calls are rewritten (LibreOffice rejects them with
Err:502); the rewritten copy lives in a scratch directory and is never shipped.

No VBA is executed: documents are opened with MacroExecutionMode=NEVER_EXECUTE
and the macros (Module1.Debt_Sizing, Module2.Refi_Debt_Sizing,
Module7.DSRA_resize_3) are re-implemented below from their audited source.

Usage (system python3 with the LibreOffice UNO bridge):
    python3 tools/phase0/lo_scenarios.py WORKBOOK.xlsm scenarios.json OUT_DIR [--dump-cells]

--dump-cells also writes <id>_cells.json with every formula cell's value
("Sheet!row,col" -> value), used by the engine's cell-by-cell reconciliation.
"""
import json
import os
import re
import subprocess
import sys
import time

import openpyxl
import uno
from com.sun.star.beans import PropertyValue

HERE = os.path.dirname(os.path.abspath(__file__))
TIMELINE = [openpyxl.utils.get_column_letter(c) for c in range(10, 79)]  # J..BZ
MAX_LOOP = 500     # the VBA loops have no bound; we stop and report instead
MAX_ROUNDS = 25    # operator re-runs "until the checks are all Ok"

NEG_OFFSET = re.compile(r"OFFSET\((\$?[A-Z]{1,3}\$?\d+),0,0,(\d+),-(\$?[A-Z]{1,3}\$?\d+)\)")


def rewrite_offsets(src, dst):
    """OFFSET(ref,0,0,h,-w) -> OFFSET(ref,0,1-w,h,w): same cells, positive width."""
    wb = openpyxl.load_workbook(src)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "OFFSET(" in c.value:
                    c.value, k = NEG_OFFSET.subn(r"OFFSET(\1,0,1-\3,\2,\3)", c.value)
                    n += k
    wb.save(dst)
    return n


def pv(name, value):
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


class Office:
    def __init__(self, workdir):
        self.port = 2100 + os.getpid() % 1000
        profile = uno.systemPathToFileUrl(os.path.abspath(os.path.join(workdir, f"lo-profile-{self.port}")))
        self.proc = subprocess.Popen([
            "soffice", "--headless", "--invisible", "--nologo", "--norestore",
            f"--accept=socket,host=127.0.0.1,port={self.port};urp;",
            f"-env:UserInstallation={profile}"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(90):
            try:
                ctx = resolver.resolve(
                    f"uno:socket,host=127.0.0.1,port={self.port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(1)
        else:
            raise RuntimeError("LibreOffice did not start")
        self.desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)

    def open(self, path):
        return self.desktop.loadComponentFromURL(
            uno.systemPathToFileUrl(os.path.abspath(path)), "_blank", 0,
            (pv("Hidden", True), pv("MacroExecutionMode", 0)))

    def close(self):
        try:
            self.desktop.terminate()
        except Exception:
            pass
        self.proc.terminate()


class Book:
    def __init__(self, doc):
        self.doc = doc

    def cell(self, ref):
        sh, addr = ref.split("!")
        return self.doc.Sheets.getByName(sh.strip("'")).getCellRangeByName(addr)

    def get(self, ref):
        c = self.cell(ref)
        if c.getError():
            return {"error": c.getError()}
        if c.getType().value == "TEXT":
            return c.getString()
        return c.getValue()

    def set(self, ref, value):
        c = self.cell(ref)
        if isinstance(value, str):
            c.setString(value)
        else:
            c.setValue(float(value))
        self.doc.calculateAll()

    def paste_value(self, src, dst):
        self.set(dst, self.cell(src).getValue())


def debt_sizing(b, log):
    """Module1.Debt_Sizing: Do Until DSCR_Delta = 0 And Funding_Delta = 0
    copy DSCR trial (I300) -> G300, then Funding trial (I292) -> G292."""
    n = 0
    while not (b.get("Inputs!J300") == 0 and b.get("Inputs!J292") == 0):
        if n >= MAX_LOOP:
            log.append({"macro": "Debt_Sizing", "status": "no convergence", "iterations": n})
            return False
        b.paste_value("Inputs!I300", "Inputs!G300")
        b.paste_value("Inputs!I292", "Inputs!G292")
        n += 1
    log.append({"macro": "Debt_Sizing", "iterations": n})
    return True


def refi_debt_sizing(b, log):
    """Module2.Refi_Debt_Sizing: Do Until Refi_Delta = 0, copy I344 -> G344."""
    n = 0
    while b.get("Inputs!J344") != 0:
        if n >= MAX_LOOP:
            log.append({"macro": "Refi_Debt_Sizing", "status": "no convergence", "iterations": n})
            return False
        b.paste_value("Inputs!I344", "Inputs!G344")
        n += 1
    log.append({"macro": "Refi_Debt_Sizing", "iterations": n})
    return True


def dsra_resize(b, log):
    """Module7.DSRA_resize_3: copy DSRA_Copy (E371) -> DSRA_Paste (E372), once."""
    b.paste_value("Inputs!E371", "Inputs!E372")
    log.append({"macro": "DSRA_resize_3"})


# Checks the macros can act on (sizing state). The other Dashboard checks
# (equity injection, debt repaid, ...) describe the economics of the case and are
# reported as warnings: re-running the macros cannot change them.
SIZING_CHECKS = ["Inputs!N292", "Inputs!N300", "Inputs!N344", "Inputs!E373"]


def checks_ok(b):
    return all(b.get(k) == 0 for k in SIZING_CHECKS)


def solve(b, log):
    """Operator procedure from the workbook ('Key' and 'Currency' sheets):
    run debt sizing, refi sizing when the refi switch is on, then DSRA sizing,
    and repeat until the sizing checks read 0."""
    for rnd in range(1, MAX_ROUNDS + 1):
        if checks_ok(b):
            return {"converged": True, "rounds": rnd - 1}
        ok = debt_sizing(b, log)
        if ok and b.get("Inputs!G336") == 1:
            ok = refi_debt_sizing(b, log)
        if not ok:
            return {"converged": False, "rounds": rnd}
        dsra_resize(b, log)
    return {"converged": checks_ok(b), "rounds": MAX_ROUNDS}


def extract(b, spec):
    out = {"scalars": {}, "pasted_values": {}, "series": {}}
    for k, (ref, _label, _unit) in spec["scalars"].items():
        out["scalars"][k] = b.get(ref)
    for k, ref in spec["pasted_values"].items():
        out["pasted_values"][k] = b.get(ref)
    for k, (sh, row, _label) in spec["series"].items():
        out["series"][k] = [b.get(f"{sh}!{col}{row}") for col in TIMELINE]
    return out


def dump_cells(doc):
    res = {}
    for i in range(doc.Sheets.Count):
        sh = doc.Sheets.getByIndex(i)
        cur = sh.createCursor()
        cur.gotoEndOfUsedArea(False)
        addr = cur.getRangeAddress()
        rng = sh.getCellRangeByPosition(0, 0, addr.EndColumn, addr.EndRow)
        forms, data = rng.getFormulaArray(), rng.getDataArray()
        for r, row in enumerate(forms):
            for c, f in enumerate(row):
                if isinstance(f, str) and f.startswith("="):
                    res[f"{sh.Name}!{r + 1},{c + 1}"] = data[r][c]
    return res


def main():
    workbook, scenarios_path, out_dir = sys.argv[1:4]
    want_cells = "--dump-cells" in sys.argv[4:]
    spec = json.load(open(os.path.join(HERE, "outputs_spec.json")))
    scenarios = json.load(open(scenarios_path))["scenarios"]
    os.makedirs(out_dir, exist_ok=True)
    work = os.path.join(out_dir, "_work")
    os.makedirs(work, exist_ok=True)
    calc_copy = os.path.join(work, "calc_copy.xlsx")
    n = rewrite_offsets(workbook, calc_copy)
    print(f"rewrote {n} negative-width OFFSET calls in scratch copy")
    office = Office(work)
    try:
        for sc in scenarios:
            doc = office.open(calc_copy)
            b = Book(doc)
            doc.calculateAll()
            for ref, value in sc.get("set", {}).items():
                b.set(ref, value)
            log = []
            status = solve(b, log) if sc.get("replay_macros", True) else {"converged": None}
            out = extract(b, spec)
            status["failed_model_checks"] = sorted(
                k for k, v in out["scalars"].items() if k.startswith("check_") and v != 0)
            result = {
                "scenario": sc["id"],
                "engine": "LibreOffice (cross-check, not reference)",
                "inputs_changed": sc.get("set", {}),
                "solver": status,
                "macro_log": log,
                **out,
            }
            json.dump(result, open(os.path.join(out_dir, f"{sc['id']}_libreoffice.json"), "w"), indent=1)
            if want_cells:
                json.dump(dump_cells(doc), open(os.path.join(out_dir, f"{sc['id']}_cells.json"), "w"))
            print(sc["id"], status, "IRR", result["scalars"]["project_irr"])
            doc.close(True)
    finally:
        office.close()


if __name__ == "__main__":
    main()
