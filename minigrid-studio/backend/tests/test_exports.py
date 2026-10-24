"""Phase 2 exports: Studio workbook, PDF report, CrossBoundary-format fill, labels."""
import io
import json
import shutil
import subprocess
import zipfile

import openpyxl
import pytest

from app.engine.assumptions import from_dict, to_dict
from app.engine.results import run_scenario
from app.exports.crossboundary import _patch_sheet, fill_template
from app.exports.labels import FIELDS_FR, field_entries, labels, translate_message
from app.exports.pdf_report import _nice_ticks, build_pdf_report
from app.exports.xlsx_studio import build_studio_workbook
from app.services.scenarios import load_template

from .conftest import WORKBOOK, needs_workbook


@pytest.fixture(scope="module")
def demo():
    a = load_template("demo_fictitious")
    res = run_scenario(from_dict(a), per_lot=True)
    return {"project": {"name": "Fictitia", "country": "Fictitia"},
            "version": {"scenario": "base", "version": 1, "status": "draft", "assumptions": a},
            "run": {"id": "r1", "status": res.status, "inputs_hash": "h", "engine_version": "1.0.0",
                    "result": json.loads(json.dumps(res.as_dict(), default=lambda o: None))}}


def test_identities_used_by_the_studio_formulas_hold(demo):
    s = {k: v["values"] for k, v in demo["run"]["result"]["series"].items()}
    for t in range(len(s["year"])):
        assert s["total_opex"][t] == pytest.approx(s["opco_fees"][t] + s["diesel_cost"][t] + s["admin_costs"][t], abs=1e-6)
        assert s["ebitda"][t] == pytest.approx(s["revenue"][t] + s["total_opex"][t], abs=1e-6)
        assert s["ebit"][t] == pytest.approx(s["ebitda"][t] + s["depreciation"][t], abs=1e-6)
        assert s["ebt"][t] == pytest.approx(s["ebit"][t] + s["interest"][t], abs=1e-6)
        assert s["net_assets"][t] == pytest.approx(s["total_assets_current"][t] + s["fixed_assets"][t] + s["liabilities"][t], abs=1e-6)
        if s["dscr"][t] is not None:
            assert s["dscr"][t] == pytest.approx(s["cfads"][t] / s["debt_service"][t])


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_studio_workbook_structure(demo, lang):
    data = build_studio_workbook(**demo, lang=lang)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    names = wb.sheetnames
    assert len(names) == 7
    kpis = wb[names[1]]
    formulas = {row[0].value: row[6].value for row in kpis.iter_rows(min_row=2)}
    assert formulas["project_irr"].startswith("=XIRR(")
    assert formulas["project_npv"].startswith("=XNPV(")
    annual = wb[[n for n in names if n in ("Annuel", "Annual")][0]]
    recalc = [r[0].value for r in annual.iter_rows(min_row=2) if str(r[0].value).startswith("recalc.")]
    assert recalc == [f"recalc.{k}" for k in ("total_opex", "ebitda", "ebit", "ebt", "net_assets", "balance_check", "dscr")]


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_studio_formulas_recalculate_to_engine_values(demo, tmp_path):
    """LibreOffice (stand-in for Excel) evaluates the exported formulas; the gaps must be ~0."""
    src = tmp_path / "studio.xlsx"
    src.write_bytes(build_studio_workbook(**demo, lang="en"))
    subprocess.run(["soffice", "--headless", f"-env:UserInstallation=file://{tmp_path}/profile", "--convert-to", "xlsx",
                    "--outdir", str(tmp_path / "out"), str(src)], check=True, capture_output=True, timeout=180)
    wb = openpyxl.load_workbook(tmp_path / "out" / "studio.xlsx", data_only=True)
    for row in wb["KPIs"].iter_rows(min_row=2):
        if row[6].value is not None:
            assert abs(row[7].value) <= 1e-6 * max(1, abs(row[2].value)), row[0].value
    for row in wb["Annual"].iter_rows(min_row=2):
        if str(row[0].value).startswith("recalc."):
            assert abs(row[4].value) < 1e-6, row[0].value


@pytest.mark.parametrize("lang", ["fr", "en"])
def test_pdf_report(demo, lang):
    pdf = build_pdf_report(**demo, lang=lang)
    assert pdf.startswith(b"%PDF") and len(pdf) > 10_000
    failed = {**demo["run"], "status": "failed", "result": {}, "error": "lots: No included lot with grids (x)."}
    assert build_pdf_report(project=demo["project"], version=demo["version"], run=failed, lang=lang).startswith(b"%PDF")


def test_axis_reaches_data_maximum():
    for lo, hi in ((0, 4.33e6), (0, 1.2), (-3.8e6, 3.2e6), (0, 0), (0.37, 1.2)):
        t = _nice_ticks(min(lo, 0), max(hi, 0))
        assert t[0] <= min(lo, 0) and t[-1] >= hi


def test_labels_complete_and_messages_translated():
    assert all(e["path"].replace("<i>.", "") in FIELDS_FR for e in field_entries())
    fr, en = labels("fr"), labels("en")
    assert set(fr["kpis"]) == set(en["kpis"]) and set(fr["series"]) == set(en["series"])
    assert translate_message("Derived DSCR falls to 0.37x in a repayment year, below the 1.20x target "
                             "(the workbook has no DSCR covenant test).", "fr").startswith("Le DSCR dérivé descend à 0.37x")
    assert translate_message("lots.0.included: Included flag must be 0 or 1.; debt.sizing_method: Must be 'Gearing' or 'DSCR'.",
                             "fr") == ("lots.0.included: L'indicateur « inclus » doit valoir 0 ou 1.; "
                                       "debt.sizing_method: Valeur attendue : 'Gearing' or 'DSCR'.")
    assert translate_message("Something new.", "fr") == "Something new."


def test_patch_sheet_keeps_other_cells_and_refuses_shared_masters():
    xml = ('<sheetData><row r="1"><c r="A1" s="3"><v>1</v></c><c r="B1" s="2" t="s"><v>0</v></c>'
           '<c r="C1" s="4"/><c r="D1"><f t="shared" ref="D1:D9" si="0">A1*2</f><v>2</v></c>'
           '<c r="E1"><f t="shared" si="0"/><v>2</v></c></row></sheetData>')
    problems = []
    out = _patch_sheet(xml, {"A1": 2.5, "B1": "Pari Passu", "C1": None, "E1": 7, "D1": 1, "Z9": 1}, "S", problems)
    assert '<c r="A1" s="3"><v>2.5</v></c>' in out
    assert '<c r="B1" s="2" t="inlineStr"><is><t xml:space="preserve">Pari Passu</t></is></c>' in out
    assert '<c r="C1" s="4"/>' in out and '<c r="E1"><v>7</v></c>' in out
    assert '<f t="shared" ref="D1:D9" si="0">A1*2</f>' in out          # shared master untouched
    assert any("D1" in p for p in problems) and any("Z9" in p for p in problems)


@needs_workbook()
@pytest.mark.parametrize("template", ["demo_fictitious", "crossboundary_reference"])
def test_crossboundary_fill_round_trip(template):
    """Assumptions written into the workbook read back identically, and the file keeps its macros."""
    from app.engine.loader import load_assumptions
    from app.engine.solver import solve
    src = WORKBOOK.read_bytes()
    a = from_dict(load_template(template))
    solved, _L, _ = solve(a)
    p = solved.pasted
    out = fill_template(src, a, {"funding": p.funding, "dscr_limit": p.dscr_limit, "refi_limit": p.refi_limit,
                                 "dsra_initial": p.dsra_initial})
    b = load_assumptions(io.BytesIO(out))
    da, db = to_dict(a), to_dict(b)
    db["pasted"] = da["pasted"] = None
    db["lots"] = db["lots"][:len(da["lots"])]
    assert da == db
    with zipfile.ZipFile(io.BytesIO(out)) as z, zipfile.ZipFile(io.BytesIO(src)) as zs:
        assert z.read("xl/vbaProject.bin") == zs.read("xl/vbaProject.bin")
        assert 'fullCalcOnLoad="1"' in z.read("xl/workbook.xml").decode()
        assert "xl/calcChain.xml" not in z.namelist()
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb["Inputs"]["G292"].value == pytest.approx(p.funding)


@needs_workbook()
def test_crossboundary_fill_refuses_basis_change():
    from app.exports.crossboundary import TemplateMismatch
    a = from_dict(load_template("crossboundary_reference"))
    a.currency.tariff.basis = "usd" if a.currency.tariff.basis != "usd" else "lcy_contracted"
    with pytest.raises(TemplateMismatch) as e:
        fill_template(WORKBOOK.read_bytes(), a, {})
    assert any("currency.tariff" in p for p in e.value.problems)

