"""Studio workbook: a calculation run exported to .xlsx with auditable formulas.

Engine values are written as numbers, each labelled with the workbook cell or
row it reproduces. Below them, Excel formulas recompute what can be recomputed
from the exported rows alone (opex total, EBITDA, EBIT, EBT, net assets,
balance check, DSCR, project IRR and NPV), next to a gap column, so a reviewer
can check the engine with Excel's own functions. The full CrossBoundary
formula set is available with the CrossBoundary-format export instead.
"""
from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engine import ENGINE_VERSION
from .labels import CHECKS_EN, CHECKS_FR, field_entries, kpi_label, labels, series_label

INPUT_FONT = Font(color="1F4E9E")            # typed assumptions
FORMULA_FONT = Font(color="000000")
ENGINE_FONT = Font(color="2E6B30")           # values computed by the engine
HEAD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E8EEF4")
FIRST_YEAR_COL = 6                           # F: first year; E is the XIRR anchor column

T = {
    "fr": dict(readme="Lisez-moi", kpis="Indicateurs", annual="Annuel", assumptions="Hypothèses", lots="Lots",
               checks="Contrôles", alerts="Alertes", key="Clé", label="Libellé", value="Valeur moteur", unit="Unité",
               source="Source (classeur)", status="Statut", recalc="Recalcul Excel", gap="Écart", anchor="Ancrage",
               path="Chemin", cell="Cellule", input="Valeur", doc="Note", recalc_title="Recalcul par formules Excel (contrôle)",
               total_gap="Écart absolu total", ok="OK", fail="ÉCHEC", derived="dérivé", experimental="expérimental",
               none="Aucune"),
    "en": dict(readme="ReadMe", kpis="KPIs", annual="Annual", assumptions="Assumptions", lots="Lots",
               checks="Checks", alerts="Warnings", key="Key", label="Label", value="Engine value", unit="Unit",
               source="Source (workbook)", status="Status", recalc="Excel recalculation", gap="Gap", anchor="Anchor",
               path="Path", cell="Cell", input="Value", doc="Note", recalc_title="Recalculated with Excel formulas (audit)",
               total_gap="Total absolute gap", ok="OK", fail="FAIL", derived="derived", experimental="experimental",
               none="None"),
}

README = {
    "fr": [
        "Export du Mini-Grid Financial Model Studio.",
        "Les valeurs en vert sont calculées par le moteur ; chacune nomme la cellule ou la ligne du classeur",
        "CrossBoundary Access Open Source Model 1.0 qu'elle reproduit. Les hypothèses saisies sont en bleu.",
        "Les lignes « Recalcul » sont des formules Excel vivantes, à côté d'une colonne d'écart :",
        "elles permettent de contrôler le moteur avec les fonctions d'Excel (XIRR, XNPV, sommes).",
        "Pour disposer de toutes les formules d'origine, utilisez l'export au format CrossBoundary.",
        "Outil d'analyse : ni certification BAD, Banque mondiale ou CrossBoundary, ni avis juridique,",
        "fiscal ou d'investissement. ACCESS360 Africa n'est ni affiliée à CrossBoundary ni approuvée par elle.",
    ],
    "en": [
        "Export of Mini-Grid Financial Model Studio.",
        "Green values are computed by the engine; each names the cell or row of the CrossBoundary Access",
        "Open Source Model 1.0 workbook it reproduces. Typed assumptions are blue.",
        "'Recalculated' rows are live Excel formulas next to a gap column:",
        "they let a reviewer check the engine with Excel's own functions (XIRR, XNPV, sums).",
        "For every original formula, use the CrossBoundary-format export.",
        "Analysis tool: not a certification by AfDB, the World Bank or CrossBoundary, nor legal, tax",
        "or investment advice. ACCESS360 Africa is not affiliated with or endorsed by CrossBoundary.",
    ],
}

NUMBER_FORMATS = {"%": "0.00%", "% p.a.": "0.00%", "x": "0.00", "currency": "#,##0;-#,##0",
                  "kWh": "#,##0", "year": "0", "date": "dd/mm/yyyy"}

# identities verified on the engine's rows (tests/test_exports.py)
RECALC = [  # (key of the engine row, formula over engine rows)
    ("total_opex", "{opco_fees}+{diesel_cost}+{admin_costs}"),
    ("ebitda", "{revenue}+{total_opex}"),
    ("ebit", "{ebitda}+{depreciation}"),
    ("ebt", "{ebit}+{interest}"),
    ("net_assets", "{total_assets_current}+{fixed_assets}+{liabilities}"),
    ("balance_check", "{net_assets}-{equity}"),
    ("dscr", "IF({debt_service}>0.000000001,{cfads}/{debt_service},0)"),   # 0 where the engine leaves a blank
]


def _fmt(unit: str) -> str:
    return NUMBER_FORMATS.get(unit, "#,##0.00" if unit.startswith("currency") else "General")


def _flatten(d, prefix=""):
    if isinstance(d, dict):
        for k, v in d.items():
            yield from _flatten(v, f"{prefix}{k}.")
    elif isinstance(d, list):
        for i, v in enumerate(d):
            yield from _flatten(v, f"{prefix}{i}.")
    else:
        yield prefix[:-1], d


def build_studio_workbook(*, project: dict, version: dict, run: dict, lang: str = "fr") -> bytes:
    project, version, run = _clean(project), _clean(version), _clean(run)
    t, lab = T[lang], labels(lang)
    res = run["result"]
    wb = Workbook()
    generated: set[tuple[str, str]] = set()

    def formula(ws, row, col, text):
        c = ws.cell(row, col, text)
        generated.add((ws.title, c.coordinate))
        return c

    # ---- ReadMe
    ws = wb.active
    ws.title = t["readme"]
    ws["A1"] = "Mini-Grid Financial Model Studio"
    ws["A1"].font = Font(bold=True, size=14)
    meta = [
        ("Projet" if lang == "fr" else "Project", project.get("name", "")),
        ("Pays" if lang == "fr" else "Country", project.get("country", "")),
        ("Scénario" if lang == "fr" else "Scenario", f"{version['scenario']} v{version['version']}"),
        ("Statut de la version" if lang == "fr" else "Version status", version.get("status", "")),
        ("Statut du calcul" if lang == "fr" else "Run status", run.get("status", "")),
        ("Options", ", ".join(res.get("options") or []) or t["none"]),
        ("Solveur" if lang == "fr" else "Solver", res.get("solver", {}).get("method", "")),
        ("Version du moteur" if lang == "fr" else "Engine version", run.get("engine_version", ENGINE_VERSION)),
        ("Empreinte des hypothèses" if lang == "fr" else "Inputs hash", run.get("inputs_hash", "")),
        ("Calcul" if lang == "fr" else "Run", run.get("id", "")),
        ("Exporté le" if lang == "fr" else "Exported on", dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(i, 1, k).font = HEAD
        ws.cell(i, 2, v)
    for i, line in enumerate(README[lang], start=len(meta) + 4):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70

    # ---- Assumptions (written first so formulas can reference the WACC)
    wa = wb.create_sheet(t["assumptions"])
    entries = {e["path"]: e for e in field_entries()}
    for c, h in enumerate((t["path"], t["label"], t["input"], t["unit"], t["cell"], t["doc"]), 1):
        wa.cell(1, c, h).font = HEAD
        wa.cell(1, c).fill = HEAD_FILL
    wacc_ref = None
    r = 2
    for path, value in _flatten({k: v for k, v in version["assumptions"].items() if k != "lots"}):
        base = path
        parts = path.split(".")
        if len(parts) > 2:
            base = ".".join(parts[:2])
        e = entries.get(base, {})
        sub = parts[2] if len(parts) > 2 and not parts[2].isdigit() else ""
        unit, cell_ref = e.get("unit", ""), e.get("cell", "")
        if sub in ("basis", "vat"):          # MoneyInput descriptors: no unit, same workbook row
            unit = ""
        elif sub == "lcy_value" and cell_ref.startswith("Currency!H"):
            cell_ref = "Currency!I" + cell_ref[len("Currency!H"):]
        if unit == "date" and isinstance(value, str):
            try:
                value = dt.date.fromisoformat(value[:10])
            except ValueError:
                pass
        wa.cell(r, 1, path)
        wa.cell(r, 2, lab["fields"].get(base, "") + (f" ({sub})" if sub else ""))
        cell = wa.cell(r, 3, value)
        cell.font = INPUT_FONT
        if isinstance(value, dt.date):
            cell.number_format = "dd/mm/yyyy"
        elif isinstance(value, (int, float)) and unit.startswith("%"):
            cell.number_format = "0.00%"
        wa.cell(r, 4, unit)
        wa.cell(r, 5, cell_ref)
        wa.cell(r, 6, e.get("doc", ""))
        if path == "funding.wacc":
            wacc_ref = f"'{t['assumptions']}'!$C${r}"
        r += 1
    for col, w in zip("ABCDEF", (38, 52, 18, 22, 18, 50)):
        wa.column_dimensions[col].width = w

    # ---- Annual
    wn = wb.create_sheet(t["annual"])
    series = res["series"]
    years = series["year"]["values"]
    n = len(years)
    last_col = get_column_letter(FIRST_YEAR_COL + n - 1)
    heads = (t["key"], t["label"], t["unit"], t["source"], t["anchor"])
    for c, h in enumerate(heads, 1):
        wn.cell(1, c, h).font = HEAD
        wn.cell(1, c).fill = HEAD_FILL
    for j, y in enumerate(years):
        wn.cell(1, FIRST_YEAR_COL + j, y).font = HEAD
        wn.cell(1, FIRST_YEAR_COL + j).fill = HEAD_FILL
    row_of: dict[str, int] = {}
    r = 2
    for key, s in series.items():
        if key == "year":
            continue
        row_of[key] = r
        wn.cell(r, 1, key)
        wn.cell(r, 2, series_label(key, s["label"], lang))
        wn.cell(r, 3, s["unit"])
        wn.cell(r, 4, s["source"])
        fmt = _fmt(s["unit"])
        for j, v in enumerate(s["values"]):
            if key == "period_end" and v is not None:
                v = dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
            c = wn.cell(r, FIRST_YEAR_COL + j, v)
            c.font = ENGINE_FONT
            c.number_format = fmt
        r += 1
    # XIRR anchor (Names!F19 = 1e-7 at the model start date), as in Model!E375
    start = version["assumptions"].get("timing", {}).get("model_start")
    if "period_end" in row_of and start:
        wn.cell(row_of["period_end"], 5, dt.date.fromisoformat(str(start)[:10])).number_format = "dd/mm/yyyy"
    if "irr_cash_flow" in row_of:
        wn.cell(row_of["irr_cash_flow"], 5, -1e-7)
    r += 1
    wn.cell(r, 1, t["recalc_title"]).font = HEAD
    wn.cell(r, 5, t["total_gap"]).font = HEAD
    r += 1
    recalc_rows = {}
    for key, expr in RECALC:
        needed = [k for k in _names(expr)]
        if key not in row_of or any(k not in row_of for k in needed):
            continue
        wn.cell(r, 1, f"recalc.{key}")
        wn.cell(r, 2, f"{t['recalc']} : {series_label(key, series[key]['label'], lang)}")
        wn.cell(r, 3, series[key]["unit"])
        wn.cell(r, 4, expr.replace("{", "").replace("}", ""))
        for j in range(n):
            col = get_column_letter(FIRST_YEAR_COL + j)
            f = expr.format(**{k: f"{col}{row_of[k]}" for k in needed})
            c = formula(wn, r, FIRST_YEAR_COL + j, f"={f}")
            c.font = FORMULA_FONT
            c.number_format = _fmt(series[key]["unit"])
        rng_new = f"{get_column_letter(FIRST_YEAR_COL)}{r}:{last_col}{r}"
        rng_eng = f"{get_column_letter(FIRST_YEAR_COL)}{row_of[key]}:{last_col}{row_of[key]}"
        formula(wn, r, 5, f"=SUMPRODUCT(ABS({rng_new}-{rng_eng}))").number_format = "0.000E+00"
        recalc_rows[key] = r
        r += 1
    wn.column_dimensions["A"].width = 24
    wn.column_dimensions["B"].width = 46
    wn.column_dimensions["D"].width = 22
    wn.column_dimensions["E"].width = 16
    wn.freeze_panes = wn.cell(2, FIRST_YEAR_COL)

    # ---- KPIs
    wk = wb.create_sheet(t["kpis"], 1)
    heads = (t["key"], t["label"], t["value"], t["unit"], t["source"], t["status"], t["recalc"], t["gap"])
    for c, h in enumerate(heads, 1):
        wk.cell(1, c, h).font = HEAD
        wk.cell(1, c).fill = HEAD_FILL
    an = f"'{t['annual']}'!"
    first = get_column_letter(FIRST_YEAR_COL)

    def rng(key, from_anchor=False):
        c0 = "E" if from_anchor else first
        return f"{an}${c0}${row_of[key]}:${last_col}${row_of[key]}"

    formulas = {}
    if {"irr_cash_flow", "period_end"} <= row_of.keys():
        formulas["project_irr"] = f"=XIRR({rng('irr_cash_flow', True)},{rng('period_end', True)})"
    if {"npv_cash_flow", "period_end"} <= row_of.keys() and wacc_ref:
        formulas["project_npv"] = f"=XNPV({wacc_ref},{rng('npv_cash_flow')},{rng('period_end')})"
    # min/avg DSCR depend on the repayment flag, which is not exported: no formula rather than a wrong one
    for i, (key, k) in enumerate(res["kpis"].items(), start=2):
        wk.cell(i, 1, key)
        wk.cell(i, 2, kpi_label(key, k["label"], lang))
        c = wk.cell(i, 3, k["value"])
        c.font = ENGINE_FONT
        c.number_format = _fmt(k["unit"])
        wk.cell(i, 4, k["unit"])
        wk.cell(i, 5, k["source"])
        wk.cell(i, 6, t["experimental"] if k.get("experimental") else t["derived"] if k.get("derived") else "")
        if key in formulas:
            formula(wk, i, 7, formulas[key]).number_format = _fmt(k["unit"])
            formula(wk, i, 8, f"=G{i}-C{i}").number_format = "0.000E+00"
    for col, w in zip("ABCDEFGH", (30, 52, 18, 18, 30, 14, 18, 14)):
        wk.column_dimensions[col].width = w

    # ---- Lots
    wl = wb.create_sheet(t["lots"])
    lots = version["assumptions"].get("lots", [])
    wl.cell(1, 1, t["key"]).font = HEAD
    wl.cell(1, 2, t["label"]).font = HEAD
    for j, lot in enumerate(lots):
        wl.cell(1, 3 + j, lot.get("name") or f"Lot {j + 1}").font = HEAD
    for i, name in enumerate(lots[0].keys() if lots else [], start=2):
        wl.cell(i, 1, name)
        wl.cell(i, 2, lab["fields"].get(f"lots.<i>.{name}", name))
        for j, lot in enumerate(lots):
            wl.cell(i, 3 + j, lot.get(name)).font = INPUT_FONT
    wl.column_dimensions["A"].width = 24
    wl.column_dimensions["B"].width = 44

    # ---- Checks and warnings
    wc = wb.create_sheet(t["checks"])
    names = CHECKS_FR if lang == "fr" else CHECKS_EN
    for c, h in enumerate((t["key"], t["label"], t["value"], t["status"]), 1):
        wc.cell(1, c, h).font = HEAD
    failed = set(res.get("failed_checks", []))
    for i, (k, v) in enumerate(res.get("checks", {}).items(), start=2):
        wc.cell(i, 1, k)
        wc.cell(i, 2, names.get(k, k))
        wc.cell(i, 3, v)
        wc.cell(i, 4, t["fail"] if k in failed else t["ok"])
    wc.column_dimensions["B"].width = 44
    ww = wb.create_sheet(t["alerts"])
    for i, w in enumerate(res.get("warnings") or [t["none"]], start=1):
        ww.cell(i, 1, w).alignment = Alignment(wrap_text=True)
    ww.column_dimensions["A"].width = 120

    # Only the formulas written above may be formulas: any other text that Excel would read as one
    # (a lot or project name starting with '=', '+', '-', '@') is forced to plain text.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    if (ws.title, c.coordinate) in generated:
                        continue
                    if c.data_type == "f" or c.value[:1] in ("=", "+", "-", "@", "\t", "\r"):
                        c.data_type = "s"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _clean(o):
    """Drop characters that are illegal in worksheet XML from every text (user names, notes)."""
    if isinstance(o, str):
        return ILLEGAL_CHARACTERS_RE.sub("", o)
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, list):
        return [_clean(v) for v in o]
    return o


def _names(expr: str) -> list[str]:
    import re
    return re.findall(r"\{(\w+)\}", expr)
