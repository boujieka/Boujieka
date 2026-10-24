"""Cell-by-cell reconciliation of an engine run against workbook values.

The reference is either the Excel values cached in a saved workbook (scenario
S1 = the published file) or a JSON dump of cell values ("Sheet!row,col" ->
value, as written by tools/phase0 LibreOffice runs).
"""
from __future__ import annotations

import datetime as dt
import json
import math
from dataclasses import dataclass, field

from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as CL

from .model import T, Ledger

EPOCH = dt.datetime(1899, 12, 30)
FIRST_COL = 10  # J
TABLE_FIRST_COL = 9  # I (Inputs tables start in column I)


def _num(v):
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, dt.datetime):
        return (v - EPOCH).total_seconds() / 86400
    if isinstance(v, dt.time):
        return 0.0
    return v


@dataclass
class _Missing:
    pass


MISSING = _Missing()  # the reference holds no value for this cell (not compared)


@dataclass
class Mismatch:
    key: str
    cell: str
    expected: object
    actual: object


@dataclass
class Report:
    compared: int = 0
    equal: int = 0
    mismatches: list[Mismatch] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.mismatches

    def by_row(self) -> dict[str, list[Mismatch]]:
        out: dict[str, list[Mismatch]] = {}
        for m in self.mismatches:
            out.setdefault(m.key, []).append(m)
        return out

    def summary(self, limit: int = 40) -> str:
        lines = [f"compared {self.compared} cells, equal {self.equal}, mismatching {len(self.mismatches)} "
                 f"in {len(self.by_row())} rows/cells"]
        for key, ms in list(self.by_row().items())[:limit]:
            m = ms[0]
            lines.append(f"  {key}: {len(ms)} cell(s), first {m.cell} expected {m.expected!r} got {m.actual!r}")
        return "\n".join(lines)


def close(expected, actual, rel=1e-9, abs_=1e-6) -> bool:
    e, a = _num(expected), _num(actual)
    if e is None:
        e = 0.0
    if isinstance(a, bool):
        a = float(a)
    if isinstance(e, (int, float)) and isinstance(a, (int, float)):
        if math.isnan(a):
            return False
        return abs(e - a) <= max(abs_, rel * abs(e))
    return str(e) == str(a)


def reconcile(L: Ledger, value_of, sheets=("Model", "Debt", "FinStat"), rel=1e-9, abs_=1e-6,
              skip_rows: set[str] = frozenset()) -> Report:
    """``value_of(sheet, row, col)`` returns the reference value of a cell."""
    rep = Report()

    def cmp(key, cell, expected, actual):
        if expected is MISSING:
            return
        rep.compared += 1
        if close(expected, actual, rel, abs_):
            rep.equal += 1
        else:
            rep.mismatches.append(Mismatch(key, cell, expected, actual))

    for key, arr in L.v.items():
        sh, row = key.split("!")
        if sh not in sheets or key in skip_rows:
            continue
        row = int(row)
        for t in range(T):
            cmp(key, f"{sh}!{CL(FIRST_COL + t)}{row}", value_of(sh, row, FIRST_COL + t), arr[t])
    for key, val in L.totals.items():
        sh, row = key.split("!")
        if sh not in sheets or key in skip_rows:
            continue
        cmp(key + " (col I)", f"{sh}!I{row}", value_of(sh, int(row), 9), val)
    for key, val in L.s.items():
        if "!" not in key or key.split("!")[0] not in sheets + ("Inputs", "Currency", "Portfolio", "Dashboard"):
            continue
        sh, addr = key.split("!")
        col = "".join(c for c in addr if c.isalpha())
        row = "".join(c for c in addr if c.isdigit())
        if not row or not col or isinstance(val, bool) and key.endswith("pending"):
            continue
        cmp(key, key, value_of(sh, int(row), CI(col)), val)
    return rep


def reconcile_tables(tables: dict, value_of, rel=1e-9, abs_=1e-6) -> Report:
    rep = Report()
    for key, arr in tables.items():
        sh, row = key.split("!")
        for i, val in enumerate(arr):
            exp = value_of(sh, int(row), TABLE_FIRST_COL + i)
            rep.compared += 1
            if close(exp, val, rel, abs_):
                rep.equal += 1
            else:
                rep.mismatches.append(Mismatch(key, f"{sh}!{CL(TABLE_FIRST_COL + i)}{row}", exp, val))
    return rep


def workbook_values(path: str):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    def value_of(sh, row, col):
        return _num(wb[sh].cell(row, col).value)
    return value_of


def json_values(path: str):
    d = json.load(open(path))

    def value_of(sh, row, col):
        k = f"{sh}!{row},{col}"
        if k not in d:
            return MISSING  # constants are not in the dump
        v = d[k]
        return None if v == "" else v
    return value_of
